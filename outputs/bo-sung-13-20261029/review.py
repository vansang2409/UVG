"""Rà độc lập file bổ sung với nguồn, kế hoạch đã chốt và liên kết mới."""
from pathlib import Path
from collections import defaultdict, Counter
from decimal import Decimal
from datetime import datetime
import hashlib, json, sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
from lxml import etree

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_invoice_reports import rows, invoice_key, amount, invoice_day
HERE=Path(__file__).resolve().parent
TARGET=HERE/'IMPORT_DIEU_CHINH_BO_SUNG_13_HD.xlsx'
OLD=ROOT/'scripts/data/apply-13-20261019'
plan=json.loads((OLD/'plan.json').read_text(encoding='utf-8'))
src=json.loads((OLD/'sources.json').read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
before=sha(TARGET)
for path,h in src['hashes'].items():assert sha(ROOT/path)==h,('Nguồn thay đổi',path)
wb=load_workbook(TARGET,read_only=True,data_only=False)
ws=wb['Import dieu chinh'];groups=defaultdict(list);allrows=[]
for rn,row in enumerate(ws.iter_rows(min_row=10,max_col=24),10):
    assert all(c.data_type not in ('e','f') for c in row),('Lỗi/công thức',rn)
    v=[c.value for c in row];groups[v[0]].append((rn,v));allrows.append(v)
assert set(groups)==set(range(2033,2046)) and len(allrows)==206
actions={a['seq']:a for a in plan['actions']}
assert set(groups)==set(actions)
rounding=[];roots=defaultdict(list)
for seq,g in groups.items():
    a=actions[seq];root=a['root'];target=a['target'];roots[root].append(seq)
    assert len(g)==len(a['lines'])
    assert [g[0][1][c] for c in [19,21,22]]==a['money']
    assert sum(Decimal(str(v[18])) for _,v in g if v[18] is not None)==Decimal(str(a['money'][0]))
    assert a['money'][0]+a['money'][1]==a['money'][2]
    for index,(rn,v) in enumerate(g):
        assert v[1]==datetime(2026,10,19)
        assert '/'.join(v[9:11])==root and isinstance(v[10],str) and len(v[10])==8
        assert invoice_day(v[11]).isoformat()==src['headers'][root]['date'][:10]
        assert v[2]==v[5] and v[6] is None and v[7]=='TM/CK'
        assert v[17] is None or v[17]>0
        if index:assert all(v[c] is None for c in [19,20,21,22])
        else:assert v[20]==8
        assert v[13].startswith('Điều chỉnh tăng' if a['money'][2]>0 else 'Điều chỉnh giảm')
        line=a['lines'][index]
        for c,k in [(14,'name'),(15,'unit'),(16,'qty'),(17,'price'),(18,'goods')]:assert v[c]==line.get(k),(seq,rn,k)
        if target==root:
            original=src['details'][target][index]
            assert v[14]==original['name'] and v[15]==original['unit']
            assert v[16]==-original['qty'] and v[17]==original['price']
            assert v[18]==(-original['goods'] if original['goods'] is not None else None)
        else:assert len(g)==1 and v[16:18]==[None,None]
        if all(v[c] is not None for c in [16,17,18]):
            delta=Decimal(str(v[16]))*Decimal(str(v[17]))-Decimal(str(v[18]))
            if delta:rounding.append({'seq':seq,'row':rn,'delta':str(delta),'quantity':v[16],'price':v[17],'line_amount':v[18]})
assert len(rounding)==2 and {r['row'] for r in rounding}=={75,125}
# So toàn bộ 24 cột với block kế hoạch lịch sử, khác trống/chuỗi rỗng được chuẩn hóa.
block=next(b for f in plan['files'] if f['kind']=='compact' for b in f['sheets']['Import dieu chinh']['blocks'] if b['range']=='A41800:X42005')
expected=[]
for r in block['values']:
    r=[None if x=='' else x for x in r]
    r[1]=datetime.fromisoformat(r[1]);expected.append(r)
assert allrows==expected
# Đọc lại trực tiếp quan hệ mới; đếm một lần cho từng số hóa đơn.
children=defaultdict(set);newmoney={};link_observations=0
link_files=sorted((ROOT/'data/Meinv/HD_bi_dieu_chinh_2024_2026').glob('*.xls'))
for p in link_files:
    for rn,r in rows(p):
        if rn<7 or amount(r.get('A')) is None:continue
        old=invoice_key(r.get('B'),r.get('D'));new=invoice_key(r.get('M'),r.get('O'))
        if not old or not new:continue
        children[old].add(new);link_observations+=1
        for key,cols in [(old,['J','K','L']),(new,['U','V','W'])]:
            m=[amount(r.get(c)) for c in cols]
            assert key not in newmoney or newmoney[key]==m
            newmoney[key]=m
proof=[]
for check in plan['root_checks']:
    root=tuple(check['root'].split('/'));keys=set();queue=[root]
    while queue:
        k=queue.pop()
        if k in keys:continue
        keys.add(k);queue.extend(children[k])
    assert set(check['source_invoices'])=={'/'.join(k) for k in keys}
    assert not any(k[0].startswith('1C26') for k in keys),('Đã có hóa đơn 2026 trong chuỗi',root)
    values=[];basis=[]
    for k in keys:
        text='/'.join(k);m=list(newmoney[k])
        if text in src['xmls']:
            xp=ROOT/src['xmls'][text]['file'];tree=etree.parse(str(xp))
            xm=[Decimal(tree.findtext('.//'+n)) for n in ['TgTCThue','TgTThue','TgTTTBSo']]
            assert xm==[Decimal(str(x)) for x in src['xmls'][text]['money']]
            for i in range(3):
                if m[i] is None:m[i]=xm[i];basis.append({'key':text,'field':i,'basis':'XML','value':str(xm[i])})
        if text=='1C24TUV/00000574' and m[2] is None:
            m[2]=Decimal(0);basis.append({'key':text,'field':2,'basis':'Quy ước người dùng đã chốt, ô nguồn vẫn trống','value':'0'})
        assert all(x is not None for x in m),(text,m)
        values.append(m)
    supplements=[actions[seq]['money'] for seq in check['sequences']]
    balance=[sum(v[c] for v in values+supplements) for c in range(3)]
    assert balance==[0,0,0],(root,balance)
    proof.append({'root':check['root'],'sequences':check['sequences'],'after':[str(x) for x in balance],'blank_resolution':basis})
assert sha(TARGET)==before
result={'passed':True,'file_sha256':before,'invoices':len(groups),'rows':len(allrows),'roots':len(roots),'all_4944_cells_match_approved_plan':True,'source_hashes_verified':len(src['hashes']),'raw_link_rows':link_observations,'eight_chains_zero_after_supplement':proof,'rounding_preserved':rounding,'zero_amount_invoice':2037,'positive_reversal_invoice':2034,'invoice_date':'2026-10-19','source_unchanged':True,'software_import_tested':False}
(HERE/'deep-review.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2));wb.close()
