from pathlib import Path
import json, hashlib, shutil, datetime
from openpyxl import load_workbook

root=Path(__file__).parent
source=root/'ĐỐI SOÁT ĐÃ GỘP.xlsx'
validation=json.loads((root/'validation.json').read_text(encoding='utf-8'))
assert validation['status']=='PASS'
assert validation['file_size']==source.stat().st_size
wb=load_workbook(source,read_only=True,data_only=True)
assert wb.sheetnames==['Hóa đơn gộp']
ws=wb.active
assert ws.max_row in (None,26551) and ws.max_column in (None,50)
assert ws['A2'].value==datetime.datetime(2025,7,1)
assert ws['B2'].value=='10299' and ws['H2'].value=='250626J17FHTEK'
assert ws['M2'].number_format=='#,##0'
assert ws['P2'].number_format=='#,##0.0000'
assert ws['A2'].font.name=='Arial'
wb.close()
destination=Path(r'C:\Users\thuong ngo\Downloads\ĐỐI SOÁT ĐÃ GỘP_12-10-2026.xlsx')
with source.open('rb') as src,destination.open('xb') as dest:shutil.copyfileobj(src,dest)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(source)==sha(destination)
result={'path':str(destination),'bytes':destination.stat().st_size,'rows':26550,'sheets':1,'sha256':sha(destination),'native_reader_check':'PASS'}
(root/'final_delivery.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
