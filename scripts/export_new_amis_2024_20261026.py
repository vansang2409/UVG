"""Xuất dự phòng khi Artifact Tool không hoàn tất workbook lớn trong runtime.

Dữ liệu và bố cục đã được dựng/kiểm tra bằng Artifact Tool. Dùng XlsxWriter
để đóng gói dữ liệu lớn, rồi kiểm độc lập bằng trình đọc khác.
"""
from pathlib import Path
from datetime import datetime
from io import BytesIO
from math import ceil
import json
import xlsxwriter
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/new-amis-2024-20261026'
OUT=ROOT/'outputs/amis-2024-moi-20261026'


def day(v):
    return datetime.fromisoformat(v) if v else None


def main():
    p=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
    template=load_workbook(BytesIO((ROOT/'data/Du lieu Meinvoice/Mau_hoa_don_GTGT_MTT_1_thue_suat.xls').read_bytes()),read_only=True,data_only=True).active
    template_rows=list(template.iter_rows(min_row=1,max_row=9,values_only=True))
    ready=[d for d in p['docs'] if d['category']=='Đưa vào import']
    pending=[d for d in p['docs'] if d['category']=='Cần kiểm tra']
    excluded=[d for d in p['docs'] if d['category']=='Đã loại']
    OUT.mkdir(parents=True,exist_ok=True)

    def create(filename):
        target=OUT/filename
        if target.exists():
            raise FileExistsError(target)
        w=xlsxwriter.Workbook(target,{'constant_memory':True,'strings_to_formulas':False,'strings_to_urls':False,'default_date_format':'dd/mm/yyyy'})
        w.set_calc_mode('auto')
        base={'font_name':'Arial','font_size':10,'valign':'vcenter'}
        fmt={
            'body':w.add_format(base),
            'wrap':w.add_format({**base,'text_wrap':True}),
            'text':w.add_format({**base,'num_format':'@'}),
            'money':w.add_format({**base,'num_format':'#,##0;[Red](#,##0);0'}),
            'decimal':w.add_format({**base,'num_format':'#,##0.00########;[Red](#,##0.00########);0.00'}),
            'date':w.add_format({**base,'num_format':'dd/mm/yyyy'}),
            'header':w.add_format({**base,'bold':True,'bg_color':'#24435C','font_color':'white','text_wrap':True,'align':'center'}),
            'input':w.add_format({**base,'bg_color':'#FFF2CC','text_wrap':True}),
            'total':w.add_format({**base,'bold':True,'top':1,'num_format':'#,##0;[Red](#,##0);0'}),
            'note':w.add_format({**base,'text_wrap':True}),
        }
        return w,fmt

    def write_sheet(w,f,name,headers,rows,widths,header_row=0,dates=(),money=(),decimals=(),text=(),wrap=(),inputs=(),notes=None):
        s=w.add_worksheet(name)
        s.hide_gridlines(2);s.set_default_row(24)
        for c,width in enumerate(widths):
            cf=f['date'] if c in dates else f['money'] if c in money else f['decimal'] if c in decimals else f['text'] if c in text else f['wrap'] if c in wrap else f['body']
            s.set_column(c,c,width,cf)
        if notes:
            for r,content in enumerate(notes):
                if content:
                    s.merge_range(r,0,r,len(headers)-1,content,f['note']);s.set_row(r,24)
        s.write_row(header_row,0,headers,f['header']);s.set_row(header_row,44)
        count=0
        for count,row in enumerate(rows,1):
            rr=header_row+count
            height=24
            for c,v in enumerate(row):
                cf=f['input'] if c in inputs else f['date'] if c in dates else f['money'] if c in money else f['decimal'] if c in decimals else f['text'] if c in text else f['wrap'] if c in wrap else f['body']
                if v is None or v=='':s.write_blank(rr,c,None,cf)
                elif isinstance(v,datetime):s.write_datetime(rr,c,v,cf)
                elif isinstance(v,(int,float)):s.write_number(rr,c,v,cf)
                else:s.write_string(rr,c,str(v),cf)
                if c in wrap and v:
                    height=max(height,15*ceil(len(str(v))/max(8,widths[c]))+6)
            if height>24:s.set_row(rr,height)
        s.autofilter(header_row,0,header_row+count,len(headers)-1)
        s.freeze_panes(header_row+1,2)
        return s

    w,f=create('IMPORT_HOA_DON_MOI_AMIS_2024.xlsx')
    headers=list(template_rows[8])+['Số chứng từ AMIS','Tổng tiền hàng','Tổng tiền thanh toán','Dòng nguồn AMIS','Mã hàng hóa/dịch vụ']
    notes=[r[0] for r in template_rows[:7]]
    notes[6]='- Dữ liệu AMIS 2024 sau loại trừ; ngày lập theo Ngày chứng từ, thanh toán TM/CK theo xác nhận.'
    def import_rows():
        for n,d in enumerate(ready,1):
            for j,l in enumerate(d['lines']):
                r=l['values'];first=j==0
                yield [n,day(d['date']) if first else None,d['buyer'] if first else None,d['tax'] if first else None,d['address'] if first else None,d['buyer'] if first else None,None,None,None,'TM/CK' if first else None,d['rate'] if first else None,d['sums']['vat'] if first else None,r[7],r[16],r[17],r[19],r[20],d['doc'],d['sums']['goods'] if first else None,d['sums']['payment'] if first else None,l['row'],r[12]]
    write_sheet(w,f,'Hóa đơn GTGT',headers,import_rows(),[14,15,38,18,55,38,25,20,22,23,18,23,86,12,15,23,23,27,25,26,18,27],header_row=8,dates=[1],money=[11,16,18,19],decimals=[14,15],text=[3,17,21],wrap=[2,4,5,12],notes=notes)
    guide=[
        ['Nhập khẩu','Chọn sheet Hóa đơn GTGT, dòng tiêu đề 9. Mẫu GTGT máy tính tiền, một thuế suất do bạn cung cấp.'],
        ['Gom hóa đơn','Mỗi số chứng từ AMIS là một hóa đơn; các dòng cùng số thứ tự thuộc cùng hóa đơn.'],
        ['Ngày lập','Lấy Ngày chứng từ AMIS theo yêu cầu. Ngày hạch toán và Ngày hóa đơn nguồn được giữ trong Doi chieu import của file kiểm tra.'],
        ['Thanh toán','TM/CK theo xác nhận của bạn.'],
        ['Thuế suất','8% suy ra duy nhất từ tiền hàng/VAT, chênh làm tròn không quá 1đ/dòng. Giữ nguyên số tiền nguồn.'],
        ['Dòng đầu hóa đơn','Ngày, người mua, thanh toán, thuế suất và tổng VAT chỉ ghi tại dòng đầu của mỗi hóa đơn, đúng bố cục mẫu.'],
        ['Người mua','Giữ tên, địa chỉ và MST từ AMIS, kể cả tên khách hàng dùng chung theo sàn. Email, điện thoại, CCCD trống vì nguồn không có.'],
        ['Cột bổ sung','Giữ đúng 17 cột A:Q. Cột R:V dùng đối chiếu chứng từ, tổng tiền, dòng nguồn và mã hàng.'],
        ['Loại trừ','Lọc BHH và các mã tách/gộp ghi tại tab HD cty 2024. Không coi toàn bộ tab F0/F1/F2 là danh sách loại.'],
        ['Trả lại','Giữ ngoài import phiếu trả lại và mọi đơn bán có thể liên quan. Xem file kiểm tra riêng.'],
        ['Phạm vi','Chưa kiểm chứng trạng thái mọi hóa đơn cũ trên meInvoice; đây là dữ liệu AMIS sau lọc theo danh sách bạn cung cấp.'],
        ['Nguồn','So_chi_tiet_ban_hang_2024.xlsx; HD Xuất CTY 2024.xlsx; Mau_hoa_don_GTGT_MTT_1_thue_suat.xls.'],
        ['MISA','https://helpv4.meinvoice.vn/kb/nhap-khau-hoa-don-khoi-tao-tu-may-tinh-tien-vao-phan-mem/'],
        ['Trạng thái','Chưa nhập, ký hoặc phát hành.'],
    ]
    write_sheet(w,f,'Huong dan',['Mục','Nội dung'],guide,[30,125],wrap=[1])
    w.close();print('SAVED import',flush=True)

    w,f=create('KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx')
    s=w.add_worksheet('Tong hop');s.hide_gridlines(2)
    widths=[30,20,23,28,26,25,29]
    for c,width in enumerate(widths):s.set_column(c,c,width,f['body'])
    s.write_row(0,0,['Phân nhóm','Số chứng từ','Số dòng AMIS','Doanh số bán (VND)','Trả lại (VND)','VAT (VND)','Thanh toán (VND)'],f['header']);s.set_row(0,44)
    for i,(category,values) in enumerate(p['summary'].items(),1):
        s.write_string(i,0,category,f['body'])
        for c,k in enumerate(['documents','lines','goods','returns','vat','payment'],1):s.write_number(i,c,values[k],f['money'])
        s.set_row(i,24)
    s.write_string(4,0,'Tổng sổ AMIS',f['total'])
    for c,k in enumerate(['documents','lines','goods','returns','vat','payment'],1):s.write_formula(4,c,f'=SUM({chr(65+c)}2:{chr(65+c)}4)',f['total'],sum(v[k] for v in p['summary'].values()))
    s.set_row(4,24)
    summary_notes=[
        ['Cách sử dụng','Lọc Can kiem tra theo Lý do. Điền quyết định và BHH đúng vào hai cột vàng cuối bảng.'],
        ['Trả lại','15.879 đơn bán là ứng viên có thể liên quan phiếu trả lại, chưa kết luận đã trả. Chưa đưa vào import.'],
        ['Điều kiện dò trả lại','Cùng mã KH/mã hàng/đơn giá ±1đ, bán trước hoặc cùng ngày trả, đủ SL. Không tự chọn đơn để bù trừ.'],
        ['Ngày và thanh toán','Ngày lập theo Ngày chứng từ AMIS. Thanh toán TM/CK theo xác nhận.'],
        ['Thuế suất','AMIS không có cột thuế suất. 8% ở file import suy ra từ tiền hàng/VAT, giữ nguyên tiền nguồn.'],
        ['Loại chưa rõ','HĐ thiếu BHH, mã cũ và HĐ chỉ có ở bảng ghi nhận được liệt kê riêng tại Loai chua ro.'],
        ['Nguồn','So_chi_tiet_ban_hang_2024.xlsx và HD Xuất CTY 2024.xlsx; giữ nguyên bản gốc.'],
        ['Phạm vi','Tab F0/F1/F2 chỉ là lịch sử tham khảo. Chưa kiểm chứng toàn bộ lịch sử phát hành trên meInvoice.'],
    ]
    for i,(label,content) in enumerate(summary_notes,7):
        s.write_string(i,0,label,f['wrap']);s.merge_range(i,1,i,6,content,f['note']);s.set_row(i,36)
    common=['Số chứng từ AMIS','Ngày chứng từ','Mã khách hàng','Tên khách hàng','MST','Doanh số bán (VND)','Trả lại (VND)','VAT (VND)','Thanh toán (VND)','Số dòng','Dòng AMIS','Lý do','Quyết định của bạn','BHH đúng/liên quan']
    def docs_rows(ds):
        for d in ds:
            yield [d['doc'],day(d['date']),d['customer_code'],d['buyer'],d['tax'],d['sums']['goods'],d['sums']['returns'],d['sums']['vat'],d['sums']['payment'],len(d['lines']),', '.join(map(str,d['rows'])),'; '.join(d['reasons']),None,None]
    widths=[27,18,21,42,18,25,24,24,26,13,32,110,55,36]
    write_sheet(w,f,'Can kiem tra',common,docs_rows(pending),widths,dates=[1],money=[5,6,7,8],text=[0,2,4,13],wrap=[3,11],inputs=[12,13])
    def pending_rows():
        for d in pending:
            for l in d['lines']:
                yield [l['row']]+[day(v) if i in (0,1,4) else v for i,v in enumerate(l['values'])]
    raw_widths=[15]+[70 if i in [6,7,9,10,13,14] else 18 if i in [0,1,4] else 23 for i in range(42)]
    write_sheet(w,f,'Chi tiet can kiem tra',['Dòng AMIS']+p['headers'],pending_rows(),raw_widths,dates=[1,2,5],money=[21,25,26,28,29,31,33,34,35,40],decimals=[18,19,20,27,37],text=[3,4,6,9,12,13],wrap=[7,8,10,11,14,15])
    write_sheet(w,f,'Da loai',common,docs_rows(excluded),widths,dates=[1],money=[5,6,7,8],text=[0,2,4,13],wrap=[3,11])
    unresolved=[]
    for e in p['exclusions']:
        if not e['codes']:unresolved.append([e['key'],f"HD cty 2024 dòng {e['row']}",None,e['values'][4],'Thiếu BHH: chưa tự loại đơn theo tên/MST',None])
        for code in e['codes']:
            if code in p['exclusion_codes_not_found']:
                note='Mã cũ đã gộp vào BHHN2024000078; đã loại mã mới' if code=='BHHN2024000079' else 'Mã thuộc năm 2023, không có trong sổ AMIS 2024'
                unresolved.append([e['key'],f"HD cty 2024 dòng {e['row']}",code,e['values'][4],note,None])
    seen=set()
    for row,key,buyer,tax,order in p['extra_invoices']:
        if key not in seen:
            seen.add(key);unresolved.append([key,f'HĐ ghi nhận UVG 2024 dòng {row}',None,buyer,f'Không có trong tab HD cty 2024; thiếu BHH. Mã đơn ghi nhận: {order or ""}',None])
    write_sheet(w,f,'Loai chua ro',['Hóa đơn','Vị trí nguồn','BHH','Tên trong nguồn','Nội dung cần kiểm tra','BHH đúng do bạn bổ sung'],unresolved,[29,49,28,65,100,38],wrap=[3,4],inputs=[5])
    def back_rows():
        for i,d in enumerate(ready,1):
            yield [i,d['doc'],day(d['date']),d['buyer'],len(d['lines']),', '.join(map(str,d['rows'])),d['sums']['goods'],d['sums']['vat'],d['sums']['payment'],d['rate'],day(d['accounting_date']),day(d['invoice_date_source'])]
    write_sheet(w,f,'Doi chieu import',['STT hóa đơn','Số chứng từ AMIS','Ngày chứng từ','Tên khách hàng','Số dòng','Dòng nguồn AMIS','Tiền hàng (VND)','VAT (VND)','Thanh toán (VND)','Thuế suất suy ra (%)','Ngày hạch toán AMIS','Ngày hóa đơn nguồn AMIS'],back_rows(),[17,28,19,45,14,36,26,25,27,23,22,25],dates=[2,10,11],money=[6,7,8],text=[1],wrap=[3])
    links=([doc,row,len(candidates),f'Lọc cột Lý do của Can kiem tra theo {doc}; đơn trong Da loai vẫn được giữ ngoài import.'] for doc,row,candidates in p['return_evidence'])
    write_sheet(w,f,'Phieu tra lai',['Phiếu trả lại','Dòng AMIS','Số dòng bán có thể liên quan','Cách tìm BHH bán cần kiểm tra'],links,[26,18,28,130],wrap=[3])
    w.close();print('SAVED review',flush=True)
    (WORK/'export-fallback.json').write_text(json.dumps({'authoring':'Artifact Tool đã dựng model và render đủ 7 sheet kiểm tra; xuất workbook lớn không hoàn tất, dừng tiến trình riêng sau chờ kéo dài và bộ nhớ tăng ~3.7GB','final_packaging':'Bundled XlsxWriter; kiểm độc lập bằng openpyxl chỉ đọc','template_columns':17,'template_header_row':9},ensure_ascii=False,indent=2),encoding='utf8')


if __name__=='__main__':main()
