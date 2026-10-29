"""Đối chiếu toàn bộ workbook đã lưu với nguồn ghép, chỉ đọc Excel."""
import json, sys, hashlib, re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
DIR=ROOT/'scripts/data/h2-chain-readiness-20261022'
OUT=ROOT/'outputs/HOA_DON_GOC_T7_T12_2025_F0_F1_F2.xlsx'
data=json.loads((DIR/'workbook-input.json').read_text(encoding='utf-8'))
checks=json.loads((DIR/'expected-workbook.json').read_text(encoding='utf-8'))
wb=load_workbook(OUT,read_only=True,data_only=True)
assert wb.sheetnames==['F0, F1, F2, ...','Chi tiet nguon']
main=wb.worksheets[0];source=wb.worksheets[1]
assert main.max_row==37093 and source.max_row==37764
assert len(data['chains'])==37087
blank_cells=0;flagged_chains=0;blank_current=0;keys=set()
for row,c,expected in zip(main.iter_rows(min_row=7,values_only=True),data['chains'],checks['expected'],strict=True):
    inv=c['invoices'];rt=inv[0];key=tuple(rt['key']);assert key not in keys;keys.add(key)
    assert row[0]==len(keys) and row[4]==len(inv)-1
    for f,v in enumerate(inv):
        at=[6,13,20][f]
        assert list(row[at:at+2])==v['key'],(key,f,'identifier')
        assert isinstance(row[at+1],str) and len(row[at+1])==8
        assert isinstance(row[at+2],datetime) and row[at+2].strftime('%Y-%m-%d')==v['date']
        assert list(row[at+3:at+6])==v['money'],(key,f,'money')
        if v['money'][2] is None:blank_cells+=1
    flag=''.join(f+'; ' for f in expected['blank'])
    assert (row[5] or '')==flag,(key,'flag',row[5],flag)
    if flag:flagged_chains+=1
    for m,exp in enumerate(expected['total']):
        actual=row[26+m]
        if exp is None: assert actual in (None,''),(key,'blank_total',actual)
        else: assert isinstance(actual,(int,float)) and abs(actual-exp)<0.000001,(key,'total',actual,exp)
    if row[28] in (None,''):blank_current+=1
    assert all(not isinstance(v,str) or not v.startswith(('#REF!','#VALUE!','#DIV/0!','#NAME?','#NUM!')) for v in row)
assert blank_cells==436
raw_expected=[(c,v,f) for c in data['chains'] for f,v in enumerate(c['invoices'])]
raw_keys=set()
for row,(c,v,f) in zip(source.iter_rows(min_row=7,values_only=True),raw_expected,strict=True):
    assert list(row[:2])==c['root']
    assert row[2]=='F'+str(f) and list(row[3:5])==v['key']
    assert list(row[12:15])==v['money']
    assert row[18]==v['source'].replace('\\','/') and row[19]==v['row']
    if v['relationship']: assert list(row[7:9])==v['relationship']['parent']
    raw_keys.add(tuple(v['key']))
assert len(raw_keys)==37758
wb.close()
structure=load_workbook(OUT,data_only=False,read_only=True)
assert structure.worksheets[0]['H7'].number_format=='00000000'
assert structure.worksheets[0]['I7'].number_format=='dd/mm/yyyy'
assert structure.worksheets[0]['AA7'].data_type=='f'
structure.close()
with ZipFile(OUT) as z:
    for name in ('xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml'):
        xml=z.read(name)
        assert b'topLeftCell="A7"' in xml and b'state="frozen"' in xml
        assert len(re.findall(rb'<tablePart\s',xml))==1
        assert b'<conditionalFormatting ' in xml
report=dict(root_count=len(keys),source_invoices=len(raw_keys),blank_payment_cells=blank_cells,
    flagged_chains=flagged_chains,blank_current_payment=blank_current,all_rows_verified=True,
    sha256=hashlib.sha256(OUT.read_bytes()).hexdigest(),path=str(OUT))
(DIR/'workbook-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
sys.stdout.reconfigure(encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))
