from pathlib import Path
import os
from openpyxl import load_workbook
from collections import defaultdict, Counter
from datetime import datetime
from decimal import Decimal

root=Path(os.environ.get('UVG_DATA_DIR', str(Path(__file__).resolve().parents[1] / 'data')))
out=Path(os.environ.get('UVG_OUTPUT_DIR', str(Path(__file__).resolve().parents[1] / 'outputs')))/'uvg-doi-chieu'
out.mkdir(parents=True,exist_ok=True)
def n(v):
    return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
def money(v): return f'{v:,.0f}'.replace(',','.')
amis=defaultdict(lambda: {'docs':set(),'lines':0,'sales':Decimal(0),'vat':Decimal(0),'total':Decimal(0),'returns':0})
source=[]
controls=[]
for p in sorted((root/'Du lieu Amis').glob('*.xlsx')):
    s=load_workbook(p,read_only=True,data_only=True).worksheets[0]
    totals=[Decimal(0)]*3
    for i,r in enumerate(s.values,1):
        if i<5: continue
        if isinstance(r[1],datetime) and r[2]:
            k=r[1].strftime('%Y-%m'); a=amis[k]
            a['docs'].add(str(r[2])); a['lines']+=1
            for label,col,j in [('sales',20,0),('vat',32,1),('total',34,2)]:
                a[label]+=n(r[col]); totals[j]+=n(r[col])
            a['returns']+=any(n(r[c])!=0 for c in (26,27,28))
        elif r[0]=='Tổng cộng':
            controls.append((p.name,totals[0]-n(r[20]),totals[1]-n(r[32]),totals[2]-n(r[34])))
    source.append(str(p))
used={}; detail=set(); edges=[]
for p in sorted((root/'Du lieu Meinvoice').glob('*.xls')):
    if p.name.startswith('Mau'): continue
    s=load_workbook(p.open('rb'),read_only=True,data_only=True).worksheets[0]
    source.append(str(p))
    for i,r in enumerate(s.values,1):
        if not isinstance(r[0],(int,float)): continue
        if p.name.startswith('Bang_ke') and i>=6:
            k=(str(r[1]),str(r[2])); used[k]={'row':r,'file':p.name,'line':i}
        elif p.name.startswith('Bảng kê chi tiết') and i>=6:
            detail.add((str(r[1]),str(r[2])))
        elif p.name.startswith('DS') and i>=7:
            edges.append({'type':'Điều chỉnh' if 'điều chỉnh' in p.name else 'Thay thế','old':(str(r[1]),str(r[3])),'new':(str(r[12]),str(r[14])),'base':n(r[11]),'delta':n(r[22]),'file':p.name,'line':i})
mei=defaultdict(lambda:{'count':0,'gross':Decimal(0),'excluded':Decimal(0),'net':Decimal(0),'blank_total':0})
for k,v in used.items():
    r=v['row']; m=datetime.strptime(r[3],'%d/%m/%Y').strftime('%Y-%m'); a=mei[m]; a['count']+=1
    a['gross']+=n(r[17]); a['blank_total']+=not isinstance(r[17],(int,float))
    if r[20] in ('Hóa đơn đã bị thay thế','Hoá đơn đã bị hủy'): a['excluded']+=n(r[17])
    else: a['net']+=n(r[17])
byold=defaultdict(list)
for e in edges: byold[e['old']].append(e)
negative=[]; repeat=[]; chains=[]
for k,es in byold.items():
    adj=[e for e in es if e['type']=='Điều chỉnh']
    if len(adj)>1: repeat.append((k,adj))
    if adj:
        balance=adj[0]['base']+sum((e['delta'] for e in adj),Decimal(0))
        if balance<0: negative.append((k,adj,balance))
for e in edges:
    if e['type']=='Thay thế' and e['new'] in byold: chains.append(e)
report=['# Đối chiếu UVG – bước 1','', 'Ngày phân tích: 06/10/2026. Chỉ đọc dữ liệu nguồn; chưa sửa hoặc phát hành hóa đơn.','',
'AMIS được dùng làm số liệu mục tiêu đã được khách hàng chốt theo xác nhận của người dùng. meInvoice là lịch sử phát hành trong các file cung cấp. Chưa có dữ liệu CQT hoặc mã đơn sàn.','',
'## Phạm vi và kiểm tra dữ liệu','',f'- AMIS: {sum(a["lines"] for a in amis.values()):,} dòng, {sum(len(a["docs"]) for a in amis.values()):,} số chứng từ tính riêng từng tháng. Không đồng nghĩa số đơn hàng hoặc hóa đơn cần phát hành.',
f'- meInvoice: {len(used):,} hóa đơn trong bảng tổng; {len(detail):,} khóa hóa đơn trong bảng chi tiết.',
f'- Thiếu chi tiết so với bảng tổng: {len(set(used)-detail)}; thiếu bảng tổng so với chi tiết: {len(detail-set(used))}.','',
'## Tổng hợp theo tháng','',
'AMIS phân kỳ theo ngày chứng từ. meInvoice phân kỳ theo ngày hóa đơn, bao gồm lần sửa hóa đơn kỳ cũ. Hai cột không cùng cơ sở thời gian; chênh lệch không phải số tiền cần xuất bổ sung. Tổng thanh toán meInvoice quy đổi dưới đây loại hóa đơn có trạng thái đã bị thay thế/đã hủy, giữ hóa đơn bị điều chỉnh và cộng các hóa đơn điều chỉnh hiện có. Đây là phép tính theo file, chưa xác minh XML, dữ liệu CQT hoặc chuỗi sửa đầy đủ.','',
'| Tháng | Chứng từ AMIS | Tổng thanh toán AMIS | Hóa đơn meInvoice | Tổng thanh toán meInvoice quy đổi | HĐ thiếu giá trị tổng |',
'|---|---:|---:|---:|---:|---:|']
for m in sorted(set(amis)|set(mei)):
    a=amis[m]; b=mei[m]
    report.append(f'| {m} | {len(a["docs"]):,} | {money(a["total"])} | {b["count"]:,} | {money(b["net"])} | {b["blank_total"]} |')
report+=['','Đơn vị tiền: VND. Dòng thiếu tổng thanh toán không được tự coi là hóa đơn 0 đồng; cần kiểm tra chi tiết/XML.','', '## Kiểm tra tổng dòng AMIS với dòng Tổng cộng','', '| File | Lệch doanh số | Lệch thuế | Lệch thanh toán |','|---|---:|---:|---:|']
for file,s,v,t in controls: report.append(f'| {file} | {money(s)} | {money(v)} | {money(t)} |')
report+=['','## Hóa đơn tổng tiền âm sau các điều chỉnh trong bảng liên kết','',
'Đây là danh sách cần xác minh, chưa kết luận sai phạm. Số dư = giá trị gốc tại bảng liên kết + các khoản điều chỉnh liệt kê.','',
'| Ký hiệu | Số hóa đơn | Giá trị gốc | Tổng điều chỉnh | Số dư | Nguồn và dòng |','|---|---|---:|---:|---:|---|']
for k,adj,balance in negative:
    refs='; '.join(f'{e["file"]}, dòng {e["line"]}' for e in adj)
    report.append(f'| {k[0]} | {k[1]} | {money(adj[0]["base"])} | {money(sum((e["delta"] for e in adj),Decimal(0)))} | {money(balance)} | {refs} |')
report+=['','## Hóa đơn điều chỉnh nhiều lần','']
for k,es in repeat: report.append(f'- {k[0]} / {k[1]}: '+', '.join(f'{e["new"][0]} / {e["new"][1]}' for e in es))
report+=['','## Chuỗi thay thế tiếp tục được sửa','']
for e in chains: report.append(f'- {e["old"][0]} / {e["old"][1]} → {e["new"][0]} / {e["new"][1]} → '+', '.join(f'{x["type"]}: {x["new"][0]} / {x["new"][1]}' for x in byold[e['new']]))
report+=['','## Khóa hóa đơn chưa khớp giữa các bảng','', 'Có trong tổng, không có trong chi tiết:','']
for k in sorted(set(used)-detail): report.append(f'- {k[0]} / {k[1]} ({used[k]["file"]}, dòng {used[k]["line"]})')
report+=['','Có trong chi tiết, không có trong tổng:','']
for k in sorted(detail-set(used)): report.append(f'- {k[0]} / {k[1]}')
report+=['','## Công việc tiếp theo','',
'1. Giải thích các dòng thiếu giá trị tổng và khóa hóa đơn chưa khớp trước khi dùng tổng meInvoice.',
'2. Lập bảng ánh xạ sản phẩm AMIS và meInvoice, ghép có điều kiện theo số lượng, tiền hàng, thuế và thời gian. Không coi ghép bằng số tiền đơn thuần là bằng chứng chắc chắn.',
'3. Các giao dịch có nhiều ứng viên ghép hoặc không ghép được giữ ở nhóm chưa xác định; không tự chuyển sang xuất mới.',
'4. Bổ sung dữ liệu CQT và lịch sử sửa sau tháng 6/2025 trước khi chốt phương án phát hành.','', '## Nguồn','']
for p in source: report.append('- '+p)
target=out/'UVG_DOI_CHIEU_BUOC_1.md'; target.write_text('\n'.join(report),encoding='utf-8')
print('REPORT',target)
print('CONTROLS',controls)
print('NEGATIVE',len(negative),'REPEAT',len(repeat),'CHAINS',len(chains),'BLANK_TOTALS',sum(v['blank_total'] for v in mei.values()))
