"""Đối chiếu nhóm doanh thu; chỉ đọc nguồn, không tạo chỉ định phát hành."""
from pathlib import Path
import os
from openpyxl import load_workbook
from collections import defaultdict, Counter
from datetime import datetime
from decimal import Decimal
import csv, re

BASE=Path(os.environ.get('UVG_DATA_DIR', str(Path(__file__).resolve().parents[1] / 'data'))); OUT=Path(os.environ.get('UVG_OUTPUT_DIR', str(Path(__file__).resolve().parents[1] / 'outputs')))/'uvg-doi-chieu-nhom'
OUT.mkdir(parents=True,exist_ok=True)
def n(v): return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
def fmt(v): return f'{v:,.0f}'.replace(',','.')
def family(name,sku=''):
    t=str(name or '').upper(); code=str(sku or '').upper()
    found=re.findall(r'(?:KA|KH|KM|KC)\s*[-_]?\s*\d{2,4}',code+' '+t)
    ids=sorted(set(re.sub(r'[\s_-]','',m) for m in found))
    if len(ids)!=1: return 'CHƯA XÁC ĐỊNH MODEL'
    isfilter=('FILTER' in code or bool(re.match(r'^\s*(MÀNG LỌC|LÕI LỌC|BỘ LỌC)',t)))
    return ('Lõi lọc ' if isfilter else 'Máy/thiết bị ')+ids[0]
def channel(text):
    t=str(text or '').upper()
    if 'TIKTOK' in t: return 'TikTok'
    if any(x in t for x in ('SHOPEE','SHOPPE','KHÁCH HÀNG LẺ SP')): return 'Shopee'
    if 'LAZADA' in t: return 'Lazada'
    return 'Khác/chưa xác định kênh'
def bucket(): return {'sales':Decimal(0),'vat':Decimal(0),'total':Decimal(0),'qty':Decimal(0),'unknownqty':0,'records':0}
def add(b,sales,vat,total,qty=None):
    b['sales']+=sales; b['vat']+=vat; b['total']+=total; b['records']+=1
    if qty is None: b['unknownqty']+=1
    else: b['qty']+=qty
def csvout(name,headers,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(headers); w.writerows(rows)

A=defaultdict(bucket); AC=defaultdict(bucket); amiscontrols=[]; arefs=[]; amisissues=[]
for p in sorted((BASE/'Du lieu Amis').glob('*.xlsx')):
    totals=bucket(); declared=None
    for i,r in enumerate(load_workbook(p,read_only=True,data_only=True).worksheets[0].values,1):
        if i<5: continue
        if isinstance(r[1],datetime) and r[2]:
            m=r[1].strftime('%Y-%m'); f=family(r[13],r[12]); c=channel(r[9])
            sales=n(r[20])-n(r[25])-n(r[27])-n(r[28]); vat=n(r[32]); total=n(r[34]); qty=n(r[17])-n(r[26])
            residual=total-sales-vat
            if residual!=0: amisissues.append([p.name,i,r[2],sales,vat,total,residual])
            add(A[(m,f)],sales,vat,total,qty); add(AC[(m,c,f)],sales,vat,total,qty); add(totals,sales,vat,total,qty)
            arefs.append([p.name,i,r[2],m,c,f,qty,sales,vat,total])
        elif r[0]=='Tổng cộng': declared=(n(r[20])-n(r[25])-n(r[27])-n(r[28]),n(r[32]),n(r[34]))
    amiscontrols.append((p.name,totals,declared))
print('AMIS groups loaded',len(A),flush=True)

H={}; details=defaultdict(list); edges=[]; olddates={}; problems=[]
for p in sorted((BASE/'Du lieu Meinvoice').glob('*.xls')):
    if p.name.startswith('Mau'): continue
    for i,r in enumerate(load_workbook(p.open('rb'),read_only=True,data_only=True).worksheets[0].values,1):
        if not isinstance(r[0],(int,float)): continue
        if p.name.startswith('Bang_ke') and i>=6:
            key=(str(r[1]),str(r[2])); H[key]={'date':datetime.strptime(r[3],'%d/%m/%Y'),'sales':n(r[15]),'vat':n(r[16]),'total':n(r[17]),'complete':all(isinstance(r[j],(int,float)) for j in (15,16,17)),'channel':channel(str(r[7])+' '+str(r[11])),'status':r[20],'file':p.name,'row':i}
        elif p.name.startswith('Bảng kê chi tiết') and i>=6:
            key=(str(r[1]),str(r[2])); details[key].append({'family':family(r[11]),'qty':n(r[13]) if isinstance(r[13],(int,float)) else None,'sales':n(r[15]),'vat':n(r[17]),'complete':all(isinstance(r[j],(int,float)) for j in (15,17)),'file':p.name,'row':i})
        elif p.name.startswith('DS') and i>=7:
            old=(str(r[1]),str(r[3])); new=(str(r[12]),str(r[14])); typ='Điều chỉnh' if 'điều chỉnh' in p.name else 'Thay thế'
            d=r[2] if isinstance(r[2],datetime) else datetime.strptime(str(r[2]),'%d/%m/%Y')
            olddates[old]=d
            edges.append((old,new,typ,p.name,i))
print('meInvoice loaded',len(H),len(edges),flush=True)
parents=defaultdict(set); children=defaultdict(list); kinds=defaultdict(set)
for old,new,typ,file,row in edges:
    parents[new].add(old); children[old].append((new,typ)); kinds[new].add(typ)
def root_of(k):
    visited=set(); cur=k
    while parents.get(cur):
        if cur in visited: return None,'Vòng lặp liên kết'
        visited.add(cur)
        if len(parents[cur])!=1: return None,'Nhiều hóa đơn nguồn'
        cur=next(iter(parents[cur]))
    return cur,''
M=defaultdict(bucket); MC=defaultdict(bucket); posting=defaultdict(bucket); roots=defaultdict(list)
included=0; excluded=Counter(); unidentified=Counter(); allocations=[]; invrows=[]
for key,h in H.items():
    if h['status'] in ('Hóa đơn đã bị thay thế','Hoá đơn đã bị hủy'):
        excluded[h['status']]+=1; continue
    included+=1
    root,error=root_of(key)
    flags=[]
    if error: flags.append(error)
    if not parents.get(key) and (h['status'] in ('Hoá đơn điều chỉnh','Hoá đơn thay thế') or kinds.get(key)):
        flags.append('Thiếu liên kết tới hóa đơn nguồn'); root=None
    date=(H[root]['date'] if root in H else olddates.get(root)) if root else None
    if date is None: flags.append('Chưa xác định kỳ hóa đơn gốc')
    month=date.strftime('%Y-%m') if date else 'CHƯA XÁC ĐỊNH KỲ'
    if not h['complete']: flags.append('Thiếu số tiền trên bảng tổng')
    if h['status']=='Hóa đơn đã bị điều chỉnh' and not any(t=='Điều chỉnh' for _,t in children.get(key,[])):
        flags.append('Trạng thái bị điều chỉnh nhưng thiếu hóa đơn điều chỉnh liên kết')
    if root: roots[root].append(key)
    lines=details.get(key,[])
    arithmeticok=h['sales']+h['vat']==h['total']
    if h['complete'] and not arithmeticok: flags.append('Tiền hàng cộng thuế không bằng tổng thanh toán')
    knownlines=[l for l in lines if l['complete']]
    knownmatches=arithmeticok and bool(knownlines) and sum((l['sales'] for l in knownlines),Decimal(0))==h['sales'] and sum((l['vat'] for l in knownlines),Decimal(0))==h['vat']
    monetaryok=knownmatches and len(knownlines)==len(lines)
    partialok=knownmatches and len(knownlines)<len(lines)
    if partialok: flags.append('Dòng đủ tiền khớp tổng; dòng thiếu tiền giữ riêng, chưa phân bổ lượng')
    elif not monetaryok: flags.append('Chưa phân bổ chắc chắn tiền theo sản phẩm')
    for flag in flags: unidentified[flag]+=1
    invrows.append([key[0],key[1],h['date'].strftime('%d/%m/%Y'),h['status'],root[0] if root else '',root[1] if root else '',month,h['channel'],h['sales'] if h['complete'] else '',h['vat'] if h['complete'] else '',h['total'] if h['complete'] else '', '; '.join(flags),h['file'],h['row']])
    if not h['complete']: continue
    add(posting[h['date'].strftime('%Y-%m')],h['sales'],h['vat'],h['total'])
    if monetaryok:
        components=lines
    elif partialok:
        # Chỉ phân bổ các dòng có số tiền thực tế. Phần dư tổng bằng 0 không
        # được phân bổ vào các dòng trống hoặc suy ra lượng hàng khuyến mại.
        components=knownlines+[{'family':'CHƯA PHÂN BỔ SẢN PHẨM','qty':None,'sales':Decimal(0),'vat':Decimal(0),'total':Decimal(0),'file':h['file'],'row':h['row']}]
    else:
        components=[{'family':'CHƯA PHÂN BỔ SẢN PHẨM','qty':None,'sales':h['sales'],'vat':h['vat'],'total':h['total'],'file':h['file'],'row':h['row']}]
    for l in components:
        s=l['sales']; v=l['vat']; t=l.get('total',s+v)
        # Không suy ra lượng thay đổi từ dòng hóa đơn điều chỉnh chỉ tiền.
        qty=None if 'Điều chỉnh' in kinds.get(key,set()) or h['status']=='Hoá đơn điều chỉnh' else l['qty']
        add(M[(month,l['family'])],s,v,t,qty); add(MC[(month,h['channel'],l['family'])],s,v,t,qty)
        allocations.append([key[0],key[1],month,l['family'],qty if qty is not None else '',s,v,t,l['file'],l['row'],h['status']])

group_rows=[]
for k in sorted(set(A)|set(M)):
    a=A[k]; m=M[k]
    note=[]
    if k[0]=='CHƯA XÁC ĐỊNH KỲ' or k[1].startswith('CHƯA'): note.append('Chưa đủ căn cứ so sánh sản phẩm/kỳ')
    else: note.append('Kỳ AMIS theo ngày chứng từ; kỳ meInvoice theo ngày HĐ gốc, chưa phải cùng ngày bán')
    if m['unknownqty']: note.append('Có khoản chưa xác định lượng')
    group_rows.append([*k,a['qty'],m['qty'] if not m['unknownqty'] else '',m['unknownqty'],a['sales'],m['sales'],a['sales']-m['sales'],a['vat'],m['vat'],a['vat']-m['vat'],a['total'],m['total'],a['total']-m['total'],'; '.join(note)])
csvout('DOI_CHIEU_THANG_MODEL.csv',['Tháng/kỳ','Họ model','SL AMIS sau trả hàng','SL meInvoice từ dòng phân bổ được (chưa khẳng định đầy đủ)','Số khoản meInvoice thiếu lượng','Doanh thu AMIS sau giảm trừ','Doanh thu meInvoice theo file','Chênh lệch AMIS-meInvoice','Thuế AMIS','Thuế meInvoice','Chênh lệch thuế','Thanh toán AMIS','Thanh toán meInvoice','Chênh lệch thanh toán','Giới hạn'],group_rows)
csvout('SO_LIEU_THEO_KENH.csv',['Nguồn','Tháng/kỳ','Kênh nhận diện từ tên khách hàng','Họ model','Số lượng biết được','Số khoản thiếu lượng','Doanh thu','Thuế','Thanh toán'],[[src,*k,b['qty'],b['unknownqty'],b['sales'],b['vat'],b['total']] for src,groups in [('AMIS',AC),('meInvoice',MC)] for k,b in sorted(groups.items())])
csvout('HOA_DON_VA_KY_GOC.csv',['Ký hiệu','Số hóa đơn','Ngày phát hành','Trạng thái','Ký hiệu gốc','Số HĐ gốc','Kỳ HĐ gốc','Kênh nhận diện','Doanh thu','Thuế','Thanh toán','Cần kiểm tra','File nguồn','Dòng'],invrows)
csvout('CHI_TIET_AMIS_DA_TINH_GIAM_TRU.csv',['File','Dòng','Số chứng từ','Tháng','Kênh nhận diện','Họ model','Lượng sau trả hàng','Doanh thu sau giảm trừ','Thuế','Thanh toán'],arefs)
csvout('PHAN_BO_MEINVOICE_THEO_SAN_PHAM.csv',['Ký hiệu','Số HĐ','Kỳ HĐ gốc','Họ model','Lượng biết được','Doanh thu','Thuế','Thanh toán','File nguồn','Dòng','Trạng thái'],allocations)
csvout('AMIS_CAN_KIEM_TRA_CONG_THUC.csv',['File','Dòng','Số chứng từ','Doanh thu sau giảm trừ','Thuế','Thanh toán','Thanh toán trừ doanh thu và thuế'],amisissues)
chainrows=[]
for root,keys in sorted(roots.items()):
    b=bucket(); missing=0
    for k in keys:
        h=H[k]
        if not h['complete']: missing+=1
        else: add(b,h['sales'],h['vat'],h['total'])
    chainrows.append([*root,len(keys),b['sales'],b['vat'],b['total'],missing,'; '.join(f'{k[0]}/{k[1]}' for k in keys)])
csvout('SO_DU_THEO_CHUOI_HOA_DON.csv',['Ký hiệu gốc','Số HĐ gốc','Số HĐ đóng góp số dư','Doanh thu theo file','Thuế theo file','Thanh toán theo file','Số HĐ thiếu giá trị','Hóa đơn đóng góp'],chainrows)
asum=sum((b['total'] for b in A.values()),Decimal(0)); msum=sum((b['total'] for b in M.values()),Decimal(0))
assert msum==sum((b['total'] for b in posting.values()),Decimal(0))
assert sum((b['total'] for b in MC.values()),Decimal(0))==msum
assert sum((b['total'] for b in AC.values()),Decimal(0))==asum
for file,t,d in amiscontrols:
    assert d is not None and (t['sales'],t['vat'],t['total'])==d,('AMIS control',file)
months=sorted(set(k[0] for k in A)|set(k[0] for k in M))
report=['# UVG – đối chiếu doanh thu theo nhóm','',
'Dữ liệu AMIS đã được khách hàng chốt là mục tiêu. Chỉ đọc nguồn, không dùng mã đơn sàn và không phát hành hóa đơn. Phạm vi nguồn: năm 2024 và tháng 1–6/2025; chưa bao gồm các lần sửa sau kỳ xuất dữ liệu.','',
'## Kết quả tổng thể','',f'- Tổng thanh toán AMIS sau giảm trừ: **{fmt(asum)} đồng**.',f'- Tổng thanh toán meInvoice tính được theo trạng thái và liên kết trong file: **{fmt(msum)} đồng**.',f'- Chênh lệch AMIS − meInvoice: **{fmt(asum-msum)} đồng**. Đây là chênh lệch cần giải thích, không phải số tiền cần xuất bổ sung hoặc điều chỉnh.',
f'- {included:,} hóa đơn đóng góp hoặc cần kiểm tra; loại khỏi tổng {sum(excluded.values())} hóa đơn có trạng thái đã hủy/đã bị thay thế. Hóa đơn thiếu số tiền được giữ riêng, không được coi là giá trị bằng 0.',
'','## Tổng hợp theo kỳ gốc','',
'AMIS lấy ngày chứng từ. meInvoice lần theo bảng liên kết về ngày hóa đơn gốc; đây chỉ là quy đổi về kỳ hóa đơn gốc, chưa xác định được kỳ bán hàng của từng giao dịch trong hóa đơn gộp. Khoản không có liên kết được giữ ở dòng chưa xác định kỳ.','',
'| Kỳ | Thanh toán AMIS | Thanh toán meInvoice theo file | Chênh lệch AMIS − meInvoice |','|---|---:|---:|---:|']
for month in months:
    aa=sum((b['total'] for (m,f),b in A.items() if m==month),Decimal(0)); mm=sum((b['total'] for (m,f),b in M.items() if m==month),Decimal(0))
    report.append(f'| {month} | {fmt(aa)} | {fmt(mm)} | {fmt(aa-mm)} |')
report+=['','## Phần còn thiếu căn cứ','']
for label,count in unidentified.items(): report.append(f'- {label}: {count:,} hóa đơn. Một hóa đơn có thể có nhiều vấn đề.')
unknownprod=sum((b['total'] for (m,f),b in M.items() if f=='CHƯA PHÂN BỔ SẢN PHẨM'),Decimal(0))
unknownperiod=sum((b['total'] for (m,f),b in M.items() if m=='CHƯA XÁC ĐỊNH KỲ'),Decimal(0))
unknownmodel=sum((b['total'] for (m,f),b in M.items() if f=='CHƯA XÁC ĐỊNH MODEL'),Decimal(0))
report += [f'- Tổng thanh toán chưa phân bổ chắc chắn theo sản phẩm: {fmt(unknownprod)} đồng (số cộng đại số, gồm cả khoản tăng/giảm).',f'- Tổng thanh toán chưa xác định kỳ hóa đơn gốc: {fmt(unknownperiod)} đồng.',f'- AMIS có {len(amisissues)} dòng mà công thức doanh thu sau giảm trừ + thuế chưa khớp thanh toán; chi tiết ở file kiểm tra.','',
'## Cách tính và giới hạn','',
'- AMIS: doanh thu = doanh số bán − chiết khấu − giá trị trả lại − giá trị giảm giá; số lượng = lượng bán − lượng trả lại. Thuế và thanh toán dùng dấu có sẵn trong báo cáo. Đã kiểm tra tổng dòng với dòng Tổng cộng của hai file.',
'- meInvoice: loại hóa đơn có trạng thái đã bị thay thế/đã hủy; giữ hóa đơn gốc bị điều chỉnh, cộng khoản điều chỉnh, giữ bản thay thế còn hiện hành theo trạng thái file. Các liên kết đưa khoản sửa về cùng chuỗi; trường hợp thiếu liên kết hoặc vòng lặp được đánh dấu.',
'- Phân bổ tiền meInvoice vào model từ các dòng có số tiền và thuế thực tế khi tổng các dòng đó khớp bảng tổng. Nếu còn dòng trống tiền, giữ chúng ở khoản chưa phân bổ lượng, không tự coi từng dòng là 0 đồng. Nếu không khớp tổng hoặc thiếu chi tiết, giữ toàn bộ giá trị hóa đơn ở nhóm chưa phân bổ sản phẩm. Không phân bổ tùy ý khoản điều chỉnh chung vào model.',
'- Họ model phân biệt máy/thiết bị và lõi lọc. Không suy ra biến thể Pro/Light, không tự kết luận cùng SKU. Số lượng trên dòng điều chỉnh cần kiểm tra bản chất: số lượng diễn giải không luôn là lượng tăng/giảm thực tế.',
'- Kênh chỉ nhận diện khi tên khách hàng/người mua ghi rõ TikTok, Shopee/Shoppe/SP hoặc Lazada. Phần còn lại giữ Khác/chưa xác định kênh, không tự coi là ngoài sàn. File kênh trình bày hai nguồn riêng, không lấy kênh chưa xác định để bù sang một sàn.',
'- Dữ liệu có thể chưa đủ CQT và các lần sửa sau tháng 6/2025; số dư là theo bộ file hiện có, chưa phải số dư đã xác minh đầy đủ.',
'- Bảng chênh lệch không phải file import hoặc hướng dẫn giảm toàn bộ hóa đơn. Cần giải thích lệch kỳ, lệch sản phẩm, giảm trừ và lịch sử trước khi chốt hướng xử lý.','',
'## File để kiểm tra','',
'1. DOI_CHIEU_THANG_MODEL.csv: doanh thu, thuế, thanh toán và lượng theo tháng/họ model của hai nguồn.',
'2. SO_LIEU_THEO_KENH.csv: số liệu theo kênh nhận diện, giữ phần không xác định riêng.',
'3. HOA_DON_VA_KY_GOC.csv: từng hóa đơn, trạng thái, chuỗi nguồn, kỳ gốc và vấn đề cần kiểm tra.',
'4. SO_DU_THEO_CHUOI_HOA_DON.csv: số dư theo chuỗi, liệt kê các hóa đơn đóng góp.',
'5. Các file chi tiết AMIS/meInvoice chứa file nguồn và số dòng để truy ngược.','',
'## Việc tiếp theo','',
'Ưu tiên giải thích hóa đơn thiếu liên kết và khoản chưa phân bổ theo sản phẩm. Sau đó rà soát từng nhóm tháng/model lớn, kiểm tra bảng ánh xạ hàng và thời điểm ghi nhận. Phần không đủ căn cứ giữ lại; chưa phát hành theo chênh lệch tổng.']
(OUT/'UVG_DOI_CHIEU_NHOM.md').write_text('\n'.join(report),encoding='utf-8')
reportpath=OUT/'UVG_DOI_CHIEU_NHOM.md'
reportpath.write_text(reportpath.read_text(encoding='utf-8').replace('## Cách tính và giới hạn',f'- Ngoài nhóm chưa phân bổ sản phẩm, còn {fmt(unknownmodel)} đồng thuộc nhóm chưa nhận diện được họ model từ tên hàng.\n\n## Cách tính và giới hạn'),encoding='utf-8')
print('TOTALS',asum,msum,'DIFF',asum-msum,flush=True)
print('ISSUES',dict(unidentified),'AMIS_ARITHMETIC',len(amisissues),'UNKNOWNPROD',unknownprod,'UNKNOWNPERIOD',unknownperiod,flush=True)
print('OUTPUT',OUT,flush=True)
