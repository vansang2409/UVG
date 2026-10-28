"""Dự tính doanh thu theo hóa đơn gốc 2024, đọc nguồn/import hiện hành, không phát hành."""
from pathlib import Path
from collections import defaultdict,Counter
from datetime import date,datetime
from decimal import Decimal
import json,hashlib
from openpyxl import load_workbook
from prepare_invoice_reports import rows,invoice_key,invoice_day,amount
from filter_invoice_exceptions import normal

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/revenue-forecast-20261028'
SOURCES=ROOT/'data/Du lieu Meinvoice'
ZERO_APPROVED={('1C24TUV',n) for n in ['00000574','00000671','00000744']}
def add(a,b):return [x+y for x,y in zip(a,b)]
def zero():return [Decimal(0)]*3
def native(v):
    if isinstance(v,Decimal):return int(v) if v==int(v) else float(v)
    if isinstance(v,(date,datetime)):return v.isoformat()
    if isinstance(v,set):return sorted(v)
    raise TypeError(type(v))

def main():
    WORK.mkdir(parents=True,exist_ok=True)
    records={};parents=defaultdict(set);children=defaultdict(set);root_hints=defaultdict(set);replaced=set();manifest=[];issues=[]
    def source(p):
        rel=str(p.relative_to(ROOT))
        if not any(m['path']==rel for m in manifest):manifest.append({'path':rel,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    # Bảng kê hóa đơn đã sử dụng: một bản ghi cho mỗi ký hiệu + số HĐ.
    for p in sorted((SOURCES/'Bang_ke_hoa_don_da_su_dung_2024_2025').glob('*.xlsx')):
        source(p)
        for rn,r in rows(p):
            if rn<6 or amount(r.get('A')) is None:continue
            key=invoice_key(r.get('B'),r.get('C'))
            if not key:continue
            assert key not in records,key
            records[key]={'key':key,'date':invoice_day(r.get('D')),'money':[amount(r.get(c)) for c in ['P','Q','R']],'status':r.get('U') or '', 'buyer':r.get('H') or r.get('L'),'source':str(p.relative_to(ROOT)), 'row':rn,'header':True}
    # Đọc quan hệ cả các nguồn 2026 để không bỏ sót điều chỉnh muộn cho gốc 2024.
    link_files=sorted(set(list(SOURCES.glob('*dieu_chinh_xuat*.xls'))+list(SOURCES.glob('*thay_the_xuat*.xls'))+list((SOURCES/'Xuat_lai_2026-10-09').glob('*dieu_chinh*.xls'))+list((SOURCES/'Xuat_lai_2026-10-09').glob('*thay_the*.xls'))+list((SOURCES/'Xuat_lai_2026-10-22').glob('*dieu_chinh*.xls'))+list((SOURCES/'Xuat_lai_2026-10-22').glob('*thay_the*.xls'))))
    for p in link_files:
        source(p);replacement='thay_the' in p.name
        for rn,r in rows(p):
            if rn<7 or amount(r.get('A')) is None:continue
            old=invoice_key(r.get('B'),r.get('D'));new=invoice_key(r.get('M'),r.get('O'))
            if not old or not new:continue
            parents[new].add(old);children[old].add(new)
            if replacement:replaced.add(old)
            for key,dc,cols,status,buyer in [(old,'C',['J','K','L'],'',r.get('F')),(new,'N',['U','V','W'],'Hoá đơn thay thế' if replacement else 'Hoá đơn điều chỉnh',r.get('Q'))]:
                if key not in records:records[key]={'key':key,'date':invoice_day(r.get(dc)),'money':[amount(r.get(c)) for c in cols],'status':status,'buyer':buyer,'source':str(p.relative_to(ROOT)),'row':rn,'header':False}
    # Giữ liên kết lịch sử đã xác minh trong các workbook hiện hành (kể cả 8 liên kết thiếu ở bản xuất).
    for name in ['HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx','HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','NGOAI_LE_F0_2024_T6_2025.xlsx']:
        p=ROOT/'outputs'/name;source(p);w=load_workbook(p,read_only=True,data_only=True)
        if 'Chi tiet nguon' not in w.sheetnames:continue
        s=w['Chi tiet nguon']
        for r in s.iter_rows(min_row=2,values_only=True):
            if not r[0] or not r[2]:continue
            root=invoice_key(r[0],r[1]);key=invoice_key(r[2],r[3])
            root_hints[key].add(root)
            if key!=root:parents[key].add(root);children[root].add(key)
            if key not in records:records[key]={'key':key,'date':invoice_day(r[4]),'money':[amount(r[c]) for c in [9,10,11]],'status':r[8] or '', 'buyer':r[5],'source':name+'/Chi tiet nguon','header':False}
    cache={}
    def roots(k,stack=frozenset()):
        if k in cache:return cache[k]
        assert k not in stack,('cycle',k)
        result={k} if not parents[k] else set.union(*(roots(p,stack|{k}) for p in parents[k]))
        cache[k]=result;return result
    cohort={};all_2024=[];unmapped=[]
    for key,r in list(records.items()):
        rr=roots(key)
        if len(rr)!=1:issues.append({'kind':'multiple_roots','invoice':key,'roots':sorted(rr)});continue
        root=next(iter(rr));root_record=records.get(root)
        if not root_record or not root_record['date']:continue
        if root_record['date'].year!=2024:continue
        if not r['date']:issues.append({'kind':'missing_date','invoice':key});continue
        if r['date']>date(2026,10,28):continue
        if key==root and 'dieu chinh' in normal(r['status']) and normal(r['status'])!='hoa don da bi dieu chinh':
            unmapped.append(r);continue
        if key in replaced or 'bi thay the' in normal(r['status']) or 'huy' in normal(r['status']):continue
        r={**r,'root':root,'root_date':root_record['date']}
        if r['money'][2] is None and key in ZERO_APPROVED:r['money']=r['money'][:2]+[Decimal(0)];r['zero_basis']='Người dùng đã xác nhận ô thanh toán trống tạm tính 0, chỉ đúng hóa đơn này'
        cohort[key]=r
    # Thu tiền 0 tại gốc63 đã có rõ trong nguồn chi tiết: dùng chứng cứ số0, không đoán từ ô trống.
    missing={key for key,r in cohort.items() if None in r['money']}
    if missing:
        p=SOURCES/'Bang_ke_chi_tiet_HD_da_su_dung_2024.xls';source(p);details=defaultdict(list)
        for rn,r in rows(p):
            key=invoice_key(r.get('D'),r.get('B'))
            if rn>=7 and key in missing:details[key].append([amount(r.get(c)) for c in ['P','T','U']])
        for key in missing:
            rs=details[key]
            if rs and all(None not in r for r in rs):
                sums=[sum(r[c] for r in rs) for c in range(3)]
                for c,v in enumerate(cohort[key]['money']):
                    if v is None:cohort[key]['money'][c]=sums[c];cohort[key]['detail_basis']='Giá trị đọc trực tiếp từ nguồn chi tiết 2024'
    for key,r in cohort.items():
        if None in r['money']:issues.append({'kind':'missing_money','invoice':key,'money':r['money']})
    assert not [i for i in issues if i['kind'] in ['multiple_roots','missing_money']],issues[:10]
    months={m:{'before':zero(),'adjustments':zero(),'new':zero(),'roots':set(),'old_invoices':0,'adjustment_invoices':0,'new_invoices':0} for m in range(1,13)}
    for key,r in cohort.items():
        m=months[r['root_date'].month];m['before']=add(m['before'],r['money']);m['roots'].add(r['root']);m['old_invoices']+=1
    adjustment_records=[];processed_roots=set();adjustment_dates=Counter()
    for name in ['HOA_DON_CHUA_DIEU_CHINH_IMPORT_2_SHEET.xlsx','HOA_DON_DA_DIEU_CHINH_IMPORT_2_SHEET.xlsx']:
        p=ROOT/'outputs'/name;source(p);s=load_workbook(p,read_only=True,data_only=True)['Import dieu chinh'];seen=set()
        for rownum,r in enumerate(s.iter_rows(min_row=10,values_only=True),10):
            if not isinstance(r[0],(int,float)) or r[0] in seen:continue
            seen.add(r[0]);root=invoice_key(r[9],r[10]);rd=invoice_day(r[11])
            if not rd or rd.year!=2024:continue
            money=[amount(r[c]) for c in [19,21,22]]
            assert None not in money,(name,rownum,money)
            assert root in records and records[root]['date']==rd,(root,rd)
            processed_roots.add(root);m=months[rd.month];m['adjustments']=add(m['adjustments'],money);m['adjustment_invoices']+=1
            adjustment_dates[str(invoice_day(r[1]))]+=1
            adjustment_records.append({'file':name,'row':rownum,'index':r[0],'root':root,'root_date':rd,'date':invoice_day(r[1]),'money':money})
    newfile=ROOT/'outputs/amis-2024-theo-btl-bhh-20261026/IMPORT_HOA_DON_MOI_AMIS_2024.xlsx';source(newfile)
    s=load_workbook(newfile,read_only=True,data_only=True).active;invoice_group={};new_records=[]
    for rownum,r in enumerate(s.iter_rows(min_row=10,values_only=True),10):
        if r[0] not in invoice_group:
            invoice_group[r[0]]={'index':r[0],'date':invoice_day(r[1]),'bhh':r[18],'goods':Decimal(0),'vat':amount(r[11]),'row':rownum}
        invoice_group[r[0]]['goods']+=amount(r[16])
    for r in invoice_group.values():
        assert r['date'].year==2024
        r['money']=[r['goods'],r['vat'],r['goods']+r['vat']];m=months[r['date'].month];m['new']=add(m['new'],r['money']);m['new_invoices']+=1;new_records.append(r)
    assert len(new_records)==34739
    # Phương án mới: bán gốc và điều chỉnh hàng trả lại tách riêng.
    returnfile=ROOT/'outputs/amis-2024-theo-btl-bhh-20261026/KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx';source(returnfile)
    returned=[];seen_returns=set();return_sheet=load_workbook(returnfile,read_only=True,data_only=True)['Import dieu chinh']
    for rownum,r in enumerate(return_sheet.iter_rows(min_row=10,values_only=True),10):
        if r[0] in seen_returns:continue
        seen_returns.add(r[0]);money=[amount(r[c]) for c in [19,21,22]]
        assert None not in money and money[0]+money[1]==money[2]
        returned.append({'index':r[0],'row':rownum,'bhh':r[25],'reason':r[13],'money':money})
    assert len(returned)==163
    returns_total=[sum(r['money'][i] for r in returned) for i in range(3)]
    assert returns_total==[Decimal(-225865926),Decimal(-18069275),Decimal(-243935201)]
    by_root=defaultdict(lambda:{'before':zero(),'adjustments':zero(),'after':zero(),'nodes':[]})
    for key,r in cohort.items():
        q=by_root[r['root']];q['before']=add(q['before'],r['money']);q['nodes'].append(key)
    for r in adjustment_records:by_root[r['root']]['adjustments']=add(by_root[r['root']]['adjustments'],r['money'])
    for root,r in by_root.items():r['after']=add(r['before'],r['adjustments'])
    processed_nonzero=[{'root':k,**v} for k,v in by_root.items() if k in processed_roots and any(v['after'])]
    total={k:[sum(months[m][k][i] for m in months) for i in range(3)] for k in ['before','adjustments','new']}
    total['retained']=add(total['before'],total['adjustments']);total['forecast']=add(total['retained'],total['new'])
    for r in months.values():r['retained']=add(r['before'],r['adjustments']);r['forecast']=add(r['retained'],r['new'])
    retained=[{'root':k,'date':records[k]['date'],'buyer':records[k]['buyer'],**v} for k,v in by_root.items() if any(v['after'])]
    by_year=defaultdict(lambda:{'count':0,'money':zero()})
    for r in cohort.values():
        v=by_year[r['date'].year];v['count']+=1;v['money']=add(v['money'],r['money'])
    report={'scope':'Theo ngày hóa đơn gốc năm 2024, cộng các điều chỉnh liên quan năm sau; không phải báo cáo theo ngày phát hành 2024','totals':total,'months':months,'counts':{'old_roots':len(by_root),'old_invoices':len(cohort),'adjustment_invoices':len(adjustment_records),'processed_roots':len(processed_roots),'new_invoices':len(new_records),'retained_roots':len(retained)},'existing_by_issue_year':by_year,'adjustment_dates':adjustment_dates,'retained_roots':retained,'processed_nonzero':processed_nonzero,'issues':issues,'unmapped_adjustments':[r['key'] for r in unmapped],'manifest':manifest,'old_invoices':list(cohort.values()),'adjustment_records':adjustment_records,'new_records':new_records}
    total['new_gross']=total['new'][:];total['new_returns']=returns_total
    total['forecast_before_manual']=add(total['forecast'],returns_total)
    manual_anchors={('1C24TUV',n) for n in ['00000188','00000189','00000190','00000191','00000192','00000216']}
    # HĐ 188...216 là mắt xích giữa; khi có đủ liên kết phải quy về gốc thật.
    assert all(len(roots(k))==1 for k in manual_anchors)
    manual_roots={next(iter(roots(k))) for k in manual_anchors}
    manual_records=[r for r in retained if tuple(r['root']) in manual_roots]
    for r in manual_records:r['manual_anchor']=[k for k in sorted(manual_anchors) if tuple(r['root']) in roots(k)]
    assert len(manual_records)==6
    manual=[-sum(r['after'][i] for r in manual_records) for i in range(3)]
    assert manual==[Decimal(-371344277),Decimal(-29707543),Decimal(-401051820)]
    total['manual_replacement_reduction']=manual
    total['forecast_after_manual']=add(total['forecast_before_manual'],manual)
    total['change_after_all_vs_before']=[a-b for a,b in zip(total['forecast_after_manual'],total['before'])]
    ledger=ROOT/'data/Du lieu Amis/So_chi_tiet_ban_hang_2024.xlsx';source(ledger)
    amis=zero()
    for r in load_workbook(ledger,read_only=True,data_only=True).active.iter_rows(min_row=5,values_only=True):
        if not isinstance(r[0],datetime) or not r[2]:continue
        amis=add(amis,[Decimal(str(r[20] or 0))-Decimal(str(r[27] or 0)),Decimal(str(r[32] or 0)),Decimal(str(r[34] or 0))])
    total['amis_net']=amis;total['after_all_minus_amis']=[a-b for a,b in zip(total['forecast_after_manual'],amis)]
    report['returns_records']=returned;report['manual_records']=manual_records
    report['manual_status']='Kế hoạch đưa 6 chuỗi thay thế về 0 theo xác nhận 27/10; chưa xác minh KT đã thực hiện'
    report['status']='Dự tính từ file đã chuẩn bị; chưa xác minh import/ký/phát hành hoặc báo cáo meInvoice thực tế'
    report['counts']['new_return_invoices']=len(returned)
    assert not processed_nonzero and not issues and not unmapped
    assert all(total[k][0]+total[k][1]==total[k][2] for k in ['forecast_before_manual','forecast_after_manual','new_gross','new_returns','manual_replacement_reduction','amis_net'])
    for s in manifest:assert hashlib.sha256((ROOT/s['path']).read_bytes()).hexdigest()==s['sha256']
    (WORK/'prepared.json').write_text(json.dumps(report,ensure_ascii=False,default=native),encoding='utf8')
    summary={k:v for k,v in report.items() if k not in ['old_invoices','adjustment_records','new_records','months','manifest','retained_roots','returns_records','manual_records']}
    (WORK/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=native),encoding='utf8')
    print(json.dumps(summary,ensure_ascii=False,indent=2,default=native))

if __name__=='__main__':main()
