"""Đọc lại file đã lưu: mẫu, sáu tham chiếu và toàn bộ tiền/số lượng về 0."""
from pathlib import Path
from io import BytesIO
from datetime import datetime
import json,hashlib
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/replacement-zero-20261029'
p=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
out=ROOT/'outputs/amis-2024-theo-btl-bhh-20261026/IMPORT_THAY_THE_VE_0_6_CHUOI_2024.xlsx'
w=load_workbook(out,read_only=True,data_only=False)
assert w.sheetnames==['Hóa đơn GTGT']
s=w.active
assert list(next(s.iter_rows(min_row=9,max_row=9,max_col=22,values_only=True)))==p['headers']
actual=list(s.iter_rows(min_row=10,max_row=9+len(p['rows']),max_col=22,values_only=True))
for idx,(got,expected) in enumerate(zip(actual,p['rows']),10):
    got=list(got)
    if isinstance(got[13],datetime):got[13]=(got[13]-datetime(1899,12,30)).days
    assert got==expected,(idx,got,expected)
    assert all(got[j] is None for j in [1,6,7,8,10,14])
    assert got[17:20]==[0,0,0]
    assert not any(isinstance(c,str) and c.startswith('=') for c in got)
assert len(actual)==181 and set(r[0] for r in actual)==set(range(1,7))
head=[r for r in actual if r[11]]
assert len(head)==6
assert {r[12] for r in head}=={'00002601','00002610','00002602','00002603','00002604','00002617'}
assert all(r[11]=='1C25TUV' and r[20]==8 and r[21]==0 and isinstance(r[12],str) for r in head)
assert not any(any(v is not None for v in r) for r in s.iter_rows(min_row=191,values_only=True))
assert hashlib.sha256(Path(p['template']).read_bytes()).hexdigest()==p['template_sha256']
assert hashlib.sha256(Path(p['source']).read_bytes()).hexdigest()==p['source_sha256']
result={'passed':True,'invoices':6,'lines':len(actual),'quantity_price_goods_tax_zero':True,'new_dates_blank':True,'replaced_dates':'2025-05-31','reference_series':'1C25TUV','reference_numbers':[r[12] for r in head],'template_headers_exact':True,'source_files_unchanged':True,'before':p['totals_before'],'after':[0,0,0],'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'not_imported_to_meinvoice':True}
(WORK/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
