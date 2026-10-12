"""Apply confirmed review classifications and audit whole-chain buyer exclusions.

Reads source workbooks/XML only. Does not rewrite any import workbook.
"""
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import io
import json
import re
import shutil
import xml.etree.ElementTree as ET

import openpyxl

from prepare_invoice_reports import rows, invoice_key, invoice_day, amount
from filter_invoice_exceptions import normal, is_company, tax_shaped

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / 'scripts/data/check-review/kt-workbook-20261011'
SOURCE = Path(r'C:\Users\thuong ngo\OneDrive\Desktop\JOB UVG\JOB HOA DON\UVG_CAC_VAN_DE_CAN_KE_TOAN_XAC_NHAN_20261011.xlsx')
manifest = json.loads((HERE / 'review-manifest.json').read_text(encoding='utf-8'))
baseline = json.loads((ROOT / 'scripts/data/invoice_reports.json').read_text(encoding='utf-8'))
digest = sha256(SOURCE.read_bytes()).hexdigest()
wb = openpyxl.load_workbook(SOURCE, data_only=True)
observations = defaultdict(list)
graph = defaultdict(set)
headers = {}
invoice_dates = {}
source_hashes = {}
seeds = defaultdict(list)
edges = set()

def encode(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    raise TypeError(type(value))

def company_or_household(text):
    return is_company(text) or bool(re.search(r'\b(ho kinh doanh|hkd)\b', normal(text)))

def observe(key, name, tax, source, family, note=False):
    if not key:
        return
    name = str(name or '').strip()
    tax = str(tax or '').strip()
    obs = {'name': name, 'tax': tax, 'source': source, 'family': family}
    if obs in observations[key]:
        return
    observations[key].append(obs)
    if company_or_household(name) or (not note and tax_shaped(tax)):
        seeds[key].append({**obs, 'reason': 'Đổi người mua thành công ty/hộ kinh doanh' if note
                           else 'Tên công ty/hộ kinh doanh hoặc định danh thuộc quy tắc ngoại lệ hiện có'})

def add_edge(parent, child, source):
    if parent and child and parent != child:
        graph[parent].add(child)
        graph[child].add(parent)
        edges.add((parent, child))

def buyer_note(text):
    n = normal(text)
    return bool(re.search(r'(thong tin.*(khach hang|nguoi mua)|ten.*(khach hang|nguoi mua|don vi)|mst.*(cong ty|ho kinh doanh))', n))

# Preserve all previously protected chains and their historical relations.
baseline_roots = {tuple(g['root']) for g in baseline['groups']['exceptions']}
baseline_invoices = set()
for kind, groups in baseline['groups'].items():
    for g in groups:
        root = tuple(g['root'])
        for row in g['invoices']:
            key = tuple(row[:2])
            add_edge(root, key, 'Báo cáo chuỗi hiện có')
            if kind == 'exceptions':
                baseline_invoices.add(key)
                seeds[key].append({'reason': 'Ngoại lệ đã tách trước đợt phản hồi', 'source': 'Báo cáo ngoại lệ hiện có'})
for edge in baseline['historical_edges']:
    add_edge(tuple(edge['old']), tuple(edge['new']), 'Liên kết lịch sử đã kiểm chứng')

# Read current exports, including 2026 relations: exclusion has no 2025 cutoff.
NEW = ROOT / 'data/Du lieu Meinvoice/Xuat_lai_2026-10-09'
for path in sorted(NEW.iterdir()):
    if path.suffix not in ('.xls', '.xlsx', '.zip') or '1791532943641' in path.name:
        continue
    content = path.read_bytes()
    source_hashes[path.relative_to(ROOT).as_posix()] = sha256(content).hexdigest()
    if path.suffix == '.zip':
        with ZipFile(io.BytesIO(content)) as archive:
            books = [(n, archive.read(n)) for n in archive.namelist() if n.endswith(('.xls', '.xlsx'))]
    else:
        books = [(path.name, content)]
    for name, book in books:
        family = 'detail' if 'chi_tiet' in name else 'link' if any(t in name for t in ('dieu_chinh_xuat', 'thay_the_xuat')) else 'total'
        label = path.relative_to(ROOT).as_posix() + (' / ' + name if path.suffix == '.zip' else '')
        for rn, row in rows(io.BytesIO(book)):
            if amount(row.get('A')) is None:
                continue
            source = label + f' | dòng {rn}'
            if family == 'total':
                key = invoice_key(row.get('B'), row.get('C'))
                headers[key] = {'date': invoice_day(row.get('D')), 'name': row.get('H'), 'buyer': row.get('L'),
                                'tax': row.get('J'), 'money': [amount(row.get(c)) for c in ('P', 'Q', 'R')]}
                invoice_dates[key] = invoice_day(row.get('D'))
                observe(key, row.get('H'), row.get('J'), source, 'Bảng tổng: tên khách hàng')
                observe(key, row.get('L'), row.get('J'), source, 'Bảng tổng: người mua')
            elif family == 'detail':
                key = invoice_key(row.get('B'), row.get('C'))
                observe(key, row.get('E'), row.get('G'), source, 'Chi tiết: tên khách hàng')
                observe(key, row.get('H'), row.get('G'), source, 'Chi tiết: người mua')
            else:
                parent = invoice_key(row.get('B'), row.get('D'))
                child = invoice_key(row.get('M'), row.get('O'))
                add_edge(parent, child, source)
                if parent:
                    invoice_dates.setdefault(parent, invoice_day(row.get('C')))
                if child:
                    invoice_dates.setdefault(child, invoice_day(row.get('N')))
                observe(parent, row.get('F'), row.get('G'), source, 'Liên kết: hóa đơn gốc')
                observe(child, row.get('Q'), row.get('R'), source, 'Liên kết: hóa đơn điều chỉnh/thay thế')
                if buyer_note(row.get('X')):
                    observe(child, row.get('X'), '', source, 'Lý do điều chỉnh thông tin người mua', note=True)

def first(node, name):
    return next((e for e in node.iter() if e.tag.rsplit('}', 1)[-1] == name), None) if node is not None else None

def value(node, name):
    return next((e.text for e in node if e.tag.rsplit('}', 1)[-1] == name), None) if node is not None else None

xml_count = 0
xml_keys = set()
for path in sorted((ROOT / 'data/Bo sung/meInvoice').rglob('*.xml')):
    content = path.read_bytes()
    source = path.relative_to(ROOT).as_posix()
    source_hashes[source] = sha256(content).hexdigest()
    doc = ET.fromstring(content)
    tt = first(doc, 'TTChung')
    key = invoice_key((value(tt, 'KHMSHDon') or '') + (value(tt, 'KHHDon') or ''), value(tt, 'SHDon'))
    if not key:
        continue
    xml_count += 1
    xml_keys.add(key)
    invoice_dates[key] = invoice_day(value(tt, 'NLap'))
    buyer = first(doc, 'NMua')
    observe(key, value(buyer, 'Ten'), value(buyer, 'MST'), source, 'XML: tên đơn vị')
    observe(key, value(buyer, 'HVTNMHang'), value(buyer, 'MST'), source, 'XML: người mua')
    related = first(tt, 'TTHDLQuan')
    if related is not None:
        parent = invoice_key((value(related, 'KHMSHDCLQuan') or '') + (value(related, 'KHHDCLQuan') or ''), value(related, 'SHDCLQuan'))
        add_edge(parent, key, source)
        note = value(related, 'GChu')
        if buyer_note(note):
            observe(key, note, '', source, 'XML: điều chỉnh thông tin người mua', note=True)

# Direct user confirmation supersedes earlier 3/508 preparation decision.
explicit_hold = []
for case in manifest['resolved'][:6]:
    root, child = tuple(case['root']), tuple(case['focus'])
    add_edge(root, child, 'Hồ sơ kế toán')
    seeds[root].append({'reason': 'Người dùng chốt giữ cả cặp ngoài đợt xử lý (điểm 6)', 'source': 'Tin nhắn người dùng 11/10/2026', 'case': case['id']})
    explicit_hold.append(root)

components = []
component_for = {}
for start in sorted(set(observations) | set(graph) | set(seeds)):
    if start in component_for:
        continue
    nodes, queue = set(), [start]
    while queue:
        key = queue.pop()
        if key in nodes:
            continue
        nodes.add(key)
        queue.extend(graph[key] - nodes)
    index = len(components)
    direct = sorted(nodes & set(seeds))
    components.append({'nodes': sorted(nodes), 'seeds': direct})
    for key in nodes:
        component_for[key] = index
protected = {key for component in components if component['seeds'] for key in component['nodes']}

review_cases = []
for case in manifest['cases']:
    ws = wb[case['group']]
    row = next(r for r in range(8, ws.max_row + 1) if ws.cell(r, 1).value == case['id'])
    root = tuple(case['root'])
    assert ws.cell(row, 2).value == '/'.join(root)
    first_col = 12 if case['group'] == 'Lech tien' else 8
    reply = {ws.cell(7, c).value: ws.cell(row, c).value for c in range(first_col, first_col + 6)}
    review_cases.append({**case, 'row': row, 'reply': reply, 'protected': root in protected,
                         'component': components[component_for[root]] if root in component_for else None})

planned = []
plan_impacts = []
for filename in ('HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx', 'HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx'):
    path = ROOT / 'outputs' / filename
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book['Import dieu chinh']
    seqs = {}
    for row in ws.iter_rows(min_row=10, max_col=24, values_only=True):
        if row[0] is None:
            continue
        seq, target = row[0], invoice_key(row[9], row[10])
        seqs.setdefault(seq, {'file': filename, 'seq': seq, 'target': target, 'lines': 0})['lines'] += 1
    planned.extend(seqs.values())
    plan_impacts.extend(x for x in seqs.values() if x['target'] in protected)
    book.close()

scope_roots = set(baseline_roots)
scope_roots.update(tuple(g['root']) for g in baseline['groups']['adjusted'])
untreated = json.loads((ROOT / 'scripts/data/untreated/untreated.json').read_text(encoding='utf-8'))
scope_roots.update(tuple(g['root']) for g in untreated['groups'])
new_protected_roots = sorted((scope_roots & protected) - baseline_roots)
held = []
for root in sorted(scope_roots & protected):
    comp = components[component_for[root]]
    evidence = [{'invoice': key, **reason} for key in comp['seeds'] for reason in seeds[key]]
    held.append({'root': root, 'date': headers.get(root, {}).get('date'), 'invoices': comp['nodes'],
                 'already_in_baseline_exceptions': root in baseline_roots, 'evidence': evidence})

destination = ROOT / 'data/Bo sung/Ke toan/2026-10-11' / digest[:12]
destination.mkdir(parents=True, exist_ok=True)
saved = destination / SOURCE.name
if saved.exists():
    assert sha256(saved.read_bytes()).hexdigest() == digest
else:
    shutil.copy2(SOURCE, saved)
(destination / 'PROVENANCE.json').write_text(json.dumps({'source': str(SOURCE), 'preserved_file': saved.name,
    'sha256': digest, 'bytes': len(SOURCE.read_bytes()), 'received_date': '2026-10-11',
    'purpose': 'Revised accountant replies with direct user confirmations 1-6'}, ensure_ascii=False, indent=2), encoding='utf-8')
summary = {'current_total_invoices': len(headers), 'xml_files_read': xml_count, 'unique_xml_invoices': len(xml_keys),
           'relations_read': len(edges), 'baseline_exception_roots': len(baseline_roots),
           'protected_scope_roots': len(scope_roots & protected), 'new_protected_roots': len(new_protected_roots),
           'explicit_keep_pairs': len(explicit_hold), 'existing_planned_invoices': len(planned),
           'planned_invoices_to_exclude': len(plan_impacts), 'planned_lines_to_exclude': sum(x['lines'] for x in plan_impacts),
           'protected_review_cases': sum(c['protected'] for c in review_cases)}
summary['protected_scope_invoices'] = len({tuple(k) for h in held for k in h['invoices']})
result = {'date': '2026-10-11', 'source': str(SOURCE), 'source_sha256': digest, 'preserved_path': str(saved),
          'policy': 'Keep the whole connected chain if any invoice was issued to or adjusted to a company/household business. Keep all previous tax-identity exceptions and six expressly excluded review pairs.',
          'summary': summary, 'review_cases': review_cases, 'held_chains': held, 'new_protected_roots': new_protected_roots,
          'existing_import_impacts': plan_impacts, 'source_hashes': source_hashes,
          'invoice_dates': {'/'.join(k):v for k,v in invoice_dates.items()},
          'application_status': 'Classification and numeric preparation only; no source or import workbook rewritten.'}
for filename, expected in manifest['workbookHashes'].items():
    assert sha256((ROOT / 'outputs' / filename).read_bytes()).hexdigest() == expected
for name, expected in source_hashes.items():
    assert sha256((ROOT / name).read_bytes()).hexdigest() == expected
assert sha256(SOURCE.read_bytes()).hexdigest() == digest
(HERE / 'confirmed-decisions-audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=encode), encoding='utf-8')
wb.close()
print(json.dumps({'source_sha256': digest, **summary, 'new_protected_root_keys': new_protected_roots,
                  'import_impact_root_keys': sorted({x['target'] for x in plan_impacts})}, ensure_ascii=False))
