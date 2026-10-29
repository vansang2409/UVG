"""Chuẩn bị sáu hóa đơn thay thế về 0, không nhập hay phát hành."""
from pathlib import Path
from io import BytesIO
from collections import defaultdict
from datetime import date
import json,hashlib
from openpyxl import load_workbook
from prepare_invoice_reports import rows,invoice_key,amount,invoice_day

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/replacement-zero-20261029'
TEMPLATE=ROOT/'data/Mau_hoa_don_GTGT_MTT_1_thue_suat_Thay_the.xls'
SOURCE=ROOT/'data/Du lieu Meinvoice/Bang_ke_chi_tiet_HD_da_su_dung_2025.xls'
CHAIN=[(80,188,2601),(87,189,2610),(116,190,2602),(117,191,2603),(118,192,2604),(184,216,2617)]

def main():
    WORK.mkdir(parents=True,exist_ok=True)
    f=json.loads((ROOT/'scripts/data/revenue-forecast-20261028/prepared.json').read_text(encoding='utf8'))
    headers={int(x['key'][1]):x for x in f['old_invoices'] if x['key'][0]=='1C25TUV' and int(x['key'][1]) in {c[2] for c in CHAIN}}
    detail=defaultdict(list)
    for rn,r in rows(SOURCE):
        k=invoice_key(r.get('D'),r.get('B'))
        if rn>=7 and k and k[0]=='1C25TUV' and int(k[1]) in headers:detail[int(k[1])].append((rn,r))
    t=load_workbook(BytesIO(TEMPLATE.read_bytes()),read_only=True,data_only=True)
    template=t.active
    template_headers=list(next(template.iter_rows(min_row=9,max_row=9,max_col=22,values_only=True)))
    data=[];audit=[]
    for seq,(root,middle,num) in enumerate(CHAIN,1):
        h=headers[num];ls=detail[num]
        assert ls and int(h['root'][1])==root
        assert h['date']=='2025-05-31'
        sums=[sum((amount(r.get(c)) or 0) for _,r in ls) for c in ['P','T','U']]
        assert sums==h['money'],(num,sums,h['money'])
        first=len(data)+10
        for i,(rn,r) in enumerate(ls):
            assert r['B']==f'{num:08}' and r['D']=='1C25TUV'
            assert invoice_day(r['C'])==date(2025,5,31)
            assert amount(str(r['S']).replace('%',''))==8
            assert r['L'] and not r.get('G')
            row=[None]*22
            row[0]=seq
            if i==0:
                # Ô không có dữ liệu nguồn giữ trống, không gán khách lẻ vào tên đơn vị.
                row[2]=r.get('E') or None;row[3]=r.get('G') or None;row[4]=r.get('F') or None
                row[5]=r.get('H') or None;row[9]=r.get('I') or 'TM/CK'
                row[11]='1C25TUV';row[12]=f'{num:08}'
                row[13]=(date(2025,5,31)-date(1899,12,30)).days
                row[20]=8;row[21]=0
            # Thay thế toàn bộ về 0 cả SL và tiền; giữ đúng tên hàng, ĐVT từng dòng.
            row[15]=r['L'];row[16]=r.get('M') or None
            row[17]=0;row[18]=0;row[19]=0
            data.append(row)
        audit.append({'root':f'1C24TUV/{root:08}','middle':f'1C24TUV/{middle:08}','replaced':f'1C25TUV/{num:08}','date':h['date'],'source_rows':[n for n,_ in ls],'output_rows':[first,len(data)+9],'lines':len(ls),'before':h['money'],'after':[0,0,0]})
    assert len(headers)==len(detail)==6
    totals=[sum(x['before'][i] for x in audit) for i in range(3)]
    assert totals==[371344277,29707543,401051820],totals
    p={'headers':template_headers,'rows':data,'audit':audit,'totals_before':totals,'totals_after':[0,0,0],'template':str(TEMPLATE),'source':str(SOURCE),'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'new_date_blank':True,'quantity_and_money_zero':True}
    (WORK/'prepared.json').write_text(json.dumps(p,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'invoices':6,'lines':len(data),'by_invoice':[(x['replaced'],x['lines']) for x in audit],'before':totals,'after':[0,0,0]},ensure_ascii=False))

if __name__=='__main__':main()
