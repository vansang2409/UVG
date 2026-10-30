"""Đọc lại quan hệ nguồn để kiểm tra thiếu/trùng độc lập với bước dựng Excel."""
from pathlib import Path
from collections import defaultdict,Counter
from datetime import datetime,date
from io import BytesIO
import json,sys,hashlib
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_invoice_reports import invoice_key,invoice_day,amount
DIR=ROOT/'outputs/chuoi-sau-xu-ly-20261030'
sys.stdout.reconfigure(encoding='utf-8')
p=json.loads((DIR/'prepared.json').read_text(encoding='utf-8'))
a=json.loads((ROOT/'scripts/data/post-issue-20261029/audit.json').read_text(encoding='utf-8'))
cached={tuple(v['key']):v for v in a['invoices']}
children=defaultdict(set);repparents=defaultdict(set);raw_edges=set();linkrows=0

def read(path):
    w=load_workbook(BytesIO(path.read_bytes()),read_only=True,data_only=True)
    s=w.worksheets[0]
    for rn,r in enumerate(s.iter_rows(max_col=24,values_only=True),1):yield rn,r
    w.close()

link_paths=[ROOT/v['file'] for v in a['files'] if v['kind']!='total']
historical=ROOT/'data/Du lieu Meinvoice/Bang_ke_hoa_don_thay_the_xuat_theo_HD_bi_thay_the_1793158686512.xls'
for path in link_paths+[historical]:
    replacement=path.parent.name=='THAY_THE' or path==historical
    for rn,r in read(path):
        if rn<7 or amount(r[0]) is None:continue
        old=invoice_key(r[1],r[3]);new=invoice_key(r[12],r[14])
        if not old or not new:continue
        if path!=historical:linkrows+=1
        children[old].add(new);raw_edges.add((old,new))
        if replacement:repparents[new].add(old)

def reach(root):
    found=set();todo=[root]
    while todo:
        key=todo.pop()
        if key in found:continue
        found.add(key);todo.extend(children[key])
    return found

planroots=set()
for name in ['HOA_DON_DA_DIEU_CHINH_IMPORT_2_SHEET.xlsx','HOA_DON_CHUA_DIEU_CHINH_IMPORT_2_SHEET.xlsx']:
    w=load_workbook(ROOT/'outputs'/name,read_only=True,data_only=True)
    for r in w['Import dieu chinh'].iter_rows(min_row=10,max_col=11,values_only=True):
        if r[0] is not None:planroots.add(invoice_key(r[9],r[10]))
    w.close()
adjustment=p['sheets'][0];replacement=p['sheets'][1]
assert {tuple(r['keys'][0]) for r in adjustment['rows']}==planroots
issues=[];occurrences=Counter();needed=set();roots_seen=set()
for sheet in p['sheets']:
    for rownum,r in enumerate(sheet['rows'],10):
        keys=[tuple(k) for k in r['keys']];root=keys[0]
        assert root not in roots_seen;roots_seen.add(root)
        if len(keys)!=len(set(keys)):issues.append({'kind':'duplicate_within_chain','root':root})
        expected=reach(root)
        if set(keys)!=expected:issues.append({'kind':'missing_or_extra_invoice','root':root,'missing':sorted(expected-set(keys)),'extra':sorted(set(keys)-expected)})
        if r['replacement']:
            assert len(keys)==4
            for old,new in zip(keys,keys[1:]):assert repparents[new]=={old},(old,new,'Sai chặng thay thế')
        occurrences.update(keys);needed.update(keys)
        for i,k in enumerate(keys):
            assert r['values'][8+6*i:10+6*i]==list(k)
assert not issues,issues
assert all(v==1 for v in occurrences.values()),[k for k,v in occurrences.items() if v>1]

# Dùng metadata chỉ để định tuyến file; dữ liệu ngày/tiền đọc lại bằng openpyxl từ nguồn gốc.
file_routes={cached[k]['source']['file'] for k in needed if k in cached}
totals=[v for v in a['files'] if v['kind']=='total']
file_routes.add(min(totals,key=lambda v:v['date_min'])['file'])
raw={};opened=set();duplicate_source=[]
def load_total(rel):
    if rel in opened:return
    opened.add(rel)
    for rn,r in read(ROOT/rel):
        if rn<6 or amount(r[0]) is None:continue
        key=invoice_key(r[1],r[2])
        if key not in needed:continue
        if key in raw:duplicate_source.append(key)
        raw[key]={'date':invoice_day(r[3]),'money':[amount(r[c]) for c in [15,16,17]],'cqt':r[18],'status':r[20]}
for rel in sorted(file_routes):load_total(rel)
if set(raw)!=needed:
    for v in totals:load_total(v['file'])
assert set(raw)==needed,sorted(needed-set(raw))
assert not duplicate_source,duplicate_source
payment_blanks=0;new_adjustments=set();new_replacements=set();data_fields=0
for sheet in p['sheets']:
    for rn,r in enumerate(sheet['rows'],10):
        keys=[tuple(k) for k in r['keys']]
        for i,k in enumerate(keys):
            v=r['values'][8+6*i:14+6*i];current=raw[k]
            from openpyxl.utils.datetime import to_excel
            assert v[2]==to_excel(datetime.combine(current['date'],datetime.min.time())),(k,'Sai ngày')
            for c in range(3):
                if current['money'][c] is None:
                    assert c==2 and v[3+c]==0,(k,'Ô trống chưa có căn cứ')
                    payment_blanks+=1
                else:assert current['money'][c]==amount(v[3+c]),(k,'Sai tiền',c)
            data_fields+=6
            if current['date'].year==2026:
                assert current['cqt']=='HĐ hợp lệ',(k,current['cqt'])
                (new_replacements if r['replacement'] else new_adjustments).add(k)
        assert keys==[keys[0]]+sorted(keys[1:],key=lambda k:(raw[k]['date'],k))
assert len(new_adjustments)==10454 and len(new_replacements)==6
assert {('1C26MUV',str(n).zfill(8)) for n in range(110692,110705)}<=new_adjustments
assert sum(r['excluded'] for s in p['sheets'] for r in s['rows'])==1
assert [r['keys'][0] for r in adjustment['rows'] if r['excluded']]==[['1C25TUV','00001943']]
for rel,h in p['source_hashes'].items():assert hashlib.sha256((ROOT/rel).read_bytes()).hexdigest()==h
result={'passed':True,'scope':'Các chuỗi điều chỉnh/thay thế trong đợt xử lý 29/10/2026, gồm một gốc KT loại trừ.',
 'adjustment_rows':len(adjustment['rows']),'replacement_rows':len(replacement['rows']),'unique_chain_invoice_keys':len(needed),
 'duplicate_within_chain':0,'duplicate_across_chains':0,'missing_from_source_graph':0,'extra_against_source_graph':0,
 'raw_link_rows_read':linkrows,'raw_total_files_read':len(opened),'raw_total_invoice_rows_matched':len(raw),'invoice_fields_matched':data_fields,
 'issued_2026_adjustments':len(new_adjustments),'issued_2026_replacements':len(new_replacements),'supplemental_13_all_present_once':True,
 'blank_payment_fields_with_zero_shown':payment_blanks,'accountant_excluded_invoice':'1C25TUV/00001943','sources_unchanged':True}
(DIR/'completeness-review.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
