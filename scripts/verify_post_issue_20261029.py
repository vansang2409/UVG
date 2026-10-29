"""Verify the saved reconciliation using a separate workbook reader."""
from pathlib import Path
from io import BytesIO
from collections import defaultdict, Counter
from decimal import Decimal
import csv
import hashlib
import json
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'scripts/data/post-issue-20261029'


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    a = json.loads((WORK / 'audit.json').read_text(encoding='utf-8'))
    byfile = defaultdict(dict)
    for rec in a['invoices']:
        byfile[rec['source']['file']][rec['source']['row']] = rec
    verified = 0
    for name, wanted in sorted(byfile.items()):
        w = load_workbook(BytesIO((ROOT / name).read_bytes()), read_only=True, data_only=True)
        sheet = w.active
        header = list(next(sheet.iter_rows(min_row=5, max_row=5, values_only=True)))
        col = {v: i for i, v in enumerate(header) if v}
        seen = set()
        for rn, raw in enumerate(sheet.iter_rows(min_row=6, max_row=max(wanted), values_only=True), 6):
            if rn not in wanted:
                continue
            rec = wanted[rn]
            get = lambda label: raw[col[label]]
            key = [str(get('Ký hiệu')).strip().upper(), str(get('Số hóa đơn')).zfill(8)]
            assert key == rec['key'], (name, rn, 'key')
            for i, label in enumerate(['Doanh số bán chưa thuế', 'Tiền thuế GTGT', 'Tổng tiền']):
                val = get(label)
                num = None if val in (None, '') else Decimal(str(val))
                expected = rec['money'][i]
                assert num == (None if expected is None else Decimal(str(expected))), (name, rn, label)
            assert (get('Trạng thái gửi CQT') or '') == rec['cqt_status']
            assert (get('Trạng thái hóa đơn') or '') == rec['status']
            assert bool(get('Mã của CQT')) == rec['cqt_code_present']
            seen.add(rn)
            verified += 1
        assert seen == set(wanted)
        w.close()
    inv = {tuple(r['key']): r for r in a['invoices']}
    policy = json.loads((WORK / 'confirmed-blank-payment-zero.json').read_text(encoding='utf-8'))
    confirmed_keys = {tuple(k) for k in policy['invoice_keys']}
    normalized_keys = {tuple(r['key']) for r in a['invoices'] if r.get('payment_basis') == 'user_confirmed_blank_zero'}
    assert normalized_keys == confirmed_keys and len(confirmed_keys) == 24 and len(policy['roots']) == 15
    for key in confirmed_keys:
        assert inv[key]['money'] == [0, 0, None]
        assert inv[key]['resolved_money'] == [0, 0, 0]
    zeros = 0
    for group in ['adjusted', 'untreated', 'replacement_six']:
        for r in a['groups'][group]:
            # Independently add every known field; no whole-row skip on blank.
            for i in range(3):
                vals = [inv[tuple(k)]['resolved_money'][i] for k in r['active']]
                calculated = None if any(v is None for v in vals) else sum(Decimal(str(v)) for v in vals)
                assert calculated == r['balance'][i], (group, r['root'], i)
            if r['status'] == 'zero':
                assert r['balance'] == [0, 0, 0] and not r['issues']
                zeros += 1
    assert zeros == 9430
    assert a['summary']['linked_2026_cqt_statuses'] == {'HĐ hợp lệ': 10454}
    assert a['summary']['supplemental_matching'] == {'unique': 13}
    assert Counter(r['money_match'] for r in a['supplemental']) == {'exact': 13}
    nonzero = [r for g in ['adjusted', 'untreated'] for r in a['groups'][g] if any(v not in (None, 0) for v in r['balance'])]
    assert not nonzero
    excluded = a['accountant_exclusions']
    assert len(excluded) == 1 and excluded[0]['decision']['root'] == ['1C25TUV', '00001943']
    assert excluded[0]['source_result']['balance'] == [597269, 47782, 645051]
    assert a['summary']['groups']['untreated']['roots'] == 8409
    assert a['summary']['groups']['untreated']['planned_invoices'] == 8409
    unknown = [r for g in ['adjusted', 'untreated'] for r in a['groups'][g] if r['balance'] == [0, 0, None]]
    assert not unknown
    candidates = a['summary']['replacement_zero_candidates']
    assert len(candidates) == 6
    for key in candidates:
        rec = inv[tuple(key)]
        assert rec['money'] == [0, 0, None] and rec['resolved_money'] == [0, 0, 0] and rec['cqt_status'] == 'HĐ hợp lệ'
        assert sum(edge['new'] == key and edge['kind'] == 'replacement' for edge in a['edges']) == 1
    expected_pairs = {'00002601': '00110705', '00002610': '00110706', '00002602': '00110707', '00002603': '00110708', '00002604': '00110709', '00002617': '00110710'}
    assert {r['root'][1]: r['direct_replacements'][0][1] for r in a['groups']['replacement_six']} == expected_pairs
    assert all(r['balance'] == [0, 0, 0] and r['status'] == 'zero' and not r['issues'] for r in a['groups']['replacement_six'])
    p = ROOT / 'data/Meinv_SAU_PHAT_HANH_20261029/THAY_THE/Bang_ke_hoa_don_thay_the_xuat_theo_HD_thay_the_1793293386111.xls'
    w = load_workbook(BytesIO(p.read_bytes()), read_only=True, data_only=True)
    pairs = {}
    old_sums = [Decimal(0)] * 3
    for raw in w.active.iter_rows(min_row=7, max_col=23, values_only=True):
        if not isinstance(raw[0], (int, float)):
            continue
        assert raw[1] == '1C25TUV' and raw[12] == '1C26MUV'
        old, new = str(raw[3]).zfill(8), str(raw[14]).zfill(8)
        assert old not in pairs
        pairs[old] = new
        assert raw[20] == raw[21] == 0 and raw[22] in (None, '')
        for i, c in enumerate([9, 10, 11]):
            assert Decimal(str(raw[c])) == inv[('1C25TUV', old)]['money'][i]
            old_sums[i] += Decimal(str(raw[c]))
    w.close()
    assert pairs == expected_pairs
    assert old_sums == [371344277, 29707543, 401051820]
    assert not a['summary']['duplicate_total_keys'] and not a['summary']['source_conflicts']
    for path, digest in a['manifest'].items():
        assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest, path
    dest = ROOT / 'outputs/kiem-tra-sau-phat-hanh-20261029'
    expected_counts = {'TAT_CA_CHUOI.csv': 9431, 'CAC_TRUONG_HOP_CAN_KIEM_TRA.csv': 0, 'LOAI_TRU_THEO_XAC_NHAN_KT.csv': 1}
    for name, count in expected_counts.items():
        with (dest / name).open(encoding='utf-8-sig', newline='') as f:
            rs = list(csv.reader(f))
        assert len(rs) - 1 == count
        assert all(len(r) == 11 for r in rs)
        if name == 'TAT_CA_CHUOI.csv':
            assert sum(r[5] == 'Đã về 0 cả 3 khoản' for r in rs[1:]) == 9430
            assert sum(r[5] == 'Đã loại theo xác nhận KT' for r in rs[1:]) == 1
    result = {'passed': True, 'source_reader': 'openpyxl independent of extraction parser', 'invoice_source_rows_verified': verified,
              'total_files_reopened': len(byfile), 'zero_chains_verified': zeros, 'unknown_payment_chains': len(unknown),
              'actual_nonzero_roots_in_scope': 0, 'accountant_excluded_roots': len(excluded),
              'replacement_pairs_verified': len(pairs), 'replacement_candidates_without_verified_links': 0,
              'replacement_chains_unknown_payment': 0,
              'payment_fields_zero_by_user_confirmation': len(confirmed_keys), 'confirmation_date': policy['confirmed_on'],
              'report_csv_rows': expected_counts, 'source_and_plan_hashes_unchanged': True,
              'limits': 'No live CQT query, signature validation or new invoice line-item verification.'}
    (WORK / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
