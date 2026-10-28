"""Phân nhóm sổ AMIS 2024, không sửa dữ liệu nguồn."""
from pathlib import Path
from collections import defaultdict, Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import re
import unicodedata
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'scripts/data/new-amis-2024-20261026'
OUT.mkdir(parents=True, exist_ok=True)
AMIS = ROOT / 'data/Du lieu Amis/So_chi_tiet_ban_hang_2024.xlsx'
EXCL = ROOT / 'data/HD Xuất CTY 2024.xlsx'


def txt(v):
    return '' if v is None else str(v).strip()


def norm(v):
    return ' '.join(unicodedata.normalize('NFC', txt(v)).upper().split())


def num(v):
    return Decimal(str(v)) if isinstance(v, (float, int)) else None


def rounded(v):
    return v.quantize(Decimal(1), rounding=ROUND_HALF_UP)


def native(v):
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    return v


def main():
    manifest = [{'file': str(p.relative_to(ROOT)), 'bytes': p.stat().st_size,
                 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in (AMIS, EXCL)]
    w = load_workbook(EXCL, read_only=True, data_only=True)
    exclusions = []
    by_bhh = defaultdict(list)
    buyer_pending = defaultdict(list)
    main_numbers = set()
    for i, raw in enumerate(w['HD cty 2024'].iter_rows(min_row=2, values_only=True), 2):
        r = list(raw) + [None] * max(0, 26-len(raw))
        if not r[0] or not r[1]:
            continue
        inv = str(int(r[1])).zfill(8) if isinstance(r[1], (int, float)) else txt(r[1]).zfill(8)
        key = txt(r[0]) + '/' + inv
        main_numbers.add(int(inv))
        # Lọc cả mã gốc và mã mới ghi trong hướng xử lý tách/gộp.
        codes = sorted(set(re.findall(r'BH[A-Z]*\d+(?:\.\d+)?', txt(r[12])+' '+txt(r[23]))))
        entry = {'row': i, 'key': key, 'codes': codes, 'values': r[:24]}
        exclusions.append(entry)
        for code in codes:
            by_bhh[code].append(key)
        if not codes:
            for ident in (norm(r[3]), norm(r[4])):
                if ident:
                    buyer_pending[ident].append(f'Thiếu BHH ở danh sách loại: {key}, dòng {i}')
    # Các số bổ sung chỉ nằm trong bảng ghi nhận, chưa có BHH xác nhận.
    extra_invoices = []
    for i, raw in enumerate(w['HĐ ghi nhận UVG 2024'].iter_rows(min_row=6, values_only=True), 6):
        r = list(raw) + [None] * max(0, 16-len(raw))
        if isinstance(r[1], (int, float)) and int(r[1]) not in main_numbers:
            key = f'1C24TUV/{int(r[1]):08d}'
            extra_invoices.append([i, key, r[3], r[5], r[14]])
            if norm(r[3]):
                buyer_pending[norm(r[3])].append(f'Có trong HĐ ghi nhận UVG 2024 nhưng thiếu ánh xạ BHH: {key}')
    s = load_workbook(AMIS, read_only=True, data_only=True).active
    groups = defaultdict(list)
    headers = None
    source_total = None
    for i, raw in enumerate(s.iter_rows(values_only=True), 1):
        r = list(raw)
        if i == 4:
            headers = r
        if i < 5:
            continue
        if r[0] == 'Tổng cộng':
            source_total = r
            continue
        if isinstance(r[0], datetime) and r[2]:
            groups[txt(r[2])].append((i, r))
    # Các phiếu trả lại không có mã chứng từ bán gốc: giữ mọi ứng viên để KT xác nhận.
    sales_index = defaultdict(list)
    returns = []
    for doc, lines in groups.items():
        for i, r in lines:
            if doc.startswith('BTL') or (num(r[27]) or 0) > 0:
                returns.append((doc, i, r))
            elif num(r[17]) is not None and num(r[17]) > 0:
                sales_index[(norm(r[8]), norm(r[12]))].append((doc, i, r))
    return_links = defaultdict(set)
    return_evidence = []
    for doc, i, r in returns:
        candidates = []
        for sale, si, sr in sales_index[(norm(r[8]), norm(r[12]))]:
            if sr[1] <= r[1] and num(sr[19]) is not None and num(r[19]) is not None:
                if abs(num(sr[19])-num(r[19])) <= 1 and (num(sr[17]) or 0) >= (num(r[26]) or 0):
                    candidates.append([sale, si])
                    return_links[sale].add(doc)
        return_evidence.append([doc, i, candidates])
    docs = []
    line_records = []
    reason_counts = Counter()
    for doc, lines in sorted(groups.items(), key=lambda x: (x[1][0][1][1], x[0])):
        r0 = lines[0][1]
        reasons = []
        sums = {label: sum((num(r[c]) or Decimal(0) for _, r in lines), Decimal(0))
                for label, c in [('goods',20),('discount',25),('returns',27),('reduction',28),('export_tax',30),('vat',32),('payment',34)]}
        if not isinstance(r0[1], datetime) or r0[1].year != 2024:
            reasons.append('Ngày chứng từ không thuộc năm 2024 hoặc không hợp lệ')
        if any(len({norm(r[c]) for _, r in lines}) > 1 for c in [0,1,8,9,10,11]):
            reasons.append('Cùng số chứng từ nhưng khác ngày/thông tin người mua')
        if not txt(r0[9]):
            reasons.append('Thiếu tên khách hàng')
        if doc.startswith('BTL') or sums['returns'] or any((num(r[26]) or 0) != 0 for _, r in lines):
            reasons.append('Phiếu trả lại: chưa xác định chứng từ bán gốc để bù trừ')
        if doc in return_links:
            reasons.append('Có thể liên quan trả lại: '+', '.join(sorted(return_links[doc])))
        if any(r[3] or r[5] for _, r in lines):
            reasons.append('Có ký hiệu/số hóa đơn tại AMIS, cần xác nhận trước khi lập mới')
        for ident in {norm(r0[9]), norm(r0[11])}:
            reasons.extend(buyer_pending.get(ident, []))
        rates = set()
        line_rates = {}
        for i, r in lines:
            if any(num(r[c]) is None for c in [17,19,20,25,27,28,30,32,34]):
                reasons.append('Thiếu giá trị số bắt buộc trong dòng hàng')
                continue
            if not txt(r[7]) or not txt(r[16]) or not txt(r[12]):
                reasons.append('Thiếu tên hàng/đơn vị tính/mã hàng')
            if any(num(r[c]) < 0 for c in [17,19,20,32,34]):
                reasons.append('Có số lượng/đơn giá/tiền âm')
            if num(r[17]) <= 0 or num(r[20]) <= 0:
                reasons.append('Dòng số lượng hoặc thành tiền bằng 0, chưa rõ hàng khuyến mại/diễn giải')
            if num(r[25]) or num(r[28]) or num(r[30]):
                reasons.append('Có chiết khấu/giảm giá/thuế xuất khẩu cần kiểm tra cách nhập')
            if abs(rounded(num(r[17])*num(r[19]))-num(r[20])) > 1:
                reasons.append('Thành tiền lệch số lượng nhân đơn giá quá 1 đồng')
            if num(r[20]) > 0:
                candidates = [rate for rate in (5,8,10) if abs(rounded(num(r[20])*Decimal(rate)/100)-num(r[32])) <= 1]
                if len(candidates) == 1:
                    rates.add(candidates[0]); line_rates[i] = candidates[0]
                else:
                    reasons.append('Không suy ra duy nhất thuế suất 5/8/10% từ tiền hàng và VAT')
            if abs(num(r[20])-num(r[25])-num(r[27])-num(r[28])+num(r[30])+num(r[32])-num(r[34])) > 1:
                reasons.append('Tiền hàng + VAT sau giảm trừ lệch thanh toán quá 1 đồng')
        if len(rates) != 1:
            reasons.append('Chứng từ chưa có một thuế suất duy nhất')
        if sums['payment'] <= 0:
            reasons.append('Tổng thanh toán không dương')
        if sums['goods']+sums['vat'] != sums['payment']:
            reasons.append('Tổng tiền hàng + VAT không bằng tổng thanh toán nguồn')
        reasons = list(dict.fromkeys(reasons))
        if doc in by_bhh:
            category = 'Đã loại'
            reasons = ['Danh sách loại trừ: '+', '.join(by_bhh[doc])] + reasons
        elif reasons:
            category = 'Cần kiểm tra'
        else:
            category = 'Đưa vào import'
        reason_counts.update(reasons if category == 'Cần kiểm tra' else [])
        item = {'doc':doc,'category':category,'reasons':reasons,'date':r0[1], 'accounting_date':r0[0],
                'invoice_date_source':r0[4], 'customer_code':r0[8],'buyer':r0[9],'address':r0[10],'tax':r0[11],
                'sums':sums,'rate':next(iter(rates)) if len(rates)==1 else None,'rows':[i for i,_ in lines],
                'lines':[{'row':i,'values':r,'rate':line_rates.get(i)} for i,r in lines]}
        docs.append(item)
        for i,r in lines:
            line_records.append((i,category,r))
    summary = {}
    for category in ['Đưa vào import','Đã loại','Cần kiểm tra']:
        ds = [d for d in docs if d['category']==category]
        summary[category] = {'documents':len(ds),'lines':sum(len(d['lines']) for d in ds),
                             **{k:sum((d['sums'][k] for d in ds),Decimal(0)) for k in ['goods','returns','vat','payment']}}
    assert len(line_records)==44578 and len({i for i,_,_ in line_records})==44578
    assert sum(x['documents'] for x in summary.values())==35017
    for key,c in [('goods',20),('returns',27),('vat',32),('payment',34)]:
        assert sum(d['sums'][key] for d in docs)==num(source_total[c]),key
    payload={'manifest':manifest,'headers':headers,'source_total':source_total,'summary':summary,
             'reason_counts':dict(reason_counts),'docs':docs,'exclusions':exclusions,
             'exclusion_codes_not_found':sorted(set(by_bhh)-set(groups)),
             'extra_invoices':extra_invoices,'return_evidence':return_evidence,
             'rules':{'date':'Ngày chứng từ AMIS (người dùng xác nhận lấy theo AMIS)',
                      'payment':'TM/CK (người dùng xác nhận)',
                      'tax_rate':'Suy ra duy nhất trong 5/8/10% khi VAT khớp sau làm tròn nguyên đồng, dung sai 1đ/dòng; không tự cân số tiền',
                      'exclusions':'HD cty 2024 cột BHH và mã BHH tại Hướng xử lý; tab F0/F1/F2 chỉ là tham khảo, không loại toàn tab',
                      'returns':'Giữ riêng phiếu trả lại và mọi ứng viên bán cùng mã KH/mã hàng/đơn giá (±1đ), bán trước hoặc cùng ngày, đủ SL. Không tự chọn đơn để bù trừ.'}}
    (OUT/'prepared.json').write_text(json.dumps(payload,ensure_ascii=False,default=native),encoding='utf8')
    report={k:v for k,v in payload.items() if k not in ['docs','return_evidence','exclusions','headers','source_total']}
    (OUT/'summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=native),encoding='utf8')
    print(json.dumps(report,ensure_ascii=False,indent=2,default=native))


if __name__=='__main__':
    main()
