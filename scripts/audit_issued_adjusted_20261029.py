"""Đối chiếu chỉ đọc bộ đã điều chỉnh với bảng kê phát hành mới."""
from pathlib import Path
from collections import defaultdict, Counter
from decimal import Decimal
import json, sys, hashlib
from openpyxl import load_workbook
from prepare_invoice_reports import rows, invoice_key, invoice_day, amount
from filter_invoice_exceptions import normal

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'data/Meinv'
OUT = ROOT / 'scripts/data/issued-adjusted-20261029'
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding='utf-8')
planpath = ROOT / 'outputs/HOA_DON_DA_DIEU_CHINH_IMPORT_2_SHEET.xlsx'
manifest = {}
def source(p):
    manifest[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
source(planpath)
w = load_workbook(planpath, read_only=True, data_only=True)
plans = defaultdict(list)
seen = set()
for r in w['Import dieu chinh'].iter_rows(min_row=10, max_col=24, values_only=True):
    if r[0] is None or r[0] in seen: continue
    seen.add(r[0])
    plans[invoice_key(r[9], r[10])].append({'seq': r[0], 'money': [amount(r[c]) for c in [19,21,22]]})
baseline = {}
for rn, r in enumerate(w['F0, F1, F2, ...'].iter_rows(min_row=3, max_col=55, values_only=True), 3):
    key = invoice_key(r[0],r[1])
    if key in plans:
        baseline[key] = {'row': rn, 'before': [amount(r[c]) for c in [24,25,26]],
                         'old': [invoice_key(r[c],r[c+1]) for c in [0,6,12,18] if r[c] and r[c+1]]}
w.close()
print('Plan loaded',len(plans),len(seen),flush=True)
records = {}; conflicts=[]; parents=defaultdict(set); children=defaultdict(set); edges=set(); linkcounts={}
link_records={}
for p in sorted((BASE/'HD_bi_dieu_chinh_2024_2026').glob('*.xls')):
    source(p); count=0
    for rn,r in rows(p):
        if rn<7 or amount(r.get('A')) is None: continue
        old=invoice_key(r.get('B'),r.get('D')); new=invoice_key(r.get('M'),r.get('O'))
        if not old or not new: continue
        count+=1; edges.add((old,new)); parents[new].add(old); children[old].add(new)
        for key, cols, dc in [(old,['J','K','L'],'C'),(new,['U','V','W'],'N')]:
            money=[amount(r.get(c)) for c in cols]
            entry={'money':money,'date':invoice_day(r.get(dc)), 'source':str(p.relative_to(ROOT)), 'row':rn}
            if key in link_records and link_records[key]['money'] != money:
                conflicts.append({'kind':'link_amount_conflict','key':key,'a':link_records[key],'b':entry})
            else: link_records[key]=entry
    linkcounts[p.name]=count
# Tổng phát hành được ưu tiên; không suy tiền trống thành 0.
totals = [BASE/'2024.xls'] + sorted((BASE/'2025').glob('*.xls')) + sorted((BASE/'2026').glob('*.xls'))
duplicate_totals=[]
for p in totals:
    source(p)
    for rn,r in rows(p):
        if rn<6 or amount(r.get('A')) is None: continue
        key=invoice_key(r.get('B'),r.get('C'))
        if not key: continue
        entry={'money':[amount(r.get(c)) for c in ['P','Q','R']], 'date':invoice_day(r.get('D')), 'status':r.get('U') or '', 'cqt':r.get('S') or '', 'source':str(p.relative_to(ROOT)), 'row':rn}
        if key in records: duplicate_totals.append(key)
        records[key]=entry
print('Sources loaded',len(records),len(edges),flush=True)
def reach(root):
    found=set(); queue=[root]
    while queue:
        key=queue.pop()
        if key in found: continue
        found.add(key); queue.extend(children[key])
    return found
results=[]
for root, ps in sorted(plans.items()):
    keys=reach(root); issues=[]
    expected_old=set(baseline.get(root,{}).get('old',[]))
    if not expected_old <= keys: issues.append({'kind':'old_link_missing','keys':sorted(expected_old-keys)})
    active=[]; balances=[Decimal(0)]*3
    for key in sorted(keys):
        if key not in records:
            issues.append({'kind':'missing_issued_record','key':key}); continue
        r=records[key]; st=normal(r['status'])
        if 'bi thay the' in st or 'huy' in st:
            issues.append({'kind':'replacement_or_cancel_requires_review','key':key}); continue
        if len(parents[key])>1: issues.append({'kind':'multiple_parents','key':key})
        if key in link_records and r['money']!=link_records[key]['money']:
            issues.append({'kind':'total_link_money_difference','key':key,'total':r['money'],'link':link_records[key]['money']})
        if None in r['money']:
            issues.append({'kind':'blank_money','key':key,'money':r['money']}); continue
        active.append(key)
        balances=[a+b for a,b in zip(balances,r['money'])]
    # So số tiền và số lượng HĐ mới theo từng gốc, không ghép chỉ dựa tổng tiền.
    actual_new=[k for k in active if records[k]['date'] and records[k]['date'].year==2026]
    actual_money=Counter(tuple(records[k]['money']) for k in actual_new)
    expected_money=Counter(tuple(p['money']) for p in ps)
    missing=expected_money-actual_money; extra=actual_money-expected_money
    if missing: issues.append({'kind':'planned_amounts_missing','items':list(missing.items())})
    if extra: issues.append({'kind':'additional_2026_amounts','items':list(extra.items())})
    status='needs_review' if issues else 'zero' if all(x==0 for x in balances) else 'nonzero'
    results.append({'root':root,'status':status,'balance':balances,'issues':issues,'planned_count':len(ps),'issued_2026_count':len(actual_new),'invoices':[{'key':k,**records[k]} for k in sorted(keys) if k in records], 'baseline':baseline.get(root)})
def native(v):
    if isinstance(v,Decimal):return int(v) if v==int(v) else float(v)
    if hasattr(v,'isoformat'):return v.isoformat()
    raise TypeError(type(v))
summary={'read_only':True,'roots':len(plans),'planned_invoices':len(seen),'link_counts':linkcounts,'unique_edges':len(edges),'issued_records':len(records),'duplicate_total_keys':duplicate_totals,'source_conflicts':conflicts,'statuses':dict(Counter(r['status'] for r in results)),'issue_kinds':dict(Counter(i['kind'] for r in results for i in r['issues'])),'balance_totals':[sum(r['balance'][c] for r in results) for c in range(3)]}
payload={'summary':summary,'results':results,'manifest':manifest}
(OUT/'audit.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=native),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,default=native),flush=True)
for r in results:
    if r['status']!='zero':print(json.dumps({k:v for k,v in r.items() if k not in ['invoices','baseline']},ensure_ascii=False,default=native),flush=True)
