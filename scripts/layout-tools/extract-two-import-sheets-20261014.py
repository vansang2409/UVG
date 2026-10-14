from pathlib import Path, PurePosixPath
from zipfile import ZipFile
from lxml import etree as E
from openpyxl import load_workbook
from copy import copy
from collections import Counter
import hashlib, json, posixpath, sys

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
untreated = '--untreated' in sys.argv
stem = 'HOA_DON_CHUA_DIEU_CHINH' if untreated else 'HOA_DON_DA_DIEU_CHINH'
source = ROOT / 'outputs' / (stem + '_F0_2024_T6_2025.xlsx')
output = ROOT / 'outputs' / (stem + '_IMPORT_2_SHEET.xlsx')
audit = ROOT / 'scripts/data' / ('import-2-sheet-untreated-20261014' if untreated else 'import-2-sheet-20261014')
audit.mkdir(parents=True, exist_ok=True)
assert not output.exists(), 'Không ghi đè file đã tồn tại'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
before_hash = sha(source)
ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
tag = lambda x: '{' + ns + '}' + x
rid_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
mapping = [('Import dieu chinh', 'Import dieu chinh'), ('Chua dieu chinh', 'F0')] if untreated else [('Import dieu chinh', 'Import dieu chinh'), ('Dieu chinh', 'F0, F1, F2, ...')]
with ZipFile(source) as z:
    parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    wb = E.fromstring(parts['xl/workbook.xml'])
    sheets = wb.find(tag('sheets'))
    old_sheets = list(sheets)
    selected = []
    index_map = {}
    for old, new in mapping:
        s = next(s for s in old_sheets if s.get('name') == old)
        index_map[old_sheets.index(s)] = len(selected)
        s.set('name', new)
        selected.append(s)
    sheets[:] = selected
    for view in wb.iter(tag('workbookView')):
        view.set('activeTab', '0'); view.set('firstSheet', '0')
    for names in wb.findall(tag('definedNames')):
        for item in list(names):
            if item.get('localSheetId') is not None:
                old_id = int(item.get('localSheetId'))
                if old_id not in index_map:
                    names.remove(item); continue
                item.set('localSheetId', str(index_map[old_id]))
            if item.text:
                for old, new in mapping:
                    item.text = item.text.replace("'" + old + "'!", "'" + new + "'!")
    def xml(root):
        return E.tostring(root, xml_declaration=True, encoding='utf-8', standalone=True)
    parts['xl/workbook.xml'] = xml(wb)
    rels_path = 'xl/_rels/workbook.xml.rels'
    rels = E.fromstring(parts[rels_path])
    keep_ids = {s.get(rid_attr) for s in selected}
    for r in list(rels):
        if (r.get('Type').endswith('/worksheet') and r.get('Id') not in keep_ids) or r.get('Type').endswith('/calcChain'):
            rels.remove(r)
    parts[rels_path] = xml(rels)
    # Giữ nguyên XML hai sheet và các tài nguyên mà chúng sử dụng.
    def relpath(part):
        if not part: return '_rels/.rels'
        p = PurePosixPath(part)
        return str(p.parent / '_rels' / (p.name + '.rels'))
    keep = {'[Content_Types].xml'}
    pending = ['']
    seen = set()
    while pending:
        part = pending.pop()
        if part in seen: continue
        seen.add(part)
        if part: keep.add(part)
        rp = relpath(part)
        if rp not in parts: continue
        keep.add(rp)
        for r in E.fromstring(parts[rp]):
            if r.get('TargetMode') == 'External': continue
            t = r.get('Target')
            target = t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), t))
            assert target in parts, ('Thiếu tài nguyên', target)
            pending.append(target)
    # Cập nhật thông tin số sheet trong thuộc tính tài liệu.
    if 'docProps/app.xml' in keep:
        app = E.fromstring(parts['docProps/app.xml'])
        apns = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
        vt = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
        titles = app.find('{' + apns + '}TitlesOfParts')
        if titles is not None:
            vector = titles[0]; vector[:] = []
            vector.set('size', '2')
            for _, name in mapping: E.SubElement(vector, '{' + vt + '}lpstr').text = name
        headings = app.find('{' + apns + '}HeadingPairs')
        if headings is not None:
            variants = list(headings[0])
            for i in range(len(variants) - 1):
                if ''.join(variants[i].itertext()) == 'Worksheets':
                    variants[i + 1][0].text = '2'
        parts['docProps/app.xml'] = xml(app)
    ct = E.fromstring(parts['[Content_Types].xml'])
    for item in list(ct):
        if item.get('PartName') and item.get('PartName').lstrip('/') not in keep:
            ct.remove(item)
    parts['[Content_Types].xml'] = xml(ct)
    with ZipFile(output, 'w') as dest:
        for item in z.infolist():
            if item.filename in keep: dest.writestr(item, parts[item.filename])

# Kiểm độc lập toàn bộ giá trị/công thức/định dạng hai sheet sau khi tách.
a, b = load_workbook(source, data_only=False), load_workbook(output, data_only=False)
assert b.sheetnames == [new for _, new in mapping]
checked = 0
for old, new in mapping:
    sa, sb = a[old], b[new]
    assert (sa.max_row, sa.max_column) == (sb.max_row, sb.max_column)
    assert sa.freeze_panes == sb.freeze_panes
    assert set(sa.merged_cells.ranges) == set(sb.merged_cells.ranges)
    assert {k: sa.tables[k].ref for k in sa.tables} == {k: sb.tables[k].ref for k in sb.tables}
    checked_styles = set()
    for ra, rb in zip(sa, sb):
        for ca, cb in zip(ra, rb):
            assert ca.value == cb.value, (old, ca.coordinate)
            if ca.data_type == 'f': assert '!' not in ca.value, ('Cần kiểm tham chiếu ngoài sheet', ca.coordinate)
            key = (ca.style_id, cb.style_id)
            if key not in checked_styles:
                for attr in ['font', 'fill', 'border', 'alignment', 'protection', 'number_format']:
                    assert copy(getattr(ca, attr)) == copy(getattr(cb, attr)), (old, ca.coordinate, attr)
                checked_styles.add(key)
            checked += 1
heads = {r[0] for r in b['Import dieu chinh'].iter_rows(min_row=10, values_only=True) if isinstance(r[0], (int, float))}
assert len(heads) == (8410 if untreated else 2032)
assert sha(source) == before_hash
import_lines = sum(1 for r in b['Import dieu chinh'].iter_rows(min_row=10, values_only=True) if isinstance(r[0], (int, float)))
report = {'passed': True, 'source_sha256': before_hash, 'output_sha256': sha(output), 'sheets': b.sheetnames, 'invoices': len(heads), 'import_lines': import_lines, 'checked_cells': checked, 'source_unchanged': True}
(audit / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
