from pathlib import Path
from openpyxl import load_workbook
from itertools import zip_longest
from copy import copy
import json, hashlib, shutil, sys

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'scripts/data/import-reasons-20261014'
plan = json.loads((OUT / 'plan.json').read_text(encoding='utf-8'))
index = int(sys.argv[1])
f = plan['files'][index]
oldpath = OUT / 'before' / f['name']
newpath = ROOT / 'outputs' / f['name'] if '--final' in sys.argv else OUT / 'staged' / f['name']
a, b = [load_workbook(p, data_only=False) for p in [oldpath, newpath]]
assert a.sheetnames == b.sheetnames
expected = {c['row']: c for c in f['changes']}
changes = 0
styles = {}
for sa, sb in zip(a, b):
    assert sa.freeze_panes == sb.freeze_panes, (sa.title, 'freeze')
    assert set(sa.merged_cells.ranges) == set(sb.merged_cells.ranges), (sa.title, 'merges')
    assert {n:sa.tables[n].ref for n in sa.tables} == {n:sb.tables[n].ref for n in sb.tables}, (sa.title, 'tables')
    assert (sa.max_row, sa.max_column) == (sb.max_row, sb.max_column), (sa.title, 'dimensions')
    for ra, rb in zip(sa, sb):
        for ca, cb in zip(ra, rb):
            allowed = sa.title == 'Import dieu chinh' and ca.column == 14 and ca.row in expected
            if allowed:
                assert ca.value == expected[ca.row]['old']
                assert cb.value == expected[ca.row]['new']
                changes += 1
            else:
                assert ca.value == cb.value, (sa.title, ca.coordinate, ca.value, cb.value)
            if ca.has_style or cb.has_style:
                key = (ca.style_id, cb.style_id)
                if key not in styles:
                    for attr in ['font', 'fill', 'border', 'alignment', 'protection', 'number_format']:
                        assert copy(getattr(ca, attr)) == copy(getattr(cb, attr)), (sa.title, ca.coordinate, attr)
                    styles[key] = True
assert changes == len(expected)
target = ROOT / 'outputs' / f['name']
if '--final' not in sys.argv:
    assert hashlib.sha256(target.read_bytes()).hexdigest() == f['sha256_before'], 'File đã thay đổi trong lúc xử lý'
report = {'passed': True, 'name': f['name'], 'edited_cells': changes, 'invoices': f['invoices'], 'counts': f['counts'], 'sha256_after': hashlib.sha256(newpath.read_bytes()).hexdigest(), 'other_cell_values_and_styles_unchanged': True}
(OUT / f'verified-{index}.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
if '--promote' in sys.argv:
    a.close(); b.close()
    shutil.copy2(newpath, target)
    assert hashlib.sha256(target.read_bytes()).hexdigest() == report['sha256_after']
print(json.dumps(report, ensure_ascii=False))
