"""Read-only verification of report exports against current invoice totals."""
import hashlib
import json
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from filter_invoice_exceptions import normal

data = json.loads((ROOT / 'scripts/data/invoice_reports.json').read_text(encoding='utf-8'))
manifest = json.loads((ROOT / 'SOURCE_MANIFEST.json').read_text(encoding='utf-8'))
assert manifest == data['manifest'] and len(manifest) == 18
for entry in manifest:
    p = ROOT / entry['path']
    assert p.stat().st_size == entry['bytes']
    assert hashlib.sha256(p.read_bytes()).hexdigest() == entry['sha256'], p

def key(sig, no):
    return str(sig).strip(), str(no).removesuffix('.0').zfill(8)

def day(value):
    if isinstance(value, datetime):
        return value.date()
    return datetime.strptime(str(value), '%d/%m/%Y').date()

def value(v):
    return None if v == '' else v

totals = {}
for p in (ROOT / 'data/Du lieu Meinvoice/Bang_ke_hoa_don_da_su_dung_2024_2025').glob('*.xlsx'):
    w = load_workbook(p, read_only=True, data_only=True)
    for r in w.worksheets[0].iter_rows(min_row=6, values_only=True):
        if not isinstance(r[0], (int, float)) or not r[1] or r[2] is None:
            continue
        k = key(r[1], r[2])
        assert k not in totals
        totals[k] = (day(r[3]), [value(v) for v in r[15:18]], r[20])
    w.close()
assert len(totals) == 48874

all_keys = []
for name, filename, sheet_name, offset, freeze in [
    ('adjusted', 'HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx', 'Dieu chinh', 0, 'C3'),
    ('exceptions', 'NGOAI_LE_F0_2024_T6_2025.xlsx', 'Ngoai le', 2, 'E3'),
]:
    groups = data['groups'][name]
    blocks = data['summary'][name]['max_blocks']
    w = load_workbook(ROOT / 'outputs' / filename, data_only=False)
    assert w.sheetnames == [sheet_name, 'Huong dan', 'Chi tiet nguon']
    s = w.worksheets[0]
    assert s.max_row == len(groups) + 2 and s.max_column == offset + 6 * blocks + 4
    assert s.freeze_panes == freeze and s.sheet_view.showGridLines is False
    assert len(s.tables) == 1
    table = next(iter(s.tables.values()))
    assert table.ref == f'A2:{s.cell(2,s.max_column).column_letter}{s.max_row}'
    assert table.autoFilter is not None and table.tableStyleInfo.name == 'TableStyleMedium2'
    assert s.row_dimensions[1].height == 26 and s.row_dimensions[2].height == 42
    assert s['A1'].font.name == 'Arial' and s['A1'].font.sz == 10
    if offset:
        assert [s['A2'].value, s['B2'].value] == ['Tên người mua', 'MST/CCCD chủ hộ nguyên bản']
    invoice_keys = set()
    for rn, group in enumerate(groups, 3):
        expected = [group['name'], group['tax']] if offset else []
        for block in range(blocks):
            if block < len(group['invoices']):
                r = group['invoices'][block]
                expected.extend([r[0], r[1], datetime.combine(day(r[2]), datetime.min.time()), *r[4:]])
            else:
                expected.extend([None] * 6)
        expected.extend([*group['sums'], group['notes']])
        actual = [s.cell(rn,c).value for c in range(1, s.max_column + 1)]
        assert list(map(value, actual)) == list(map(value, expected)), (name, rn, group['root'])
        root_day = day(group['invoices'][0][2])
        assert day('01/01/2024') <= root_day <= day('30/06/2025')
        if offset:
            assert s.cell(rn,2).number_format == '@'
            assert isinstance(s.cell(rn,2).value, str) or s.cell(rn,2).value is None
        active = []
        replaced = set()
        for r in group['detail']:
            if r[7].startswith('Thay thế cho '):
                replaced.add(tuple(r[7].removeprefix('Thay thế cho ').split('/')))
        for r in group['invoices']:
            k = tuple(r[:2])
            invoice_keys.add(k)
            assert day(r[2]) <= day('31/12/2025')
            confirmation = next((c for c in data.get('user_confirmed_dates',[]) if tuple(c['invoice']) == k),None)
            expected_day = day(confirmation['date']) if confirmation else totals[k][0]
            assert (day(r[2]),r[4:]) == (expected_day,totals[k][1]), (k,'source totals/date mismatch')
            d = next(t for t in group['detail'] if tuple(t[2:4]) == k)
            assert d[8] == totals[k][2]
            if k in replaced or 'huy' in normal(d[8]):
                continue
            amounts = r[4:].copy()
            if k == ('1C24TUV','00000574'):
                assert amounts[2] is None
                amounts[2] = 0
            active.append(amounts)
        cumulative = [sum(Decimal(str(r[i])) for r in active) if all(r[i] is not None for r in active) else None for i in range(3)]
        assert cumulative == group['sums'], (group['root'],'cumulative mismatch')
    assert len(invoice_keys) == data['summary'][name]['invoices']
    all_keys.append(invoice_keys)
    t = w['Chi tiet nguon']
    detail = [r for g in groups for r in g['detail']]
    assert t.max_row == len(detail) + 1 and t.max_column == 13
    assert t.freeze_panes == 'A2' and len(t.tables) == 1
    for rn,r in enumerate(detail,2):
        expected = [*r[:4],datetime.combine(day(r[4]),datetime.min.time()),*r[5:]]
        actual = [t.cell(rn,c).value for c in range(1,14)]
        assert list(map(value,actual)) == list(map(value,expected)), (name,'detail',rn)
    assert not any(c.data_type == 'f' for sh in w for row in sh for c in row)
    assert not any(c.value in ('#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!') for sh in w for row in sh for c in row)
    w.close()
    print(json.dumps({'report':name, 'verified_roots':len(groups), 'verified_invoices':len(invoice_keys),
                      'checks':'source totals, cumulative amounts, source blanks, scope, format, filters, panes'},ensure_ascii=False))
assert not (all_keys[0] & all_keys[1])
assert {p.name for p in (ROOT/'outputs').iterdir()} == {
    'HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','NGOAI_LE_F0_2024_T6_2025.xlsx'}
print('PASS: 18 unchanged source hashes; outputs has exactly two Excel files, no subfolders.')
