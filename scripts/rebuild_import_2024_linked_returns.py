"""Lập lại dữ liệu nhập mới từ liên kết BTL–BHH xác nhận, không dò ứng viên."""
from pathlib import Path
from collections import defaultdict, Counter
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from io import BytesIO
import hashlib
import json
import gzip
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT/'scripts/data/amis-2024-linked-20261026'
OUT = ROOT/'outputs/amis-2024-theo-btl-bhh-20261026'
LINK = ROOT/'outputs/btl-bhh-20261026/LIEN_KET_BTL_BHH_2024.xlsx'
CONFIRMED_ROUNDING = {
    'BHHT2024003185','BHHT2024011030','BHHT2024015896','BHHT2024017438',
    'BHHN2024000179','BHHT2024023472','BHHT2024025057',
}
CONFIRMED_PROMOTIONS = {'BHHT2024013153','BHHT2024013026'}
# Kế toán xác nhận ngày 27/10: lập bán gốc trước, trả lại lập điều chỉnh riêng.
SPLIT_RETURNS = True
PROMOTION_REASONS = {
    'Dòng số lượng hoặc thành tiền bằng 0, chưa rõ hàng khuyến mại/diễn giải',
    'Chứng từ chưa có một thuế suất duy nhất','Tổng thanh toán không dương',
}

def dec(v): return Decimal(str(v or 0))
def rnd(v): return v.quantize(Decimal(1), rounding=ROUND_HALF_UP)
def native(v):
    if isinstance(v, Decimal): return int(v) if v == int(v) else float(v)
    if isinstance(v, datetime): return v.isoformat()
    raise TypeError(type(v))
def triple(r, cols): return [dec(r[c]) for c in cols]
def aggregate(lines, cols, keycols):
    result=defaultdict(lambda:[Decimal(0)]*3)
    for r in lines:
        k=tuple(r[c] for c in keycols)
        result[k]=[a+b for a,b in zip(result[k],triple(r,cols))]
    return dict(result)

def vn_number(v):
    s=format(dec(v),'f');whole,_,fraction=s.partition('.')
    text=f'{int(whole):,}'.replace(',','.')
    return text+(','+fraction.rstrip('0') if fraction.rstrip('0') else '')

def detailed_reasons(code, lines, reasons, return_rows):
    """Giải thích bằng số nguồn; chênh lệch = AMIS - ROUND(SL × đơn giá, 0)."""
    result=[]
    for reason in reasons:
        if reason=='Thành tiền lệch số lượng nhân đơn giá quá 1 đồng':
            for l in lines:
                r=l['values'];product=dec(r[17])*dec(r[19]);expected=rnd(product);delta=dec(r[20])-expected
                if abs(delta)>1:
                    rounding=f' (làm tròn {vn_number(expected)}đ)' if expected!=product else ''
                    result.append(f"Dòng AMIS {l['row']} / {r[12]}: SL {vn_number(r[17])} × đơn giá {vn_number(r[19])}đ = {vn_number(product)}đ{rounding}; thành tiền AMIS {vn_number(r[20])}đ; AMIS {'thấp' if delta<0 else 'cao'} hơn {vn_number(abs(delta))}đ (chênh {'+' if delta>0 else '-'}{vn_number(abs(delta))}đ).")
        elif reason.startswith('BTL') and ('giá/tiền trả lại khác giá dòng bán' in reason or 'SL hoặc tiền trả lại vượt dòng bán' in reason):
            prefix=reason.split(':',1)[0]
            rr=[r for r in return_rows if f'{r[0]} / {r[4]}'==prefix]
            if len(rr)!=1:result.append(reason);continue
            ret=rr[0];sales=[l for l in lines if l['values'][12]==ret[4] and l['values'][16]==ret[6]]
            if len(sales)!=1:result.append(reason);continue
            l=sales[0];r=l['values'];price_delta=dec(ret[8])-dec(r[19]);expected=rnd(dec(ret[7])*dec(r[19]));money_delta=dec(ret[9])-expected
            result.append(f"{prefix}: đơn giá bán {vn_number(r[19])}đ, giá trả {vn_number(ret[8])}đ (chênh {'+' if price_delta>=0 else '-'}{vn_number(abs(price_delta))}đ/{r[16]}); SL trả {vn_number(ret[7])}; tiền theo giá bán {vn_number(expected)}đ, tiền trả thực tế {vn_number(ret[9])}đ (trả {'cao' if money_delta>=0 else 'thấp'} hơn {vn_number(abs(money_delta))}đ). Dòng bán AMIS {l['row']}; chưa bù trừ, chưa tự chỉnh đơn giá.")
        else:result.append(reason)
    return result

def main():
    WORK.mkdir(parents=True,exist_ok=True)
    # Chỉ tái dùng bản đọc giá trị nguồn đã lưu; không dùng phân nhóm hay ứng viên cũ.
    snapshot=ROOT/'scripts/data/new-amis-2024-20261026/prepared.json'
    old=json.loads(snapshot.read_text(encoding='utf8') if snapshot.exists() else gzip.decompress(snapshot.with_suffix('.json.gz').read_bytes()).decode('utf8'))
    link_check=json.loads((ROOT/'scripts/data/btl-bhh-live-20261026/verification.json').read_text(encoding='utf8'))
    assert link_check['passed']
    manifest=old['manifest']+[{'file':str(LINK.relative_to(ROOT)), 'sha256':link_check['sha256']}]
    confirmed_link_file=ROOT/'scripts/data/btl-bhh-live-20261026/user-confirmed-links.json'
    if confirmed_link_file.exists():manifest.append({'file':str(confirmed_link_file.relative_to(ROOT)),'sha256':hashlib.sha256(confirmed_link_file.read_bytes()).hexdigest()})
    template=ROOT/'data/Du lieu Meinvoice/Mau_hoa_don_GTGT_MTT_1_thue_suat.xls'
    manifest.append({'file':str(template.relative_to(ROOT)), 'sha256':'64346fb386624e1f900ac9130303737f79348491951eb17e53133190514847be'})
    for m in manifest: assert hashlib.sha256((ROOT/m['file']).read_bytes()).hexdigest()==m['sha256'],m['file']
    docs={d['doc']:d for d in old['docs']}
    excluded=defaultdict(list)
    for e in old['exclusions']:
        for code in e['codes']:excluded[code].append(e['key'])
    links=[list(r) for r in load_workbook(LINK,read_only=True,data_only=True)['Chi tiet BTL-BHH'].iter_rows(min_row=2,values_only=True)]
    assert len(links)==219
    by_bhh=defaultdict(list);by_btl=defaultdict(list)
    for r in links:
        by_btl[r[0]].append(r)
        if r[2]:by_bhh[r[2]].append(r)
    # Từng dòng liên kết phải khớp chính xác dòng trả lại trong sổ AMIS.
    assert set(by_btl)=={k for k in docs if k.startswith('BTL')}
    for btl,rs in by_btl.items():
        a=Counter((r[4],dec(r[7]),dec(r[8]),dec(r[9]),dec(r[11])) for r in rs)
        b=Counter((l['values'][12],dec(l['values'][26]),dec(l['values'][19]),dec(l['values'][27]),-dec(l['values'][32])) for l in docs[btl]['lines'])
        assert a==b,btl
    records=[];net_audit=[];link_issues={};return_status={}
    for d in old['docs']:
        code=d['doc'];original=d['lines'];net=deepcopy(original)
        reasons=[];btls=sorted({r[0] for r in by_bhh.get(code,[])})
        full=False;applied=False
        if code.startswith('BTL'):
            missing=any(not r[2] for r in by_btl[code])
            category='Cần kiểm tra' if missing else 'Phiếu trả lại đã liên kết'
            reasons=['AMIS đang trống Số CT bán hàng gốc; chưa bù trừ, không suy đoán BHH'] if missing else ['Dùng liên kết BHH gốc thực tế; không lập hóa đơn mới cho phiếu BTL']
            net=[]
        else:
            if code in by_bhh and SPLIT_RETURNS:
                rs=by_bhh[code]
                return_status[code]='Tách điều chỉnh sau khi có hóa đơn bán gốc'
                # Giá trả khác bán chỉ giữ kiểm tra ở phần trả, không chặn bán gốc.
                mismatch=[]
                for r in rs:
                    sales=[l for l in original if l['values'][12]==r[4] and l['values'][16]==r[6]]
                    if len(sales)==1 and abs(dec(sales[0]['values'][19])-dec(r[8]))>1:
                        mismatch.append(f'{r[0]} / {r[4]}: giá/tiền trả lại khác giá dòng bán')
                if mismatch:
                    link_issues[code]='; '.join(detailed_reasons(code,original,mismatch,rs)).replace('chưa bù trừ, chưa tự chỉnh đơn giá.','chờ kế toán chốt phần điều chỉnh; giữ nguyên bán gốc và số tiền BTL, chưa tự coi là giảm giá.')
            if code in by_bhh and not SPLIT_RETURNS:
                rs=by_bhh[code]
                sale=aggregate([l['values'] for l in original],[17,20,32],[12,16])
                ret=aggregate(rs,[7,9,11],[4,6])
                full=sale==ret
                if full:
                    net=[];applied=True
                else:
                    # Ưu tiên khớp nguyên dòng theo SKU/ĐVT/SL/tiền hàng/VAT.
                    # Nếu chỉ trả một phần một dòng duy nhất, giá chỉ được lệch do làm tròn <=1đ.
                    for r in rs:
                        qty,goods,vat=triple(r,[7,9,11])
                        cand=[l for l in net if l['values'][12]==r[4] and l['values'][16]==r[6] and dec(l['values'][17])>0]
                        exact=[l for l in cand if triple(l['values'],[17,20,32])==[qty,goods,vat]]
                        if len(exact)==1:chosen=exact[0]
                        elif len(cand)==1:chosen=cand[0]
                        else:
                            reasons.append(f'{r[0]} / {r[4]}: không xác định duy nhất dòng bán để bù trừ');break
                        v=chosen['values'];sq,sg,sv=triple(v,[17,20,32])
                        if qty>sq or goods>sg or vat>sv:
                            reasons.append(f'{r[0]} / {r[4]}: SL hoặc tiền trả lại vượt dòng bán (bán {sg}đ, trả {goods}đ)');break
                        if abs(dec(v[19])-dec(r[8]))>1 or abs(rnd(qty*dec(v[19]))-goods)>1:
                            reasons.append(f'{r[0]} / {r[4]}: giá/tiền trả lại khác giá dòng bán (bán {v[19]}đ, trả {r[8]}đ); chưa tự điều chỉnh đơn giá');break
                        nq,ng,nv=sq-qty,sg-goods,sv-vat
                        if nq==0 and (ng!=0 or nv!=0):
                            reasons.append(f'{r[0]} / {r[4]}: hết SL nhưng còn dư tiền hàng/VAT');break
                        if nq>0 and (ng<=0 or nv<0 or abs(rnd(nq*dec(v[19]))-ng)>1):
                            reasons.append(f'{r[0]} / {r[4]}: phần còn lại không khớp SL nhân đơn giá');break
                        v[17]=native(nq);v[20]=native(ng);v[32]=native(nv);v[34]=native(ng+nv)
                    if not reasons:
                        net=[l for l in net if dec(l['values'][17])>0];applied=True
                    else:
                        link_issues[code]='; '.join(reasons);net=deepcopy(original)
                return_status[code]='Trả hết' if full else 'Đã bù trừ phần trả' if applied else 'Cần kiểm tra bù trừ'
            if code in excluded:
                category='Đã loại theo danh sách';reasons=['Danh sách loại trừ: '+', '.join(excluded[code])]
            elif full:
                category='Đã trả hết';reasons=['Khớp hết SL, tiền hàng và VAT theo từng mã hàng/ĐVT với '+', '.join(btls)]
            else:
                # Các kiểm tra nguồn độc lập với liên kết trả lại; loại bỏ toàn bộ nhãn dò cũ.
                reasons += [r for r in d['reasons'] if not r.startswith('Có thể liên quan trả lại:') and not (code=='BHHT2024004546' and r=='Có ký hiệu/số hóa đơn tại AMIS, cần xác nhận trước khi lập mới') and not (code in CONFIRMED_ROUNDING and r=='Thành tiền lệch số lượng nhân đơn giá quá 1 đồng') and not (code in CONFIRMED_PROMOTIONS and r in PROMOTION_REASONS)]
                if code in CONFIRMED_PROMOTIONS:
                    assert len(net)==1 and net[0]['values'][12]=='KM300FULL'
                    assert dec(net[0]['values'][17])==1 and all(dec(net[0]['values'][c])==0 for c in [19,20,25,27,28,30,32,34])
                    net[0]['rate']=8
                if applied:
                    for l in net:
                        r=l['values'];rates=[k for k in (5,8,10) if abs(rnd(dec(r[20])*k/100)-dec(r[32]))<=1]
                        if len(rates)!=1:reasons.append('Phần còn lại không xác định duy nhất thuế suất')
                        else:l['rate']=rates[0]
                    if len({l['rate'] for l in net})!=1:reasons.append('Phần còn lại không có một thuế suất duy nhất')
                reasons=list(dict.fromkeys(reasons))
                category='Cần kiểm tra' if reasons else 'Đưa vào import'
        reasons=detailed_reasons(code,original,reasons,by_bhh.get(code,[]))
        if code in link_issues and not SPLIT_RETURNS:link_issues[code]='; '.join(reasons)
        sums={k:sum((dec(l['values'][c]) for l in net),Decimal(0)) for k,c in [('goods',20),('vat',32),('payment',34)]}
        record={**{k:d[k] for k in ['doc','date','accounting_date','invoice_date_source','customer_code','buyer','address','tax','rows','rate']},'category':category,'reasons':reasons,'source_sums':d['sums'],'sums':sums,'lines':net,'original_lines':original,'btls':btls,'return_applied':applied,'full_return':full}
        record['promotion']=code in CONFIRMED_PROMOTIONS
        if record['promotion']:record['rate']=8
        if category=='Đưa vào import':assert sums['goods']+sums['vat']==sums['payment']
        records.append(record)
        if btls:
            for l in original:
                v=l['values'];nl=next((x for x in net if x['row']==l['row']),None);nv=nl['values'] if nl else [0]*42
                net_audit.append([code,', '.join(btls),l['row'],v[12],v[7],v[16],v[17],v[20],v[32],dec(v[17])-dec(nv[17]),dec(v[20])-dec(nv[20]),dec(v[32])-dec(nv[32]),nv[17],nv[20],nv[32],category,return_status.get(code),'; '.join(reasons)])
    summary={}
    for cat in ['Đưa vào import','Cần kiểm tra','Đã loại theo danh sách','Đã trả hết','Phiếu trả lại đã liên kết']:
        ds=[d for d in records if d['category']==cat]
        summary[cat]={'documents':len(ds),'source_lines':sum(len(d['original_lines']) for d in ds),'import_lines':sum(len(d['lines']) for d in ds) if cat=='Đưa vào import' else 0,**{k:sum((dec(d['source_sums'][k]) for d in ds),Decimal(0)) for k in ['goods','returns','vat','payment']},'net_goods':sum((d['sums']['goods'] for d in ds),Decimal(0)),'net_vat':sum((d['sums']['vat'] for d in ds),Decimal(0)),'net_payment':sum((d['sums']['payment'] for d in ds),Decimal(0))}
    assert sum(s['documents'] for s in summary.values())==35017
    assert sum(s['source_lines'] for s in summary.values())==44578
    for k,c in [('goods',20),('returns',27),('vat',32),('payment',34)]:assert sum(s[k] for s in summary.values())==dec(old['source_total'][c])
    pending=[d for d in records if d['category']=='Cần kiểm tra']
    unresolved=[]
    for e in old['exclusions']:
        if not e['codes']:unresolved.append([e['key'],f"HD cty 2024 dòng {e['row']}",None,e['values'][4],'Thiếu BHH: chưa tự loại theo tên/MST',None])
        for code in e['codes']:
            if code in old['exclusion_codes_not_found']:
                unresolved.append([e['key'],f"HD cty 2024 dòng {e['row']}",code,e['values'][4],'Mã cũ đã gộp vào BHHN2024000078; đã loại mã mới' if code=='BHHN2024000079' else 'Mã năm 2023 không có trong sổ 2024',None])
    seen=set()
    for row,key,buyer,tax,order in old['extra_invoices']:
        if key not in seen:seen.add(key);unresolved.append([key,f'HĐ ghi nhận UVG 2024 dòng {row}',None,buyer,'Chỉ có ở bảng ghi nhận, chưa có BHH xác nhận; mã đơn: '+str(order or ''),None])
    template_rows=[list(r) for r in load_workbook(BytesIO(template.read_bytes()),read_only=True,data_only=True).active.iter_rows(min_row=1,max_row=9,values_only=True)]
    confirmations=[{'doc':'BHHT2024004546','decision':'Đưa vào import','basis':'Người dùng xác nhận kế toán nhập nhầm ô Số hóa đơn; thực tế chưa có số hóa đơn. Giữ nguyên ô nguồn, chỉ bỏ điều kiện giữ kiểm tra cho chứng từ này.'}]
    for code in sorted(CONFIRMED_ROUNDING):
        d=next(x for x in records if x['doc']==code)
        assert d['category']=='Đưa vào import' and d['sums']['goods']+d['sums']['vat']==d['sums']['payment']
        differences=detailed_reasons(code,d['original_lines'],['Thành tiền lệch số lượng nhân đơn giá quá 1 đồng'],[])
        confirmations.append({'doc':code,'decision':'Đưa vào import, giữ tiền AMIS','basis':'Ngày 27/10/2026 người dùng xác nhận giữ nguyên tiền hàng, VAT và tổng thanh toán theo AMIS cho riêng chứng từ này; chấp nhận lệch SL × đơn giá 2–3đ. '+'; '.join(differences)})
    for code in sorted(CONFIRMED_PROMOTIONS):
        d=next(x for x in records if x['doc']==code)
        assert d['category']=='Đưa vào import' and all(d['sums'][k]==0 for k in ['goods','vat','payment'])
        confirmations.append({'doc':code,'decision':'Đưa vào import - hàng khuyến mại','basis':'Ngày 27/10/2026 người dùng xác nhận đây là đơn khuyến mại KM300FULL, SL 1; đồng ý dùng thuế suất 8%, giữ đơn giá/tiền hàng/VAT/thanh toán 0đ. Cột Hàng khuyến mại = 1; chỉ áp dụng đúng chứng từ này.'})
    if confirmed_link_file.exists():
        for c in json.loads(confirmed_link_file.read_text(encoding='utf8')):
            confirmations.append({'doc':c['btl']+' → '+c['bhh'],'decision':'Liên kết trả lại do người dùng xác nhận','basis':c['basis']+' Ngày '+c['confirmed_on']+'. SL/giá/tiền hàng/VAT/thanh toán khớp sổ. Ghi rõ nguồn người dùng, giữ snapshot AMIS ban đầu. Theo quyết định mới: giữ bán gốc, trả lại tách chờ điều chỉnh.'})
    confirmations.append({'doc':'Toàn bộ BTL 2024','decision':'Tách bán gốc và trả lại','basis':'Ngày 27/10/2026 người dùng xác nhận theo kế toán: lập hóa đơn bán gốc trước, sau đó lấy ký hiệu/số/ngày hóa đơn thực tế để lập import điều chỉnh trả lại. Không bù trừ hoặc loại bán gốc vì trả hết; vẫn giữ danh sách loại trừ. Hai khoản lệch tiền không tự chuyển thành giảm giá.'})
    return_drafts=[]
    for r in links:
        return_drafts.append([r[0],r[2],docs[r[0]]['date'],r[4],r[5],r[6],r[7],r[8],r[9],r[10],r[11],r[12],None,None,None,None,'Chờ hóa đơn bán gốc; '+('cần chốt chênh tiền' if r[2] in link_issues else 'chưa lập import điều chỉnh'),r[16],link_issues.get(r[2])])
    payload={'return_mode':'separate_adjustment_after_original_issuance','return_drafts':return_drafts,'manifest':manifest,'summary':summary,'docs':records,'links':links,'net_audit':net_audit,'unresolved_exclusions':unresolved,'template_rows':template_rows,'source_headers':old['headers'],'link_issues':link_issues,'return_status':return_status,'confirmations':confirmations,'promotion_import_rule':{'codes':sorted(CONFIRMED_PROMOTIONS),'rate':8,'marker_column':'Hàng khuyến mại','marker':1,'source':'https://helpv4.meinvoice.vn/kb/nhap-khau-hoa-don-mot-thue-suat/'}}
    (WORK/'prepared.json').write_text(json.dumps(payload,ensure_ascii=False,default=native),encoding='utf8')
    brief={'summary':summary,'partial_return_states':{k:v for k,v in return_status.items() if v!='Trả hết'},'pending':[{'doc':d['doc'],'reason':d['reasons']} for d in pending],'unresolved_exclusions':len(unresolved)}
    (WORK/'summary.json').write_text(json.dumps(brief,ensure_ascii=False,indent=2,default=native),encoding='utf8')
    print(json.dumps(brief,ensure_ascii=False,indent=2,default=native))

if __name__=='__main__':main()
