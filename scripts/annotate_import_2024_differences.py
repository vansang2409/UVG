"""Bổ sung số chênh ngay trong workbook kiểm tra hiện hành, bảo toàn nội dung còn lại."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from collections import defaultdict
import hashlib,json,os
from lxml import etree as E
from openpyxl import load_workbook
from rebuild_import_2024_linked_returns import detailed_reasons

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/amis-2024-linked-20261026'
FILE=ROOT/'outputs/amis-2024-theo-btl-bhh-20261026/KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';Q=lambda n:'{'+NS+'}'+n
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'

def main():
    with FILE.open('r+b'):pass
    p=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
    links=defaultdict(list)
    for r in p['links']:
        if r[2]:links[r[2]].append(r)
    updates={}
    for d in p['docs']:
        old='; '.join(d['reasons']);reasons=detailed_reasons(d['doc'],d['original_lines'],d['reasons'],links[d['doc']]);new='; '.join(reasons)
        if old!=new:updates[d['doc']]=(old,new);d['reasons']=reasons
    assert len(updates)==9,len(updates)
    for code,(_,new) in updates.items():
        if code in p['link_issues']:p['link_issues'][code]=new
    for r in p['net_audit']:
        if r[0] in updates:r[17]=updates[r[0]][1]
    raw=FILE.read_bytes();changed={};edits=[]
    with ZipFile(BytesIO(raw)) as z:
        wb=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        targets={r.get('Id'):r.get('Target') for r in rels}
        strings=[''.join(s.itertext()) for s in E.fromstring(z.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in z.namelist() else []
        def value(c):
            if c is None:return ''
            if c.get('t')=='inlineStr':return ''.join(c.itertext())
            v=c.find(Q('v'))
            if v is None:return ''
            return strings[int(v.text)] if c.get('t')=='s' else (v.text or '')
        for sh in wb.find(Q('sheets')):
            name=sh.get('name')
            if name not in ['Can kiem tra','Bu tru tra lai','Lien ket BTL-BHH']:continue
            target=targets[sh.get('{'+REL+'}id')];part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
            xml=E.fromstring(z.read(part));touched=False
            key_col='C' if name=='Lien ket BTL-BHH' else 'A';note_col={'Can kiem tra':'K','Bu tru tra lai':'R','Lien ket BTL-BHH':'S'}[name]
            for row in xml.find(Q('sheetData')):
                if row.get('r')=='1':continue
                cells={''.join(filter(str.isalpha,c.get('r'))):c for c in row};code=value(cells.get(key_col))
                if code not in updates:continue
                c=cells.get(note_col);old,new=updates[code];previous=value(c)
                if old not in previous:raise ValueError(f'{name} {code}: lý do đã khác bản chuẩn; giữ nguyên ghi chú người dùng')
                if c is None:raise ValueError('Thiếu ô lý do')
                for child in list(c):c.remove(child)
                c.set('t','inlineStr');E.SubElement(E.SubElement(c,Q('is')),Q('t')).text=previous.replace(old,new)
                row.set('ht','72');row.set('customHeight','1');touched=True
                edits.append({'sheet':name,'cell':c.get('r'),'doc':code,'text':value(c)})
            if touched:changed[part]=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
        result=BytesIO()
        with ZipFile(result,'w',ZIP_DEFLATED) as dst:
            for entry in z.infolist():dst.writestr(entry,changed.get(entry.filename,z.read(entry.filename)))
    result_raw=result.getvalue()
    # So sánh từng ô thuộc ba sheet được sửa; chỉ cho phép đổi các ô lý do đã liệt kê.
    before=load_workbook(BytesIO(raw),read_only=True,data_only=False)
    after=load_workbook(BytesIO(result_raw),read_only=True,data_only=False)
    allowed={(r['sheet'],r['cell']) for r in edits}
    for name in ['Can kiem tra','Bu tru tra lai','Lien ket BTL-BHH']:
        for ra,rb in zip(before[name].iter_rows(),after[name].iter_rows()):
            for a,b in zip(ra,rb):
                if (name,a.coordinate) not in allowed:assert a.value==b.value,(name,a.coordinate)
    assert len([e for e in edits if e['sheet']=='Can kiem tra'])==9
    with ZipFile(BytesIO(raw)) as a,ZipFile(BytesIO(result_raw)) as b:
        assert set(a.namelist())==set(b.namelist())
        for name in a.namelist():
            if name not in changed:assert a.read(name)==b.read(name),name
    tmp=WORK/'_review-differences.tmp.xlsx';tmp.write_bytes(result_raw);os.replace(tmp,FILE)
    (WORK/'prepared.json').write_text(json.dumps(p,ensure_ascii=False),encoding='utf8')
    summary=json.loads((WORK/'summary.json').read_text(encoding='utf8'))
    for r in summary['pending']:
        if r['doc'] in updates:r['reason']=next(d['reasons'] for d in p['docs'] if d['doc']==r['doc'])
    (WORK/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
    report={'passed':True,'changed_documents':len(updates),'edits':edits,'unchanged':'Mọi ô khác, sheet đối chiếu import, số tiền, phân loại và dữ liệu nguồn giữ nguyên','sha256':hashlib.sha256(result_raw).hexdigest()}
    (WORK/'differences-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    verification=json.loads((WORK/'verification.json').read_text(encoding='utf8'))
    for f in verification['files']:
        path=Path(f['path']);f['sha256']=hashlib.sha256(path.read_bytes()).hexdigest();f['bytes']=path.stat().st_size
    verification['reason_details']='Đã thêm phép tính và số chênh cụ thể cho 7 chứng từ SL × đơn giá và 2 chứng từ trả lại; kiểm mọi ô khác giữ nguyên.'
    (WORK/'verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
    for code,(_,reason) in updates.items():print(code,reason)

if __name__=='__main__':main()
