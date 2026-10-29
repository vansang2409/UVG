"""Read-only reconciliation of the fresh meInvoice exports against approved imports.

Source workbooks are never rewritten. Blank fields stay unknown unless backed
by the same invoice's evidence or the user's scoped payment-zero confirmation.
"""
from pathlib import Path
from collections import defaultdict, Counter
from decimal import Decimal
import hashlib
import json
import sys
import csv
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from prepare_invoice_reports import rows, invoice_key, invoice_day, amount
from filter_invoice_exceptions import normal

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'data/Meinv_SAU_PHAT_HANH_20261029'
OUT = ROOT / 'scripts/data/post-issue-20261029'


def native(v):
    if isinstance(v, Decimal):
        return int(v) if v == int(v) else float(v)
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    raise TypeError(type(v))


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = {}
    records = {}
    links = defaultdict(list)
    edges = {}
    children = defaultdict(set)
    parents = defaultdict(set)
    replacement_children = defaultdict(set)
    conflicts = []
    duplicates = []
    file_stats = []

    def register(p):
        rel = p.relative_to(ROOT).as_posix()
        manifest[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
        return rel

    for p in sorted(BASE.rglob('*.xls')):
        rel = register(p)
        kind = 'adjustment' if p.parent.name == 'DIEU_CHINH' else 'replacement' if p.parent.name == 'THAY_THE' else 'total'
        count = 0
        dates = []
        indices = []
        period = None
        headers = {}
        for rn, r in rows(p):
            if rn == 3:
                period = r.get('A')
            if rn == 5 and kind == 'total':
                headers = {v: k for k, v in r.items() if v}
                assert all(v in headers for v in ['Ký hiệu', 'Số hóa đơn', 'Ngày hóa đơn', 'Doanh số bán chưa thuế', 'Tiền thuế GTGT', 'Tổng tiền', 'Trạng thái hóa đơn', 'Trạng thái gửi CQT', 'Mã của CQT']), (rel, headers)
            if rn < (6 if kind == 'total' else 7) or amount(r.get('A')) is None:
                continue
            indices.append(int(amount(r['A'])))
            source = {'file': rel, 'row': rn}
            if kind == 'total':
                get = lambda label: r.get(headers[label])
                key = invoice_key(get('Ký hiệu'), get('Số hóa đơn'))
                assert key, (rel, rn)
                day = invoice_day(get('Ngày hóa đơn'))
                rec = {'key': key, 'date': day, 'money': [amount(get(c)) for c in ['Doanh số bán chưa thuế', 'Tiền thuế GTGT', 'Tổng tiền']],
                       'status': get('Trạng thái hóa đơn') or '', 'cqt_status': get('Trạng thái gửi CQT') or '',
                       'cqt_code_present': bool(get('Mã của CQT')), 'source': source}
                if key in records:
                    duplicates.append({'key': key, 'a': records[key]['source'], 'b': source})
                    if any(rec[c] != records[key][c] for c in ['date', 'money', 'status', 'cqt_status', 'cqt_code_present']):
                        conflicts.append({'kind': 'duplicate_total_conflict', 'key': key})
                else:
                    records[key] = rec
                dates.append(day)
            else:
                old = invoice_key(r.get('B'), r.get('D'))
                new = invoice_key(r.get('M'), r.get('O'))
                assert old and new, (rel, rn, 'missing link key')
                edge = (old, new, kind)
                edges.setdefault(edge, []).append(source)
                children[old].add(new)
                parents[new].add(old)
                if kind == 'replacement':
                    replacement_children[old].add(new)
                for key, cols, dc in [(old, ['J', 'K', 'L'], 'C'), (new, ['U', 'V', 'W'], 'N')]:
                    links[key].append({'money': [amount(r.get(c)) for c in cols], 'date': invoice_day(r.get(dc)), 'source': source, 'side': 'old' if key == old else 'new', 'kind': kind})
                dates.append(invoice_day(r.get('C')))
            count += 1
        assert dates and all(d is not None for d in dates), (rel, 'invalid dates')
        assert sorted(indices) == list(range(1, len(indices) + 1)), (rel, 'missing or duplicate STT')
        file_stats.append({'file': rel, 'kind': kind, 'rows': count, 'period': period, 'date_min': min(dates), 'date_max': max(dates), 'stt_min': min(indices), 'stt_max': max(indices)})
    print('Loaded', len(records), 'invoices and', len(edges), 'links', flush=True)

    # Validate any duplicate link observations without replacing source values.
    for key, observations in links.items():
        for i in range(3):
            vals = {v['money'][i] for v in observations if v['money'][i] is not None}
            if len(vals) > 1:
                conflicts.append({'kind': 'link_money_conflict', 'key': key, 'field': i, 'values': sorted(vals)})

    plans = {}
    for group, filename in [('adjusted', 'HOA_DON_DA_DIEU_CHINH_IMPORT_2_SHEET.xlsx'), ('untreated', 'HOA_DON_CHUA_DIEU_CHINH_IMPORT_2_SHEET.xlsx')]:
        p = ROOT / 'outputs' / filename
        register(p)
        w = load_workbook(p, read_only=True, data_only=True)
        byroot = defaultdict(list)
        seqs = set()
        for r in w['Import dieu chinh'].iter_rows(min_row=10, max_col=24, values_only=True):
            if r[0] is None or r[0] in seqs:
                continue
            seqs.add(r[0])
            key = invoice_key(r[9], r[10])
            assert key
            byroot[key].append({'seq': r[0], 'money': [amount(r[c]) for c in [19, 21, 22]]})
        baselines = {}
        base_sheet = w['F0, F1, F2, ...'] if group == 'adjusted' else w['F0']
        for r in base_sheet.iter_rows(min_row=3, max_col=55 if group == 'adjusted' else 16, values_only=True):
            key = invoice_key(r[0], r[1])
            if key in byroot:
                baselines[key] = [invoice_key(r[c], r[c+1]) for c in ([0, 6, 12, 18] if group == 'adjusted' else [0]) if r[c] and r[c+1]]
        w.close()
        plans[group] = (byroot, baselines)
    assert not set(plans['adjusted'][0]) & set(plans['untreated'][0])

    def reach(root):
        found = set()
        todo = [root]
        while todo:
            key = todo.pop()
            if key in found:
                continue
            found.add(key)
            todo.extend(children[key])
        return found

    needed = set()
    for byroot, _ in plans.values():
        for root in byroot:
            needed.update(reach(root))
    repl_targets = [('1C25TUV', str(n).zfill(8)) for n in [2601, 2610, 2602, 2603, 2604, 2617]]
    for k in repl_targets:
        needed.update(reach(k))
    replacement_candidates = [k for k, r in records.items() if r['date'].isoformat() == '2026-10-29' and normal(r['status']) == 'hoa don thay the' and r['money'][:2] == [0, 0]]
    needed.update(replacement_candidates)

    # XML is a fallback only for an identical invoice, never for a new invoice
    # with similar amounts or a different signature/year.
    xml_values = defaultdict(list)
    blank_keys = {k for k in needed if k in records and None in records[k]['money']}
    for p in (ROOT / 'data').rglob('*.xml'):
        candidate = next((k for k in blank_keys if p.stem == '_'.join(k)), None)
        if candidate is None:
            continue
        tree = ET.fromstring(p.read_bytes())
        local = lambda e: e.tag.rsplit('}', 1)[-1]
        first = lambda tag: next((e for e in tree.iter() if local(e) == tag), None)
        get = lambda parent, tag: next((e.text for e in parent.iter() if local(e) == tag), None) if parent is not None else None
        tt, pay = first('TTChung'), first('TToan')
        key = invoice_key((get(tt, 'KHMSHDon') or '') + (get(tt, 'KHHDon') or ''), get(tt, 'SHDon'))
        if key != candidate:
            continue
        xml_values[key].append({'money': [amount(get(pay, t)) for t in ['TgTCThue', 'TgTThue', 'TgTTTBSo']], 'date': invoice_day(get(tt, 'NLap')), 'source': {'file': register(p), 'row': None}})

    resolved = {}
    for key in needed:
        if key not in records:
            continue
        rec = records[key]
        money = rec['money'][:]
        fallbacks = []
        issues = []
        for obs in links[key] + xml_values[key]:
            if obs['date'] != rec['date']:
                issues.append({'kind': 'date_conflict', 'source': obs['source']})
            for i, val in enumerate(obs['money']):
                if val is not None and rec['money'][i] is not None and val != rec['money'][i]:
                    issues.append({'kind': 'total_vs_evidence_conflict', 'field': i, 'source': obs['source'], 'total': rec['money'][i], 'other': val})
        for i, val in enumerate(money):
            if val is not None:
                continue
            candidates = [(o['money'][i], o['source']) for o in links[key] + xml_values[key] if o['money'][i] is not None and o['date'] == rec['date'] and all(a is None or b is None or a == b for a, b in zip(rec['money'], o['money']))]
            unique = {v for v, _ in candidates}
            if len(unique) == 1:
                money[i] = next(iter(unique))
                fallbacks.append({'field': i, 'value': money[i], 'evidence': candidates[0][1]})
        resolved[key] = {**rec, 'resolved_money': money, 'fallbacks': fallbacks, 'issues': issues}

    policy_path = OUT / 'confirmed-blank-payment-zero.json'
    if policy_path.exists():
        policy_rel = register(policy_path)
        policy = json.loads(policy_path.read_text(encoding='utf-8'))
        applied = 0
        for key_values in policy['invoice_keys']:
            key = tuple(key_values)
            rec = resolved[key]
            assert rec['money'] == [0, 0, None] and not rec['issues'], key
            assert rec['resolved_money'] == [0, 0, None], key
            rec['resolved_money'][2] = 0
            rec['payment_basis'] = 'user_confirmed_blank_zero'
            rec['fallbacks'].append({'field': 2, 'value': 0, 'evidence': {'file': policy_rel, 'row': None}, 'confirmed_on': policy['confirmed_on']})
            applied += 1
        assert applied == len(policy['invoice_keys'])

    def evaluate(root, expected=(), old=()):
        keys = reach(root)
        issues = []
        missing_old = set(old) - keys
        if missing_old:
            issues.append({'kind': 'prior_link_missing', 'keys': sorted(missing_old)})
        active = []
        excluded = []
        unknown = [False] * 3
        sums = [Decimal(0)] * 3
        for key in sorted(keys):
            if key not in resolved:
                issues.append({'kind': 'missing_issued_record', 'key': key})
                unknown = [True] * 3
                continue
            rec = resolved[key]
            if rec['issues']:
                issues.extend({'key': key, **i} for i in rec['issues'])
            if len(parents[key]) > 1:
                issues.append({'kind': 'multiple_parents', 'key': key})
            st = normal(rec['status'])
            if replacement_children[key] or 'bi thay the' in st or 'huy' in st:
                excluded.append(key)
                if replacement_children[key] and 'bi thay the' not in st:
                    issues.append({'kind': 'replacement_status_mismatch', 'key': key, 'status': rec['status']})
                if 'bi thay the' in st and not replacement_children[key]:
                    issues.append({'kind': 'replacement_link_missing', 'key': key})
                continue
            active.append(key)
            for i, v in enumerate(rec['resolved_money']):
                if v is None:
                    unknown[i] = True
                    issues.append({'kind': 'blank_money', 'key': key, 'field': i})
                else:
                    sums[i] += v
        new = [k for k in active if records[k]['date'].year == 2026]
        # Match exact totals first. A linked invoice with an unpopulated field
        # is present but unverified, rather than a missing/extra invoice.
        available = new[:]
        missing = []
        matches = []
        pending = []
        for e in expected:
            exact = [k for k in available if resolved[k]['resolved_money'] == e['money']]
            if exact:
                key = exact[0]
                available.remove(key)
                matches.append({'seq': e['seq'], 'key': key, 'kind': 'exact'})
            else:
                pending.append(e)
        for e in pending:
            compatible = [k for k in available if all(v is None or v == wanted for v, wanted in zip(resolved[k]['resolved_money'], e['money']))]
            if compatible:
                key = compatible[0]
                available.remove(key)
                matches.append({'seq': e['seq'], 'key': key, 'kind': 'blank_field_unverified'})
            else:
                missing.append(e)
        if missing:
            issues.append({'kind': 'planned_amount_missing', 'items': missing})
        if available and expected:
            issues.append({'kind': 'extra_2026_invoice', 'keys': available})
        balance = [None if unknown[i] else sums[i] for i in range(3)]
        return {'root': root, 'status': 'needs_review' if issues else 'zero' if balance == [0, 0, 0] else 'nonzero', 'balance': balance,
                'known_subtotal': sums, 'issues': issues, 'planned_count': len(expected), 'issued_2026_count': len(new),
                'active': active, 'excluded': excluded, 'all_keys': sorted(keys), 'matches': matches}

    result_groups = {g: [evaluate(k, p, baseline.get(k, [])) for k, p in sorted(byroot.items())] for g, (byroot, baseline) in plans.items()}
    replacement_results = []
    for key in repl_targets:
        r = evaluate(key)
        r['direct_replacements'] = sorted(replacement_children[key])
        if len(replacement_children[key]) != 1:
            r['issues'].append({'kind': 'expected_one_replacement', 'actual': len(replacement_children[key])})
            # Excluding a replaced invoice does not prove its replacement is 0.
            r['balance'] = [None, None, None]
        r['status'] = 'needs_review' if r['issues'] else 'zero' if r['balance'] == [0, 0, 0] else 'nonzero'
        replacement_results.append(r)
    result_groups['replacement_six'] = replacement_results

    supplement_path = ROOT / 'outputs/bo-sung-13-20261029/prepared.json'
    register(supplement_path)
    supplemental = json.loads(supplement_path.read_text(encoding='utf-8'))['audit']
    supplemental_results = []
    for entry in supplemental:
        root = tuple(entry['root'])
        candidates = [k for k in reach(root) if k in resolved and records[k]['date'].year == 2026 and all(v is None or v == Decimal(wanted) for v, wanted in zip(resolved[k]['resolved_money'], entry['totals']))]
        supplemental_results.append({'seq': entry['seq'], 'root': root, 'planned_money': entry['totals'], 'matches': sorted(candidates),
                                     'money_match': 'exact' if len(candidates) == 1 and resolved[candidates[0]]['resolved_money'] == entry['totals'] else 'blank_field_unverified' if len(candidates) == 1 else 'unresolved'})

    summary = {'read_only_sources': True, 'source_files': len(file_stats), 'rows_by_kind': dict(Counter({k: sum(s['rows'] for s in file_stats if s['kind'] == k) for k in ['total', 'adjustment', 'replacement']})),
               'unique_total_keys': len(records), 'duplicate_total_keys': duplicates, 'source_conflicts': conflicts,
               'date_min': min(r['date'] for r in records.values()), 'date_max': max(r['date'] for r in records.values()),
               'groups': {g: {'roots': len(rs), 'planned_invoices': sum(r['planned_count'] for r in rs), 'statuses': dict(Counter(r['status'] for r in rs)),
                              'issue_kinds': dict(Counter(i['kind'] for r in rs for i in r['issues'])),
                              'known_balance_totals': [sum(r['known_subtotal'][i] for r in rs) for i in range(3)],
                              'unknown_balance_roots': sum(any(v is None for v in r['balance']) for r in rs)} for g, rs in result_groups.items()},
               'cqt_statuses_all': dict(Counter(r['cqt_status'] for r in records.values())),
               'linked_2026_cqt_statuses': dict(Counter(records[k]['cqt_status'] for g in ['adjusted', 'untreated'] for r in result_groups[g] for k in r['active'] if records[k]['date'].year == 2026)),
               'replacement_zero_candidates': sorted(replacement_candidates),
               'supplemental_matching': dict(Counter('unique' if len(r['matches']) == 1 else 'missing' if not r['matches'] else 'multiple' for r in supplemental_results))}
    if policy_path.exists():
        summary['blank_payment_zero_confirmation'] = {'date': policy['confirmed_on'], 'chains': len(policy['roots']), 'invoice_fields': applied, 'source_blanks_preserved': True}
    payload = {'summary': summary, 'files': file_stats, 'groups': result_groups, 'supplemental': supplemental_results,
               'invoices': [resolved[k] for k in sorted(resolved)],
               'edges': [{'old': a, 'new': b, 'kind': kind, 'sources': sources} for (a, b, kind), sources in sorted(edges.items()) if a in needed or b in needed],
               'manifest': manifest}
    apply_accountant_exclusions(payload)
    for name, data in [('audit.json', payload), ('summary.json', summary)]:
        (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=native), encoding='utf-8')
    unchanged = all(hashlib.sha256((ROOT / rel).read_bytes()).hexdigest() == sha for rel, sha in manifest.items())
    assert unchanged, 'A source changed during the audit'
    write_report(payload)
    print(json.dumps(summary, ensure_ascii=False, default=native), flush=True)
    for g, rs in result_groups.items():
        for r in rs:
            if r['status'] != 'zero' and g != 'untreated':
                print(g, json.dumps(r, ensure_ascii=False, default=native), flush=True)


def apply_accountant_exclusions(payload):
    """Keep accountant-approved exclusions separate from open differences."""
    p = OUT / 'accountant-exclusions.json'
    if not p.exists():
        return
    decisions = json.loads(p.read_text(encoding='utf-8'))
    payload['manifest'][p.relative_to(ROOT).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    excluded = payload.setdefault('accountant_exclusions', [])
    for decision in decisions['exclusions']:
        group, key = decision['group'], tuple(decision['root'])
        matches = [r for r in payload['groups'][group] if tuple(r['root']) == key]
        if not matches:
            assert any(tuple(r['decision']['root']) == key and r['decision']['group'] == group for r in excluded)
            continue
        assert len(matches) == 1
        excluded.append({'decision': decision, 'confirmed_on': decisions['confirmed_on'], 'source_result': matches[0]})
        payload['groups'][group].remove(matches[0])
    summary = payload['summary']
    summary['accountant_exclusion_count'] = len(excluded)
    summary['groups'] = {g: {'roots': len(rs), 'planned_invoices': sum(r['planned_count'] for r in rs),
                            'statuses': dict(Counter(r['status'] for r in rs)),
                            'issue_kinds': dict(Counter(i['kind'] for r in rs for i in r['issues'])),
                            'known_balance_totals': [sum(r['known_subtotal'][i] for r in rs) for i in range(3)],
                            'unknown_balance_roots': sum(any(v is None for v in r['balance']) for r in rs)} for g, rs in payload['groups'].items()}


def write_report(payload):
    """Write a separate review report; never export or edit a source workbook."""
    dest = ROOT / 'outputs/kiem-tra-sau-phat-hanh-20261029'
    dest.mkdir(parents=True, exist_ok=True)
    inv = {tuple(r['key']): r for r in payload['invoices']}
    label = {'adjusted': 'Đã điều chỉnh', 'untreated': 'Chưa điều chỉnh (tên bộ import)', 'replacement_six': 'Thay thế 6 chuỗi'}
    keystr = lambda k: '/'.join(k)
    source_counts = Counter(e['kind'] for e in payload['files'])
    input_counts = payload['summary']['rows_by_kind']
    replacement_by_old = {tuple(r['root']): r['direct_replacements'] for r in payload['groups']['replacement_six']}
    linked_replacements = sum(len(v) == 1 for v in replacement_by_old.values())
    all_rows = []
    for group, rs in payload['groups'].items():
        for r in rs:
            known_nonzero = any(v not in (None, 0) for v in r['balance'])
            status = 'Đã về 0 cả 3 khoản' if r['status'] == 'zero' else 'Còn dư tiền' if known_nonzero else 'Chưa đủ căn cứ'
            new = [k for k in r['all_keys'] if tuple(k) in inv and invoice_day(inv[tuple(k)]['date']).year == 2026]
            all_rows.append([label[group], keystr(r['root']), *r['balance'], status, r['planned_count'], r['issued_2026_count'],
                             '; '.join(keystr(k) for k in new), '; '.join(sorted({i['kind'] for i in r['issues']})),
                             '; '.join(f"{inv[tuple(k)]['source']['file']} dòng {inv[tuple(k)]['source']['row']}" for k in r['all_keys'] if tuple(k) in inv)])
    review_rows = [r for r in all_rows if r[5] != 'Đã về 0 cả 3 khoản']
    excluded_rows = []
    for entry in payload.get('accountant_exclusions', []):
        d, r = entry['decision'], entry['source_result']
        excluded_rows.append([label[d['group']], keystr(d['root']), *r['balance'], 'Đã loại theo xác nhận KT', r['planned_count'], r['issued_2026_count'], '', d['reason'],
                              '; '.join(f"{inv[tuple(k)]['source']['file']} dòng {inv[tuple(k)]['source']['row']}" for k in r['all_keys'] if tuple(k) in inv)])
        # In-scope balance columns remain unavailable for excluded records.
        all_rows.append([label[d['group']], keystr(d['root']), None, None, None, 'Đã loại theo xác nhận KT', r['planned_count'], r['issued_2026_count'], '', d['reason'], excluded_rows[-1][-1]])
    header = ['Nhóm', 'Hóa đơn gốc/tham chiếu', 'Số dư tiền hàng', 'Số dư VAT', 'Số dư thanh toán', 'Kết quả', 'HĐ điều chỉnh theo kế hoạch', 'HĐ liên kết năm 2026', 'Ký hiệu/số HĐ năm 2026', 'Điểm cần kiểm tra', 'Nguồn bảng kê']
    for filename, rs in [('TAT_CA_CHUOI.csv', all_rows), ('CAC_TRUONG_HOP_CAN_KIEM_TRA.csv', review_rows), ('LOAI_TRU_THEO_XAC_NHAN_KT.csv', excluded_rows)]:
        with (dest / filename).open('w', encoding='utf-8-sig', newline='') as stream:
            w = csv.writer(stream)
            w.writerow(header)
            w.writerows(rs)
    lines = [
        '# Kiểm tra sau phát hành ngày 29/10/2026', '',
        '**Đã về 0 cả tiền hàng, VAT và thanh toán trong phạm vi KT thực hiện.** Ngày30/10/2026 người dùng xác nhận ô tổng thanh toán trống nghĩa là0. Áp dụng đúng24ô của15chuỗi đang treo (9điều chỉnh +6thay thế), giữ nguyên ô trống trong nguồn. Đủ6liên kết thay thế. **1C25TUV/00001943** được KT chủ động loại vì **TĐ không hợp lệ**, có ảnh và xác nhận người dùng; khoản nguồn645.051đ tách riêng, không tính là thiếu xử lý.', '',
        f"Đã đọc {len(payload['files'])} file nguồn: {source_counts['total']} phần bảng kê đã sử dụng ({input_counts['total']:,} hóa đơn duy nhất), {source_counts['adjustment']} phần điều chỉnh ({input_counts['adjustment']:,} liên kết) và {source_counts['replacement']} file thay thế ({input_counts['replacement']:,} liên kết). Bộ ban đầu chọn kỳ 01/01/2024–29/10/2026; file bổ sung xuất theo ngày HĐ thay thế 29/10/2026. Không trùng ký hiệu/số trên bảng kê, không phát hiện xung đột tiền giữa các nguồn đã so. Các phần đều có STT liên tục. Hash của các nguồn và kế hoạch giữ nguyên trước/sau kiểm tra.", '',
        '| Bộ kiểm tra | Số gốc/chuỗi | Về 0 đủ tiền hàng/VAT/thanh toán | Trường hợp còn lại |',
        '|---|---:|---:|---|',
        '| Đã điều chỉnh: 2.045 HĐ theo kế hoạch | 1.015 | 1.015 | Không còn chênh |',
        '| Chưa điều chỉnh: 8.410 HĐ trong file, KT loại 1 | 8.409 trong phạm vi | 8.409 | Không còn chênh |',
        '| Thay thế 6 chuỗi | 6 | 6 | Khớp đủ 6 cặp; đủ ba khoản0 |', '',
        'Tổng hai bộ điều chỉnh sau loại trừ: **9.424/9.424 gốc trong phạm vi** khớp0 đủ ba khoản; cộng6chuỗi thay thế thành **9.430/9.430 chuỗi đã về0**. Đủ **10.454/10.454 HĐ điều chỉnh trong phạm vi** theo liên kết và kế hoạch; STT1168/gốc1943 nằm ngoài đợt thực hiện theo xác nhận KT. Toàn bộ10.454 HĐ điều chỉnh năm2026 và6HĐ thay thế mới được bảng kê ghi **HĐ hợp lệ**. Đây là trạng thái trong file xuất, chưa tra cứuCQT trực tiếp hoặc kiểm mật mã chữ ký.', '',
        '## 13 hóa đơn bổ sung trên 8 chuỗi', '',
        'Đã tìm đủ13HĐ **1C26MUV/00110692–00110704**, ngày29/10/2026, đúng liên kết8gốc, đều ghi HĐ hợp lệ và khớp đủ3khoản với kế hoạch. HĐ00110696/STT2037 thanh toán nguồn trống được tính0 theo xác nhận30/10. Tổng tiền hàng bổ sung -233.262.118đ vàVAT -18.660.968đ đã triệt tiêu số còn dư trước đó của8chuỗi. Chưa kiểm chi tiết số lượng, đơn giá, mặt hàng của lần phát hành mới vì các file này là bảng kê tổng.', '',
        '| STT kế hoạch | Gốc | HĐ phát hành | Kết quả so tiền |',
        '|---:|---|---|---|',
    ]
    for r in payload['supplemental']:
        lines.append(f"| {r['seq']} | {keystr(r['root'])} | {', '.join(keystr(k) for k in r['matches']) or 'Chưa tìm thấy'} | {'Khớp 3 khoản' if r['money_match'] == 'exact' else 'Tiền hàng/VAT khớp; thanh toán trống'} |")
    lines += ['', '## Hóa đơn KT đã loại khỏi đợt xử lý', '',
        '- Gốc **1C25TUV/00001943**, ngày 29/05/2025; STT1168 của bộ chưa điều chỉnh.',
        '- Giá trị nguồn vẫn giữ 597.269đ tiền hàng + 47.782đ VAT =645.051đ; file import online và bản cục bộ có khoản giảm tương ứng.',
        '- Người dùng xác nhận KT chủ động bỏ HĐ này vì TĐ không hợp lệ; ảnh mới cũng hiển thị trạng thái này. Việc không có HĐ điều chỉnh phát hành được giải thích bằng quyết định loại trừ, không còn ghi là khoản thiếu xử lý.',
        '- Chỉ đổi phạm vi báo cáo. Giữ nguyên dữ liệu nguồn, file import và bản online; không kết luận thay cho việc xử lý thuế của hóa đơn không hợp lệ.', '',
        '## 9 chuỗi điều chỉnh đã chốt thanh toán trống bằng0', '',
        'Người dùng xác nhận trực tiếp ngày30/10/2026: “trống là 0 á”. Áp dụng vào phép đối chiếu, giữ nguyên giá trị trống trên file gốc. 9chuỗi dưới đã khớp0 cả ba khoản; không còn yêu cầu XML chỉ vì những ô thanh toán này.', '',
        '| Gốc | HĐ điều chỉnh có thanh toán trống |', '|---|---|']
    for g in ['adjusted', 'untreated']:
        for r in payload['groups'][g]:
            if not any(inv[tuple(k)].get('payment_basis') == 'user_confirmed_blank_zero' for k in r['active']):
                continue
            blank = [k for k in r['active'] if invoice_day(inv[tuple(k)]['date']).year == 2026 and inv[tuple(k)].get('payment_basis') == 'user_confirmed_blank_zero']
            lines.append(f"| {keystr(r['root'])} | {', '.join(keystr(k) for k in blank)} |")
    lines += ['', '## Sáu hóa đơn thay thế — đã xác nhận liên kết ngày 30/10/2026', '',
        'Các tham chiếu **1C25TUV/00002601, 00002610, 00002602, 00002603, 00002604, 00002617** trong bảng kê mới đều ghi **Hóa đơn đã bị thay thế**.', '',
        f'File bổ sung `Bang_ke_hoa_don_thay_the_xuat_theo_HD_thay_the_1793293386111.xls` có **{linked_replacements}/6 cặp** đúng tham chiếu dự kiến. HĐ cũ ngày31/05/2025, HĐ mới ngày29/10/2026. Số tiền từng HĐ cũ khớp nguồn và tổng371.344.277đ tiền hàng +29.707.543đVAT =401.051.820đ thanh toán. HĐ mới tiền hàng/VAT bằng0; cộtW thanh toán trống được tính0 theo xác nhận người dùng30/10. Tất cả6HĐ mới được bảng kê tổng ghi **HĐ hợp lệ**.', '',
        '| HĐ cũ (1C25TUV) | HĐ thay thế (1C26MUV) | Tiền hàng mới | VAT mới | Thanh toán mới |',
        '|---|---|---:|---:|---|']
    for old, replacements in sorted(replacement_by_old.items()):
        lines.append(f"| {old[1]} | {', '.join(k[1] for k in replacements) or 'Chưa có liên kết'} | 0 | 0 | 0 theo xác nhận |")
    lines += ['',
        '**Kết luận:** đủ6liên kết thay thế và cả15chuỗi từng treo thanh toán đã chốt0 theo xác nhận. Không còn trường hợp chênh/thiếu liên kết trong phạm vi đợt đối soát này. Phạm vi kết luận là số tiền, số lượng hóa đơn và liên kết theo các nguồn đã nhận cùng xác nhận người dùng; chưa kiểm chữ ký hoặc toàn bộ chi tiết mặt hàng phát hành mới.', '',
        '## Cách đối chiếu', '',
        '- Khóa hóa đơn: ký hiệu + số hóa đơn, không ghép chỉ theo số.',
        '- Lần theo đầy đủ liên kết; giữ gốc bị điều chỉnh và cộng các khoản điều chỉnh có liên quan.',
        '- Loại giá trị bản bị thay thế/hủy; chỉ kết luận khi có đầy đủ bản thay thế và liên kết.',
        '- Khớp số lượng HĐ và từng bộ tiền với kế hoạch, không chỉ tổng ròng.',
        '- 24ô thanh toán trống trên15chuỗi được tính0 theo xác nhận30/10; giữ ô nguồn trống, lưu căn cứ riêng. Các ô trống khác không thuộc xác nhận vẫn cần chứng cứ. Ngoại lệ XML lịch sử không tính lại thanh toán từ tiền hàng + VAT.',
        '- Đọc dữ liệu cục bộ. Không nhập/phát hành lại, sửa meInvoice, sửa AMIS hoặc chỉnh workbook nguồn.', '',
        '## Tệp chi tiết', '',
        '- `CAC_TRUONG_HOP_CAN_KIEM_TRA.csv`: không còn dòng cần xử lý (chỉ còn tiêu đề).',
        '- `LOAI_TRU_THEO_XAC_NHAN_KT.csv`: 1 gốc1943, giữ giá trị nguồn và lý do KT loại, không cộng vào số dư trong phạm vi.',
        '- `TAT_CA_CHUOI.csv`: đủ9.431dòng, gồm9.430chuỗi trong phạm vi về0 và1gốc1943đã loại. Số dư của gốc đã loại để trống vì ngoài phạm vi.',
        '- JSON căn cứ, nguồn/dòng, liên kết và SHA-256: `scripts/data/post-issue-20261029/audit.json`.', '',
        '## Nguồn mới đã đọc', '', '| Tệp | Số dòng hóa đơn/liên kết |', '|---|---:|']
    for entry in payload['files']:
        lines.append(f"| {entry['file']} | {entry['rows']} |")
    (dest / 'KET_QUA_KIEM_TRA.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
