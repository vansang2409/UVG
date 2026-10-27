"""Kiểm tra bảng liên kết BTL–BHH đọc từ giao diện AMIS, không sửa nguồn."""
from pathlib import Path
from collections import Counter,defaultdict
from decimal import Decimal
from datetime import datetime
import json, hashlib,re
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/btl-bhh-live-20261026'
SOURCE=ROOT/'data/Du lieu Amis/So_chi_tiet_ban_hang_2024.xlsx'

def number(v):
    return Decimal(v.replace('.','').replace(',','.')) if v else None

def signature(sku,qty,price,goods,vat):
    return str(sku).strip(),Decimal(str(qty)),Decimal(str(price)),Decimal(str(goods)),Decimal(str(vat))

def native(v):
    if isinstance(v,Decimal):return int(v) if v==v.to_integral_value() else float(v)
    if isinstance(v,datetime):return v.isoformat()
    raise TypeError(type(v))

def main():
    records=json.loads((WORK/'verified-live-reads.json').read_text(encoding='utf8'))
    lists=json.loads((WORK/'list-pages.json').read_text(encoding='utf8'))
    override_file=WORK/'user-confirmed-links.json'
    confirmations=json.loads(override_file.read_text(encoding='utf8')) if override_file.exists() else []
    overrides={(r['btl'],r['line'],r['sku']):r for r in confirmations}
    expected={r[2]:r for r in lists['page1']+lists['page2']}
    assert len(records)==len(expected)==163
    assert len({r['selected'][2] for r in records})==163
    assert {r['selected'][2] for r in records}==set(expected)
    before=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    source=load_workbook(SOURCE,read_only=True,data_only=True).active
    source_returns=defaultdict(list);sales=defaultdict(list)
    for i,r in enumerate(source.iter_rows(min_row=5,values_only=True),5):
        if not r[2]:continue
        if str(r[2]).startswith('BTL'):
            source_returns[str(r[2])].append((i,r))
        else:sales[str(r[2])].append((i,r))
    details=[];docs=[];issues=[]
    for rec in sorted(records,key=lambda r:(datetime.strptime(r['selected'][1],'%d/%m/%Y'),r['selected'][2])):
        head=rec['selected'];doc=head[2]
        assert head==expected[doc]
        assert rec['headers'][12]=='Số CT bán hàng'
        count=int(re.search(r'Tổng số:\s*(\d+)',rec['detailFooter']).group(1))
        assert count==len(rec['rows']) and count<20
        total=sum((number(r[8])+number(r[10]) for r in rec['rows']),Decimal(0))
        assert total==number(head[5]),doc
        live_counter=Counter(signature(r[1],number(r[5]),number(r[7]),number(r[8]),number(r[10])) for r in rec['rows'])
        src_counter=Counter(signature(r[12],r[26],r[19],r[27],-Decimal(str(r[32]))) for i,r in source_returns[doc])
        source_ok=live_counter==src_counter
        if not source_ok:issues.append([doc,'Dòng chi tiết AMIS hiện tại khác sổ xuất',str(live_counter),str(src_counter)])
        first=len(details)+2
        for r in rec['rows']:
            sig=signature(r[1],number(r[5]),number(r[7]),number(r[8]),number(r[10]))
            source_matches=[i for i,s in source_returns[doc] if signature(s[12],s[26],s[19],s[27],-Decimal(str(s[32])))==sig]
            bhh=r[12].strip()
            state='Có BHH gốc' if bhh else 'Trống BHH trên AMIS'
            override=overrides.get((doc,int(r[0]),r[1]))
            if override:
                assert bhh==override['original_amis_bhh']
                bhh=override['bhh'];state='BHH do người dùng xác nhận '+override['confirmed_on']+'; AMIS lúc trích đang trống'
            if not bhh:issues.append([doc,state,r[0],r[1]])
            details.append([doc,datetime.strptime(head[1],'%d/%m/%Y'),bhh,int(r[0]),r[1],r[2],r[3],number(r[5]),number(r[7]),number(r[8]),number(r[9]),number(r[10]),None,head[4],state,'Khớp sổ AMIS 2024' if source_ok else 'Khác sổ AMIS 2024',', '.join(map(str,source_matches))])
        last=len(details)+1
        bhhs=sorted({r[2] for r in details[first-2:] if r[0]==doc and r[2]})
        docs.append({'btl':doc,'date':datetime.strptime(head[1],'%d/%m/%Y'),'customer':head[4],'lines':len(rec['rows']),
                     'bhhs':bhhs,'first':first,'last':last,'payment':number(head[5]),'source_ok':source_ok,'captured':rec['capturedAt']})
    assert len(details)==sum(len(v) for v in source_returns.values())==219
    assert sum(d['payment'] for d in docs)==Decimal(243935201)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==before
    result={'company':lists['company'],'database':lists['database'],'filter':lists['filter'],'source_url':lists['url'],
            'source_sha256':before,'docs':docs,'details':details,'issues':issues,'confirmed_links':confirmations,
            'summary':{'documents':len(docs),'detail_lines':len(details),'unique_bhh':len({r[2] for r in details if r[2]}),
                       'missing_bhh_lines':sum(not r[2] for r in details),'source_mismatch_documents':sum(not d['source_ok'] for d in docs),
                       'goods':sum(r[9] for r in details),'vat':sum(r[11] for r in details),'payment':sum(d['payment'] for d in docs)}}
    (WORK/'prepared.json').write_text(json.dumps(result,ensure_ascii=False,default=native),encoding='utf8')
    print(json.dumps(result['summary'],ensure_ascii=False,indent=2,default=native))
    if issues:print('ISSUES',json.dumps(issues,ensure_ascii=False))

if __name__=='__main__':main()
