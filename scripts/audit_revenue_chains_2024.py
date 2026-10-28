"""Kiểm độc lập doanh thu từng chuỗi, giữ riêng cạnh điều chỉnh và thay thế."""
from pathlib import Path
from collections import defaultdict,Counter
from datetime import date,datetime
from decimal import Decimal
import json,hashlib
from openpyxl import load_workbook
from prepare_invoice_reports import rows,invoice_key,invoice_day,amount
from filter_invoice_exceptions import normal

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/revenue-chain-audit-20261028'
BASE=ROOT/'data/Du lieu Meinvoice'

def key(k):return '/'.join(k)
def money(v):return [amount(x) for x in v]
def zero():return [Decimal(0)]*3
def add(a,b):return [x+y for x,y in zip(a,b)]
def native(v):
    if isinstance(v,Decimal):return int(v) if v==int(v) else float(v)
    if isinstance(v,(date,datetime)):return v.isoformat()
    if isinstance(v,set):return sorted(v)
    raise TypeError(type(v))

def main():
    WORK.mkdir(parents=True,exist_ok=True)
    baseline=json.loads((ROOT/'scripts/data/revenue-forecast-20261028/prepared.json').read_text(encoding='utf8'))
    manifest=[];records={};observations=defaultdict(list);edge_sources=defaultdict(list);hints=defaultdict(set);history=[];issues=[]
    def source(p):
        label=str(p.relative_to(ROOT));digest=hashlib.sha256(p.read_bytes()).hexdigest()
        if not any(m['path']==label for m in manifest):manifest.append({'path':label,'sha256':digest})
        return label
    def observe(k,day,values,label,rn,status='',buyer='',primary=False):
        if not k:return
        v={'key':k,'date':day,'money':money(values),'source':label,'row':rn,'status':status or '', 'buyer':buyer or '', 'primary':primary}
        observations[k].append(v)
        if k not in records or primary:records[k]=v
    for p in sorted((BASE/'Bang_ke_hoa_don_da_su_dung_2024_2025').glob('*.xlsx')):
        label=source(p)
        for rn,r in rows(p):
            if rn<6 or amount(r.get('A')) is None:continue
            k=invoice_key(r.get('B'),r.get('C'))
            if k in records:raise ValueError(('Trùng hóa đơn trong bảng kê',k))
            observe(k,invoice_day(r.get('D')),[r.get(c) for c in ['P','Q','R']],label,rn,r.get('U'),r.get('H') or r.get('L'),True)
    links=sorted(set(list(BASE.glob('*dieu_chinh_xuat*.xls'))+list(BASE.glob('*thay_the_xuat*.xls'))+list((BASE/'Xuat_lai_2026-10-09').glob('*dieu_chinh*.xls'))+list((BASE/'Xuat_lai_2026-10-09').glob('*thay_the*.xls'))+list((BASE/'Xuat_lai_2026-10-22').glob('*dieu_chinh*.xls'))+list((BASE/'Xuat_lai_2026-10-22').glob('*thay_the*.xls'))))
    for p in links:
        label=source(p);kind='replacement' if 'thay_the' in p.name else 'adjustment'
        for rn,r in rows(p):
            if rn<7 or amount(r.get('A')) is None:continue
            old=invoice_key(r.get('B'),r.get('D'));new=invoice_key(r.get('M'),r.get('O'))
            if not old or not new:continue
            edge_sources[(old,new,kind)].append({'file':label,'row':rn})
            observe(old,invoice_day(r.get('C')),[r.get(c) for c in ['J','K','L']],label,rn)
            observe(new,invoice_day(r.get('N')),[r.get(c) for c in ['U','V','W']],label,rn,'Hoá đơn thay thế' if kind=='replacement' else 'Hoá đơn điều chỉnh',r.get('Q'))
    for name in ['HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx','HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','NGOAI_LE_F0_2024_T6_2025.xlsx']:
        p=ROOT/'outputs'/name;label=source(p);wb=load_workbook(p,read_only=True,data_only=True)
        if 'Chi tiet nguon' not in wb.sheetnames:wb.close();continue
        for rn,r in enumerate(wb['Chi tiet nguon'].iter_rows(min_row=2,values_only=True),2):
            if not r[0] or not r[2]:continue
            rt=invoice_key(r[0],r[1]);child=invoice_key(r[2],r[3]);hints[child].add(rt)
            history.append({'root':rt,'child':child,'kind_text':r[7],'status':r[8],'source':label,'row':rn,'original_source':r[12]})
            if child not in records:observe(child,invoice_day(r[4]),[r[c] for c in [9,10,11]],label,rn,r[8],r[5])
        wb.close()
    parents=defaultdict(set)
    for a,b,kind in edge_sources:parents[b].add(a)
    recovered=[]
    for h in history:
        child=h['child'];rt=h['root']
        if child==rt or parents[child]:continue
        types=normal(str(h['kind_text'] or '')+' '+str(h['status'] or ''))
        kind='replacement' if 'thay the' in types else 'adjustment' if 'dieu chinh' in types else None
        if kind and len(hints[child])==1:
            edge_sources[(rt,child,kind)].append({'file':h['source'],'row':h['row'],'basis':'Liên kết lịch sử đã xác minh trong Chi tiet nguon','original_source':h['original_source']})
            parents[child].add(rt);recovered.append({'parent':rt,'child':child,'kind':kind,'source':h['source'],'row':h['row']})
    # Giá trị thiếu chỉ bổ sung từ bằng chứng đã kiểm trước đó, không mặc định bằng 0.
    for r in baseline['old_invoices']:
        k=tuple(r['key'])
        if k in records and None in records[k]['money']:
            assert r.get('zero_basis') or r.get('detail_basis'),k
            records[k]['money']=money(r['money']);records[k]['missing_value_basis']=r.get('zero_basis') or r.get('detail_basis')
    adj=defaultdict(set);rep=defaultdict(set);allchildren=defaultdict(set)
    for a,b,kind in edge_sources:
        (adj if kind=='adjustment' else rep)[a].add(b);allchildren[a].add(b)
    def ancestors(k,stack=frozenset()):
        if k in stack:raise ValueError(('Chu trình',k))
        return {k} if not parents[k] else set.union(*(ancestors(a,stack|{k}) for a in parents[k]))
    chains=defaultdict(set)
    for k,r in records.items():
        roots=ancestors(k)
        if len(roots)!=1:
            if any(records.get(rt,{}).get('date') and records[rt]['date'].year==2024 for rt in roots):issues.append({'kind':'multiple_roots','invoice':k,'roots':roots})
            continue
        rt=next(iter(roots));rr=records.get(rt)
        if rr and rr['date'] and rr['date'].year==2024 and r['date'] and r['date']<=date(2026,10,28):chains[rt].add(k)
    result=[];contributors=set();excluded=set()
    old_by_root=defaultdict(lambda:{'money':zero(),'keys':set()})
    for r in baseline['old_invoices']:
        # So sánh cùng tập hóa đơn, quy về gốc mới khi bổ sung cạnh thay thế cũ.
        rr=ancestors(tuple(r['key']))
        assert len(rr)==1,('Gốc so sánh chưa duy nhất',r['key'],rr)
        q=old_by_root[next(iter(rr))];q['money']=add(q['money'],money(r['money']));q['keys'].add(tuple(r['key']))
    for rt,nodes in sorted(chains.items()):
        included=set();skipped={};problems=[];visited=set()
        def walk(n,stack=frozenset()):
            if n in stack:raise ValueError(('Chu trình nhánh',n))
            if n in visited:problems.append('Hóa đơn được đi tới nhiều lần: '+key(n));return
            visited.add(n);r=records[n]
            replacements=[c for c in rep[n] if c in nodes and 'huy' not in normal(records[c]['status'])]
            if len(replacements)>1:
                problems.append('Nhiều nhánh thay thế chưa xác định duy nhất: '+key(n));return
            if replacements:
                skipped[n]='Bỏ giá trị hóa đơn bị thay thế'
                for c in adj[n]:
                    skipped[c]='Điều chỉnh gắn hóa đơn đã bị thay thế; cần rà chứng cứ'
                    problems.append('Có điều chỉnh trước thay thế: '+key(n)+' → '+key(c))
                walk(replacements[0],stack|{n});return
            status=normal(r['status'])
            if 'huy' in status:skipped[n]='Hóa đơn hủy';return
            if 'bi thay the' in status:
                problems.append('Trạng thái bị thay thế nhưng thiếu hóa đơn thay thế: '+key(n));return
            if n==rt and ('hoa don dieu chinh' in status or 'hoa don thay the' in status):
                problems.append('Thiếu hóa đơn cha của '+key(n));return
            if None in r['money']:problems.append('Thiếu giá trị nguồn: '+key(n));return
            included.add(n)
            for c in sorted(adj[n]):
                if c in nodes:walk(c,stack|{n})
            if 'da bi dieu chinh' in status and not adj[n]:problems.append('Trạng thái đã điều chỉnh nhưng thiếu liên kết: '+key(n))
        walk(rt)
        unreachable=nodes-included-set(skipped)
        if unreachable:problems.append('Nhánh chưa được duyệt: '+', '.join(key(k) for k in sorted(unreachable)))
        val=zero()
        for k in included:val=add(val,records[k]['money'])
        old=old_by_root[rt];delta=[a-b for a,b in zip(val,old['money'])]
        kind='Hỗn hợp' if any(rep[n] for n in nodes) and any(adj[n] for n in nodes) else 'Thay thế' if any(rep[n] for n in nodes) else 'Điều chỉnh' if any(adj[n] for n in nodes) else 'Không phát sinh'
        result.append({'root':rt,'date':records[rt]['date'],'buyer':records[rt]['buyer'],'kind':kind,'node_count':len(nodes),'included':sorted(included),'excluded':[{'key':k,'reason':v} for k,v in sorted(skipped.items())],'before':val,'previous_before':old['money'],'difference':delta,'changed_contributors':included!=old['keys'],'issues':problems,'nodes':sorted(nodes)})
        if problems:issues.append({'kind':'chain_review','root':rt,'reasons':problems})
        if contributors&included:raise ValueError('Cộng trùng giữa hai chuỗi')
        contributors|=included;excluded|=set(skipped)
    # Xuất chứng cứ trước; không chốt tổng nếu còn chuỗi chưa rõ.
    aliases=[{'key':k,'root':next(iter(ancestors(k)))} for k in records if len(ancestors(k))==1]
    payload={'manifest':manifest,'root_aliases':aliases,'records':[v for v in records.values() if any(v['key'] in c['nodes'] for c in result)],'edges':[{'parent':a,'child':b,'kind':k,'evidence':v} for (a,b,k),v in edge_sources.items() if a in set.union(*(c['nodes'] if isinstance(c['nodes'],set) else set(c['nodes']) for c in result))],'chains':result,'recovered_edges':recovered,'issues':issues}
    total=zero()
    for c in result:total=add(total,c['before'])
    changed=[{'root':c['root'],'before':c['before'],'previous':c['previous_before'],'difference':c['difference'],'issues':c['issues']} for c in result if any(c['difference']) or c['changed_contributors']]
    summary={'provisional':bool(issues),'chain_count':len(result),'invoice_count':len(contributors),'chain_kinds':dict(Counter(c['kind'] for c in result)),'before':total,'previous_before':baseline['totals']['before'],'changed_chains':changed,'recovered_edges':len(recovered),'issues':issues,'excluded_invoice_count':len(excluded)}
    for m in manifest:assert hashlib.sha256((ROOT/m['path']).read_bytes()).hexdigest()==m['sha256']
    (WORK/'prepared.json').write_text(json.dumps(payload,ensure_ascii=False,default=native),encoding='utf8')
    (WORK/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=native),encoding='utf8')
    print(json.dumps(summary,ensure_ascii=False,indent=2,default=native))

if __name__=='__main__':main()
