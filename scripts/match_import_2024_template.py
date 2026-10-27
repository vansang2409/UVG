"""Giữ nguyên gói định dạng mẫu, chỉ thay dữ liệu; chuyển cột đối chiếu sang file kiểm tra.

Dùng OOXML trực tiếp để bảo toàn bố cục mẫu và workbook kiểm tra đã có.
Artifact Tool không xuất được workbook đầy đủ lớn trong phiên này; chỉ dùng render mẫu nhỏ.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import hashlib
import json
import math
import os
import sys
from lxml import etree as E
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/amis-2024-linked-20261026'
OUT=Path(os.environ.get('UVG_OUTPUT_DIR',ROOT/'outputs/amis-2024-theo-btl-bhh-20261026'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
Q=lambda name:'{'+NS+'}'+name
def xml(v):return E.tostring(v,encoding='UTF-8',xml_declaration=True,standalone=True)
def package(raw,changes,remove=()):
    out=BytesIO()
    with ZipFile(BytesIO(raw)) as src,ZipFile(out,'w',ZIP_DEFLATED) as dst:
        for name in src.namelist():
            if name not in remove:dst.writestr(name,changes.get(name,src.read(name)))
        for name,value in changes.items():
            if name not in src.namelist():dst.writestr(name,value)
    return out.getvalue()
def cell_value(c,strings):
    if c is None:return None
    if c.get('t')=='inlineStr':return ''.join(c.itertext())
    v=c.find(Q('v'))
    if v is None:return None
    return strings[int(v.text)] if c.get('t')=='s' else v.text
def cell(ref,value,style=0,text=False):
    c=E.Element(Q('c'),r=ref,s=str(style))
    if value is not None:
        if text:
            c.set('t','inlineStr');t=E.SubElement(E.SubElement(c,Q('is')),Q('t'));t.text=str(value)
            if str(value)!=str(value).strip():t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        else:E.SubElement(c,Q('v')).text=str(value)
    return c

def set_column_width(sheet,index,width):
    """Tách nhóm định dạng cột của mẫu, chỉ đổi đúng cột được yêu cầu."""
    cols=sheet.find(Q('cols'))
    existing=next((c for c in cols if int(c.get('min'))<=index<=int(c.get('max'))),None)
    if existing is None:
        E.SubElement(cols,Q('col'),min=str(index),max=str(index),width=str(width),customWidth='1');return
    lo,hi=int(existing.get('min')),int(existing.get('max'));pos=list(cols).index(existing);cols.remove(existing)
    pieces=[]
    if lo<index:
        left=deepcopy(existing);left.set('max',str(index-1));pieces.append(left)
    middle=deepcopy(existing);middle.set('min',str(index));middle.set('max',str(index));middle.set('width',str(width));middle.set('customWidth','1');pieces.append(middle)
    if index<hi:
        right=deepcopy(existing);right.set('min',str(index+1));pieces.append(right)
    for i,c in enumerate(pieces):cols.insert(pos+i,c)

def main():
    template=ROOT/'data/Du lieu Meinvoice/Mau_hoa_don_GTGT_MTT_1_thue_suat.xls'
    imp=OUT/'IMPORT_HOA_DON_MOI_AMIS_2024.xlsx';review=OUT/'KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx'
    for f in (imp,review):
        with f.open('r+b'):pass
    raw=imp.read_bytes();review_raw=review.read_bytes();template_raw=template.read_bytes()
    payload=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
    promotion_codes={d['doc'] for d in payload['docs'] if d.get('promotion') and d['category']=='Đưa vào import'}
    last_col='S'
    with ZipFile(BytesIO(raw)) as z:
        old=E.fromstring(z.read('xl/worksheets/sheet1.xml'))
        strings=[''.join(e.itertext()) for e in E.fromstring(z.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in z.namelist() else []
    source_rows=[r for r in old.find(Q('sheetData')) if int(r.get('r'))>=10]
    if not source_rows or not any(c.get('r','').startswith('W') for c in source_rows[0]):
        raise ValueError('Cần bản xuất có 23 cột để chuyển đối chiếu; không ghi đè file đã chuyển.')
    with ZipFile(BytesIO(template_raw)) as z:
        sheet=E.fromstring(z.read('xl/worksheets/sheet1.xml'));styles=E.fromstring(z.read('xl/styles.xml'))
        workbook=E.fromstring(z.read('xl/workbook.xml'))
    data=sheet.find(Q('sheetData'));sample=next(r for r in data if r.get('r')=='10')
    base_styles={''.join(filter(str.isalpha,c.get('r'))):int(c.get('s','0')) for c in sample}
    # Ngày trong mẫu minh họa là chuỗi; dữ liệu thật vẫn giữ kiểu ngày Excel.
    fmts=styles.find(Q('numFmts'));format_id=max([163]+[int(f.get('numFmtId')) for f in fmts])+1
    E.SubElement(fmts,Q('numFmt'),numFmtId=str(format_id),formatCode='dd/mm/yyyy');fmts.set('count',str(len(fmts)))
    xfs=styles.find(Q('cellXfs'));date_style=deepcopy(xfs[base_styles['B']]);date_style.set('numFmtId',str(format_id));date_style.set('applyNumberFormat','1')
    base_styles['B']=len(xfs);xfs.append(date_style)
    item_style=deepcopy(xfs[base_styles['M']]);alignment=item_style.find(Q('alignment'))
    if alignment is None:alignment=E.SubElement(item_style,Q('alignment'))
    alignment.set('wrapText','1');base_styles['M']=len(xfs);xfs.append(item_style);xfs.set('count',str(len(xfs)))
    for row in list(data):
        if int(row.get('r'))>=10:data.remove(row)
    if promotion_codes:
        header=next(row for row in data if row.get('r')=='9')
        style=next(c.get('s') for c in header if c.get('r')=='K9')
        header.append(cell('R9','Hàng khuyến mại',style,True))
        set_column_width(sheet,18,24.57)
    header=next(row for row in data if row.get('r')=='9')
    header.append(cell('S9','Số chứng từ',next(c.get('s') for c in header if c.get('r')=='D9'),True))
    set_column_width(sheet,19,27)
    audit=[]
    for row in source_rows:
        n=int(row.get('r'));cells={''.join(filter(str.isalpha,c.get('r'))):c for c in row}
        audit.append([cell_value(cells.get(c),strings) for c in ['A','R','U','V','W','S','T']])
        nr=E.SubElement(data,Q('row'),r=str(n),customHeight='1')
        item=cell_value(cells.get('M'),strings) or ''
        nr.set('ht',str(max(15,min(105,15*math.ceil(len(item)/27)))))
        for c in 'ABCDEFGHIJKLMNOPQ':
            oldc=cells.get(c)
            if oldc is None:nr.append(cell(c+str(n),None,base_styles[c]));continue
            newc=deepcopy(oldc);newc.set('s',str(base_styles[c]));nr.append(newc)
        if promotion_codes:nr.append(cell('R'+str(n),1 if cell_value(cells.get('R'),strings) in promotion_codes else None,base_styles['K']))
        nr.append(cell('S'+str(n),cell_value(cells.get('R'),strings),base_styles['D'],True))
    last=int(source_rows[-1].get('r'));sheet.find(Q('dimension')).set('ref',f'A1:{last_col}{last}')
    views=sheet.find(Q('sheetViews'))
    for view in views:
        view.set('topLeftCell','A1')
        for selection in view.findall(Q('selection')):selection.set('activeCell','A1');selection.set('sqref','A1')
    for name in list(workbook.findall('.//'+Q('definedName'))):
        if '#ref!' in (name.text or '').lower():name.getparent().remove(name)
    import_bytes=package(template_raw,{'xl/worksheets/sheet1.xml':xml(sheet),'xl/styles.xml':xml(styles),'xl/workbook.xml':xml(workbook)})
    # Thêm một sheet đối chiếu vào workbook kiểm tra, không đổi các sheet/cột quyết định đã có.
    with ZipFile(BytesIO(review_raw)) as z:
        rw=E.fromstring(z.read('xl/workbook.xml'));rr=E.fromstring(z.read('xl/_rels/workbook.xml.rels'));ct=E.fromstring(z.read('[Content_Types].xml'))
        existing=E.fromstring(z.read('xl/worksheets/sheet1.xml'));sample2=existing.find(Q('sheetData'))
        header_style=sample2[0][0].get('s','0');body_style=sample2[1][0].get('s','0');money_style=sample2[1][3].get('s','0')
        audit_name='Doi chieu import';sheets=rw.find(Q('sheets'))
        audit_existing=next((s for s in sheets if s.get('name')==audit_name),None)
        if audit_existing is not None:
            sid=int(audit_existing.get('sheetId'));rid=audit_existing.get('{'+REL+'}id')
            target=next(r.get('Target') for r in rr if r.get('Id')==rid)
            part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        else:
            sid=max(int(s.get('sheetId')) for s in sheets)+1;rid='UVG_ImportAudit';part=f'xl/worksheets/sheet{sid}.xml'
            assert part not in z.namelist()
    new=E.Element(Q('worksheet'),nsmap={None:NS});E.SubElement(new,Q('dimension'),ref=f'A1:H{len(audit)+1}')
    view=E.SubElement(E.SubElement(new,Q('sheetViews')),Q('sheetView'),workbookViewId='0',showGridLines='0')
    E.SubElement(view,Q('pane'),ySplit='1',state='frozen',topLeftCell='A2',activePane='bottomLeft')
    cols=E.SubElement(new,Q('cols'))
    for i,width in enumerate([18,22,29,20,29,32,27,29],1):E.SubElement(cols,Q('col'),min=str(i),max=str(i),width=str(width),customWidth='1')
    sd=E.SubElement(new,Q('sheetData'));hr=E.SubElement(sd,Q('row'),r='1',ht='44',customHeight='1')
    headers=['STT hóa đơn','Dòng Excel import','Số chứng từ AMIS','Dòng nguồn AMIS','Mã hàng','BTL đã bù trừ','Tổng tiền hàng','Tổng thanh toán']
    for i,h in enumerate(headers):hr.append(cell(chr(65+i)+'1',h,header_style,True))
    for i,r in enumerate(audit,2):
        values=[r[0],i+8,r[1],r[2],r[3],r[4],r[5],r[6]];nr=E.SubElement(sd,Q('row'),r=str(i),ht='20',customHeight='1')
        for j,value in enumerate(values):nr.append(cell(chr(65+j)+str(i),value,money_style if j in [6,7] else body_style,j in [2,4,5]))
    E.SubElement(new,Q('autoFilter'),ref=f'A1:H{len(audit)+1}')
    if audit_existing is None:
        E.SubElement(sheets,Q('sheet'),name=audit_name,sheetId=str(sid),attrib={'{'+REL+'}id':rid})
        E.SubElement(rr,'{'+PKG+'}Relationship',Id=rid,Type=REL+'/worksheet',Target=f'worksheets/sheet{sid}.xml')
        E.SubElement(ct,'{'+CT+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    review_bytes=package(review_raw,{'xl/workbook.xml':xml(rw),'xl/_rels/workbook.xml.rels':xml(rr),'[Content_Types].xml':xml(ct),part:xml(new)})
    # Xác minh đầy đủ A:Q trước khi cập nhật tại chỗ hai file chính.
    before=load_workbook(BytesIO(raw),read_only=True,data_only=True).active
    after=load_workbook(BytesIO(import_bytes),read_only=True,data_only=True).active
    assert after.max_column==19 and after.max_row==before.max_row
    count=0
    for a,b in zip(before.iter_rows(min_row=10,max_col=17,values_only=True),after.iter_rows(min_row=10,max_col=17,values_only=True)):
        assert a==b,(count,a,b)
        count+=1
    assert count==len(source_rows)
    # Lưu tạm nội bộ, chỉ thay file chính sau khi kiểm xong. Không sinh phiên bản mới ở outputs.
    temp_imp=WORK/'_format-import.tmp.xlsx';temp_review=WORK/'_format-review.tmp.xlsx'
    temp_imp.write_bytes(import_bytes);temp_review.write_bytes(review_bytes)
    os.replace(temp_review,review);os.replace(temp_imp,imp)
    # Gói xem trước nhỏ dùng Artifact Tool để kiểm bố cục, không phải file bàn giao.
    preview=deepcopy(sheet);preview_data=preview.find(Q('sheetData'))
    for row in list(preview_data):
        if int(row.get('r'))>13:preview_data.remove(row)
    preview.find(Q('dimension')).set('ref',f'A1:{last_col}13')
    (WORK/'_template-preview.tmp.xlsx').write_bytes(package(import_bytes,{'xl/worksheets/sheet1.xml':xml(preview)}))
    (WORK/'template-format-verification.json').write_text(json.dumps({'passed':True,'columns':19,'template_columns':17,'document_column':'S - Số chứng từ','promotion_column':'R - Hàng khuyến mại' if promotion_codes else None,'promotion_codes':sorted(promotion_codes),'lines':count,'same_all_import_values':True,'header_and_instructions':'Giữ nguyên A:Q dòng 1–9 và độ rộng cột từ mẫu gốc; thêm R khuyến mại, S số chứng từ theo yêu cầu','audit_sheet':audit_name,'audit_rows':len(audit),'template_sha256':hashlib.sha256(template_raw).hexdigest(),'files':[{'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in (imp,review)]},ensure_ascii=False,indent=2),encoding='utf8')
    print(f'Đã cập nhật tại chỗ 2 file; {count:,} dòng A:Q không đổi; đối chiếu chuyển sang file kiểm tra.')

def add_document_numbers():
    """Chỉ bổ sung cột S trên file đã chốt, không tái xuất hoặc đổi các ô cũ."""
    imp=OUT/'IMPORT_HOA_DON_MOI_AMIS_2024.xlsx';review=OUT/'KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx'
    with imp.open('r+b'):pass
    review_hash=hashlib.sha256(review.read_bytes()).hexdigest()
    audit=load_workbook(review,read_only=True,data_only=True)['Doi chieu import']
    mapping={int(r[1]):(int(r[0]),str(r[2])) for r in audit.iter_rows(min_row=2,values_only=True)}
    raw=imp.read_bytes()
    with ZipFile(BytesIO(raw)) as z:
        sheet=E.fromstring(z.read('xl/worksheets/sheet1.xml'))
    data=sheet.find(Q('sheetData'));rows=list(data)
    header=next(r for r in rows if r.get('r')=='9')
    if any(c.get('r')=='S9' for c in header):raise ValueError('Cột S đã có; không tự ghi đè')
    old_cells={int(r.get('r')):[xml(c) for c in r] for r in rows}
    header.append(cell('S9','Số chứng từ',next(c.get('s') for c in header if c.get('r')=='D9'),True))
    set_column_width(sheet,19,27)
    changed=0
    for row in rows:
        n=int(row.get('r'))
        if n<10:continue
        index,code=mapping[n]
        invoice=row.find(Q('c')+f'[@r="A{n}"]')
        assert invoice is not None and int(cell_value(invoice,[]))==index
        assert code and not code.startswith('BTL')
        body_style=next(c.get('s') for c in row if c.get('r')==f'D{n}')
        row.append(cell(f'S{n}',code,body_style,True));changed+=1
    assert changed==len(mapping)==43972
    sheet.find(Q('dimension')).set('ref',f'A1:S{max(mapping)}')
    filt=sheet.find(Q('autoFilter'))
    if filt is not None:filt.set('ref',f'A9:S{max(mapping)}')
    for row in rows:
        n=int(row.get('r'));actual=[xml(c) for c in row if not c.get('r').startswith('S')]
        assert actual==old_cells[n],n
    result=package(raw,{'xl/worksheets/sheet1.xml':xml(sheet)})
    with ZipFile(BytesIO(raw)) as a,ZipFile(BytesIO(result)) as b:
        for name in a.namelist():
            if name!='xl/worksheets/sheet1.xml':assert a.read(name)==b.read(name)
    check=load_workbook(BytesIO(result),read_only=True,data_only=True).active
    assert check.max_column==19 and check.cell(9,19).value=='Số chứng từ'
    count=0
    for n,r in enumerate(check.iter_rows(min_row=10,values_only=True),10):
        assert r[0]==mapping[n][0] and r[18]==mapping[n][1];count+=1
    assert count==43972 and hashlib.sha256(review.read_bytes()).hexdigest()==review_hash
    tmp=WORK/'_add-document-number.tmp.xlsx';tmp.write_bytes(result);os.replace(tmp,imp)
    report={'passed':True,'column':'S - Số chứng từ','lines':count,'unchanged':'Mọi ô A:R và toàn bộ workbook kiểm tra giữ nguyên','sha256':hashlib.sha256(result).hexdigest()}
    (WORK/'document-column-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    for name in ['verification.json','template-format-verification.json']:
        p=WORK/name;r=json.loads(p.read_text(encoding='utf8'))
        for f in r['files']:
            file=Path(f['path']);f['sha256']=hashlib.sha256(file.read_bytes()).hexdigest()
            if 'bytes' in f:f['bytes']=file.stat().st_size
        r['document_column']='S - Số chứng từ; đủ 43.972 dòng khớp Doi chieu import'
        if name=='template-format-verification.json':r['columns']=19
        p.write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf8')
    print('Đã điền Số chứng từ tại S cho đủ 43.972 dòng; A:R và file kiểm tra không đổi.')

if __name__=='__main__':
    if '--add-document-number' in sys.argv:add_document_numbers()
    else:main()
