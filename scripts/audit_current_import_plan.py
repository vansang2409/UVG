"""Audit current adjusted import against approved plans, including XML price fixes."""
from pathlib import Path
from collections import defaultdict, Counter
from decimal import Decimal
from datetime import datetime
import json
import hashlib
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
H = ROOT / 'scripts/data/check-review'
read = lambda n: json.loads((H / n).read_text(encoding='utf-8'))
base, addition, edits = read('adjusted-import-meta.json'), read('195-xml-data.json'), read('supplement-57-price-changes.json')
plan = read('adjusted-import-data.json')
p = ROOT / 'outputs/HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
before = sha(p)
expected = [list(v) for v in base['rows']] + [list(v) for v in addition['rows']]
assert len(expected) == 41268
for i, v in enumerate(expected):
    v[1] = datetime(2026, 10, 8 if i < len(base['rows']) else 9)
for c in edits['changes']:
    n = int(c['cell'][1:]) - 10
    assert expected[n][17] == c['old']
    expected[n][17] = c['new']
w = openpyxl.load_workbook(p, read_only=True, data_only=True)
groups, errors, product_diffs = defaultdict(list), [], []
rows_checked = 0
for rn, v in enumerate(w['Import dieu chinh'].iter_rows(min_row=10, max_col=24, values_only=True), 10):
    if v[0] is None:
        assert all(z is None for z in v), ('orphan', rn)
        continue
    rows_checked += 1
    assert list(v) == expected[rn - 10], ('approved row mismatch', rn)
    groups[v[0]].append((rn, v))
    if v[16] is not None and v[17] is not None and Decimal(str(v[16])) * Decimal(str(v[17])) != Decimal(str(v[18])):
        product_diffs.append({'row': rn, 'key': v[9:11], 'qty': v[16], 'price': v[17], 'goods': v[18], 'difference': str(Decimal(str(v[16])) * Decimal(str(v[17])) - Decimal(str(v[18])))})
assert rows_checked == len(expected) and set(groups) == set(range(1, 1991))
roots, reversals, reductions = set(), 0, 0
cumulative = defaultdict(lambda: [Decimal(0)] * 3)
for seq, rs in groups.items():
    v = rs[0][1]
    root = tuple(v[9:11]); roots.add(root)
    assert all(tuple(z[9:11]) == root for _, z in rs), seq
    assert sum(Decimal(str(z[18])) for _, z in rs) == Decimal(str(v[19])), seq
    money = [Decimal(str(v[c])) for c in [19, 21, 22]]
    assert money[0] + money[1] == money[2], seq
    for i, x in enumerate(money): cumulative[root][i] += x
    if len(rs) == 1 and v[16] is None: reversals += 1
    else: reductions += 1
    for rn, z in rs[1:]: assert all(z[c] is None for c in [19, 20, 21, 22]), rn
main = {tuple(v[:2]): v for v in w['Dieu chinh'].iter_rows(min_row=3, max_col=55, values_only=True) if v[0]}
for root, values in cumulative.items():
    assert all(Decimal(str(main[root][24+i])) + x == 0 for i, x in enumerate(values)), root
    assert all(main[root][i] == 0 for i in [52, 53, 54]), ('predicted balance', root)
ready = {tuple(x['key']) for x in plan['ready']} | {tuple(x['root']) for x in addition['chains']}
assert roots == ready and len(roots) == 986
pending = {tuple(v[:2]) for v in w['Cho KT xac nhan'].iter_rows(min_row=3, max_col=2, values_only=True) if v[0]}
assert len(pending) == 19 and not roots & pending
xml_verified = 0
for r in addition['chains']:
    xp = ROOT / 'data/Bo sung/meInvoice/2026-10-09-237-XML' / ('_'.join(r['adjustment']) + '.xml')
    assert sha(xp) == r['xml_sha256']; xml_verified += 1
for s in w:
    for row in s.iter_rows():
        for c in row:
            if c.data_type == 'e': errors.append((s.title, c.coordinate, c.value))
w.close()
assert not errors and before == sha(p)
result = {'passed': True, 'read_only': True, 'workbook_sha256': before, 'invoices': len(groups), 'roots': len(roots), 'rows_verified_against_approved_plan': rows_checked, 'reverse_invoices': reversals, 'reduce_invoices': reductions, 'all_986_chains_zero': True, 'pending_19_excluded': True, 'xml_adjustment_hashes_verified': xml_verified, 'product_difference_count': len(product_diffs), 'product_differences': product_diffs, 'notes': ['Quantity times unit price differences are source-derived diagnostics; approved line amount is preserved.', '57 supplemental originals are audited separately against current source data.']}
(H / 'current-import-plan-audit-20261011.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: v for k, v in result.items() if k != 'product_differences'}, ensure_ascii=False))
