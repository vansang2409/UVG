"""Read-only verification of the delivered count and 72-invoice list."""
from collections import Counter
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import json
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
d=json.loads((HERE/'data.json').read_text(encoding='utf8'))
path=ROOT/d['output']
with ZipFile(path) as z:assert z.testzip() is None
w=openpyxl.load_workbook(path,data_only=False)
c=openpyxl.load_workbook(path,data_only=True)
assert w.sheetnames==['Tong hop','Khong can dieu chinh']
s=c['Khong can dieu chinh']
assert s.max_row==78
assert [s.cell(6,col).value for col in range(1,9)]==[
 'STT','Nhóm','Ký hiệu','Số hóa đơn','Ngày hóa đơn','Hóa đơn gốc của chuỗi','Tên người mua','Lý do không cần xuất điều chỉnh']
for rn,record in enumerate(d['records'],7):
    expected=[rn-6,record['group'],*record['key'],datetime.fromisoformat(record['date']),record['root'],record['buyer'],record['reason']]
    assert [s.cell(rn,col).value for col in range(1,9)]==expected,(rn,record['key'])
    assert w[s.title].cell(rn,4).data_type=='s'
    assert s.cell(rn,4).number_format=='00000000'
    assert len(s.cell(rn,4).value)==8 and s.cell(rn,4).value.isdigit()
    assert s.cell(rn,5).number_format=='dd/mm/yyyy'
    assert s.cell(rn,8).value and len(s.cell(rn,8).value)<160
keys={tuple(r['key']) for r in d['records']}
additional={tuple(r['key']) for r in d['additional_records']}
assert len(keys)==72 and len(additional)==13 and not keys&additional
source=openpyxl.load_workbook(ROOT/d['source'],read_only=True,data_only=True)
source_keys={tuple(r[2:4]) for r in source['85 hoa don'].iter_rows(min_row=7,values_only=True) if isinstance(r[0],(int,float))}
assert source_keys==keys|additional
source.close()
counts=Counter(s.cell(row,2).value for row in range(7,79))
assert counts=={r['label']:r['count'] for r in d['groups']}
info={('1C25TUV',n) for n in ('00004203','00004204','00006108','00009186')}
assert info<=keys and not info&additional
root_keys={tuple(r['root'].split('/')) for r in d['records'] if tuple(r['key']) in info}
assert root_keys-{('1C25TUV','00009072')}<=additional
row9186=next(r for r in d['records'] if r['key'][1]=='00009186')
assert 'tên hàng' in row9186['reason'] and 'người mua' not in row9186['reason']
for key,value in {'B5':10527,'B6':72,'B7':10455,'B8':10442,'B9':13,'B13':62,'B14':6,'B15':4,'B16':72}.items():
    assert c['Tong hop'][key].value==value,(key,c['Tong hop'][key].value,value)
assert w['Tong hop']['B6'].value=="=COUNTA('Khong can dieu chinh'!D7:D78)"
assert w['Tong hop']['B7'].value=='=B5-B6'
assert w['Tong hop']['B9'].value=='=B7-B8'
table=next(iter(w[s.title].tables.values()))
assert table.ref=='A6:H78' and table.autoFilter is not None
assert w[s.title].freeze_panes=='E7'
errors=[(sheet.title,cell.coordinate,cell.value) for sheet in c for row in sheet for cell in row if cell.data_type=='e']
assert not errors,errors
for name,h in d['source_hashes'].items():assert sha256((ROOT/name).read_bytes()).hexdigest()==h,('source changed',name)
assert len(d['zero_money_checks'])==34 and all(x['goods_tax_payment']==[0,0,0] for x in d['zero_money_checks'])
result={'passed':True,'output':d['output'],'no_adjustment_invoices':72,'groups':dict(counts),
        'previous_to_issue':10442,'additional_to_issue':13,'new_to_issue':10455,
        'unique_source_invoices':10527,'distinct_keys':72,'checked_zero_chains':34,
        'information_only_invoices':4,'information_buyer_only':3,'information_product_name_only':1,
        'all_dates_and_invoice_numbers_checked':True,'original_sources_and_import_workbooks_unchanged':True,
        'formula_errors':0,'output_sha256':sha256(path.read_bytes()).hexdigest()}
(HERE/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
w.close();c.close()
