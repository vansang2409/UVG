"""Recount the 85 reconciliation invoices; read sources without changing imports."""
from collections import Counter, defaultdict
from datetime import datetime, date
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
import json
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE = ROOT/'outputs/doi-chieu-85-20261019/DANH_SACH_85_HOA_DON_KHONG_DUA_VAO_IMPORT.xlsx'
OUTPUT = ROOT/'outputs/doi-chieu-85-20261019/UVG_72_HOA_DON_KHONG_CAN_DIEU_CHINH_20261019.xlsx'
NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tracked = {}
def remember(p):
    tracked[p.relative_to(ROOT).as_posix()] = sha256(p.read_bytes()).hexdigest()
    return p

def import_count(path):
    remember(path)
    with ZipFile(path) as z:
        wb=E.fromstring(z.read('xl/workbook.xml'))
        rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        sh=next(s for s in wb.find('s:sheets',NS) if s.get('name')=='Import dieu chinh')
        rid=sh.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        target=next(r.get('Target') for r in rels if r.get('Id')==rid)
        target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        seqs=set(); rows=0
        with z.open(target) as stream:
            for _,row in E.iterparse(stream,events=('end',),tag='{'+NS['s']+'}row'):
                rn=int(row.get('r'))
                if rn>=10:
                    c=next((x for x in row if x.get('r')=='A'+str(rn)),None)
                    v=c.find('s:v',NS) if c is not None else None
                    if v is not None and v.text:
                        assert c.get('t') in (None,'n'),(path,rn,c.get('t'))
                        seqs.add(int(v.text));rows+=1
                row.clear()
                while row.getprevious() is not None: del row.getparent()[0]
        assert seqs==set(range(1,len(seqs)+1))
        return {'invoices':len(seqs),'rows':rows}

counts={stem:import_count(ROOT/'outputs'/(stem+'_F0_2024_T6_2025.xlsx'))
        for stem in ('HOA_DON_CHUA_DIEU_CHINH','HOA_DON_DA_DIEU_CHINH')}
baseline=sum(x['invoices'] for x in counts.values())
assert baseline==10442
for name in ['NGOAI_LE_F0_2024_T6_2025.xlsx','HOA_DON_CHUA_DIEU_CHINH_IMPORT_2_SHEET.xlsx','HOA_DON_DA_DIEU_CHINH_IMPORT_2_SHEET.xlsx']:
    remember(ROOT/'outputs'/name)

w=openpyxl.load_workbook(remember(SOURCE),read_only=True,data_only=True)
raw=[r for r in w['85 hoa don'].iter_rows(min_row=7,values_only=True) if isinstance(r[0],(int,float))]
assert len(raw)==85 and len({tuple(r[2:4]) for r in raw})==85
assert Counter(r[1] for r in raw)=={'31 chuỗi đã khớp':62,'6 cặp giữ riêng':12,'Ngoại lệ công ty':4,'3 chuỗi cùng sản phẩm':6,'Chỉ điều chỉnh tên':1}
w.close()

info={('1C25TUV',n) for n in ('00006108','00004203','00004204','00009186')}
info_evidence=[]
for key in sorted(info):
    p=remember(ROOT/'data/Bo sung/meInvoice/2026-10-11-pending-XML'/(key[0]+'_'+key[1]+'.xml'))
    xml=E.parse(str(p))
    values=[]
    for field in ('TgTCThue','TgTThue','TgTTTBSo'):
        matches=xml.xpath('//*[local-name()="'+field+'"]')
        assert len(matches)==1,(key,field)
        values.append(Decimal(matches[0].text))
    assert values==[0,0,0],(key,values)
    note=xml.xpath('//*[local-name()="TTHDLQuan"]/*[local-name()="GChu"]/text()')
    if not note: note=xml.xpath('//*[local-name()="GChu"]/text()')
    info_evidence.append({'key':list(key),'money':[0,0,0],'note':note})

labels=['Đã khớp chi tiết và số tiền','Cùng chuỗi sản phẩm, đã khớp','Chỉ điều chỉnh thông tin']
reasons=[
 'Đã điều chỉnh khớp chi tiết và số tiền; tổng các hóa đơn liên quan = 0.',
 'Đã điều chỉnh cùng chuỗi sản phẩm, khớp chi tiết và số tiền; tổng các hóa đơn liên quan = 0.',
 'Chỉ điều chỉnh tên/thông tin người mua, không đổi tiền hàng và tiền thuế; chỉ xử lý hóa đơn gốc.'
]
kept=[];additional=[]
for r in raw:
    key=tuple(r[2:4])
    index=0 if r[1]=='31 chuỗi đã khớp' else 1 if r[1]=='3 chuỗi cùng sản phẩm' else 2 if key in info else None
    obj={'key':list(key),'date':r[4].date().isoformat(),'root':r[5],'buyer':r[7],
         'old_group':r[1],'source_excel_row':int(r[0])+6,'accountant_locator':r[10]}
    if index is None:
        obj['reason']='MST cá nhân; bổ sung điều chỉnh trái dấu.' if r[1]=='Ngoại lệ công ty' else 'Thay quyết định giữ cả cặp; bổ sung điều chỉnh trái dấu.'
        additional.append(obj)
    else:
        reason=reasons[index]
        if key==('1C25TUV','00009186'):
            reason='Chỉ điều chỉnh tên hàng, không đổi tiền hàng và tiền thuế; giữ 00009186, chỉ xử lý gốc 00009072.'
        obj.update(group=labels[index],reason=reason,rank=index)
        kept.append(obj)
kept.sort(key=lambda r:r['rank'])
assert len(kept)==72 and len(additional)==13
assert Counter(r['group'] for r in kept)==dict(zip(labels,[62,6,4]))
assert len({r['root'] for r in additional})==8
assert {tuple(r['key']) for r in additional if r['old_group']=='Ngoại lệ công ty'}=={
 ('1C24TUV','00000478'),('1C25TUV','00002578'),('1C25TUV','00000040'),('1C25TUV','00000662')}
assert len({r['root'] for r in additional if r['old_group']=='6 cặp giữ riêng'})==6
assert {tuple(r['key']) for r in kept}&info==info

# Check the numerical basis of the retained 31+3 chains independently of the list labels.
wa=openpyxl.load_workbook(ROOT/'outputs/HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx',read_only=True,data_only=True)
by_root=defaultdict(dict)
for r in wa['Chi tiet nguon'].iter_rows(min_row=2,values_only=True):
    if r[0] and r[2]:by_root['/'.join(r[:2])][tuple(r[2:4])]=r
zero_checks=[]
for group in labels[:2]:
    roots={r['root'] for r in kept if r['group']==group}
    assert len(roots)==(31 if group==labels[0] else 3)
    for root in sorted(roots):
        rows=list(by_root[root].values())
        assert len(rows)==2,(root,len(rows))
        totals=[]
        for col in (9,10,11):
            values=[r[col] for r in rows]
            assert all(isinstance(v,(int,float)) for v in values),(root,col,values)
            totals.append(sum(Decimal(str(v)) for v in values))
        assert totals==[0,0,0],(root,totals)
        assert {tuple(r[2:4]) for r in rows}=={tuple(r['key']) for r in kept if r['root']==root}
        zero_checks.append({'root':root,'goods_tax_payment':[0,0,0]})
wa.close()

result={'date':'2026-10-19','source':SOURCE.relative_to(ROOT).as_posix(),
        'output':OUTPUT.relative_to(ROOT).as_posix(),'baseline_imports':counts,
        'source_invoices':baseline+85,'baseline_to_issue':baseline,'previous_difference':85,
        'no_adjustment':72,'additional_to_issue':13,'new_to_issue':baseline+13,
        'groups':[{'label':l,'count':n} for l,n in zip(labels,[62,6,4])],
        'records':kept,'additional_records':additional,'information_only_evidence':info_evidence,
        'zero_money_checks':zero_checks,'source_hashes':tracked,
        'quantity_price_rule':'Giảm số lượng khi điều chỉnh giảm hóa đơn gốc; đơn giá dương. Hóa đơn mới trái dấu với hóa đơn được điều chỉnh; kiểm tổng cả chuỗi bằng 0. Không lập đảo tiền cho điều chỉnh chỉ thay thông tin người mua.',
        'import_workbooks_modified':False,'software_imported':False,'signed':False,'issued':False}
(HERE/'data.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:result[k] for k in ['baseline_imports','source_invoices','no_adjustment','additional_to_issue','new_to_issue','groups']},ensure_ascii=False))
