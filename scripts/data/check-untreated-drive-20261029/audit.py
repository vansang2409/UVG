from pathlib import Path
from collections import defaultdict,Counter
from decimal import Decimal
from datetime import datetime,date
import sys,json,hashlib
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_invoice_reports import rows,invoice_key,amount,invoice_day
from filter_invoice_exceptions import normal
sys.stdout.reconfigure(encoding='utf-8')
HERE=Path(__file__).resolve().parent
def readbook(path):
 w=load_workbook(path,read_only=True,data_only=True);data={}
 for s in w:
  start=10 if s.title=='Import dieu chinh' else 3
  cols=24 if start==10 else 16
  data[s.title]=[(rn,list(r)) for rn,r in enumerate(s.iter_rows(min_row=start,max_col=cols,values_only=True),start) if any(v is not None for v in r)]
 w.close();return data
cloud=readbook(HERE/'source.xlsx');local=readbook(ROOT/'outputs/HOA_DON_CHUA_DIEU_CHINH_IMPORT_2_SHEET.xlsx')
def norm(v):
 if isinstance(v,(datetime,date)):return v.isoformat()[:10]
 return None if v=='' else v
diffs=[]
for sn in cloud:
 a=cloud[sn];b=local[sn]
 for i in range(max(len(a),len(b))):
  if i>=len(a) or i>=len(b):diffs.append({'sheet':sn,'kind':'row_count','index':i});continue
  for c,(x,y) in enumerate(zip(a[i][1],b[i][1])):
   if norm(x)!=norm(y):diffs.append({'sheet':sn,'row':a[i][0],'col':c+1,'cloud':str(x),'local':str(y)})
groups=defaultdict(list);roots=defaultdict(list);issues=[]
for rn,r in cloud['Import dieu chinh']:
 if r[0] is None:issues.append({'kind':'orphan','row':rn});continue
 groups[r[0]].append((rn,r))
for seq,g in groups.items():
 first=g[0][1];key=invoice_key(first[9],first[10]);roots[key].append(seq)
 assert all(invoice_key(r[9],r[10])==key for _,r in g)
 for c in [19,21,22]:assert first[c] is not None,(seq,c)
 assert sum(Decimal(str(r[18])) for _,r in g if r[18] is not None)==Decimal(str(first[19])),('line_total',seq)
 assert Decimal(str(first[19]))+Decimal(str(first[21]))==Decimal(str(first[22])),('tax_total',seq)
 assert all(all(r[c] is None for c in [19,20,21,22]) for _,r in g[1:]),seq
f0={};dupf0=[]
for rn,r in cloud['F0']:
 key=invoice_key(r[0],r[1])
 if key in f0:dupf0.append(key)
 f0[key]={'row':rn,'money':r[3:6]}
for key,seqs in roots.items():
 if key not in f0:continue
 for seq in seqs:
  m=[groups[seq][0][1][c] for c in [19,21,22]]
  if any(amount(f0[key]['money'][c])+amount(m[c])!=0 for c in range(3)):issues.append({'kind':'f0_not_zero','key':key,'seq':seq})
# Đọc lại trạng thái, tiền và mã CQT hiện tại của các gốc trong nhóm chưa điều chỉnh.
base=ROOT/'data/Meinv';current={};missingcode=[];moneydiff=[];status_counts=Counter()
for p in [base/'2024.xls']+sorted((base/'2025').glob('*.xls'))[:2]:
 for rn,r in rows(p):
  if rn<6 or amount(r.get('A')) is None:continue
  k=invoice_key(r.get('B'),r.get('C'))
  if k not in roots:continue
  money=[amount(r.get(c)) for c in ['P','Q','R']]
  current[k]={'file':str(p.relative_to(ROOT)),'row':rn,'status':r.get('U'),'cqt_status':r.get('S'),'cqt_code':r.get('E'),'money':money}
  status_counts[r.get('U')]+=1
  if not r.get('E'):missingcode.append(k)
  for seq in roots[k]:
   if any(money[c] is None or money[c]+amount(groups[seq][0][1][ic])!=0 for c,ic in enumerate([19,21,22])):moneydiff.append({'key':k,'seq':seq,'current':money})
links=defaultdict(set)
for p in sorted((base/'HD_bi_dieu_chinh_2024_2026').glob('*.xls')):
 for rn,r in rows(p):
  if rn<7 or amount(r.get('A')) is None:continue
  old=invoice_key(r.get('B'),r.get('D'));new=invoice_key(r.get('M'),r.get('O'))
  if old in roots:links[old].add(new)
refs=json.loads((ROOT/'scripts/data/invalid-status-20261029/refs.json').read_text(encoding='utf-8'))
overlap=[k for k in roots if any(x['group']=='Đã điều chỉnh' for x in refs.get('/'.join(k),[]))]
result={'cloud_rows':{k:len(v) for k,v in cloud.items()},'local_rows':{k:len(v) for k,v in local.items()},'cloud_local_data_diffs':diffs,'invoice_count':len(groups),'unique_roots':len(roots),'sequence_min':min(groups),'sequence_max':max(groups),'missing_sequences':sorted(set(range(1,8411))-set(groups)),'duplicate_root_sequences':[{'key':k,'seqs':s} for k,s in roots.items() if len(s)>1],'duplicate_f0':dupf0,'import_roots_not_in_f0':sorted(set(roots)-set(f0)),'f0_not_in_import':sorted(set(f0)-set(roots)),'current_missing_roots':sorted(set(roots)-set(current)),'current_status_counts':dict(status_counts),'current_money_differences':moneydiff,'current_missing_cqt_code':missingcode,'roots_with_existing_adjustments':[{'root':k,'adjustments':sorted(v)} for k,v in links.items()],'overlap_adjusted_group':overlap,'checks':issues,'totals':[sum(Decimal(str(g[0][1][c])) for g in groups.values()) for c in [19,21,22]]}
(HERE/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,default=str))
