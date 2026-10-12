from pathlib import Path
from datetime import date, datetime
from collections import Counter
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
H = ROOT/'scripts/data/check-review/kt-workbook-20261011'
plan = json.loads((H/'confirmed-plan-data.json').read_text(encoding='utf8'))
data = json.loads((H/'review-data.json').read_text(encoding='utf8'))
print('ACTIONS', json.dumps([{'case':a['case'],'root':a['root'],'target':a['target'],'action':a['action'],'money':a['money']} for a in plan['actions']],ensure_ascii=False))
print('REVIEW11',json.dumps(data['audit']['review_11'][0],ensure_ascii=False)[:3500])
for path in sorted((ROOT/'outputs').glob('*.xlsx')):
    w = openpyxl.load_workbook(path, data_only=False, read_only=False)
    print('\nWORKBOOK',path.name)
    for s in w:
        print('SHEET',s.title,s.max_row,s.max_column,'pane',s.freeze_panes,'tables',[(t.name,t.ref) for t in s.tables.values()])
        top=[list(v) for v in s.iter_rows(min_row=1,max_row=min(3,s.max_row),max_col=min(s.max_column,60),values_only=True)]
        print('TOP',json.dumps(top,ensure_ascii=False,default=str)[:2200])
        if s.title in ('Dieu chinh','Chua dieu chinh','Ngoai le'):
            targets={tuple(a['root']) for a in plan['actions']}|{tuple(c['root']) for c in json.loads((H/'review-manifest.json').read_text(encoding='utf8'))['resolved']}
            for row in s.iter_rows(min_row=3):
                if tuple(c.value for c in row[:2]) in targets:
                    print('TARGET',row[0].row,json.dumps([c.value for c in row],ensure_ascii=False,default=str))
        if s.title=='Huong dan': print('GUIDE',json.dumps([[c.row,c.value] for row in s for c in row if c.value is not None],ensure_ascii=False,default=str))
    w.close()
