"""Independently verify staged/final XLSX; never write workbook content."""
from pathlib import Path
from collections import defaultdict,Counter
from datetime import datetime
from decimal import Decimal
import hashlib,json,sys
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
d=json.loads((HERE/'edit-plan.json').read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
final='--final' in sys.argv
filter=next((x.split('=')[1] for x in sys.argv if x.startswith('--kind=')),None)
protected={tuple(k) for k in d['protected']}
kept_roots={tuple(d['cases'][ident]['root']) for ident,c in d['decisions'].items() if c['mode']=='Giữ nguyên'}
result={'passed':False,'mode':'final' if final else 'staged','files':[],
        'software_imported':False,'signed':False,'issued':False}
all_new=[];all_errors=[]
all_source_roots=set();all_import_roots=set();already_matched=set()

def allowed(f,s,r,c):
    if s=='Huong dan':return c==1 and r in ({11,13,14,15,16} if f['kind']=='untreated' else {3,5,12,13,14,15,16} if f['kind']=='adjusted' else {2,3,4,5,11,12})
    if s=='Import dieu chinh':return (r==3 or r==7 or (f['kind']=='adjusted' and r==4)) and c==1 or r>=f['append_start'] and r<=f['rows']+9 and c<=24
    if f['kind']=='untreated':
        return s==f['main'] and c==14 and r in {u['main_row'] for u in f['updates']} or s=='Kiem tra SL 0' and (r==1 and c==1 or c==17 and 3<=r<=29)
    if f['kind']=='adjusted':
        return s==f['main'] and 28<=c<=55 and r in {u['main_row'] for u in f['updates']} or s=='Cho KT xac nhan' and (r==1 and c==1 or r==2 and c in (7,8) or 3<=r<=21 and c in (7,8)) or s=='42 - Can check 11' and (r==1 and c==1 or 4<=r<=14 and c==26)
    return s=='Ngoai le' and 241<=r<=246 or s=='Chi tiet nguon' and 328<=r<=339

for f in d['files']:
    if filter and filter!=f['kind']:continue
    # Stored absolute paths describe the original run; resolve this checkout locally.
    path=ROOT/'outputs'/f['name'] if final else HERE/'staged'/f['name']
    backup=HERE/'before'/f['name']
    assert sha(backup)==f['sha256'],('baseline changed',f['name'])
    before=openpyxl.load_workbook(backup,data_only=False)
    after=openpyxl.load_workbook(path,data_only=False)
    cached=openpyxl.load_workbook(path,data_only=True)
    assert before.sheetnames==after.sheetnames,('sheets',f['kind'])
    changed=0
    for s in before:
        a=after[s.title]
        assert s.freeze_panes==a.freeze_panes,('panes',f['kind'],s.title,s.freeze_panes,a.freeze_panes)
        assert set(s.tables)==set(a.tables),('tables',s.title)
        assert set(str(x) for x in s.merged_cells.ranges)==set(str(x) for x in a.merged_cells.ranges),('merged',s.title)
        for row in s.iter_rows():
            for c in row:
                av=a.cell(c.row,c.column).value
                if c.value!=av:
                    assert allowed(f,s.title,c.row,c.column),('unauthorized change',f['kind'],s.title,c.coordinate,c.value,av)
                    changed+=1
                if c.data_type=='f' and not allowed(f,s.title,c.row,c.column): assert a.cell(c.row,c.column).data_type=='f'
        for row in a:
            for c in row:
                if c.data_type=='e':all_errors.append((f['kind'],s.title,c.coordinate,c.value))
                if c.data_type=='f' and cached[s.title].cell(c.row,c.column).data_type=='e':all_errors.append((f['kind'],s.title,c.coordinate,cached[s.title].cell(c.row,c.column).value))
        for t in a.tables.values():assert t.autoFilter is not None,('filter',s.title,t.name)
    file_result={'name':f['name'],'sha256':sha(path),'old_cells_changed_within_scope':changed,
                 'path':path.relative_to(ROOT).as_posix(),
                 'backup_path':backup.relative_to(ROOT).as_posix(),'backup_sha256':f['sha256']}
    offset=2 if f['kind']=='exceptions' else 0
    all_source_roots.update(tuple(v[offset:offset+2]) for v in cached[f['main']].iter_rows(min_row=3,values_only=True) if v[offset])
    if f['kind']!='exceptions':
        s=cached['Import dieu chinh'];groups=defaultdict(list)
        import_rows=[]
        for rn,v in enumerate(s.iter_rows(min_row=10,max_col=24,values_only=True),10):
            if v[0] is None:
                assert all(z is None for z in v),('orphan',rn)
                continue
            groups[v[0]].append((rn,v));import_rows.append(v)
            assert tuple(v[9:11]) not in protected,('protected imported',v[9:11])
            assert tuple(v[9:11]) not in kept_roots,('kept case imported',v[9:11])
            assert isinstance(v[10],str) and len(v[10])==8,('identifier',rn,v[10])
        assert len(groups)==f['invoices'] and len(import_rows)==f['rows']
        assert set(groups)==set(range(1,f['invoices']+1)),('sequence',f['kind'])
        assert len({tuple(v[9:11]) for v in import_rows})==f['roots']
        current_roots={tuple(v[9:11]) for v in import_rows}
        assert not all_import_roots&current_roots,('root in both import workbooks',all_import_roots&current_roots)
        all_import_roots.update(current_roots)
        expected=f['new_rows']
        for i,er in enumerate(expected):
            actual=list(import_rows[f['old_rows']+i]);ex=[None if v=='' else v for v in er];ex[1]=datetime(2026,10,12)
            assert actual==ex,('new row',f['kind'],f['append_start']+i,actual,ex)
            source=d['details']['/'.join(actual[9:11])][0]
            assert actual[7]==source['method'] and actual[6] is None,('buyer field mapping',f['kind'],i,actual[6:8])
        arithmetic=[]
        for a in f['new_actions']:
            rs=groups[a['seq']];v=rs[0][1]
            assert tuple(v[9:11])==tuple(a['root'])
            assert [v[c] for c in (19,21,22)]==a['money']
            goods=sum(Decimal(str(v[18])) for _,v in rs if v[18] is not None)
            assert goods==Decimal(str(a['money'][0])),('detail sum',a['id'],goods,a['money'])
            for rn,v in rs[1:]:assert all(v[c] is None for c in (19,20,21,22)),('extra header',rn)
            if a['action']=='Đảo điều chỉnh cũ':assert [x+y for x,y in zip(a['source_money'],a['money'])]==[0,0,0]
            if a['difference']:
                arithmetic.append({'case':a['case'],'seq':a['seq'],'target':a['target'],'difference':a['difference']})
            if a['case'] in ('DC02','DC03','DC04','DC05','DC06'):
                assert len(rs)==1 and rs[0][1][16]==-1
                assert (rs[0][1][17]<0 and rs[0][1][18]>0) if a['action']=='Đảo điều chỉnh cũ' else (rs[0][1][17]>0 and rs[0][1][18]<0)
            if a['case']=='SL01' and a['action']=='Giảm hóa đơn gốc':assert rs[5][1][16:19]==(None,None,-46296)
            if a['action']=='Giảm số lượng gốc':assert all(v[16]<0 and v[17] is None and v[18] is None for _,v in rs)
            if a['case'] in ('SL03','SL04','SL05','SL06'):
                notes=[v for _,v in rs if 'không thu tiền' in v[14].lower()]
                assert notes and all(v[16:19]==(None,None,None) and v[23] is None for v in notes)
            if a['case'] in ('SL07','SL08'):assert len(rs)==1
            all_new.append(a)
        if f['kind']=='adjusted':
            already_matched.update(tuple(v[:2]) for v in cached['42 - Da khop 31'].iter_rows(min_row=4,values_only=True) if v[0])
            for u in f['updates']:
                r=u['main_row'];s=after[f['main']];sv=cached[f['main']]
                if u['mode']=='Giữ nguyên':assert all(s.cell(r,c).value is None for c in range(29,53)),('kept plan',u['case'])
                else:
                    assert not any('chờ kiểm tra' in str(s.cell(r,c).value or '').lower() for c in range(29,53))
                    for i,a in enumerate(u['actions']):assert [s.cell(r,32+i*6+j).value for j in range(3)]==a['money']
                    predicted=[sv.cell(r,c).value for c in (53,54,55)]
                    assert all(x==0 for x in predicted[:2]),('prediction',u['case'],predicted)
                    assert predicted[2] in (None,'') if u['case'] in ('TT01','TT02','CC07') else predicted[2]==0,('payment prediction',u['case'],predicted)
            assert Counter(cached['Cho KT xac nhan'].cell(x['row'],7).value for x in f['pending_rows'])=={'Đã bổ sung import':13,'Giữ nguyên cả chuỗi':6}
        file_result.update(invoices=len(groups),roots=f['roots'],rows=len(import_rows),new_invoices=len(f['new_actions']),new_rows=len(expected),arithmetic_differences=arithmetic)
    else:
        roots={tuple(v[2:4]) for v in cached['Ngoai le'].iter_rows(min_row=3,values_only=True) if v[2]}
        source_keys={tuple(v[2:4]) for v in cached['Chi tiet nguon'].iter_rows(min_row=2,values_only=True) if v[2]}
        assert len(roots)==244 and len(source_keys)==338
        assert roots=={tuple(x['root']) for x in d['held_chains']}
        assert source_keys==protected
        assert after['Ngoai le'].tables[next(iter(after['Ngoai le'].tables))].ref=='A2:X246'
        file_result.update(protected_chains=244,protected_invoices=338)
    result['files'].append(file_result)
    before.close();after.close();cached.close()
assert not all_errors,all_errors[:10]
for p,h in d['source_hashes'].items():assert sha(ROOT/p)==h,('source changed',p)
assert sha(ROOT/'scripts/data/check-review/kt-workbook-20261011/confirmed-plan-data.json')==d['plan_sha256'], 'approved plan changed'
if not filter:
    assert len(all_new)==48 and len({tuple(a['root']) for a in all_new})==27
    assert len({tuple(a['target']) for a in all_new})==48
    assert sum(a['action']=='Đảo điều chỉnh cũ' for a in all_new)==21
    assert sum(a['difference']!=0 for a in all_new)==8
    held_roots={tuple(c['root']) for c in d['held_chains']}|kept_roots
    assert len(all_source_roots)==9695 and len(all_import_roots)==9417 and len(held_roots)==247 and len(already_matched)==31
    assert not (all_import_roots&held_roots or all_import_roots&already_matched or held_roots&already_matched)
    assert all_source_roots==all_import_roots|held_roots|already_matched,('scope reconciliation',len(all_source_roots))
    result['scope_reconciliation']={'source_roots':9695,'prepared_roots':9417,'protected_or_kept_roots':247,'already_matched_roots':31,'unclassified_roots':0}
result.update(passed=True,new_invoices=len(all_new),new_roots=len({tuple(a['root']) for a in all_new}),formula_errors=0,source_hashes_unchanged=True)
out=HERE/('final-verification.json' if final else 'staged-verification.json')
if filter:out=HERE/f'{filter}-verification.json'
out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
