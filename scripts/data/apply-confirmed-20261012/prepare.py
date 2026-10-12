"""Prepare authorized import edits; only read Excel, retain recoverable originals."""
from pathlib import Path
from collections import defaultdict, Counter
from datetime import date, datetime
from decimal import Decimal
from zipfile import ZipFile
import hashlib, io, json, shutil, sys
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
H = ROOT/'scripts/data/check-review/kt-workbook-20261011'
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_invoice_reports import rows, amount, invoice_key, invoice_day
read=lambda n:json.loads((H/n).read_text(encoding='utf8'))
plan, review, audit, manifest = map(read, ['confirmed-plan-data.json','review-data.json','confirmed-decisions-audit.json','review-manifest.json'])
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
encode=lambda v: v.isoformat() if isinstance(v,(date,datetime)) else int(v) if isinstance(v,Decimal) and v==v.to_integral_value() else float(v) if isinstance(v,Decimal) else v
write=lambda p,v:p.write_text(json.dumps(v,ensure_ascii=False,default=encode),encoding='utf8')
assert read('accountant-decisions.json')['remaining_points']==[]
for name,digest in plan['workbook_hashes'].items(): assert sha(ROOT/'outputs'/name)==digest,('changed input',name)
for name,digest in audit['source_hashes'].items(): assert sha(ROOT/name)==digest,('changed source',name)
cases={c['id']:c for c in audit['review_cases']}
cases.update({c['id']:c for c in manifest['resolved']})
decisions={c['id']:c for c in plan['case_plans']+plan['known_plans']}
protected={tuple(k) for c in plan['held_chains'] for k in c['invoices']}
extra_holds=[c for c in plan['held_chains'] if not c['already_in_baseline_exceptions']]
wanted={tuple(a['target']) for a in plan['actions']}|{tuple(c['root']) for c in cases.values()}|{tuple(k) for c in extra_holds for k in c['invoices']}
headers={}; detail=defaultdict(list)
for path in sorted((ROOT/'data/Du lieu Meinvoice/Xuat_lai_2026-10-09').iterdir()):
    if path.suffix not in ('.xls','.xlsx','.zip') or '1791532943641' in path.name or any(t in path.name for t in ('dieu_chinh_xuat','thay_the_xuat')): continue
    if path.suffix=='.zip':
        with ZipFile(path) as z: books=[(n,z.read(n)) for n in z.namelist() if n.endswith(('.xlsx','.xls'))]
    else: books=[(path.name,path.read_bytes())]
    for name,blob in books:
        is_detail='chi_tiet' in name
        for rn,v in rows(io.BytesIO(blob)):
            if amount(v.get('A')) is None: continue
            k=invoice_key(v.get('B'),v.get('C'))
            if k not in wanted: continue
            evidence=path.relative_to(ROOT).as_posix()+(' / '+name if path.suffix=='.zip' else '')+f' | dòng {rn}'
            if is_detail:
                detail[k].append({'name':v.get('L'),'unit':v.get('M'),'qty':amount(v.get('N')),'price':amount(v.get('O')),'goods':amount(v.get('P')),'rate':v.get('Q'),'tax':amount(v.get('R')),'payment':amount(v.get('S')),'buyer':v.get('H'),'customer':v.get('E'),'address':v.get('F'),'taxid':v.get('G'),'method':v.get('J'),'email':None,'source':evidence})
            else:
                headers[k]={'date':invoice_day(v.get('D')),'money':[amount(v.get(c)) for c in ('P','Q','R')],'buyer':v.get('L'),'customer':v.get('H'),'address':v.get('I'),'taxid':v.get('J'),'method':v.get('M'),'status':v.get('U'),'source':evidence}
assert wanted<=set(headers),(wanted-set(headers))
xml={tuple(x['key']):x for x in review['xmls']}
neg=lambda x:None if x is None else -x
def fmt_day(v): return str(v)[:10].split('-')[::-1]
def day_text(v): return '/'.join(fmt_day(v))
def rate(v):
    if v in (None,''): return None
    try:return float(str(v).rstrip('%'))
    except ValueError:return v

files=[]
for kind,name,main in [('untreated','HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx','Chua dieu chinh'),('adjusted','HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','Dieu chinh'),('exceptions','NGOAI_LE_F0_2024_T6_2025.xlsx','Ngoai le')]:
    p=ROOT/'outputs'/name;backup=HERE/'before'/name;backup.parent.mkdir(parents=True,exist_ok=True)
    if backup.exists():assert sha(backup)==sha(p)
    else:shutil.copy2(p,backup)
    w=openpyxl.load_workbook(backup,data_only=False)
    cache=openpyxl.load_workbook(backup,data_only=True)
    metadata={'kind':kind,'name':name,'main':main,'path':str(p),'backup':str(backup),'sha256':sha(p),'sheets':[],'updates':[],'new_actions':[],'new_rows':[]}
    for s in w:
        metadata['sheets'].append({'name':s.title,'rows':s.max_row,'columns':s.max_column,'pane':s.freeze_panes,'tables':[{'name':t.name,'ref':t.ref,'style':t.tableStyleInfo.name if t.tableStyleInfo else None} for t in s.tables.values()]})
    if kind!='exceptions':
        s=w['Import dieu chinh'];old=[list(v) for v in cache[s.title].iter_rows(min_row=10,max_col=24,values_only=True) if v[0] is not None]
        seq=max(v[0] for v in old);last=9+len(old)
        assert not {tuple(v[9:11]) for v in old}&protected
        existing={tuple(v[9:11]) for v in old}
        main_rows={tuple(v[:2]):rn for rn,v in enumerate(cache[main].iter_rows(min_row=3,values_only=True),3)}
        metadata.update({'old_invoices':seq,'old_rows':len(old),'append_start':last+1,'main_rows':{'/'.join(k):n for k,n in main_rows.items()}})
        for a in plan['actions']:
            root,target=tuple(a['root']),tuple(a['target'])
            if root not in main_rows:continue
            assert root not in existing and target not in protected
            seq+=1;h=headers[root];ls=detail[root];assert ls,root
            buyer=ls[0].get('buyer') or h['buyer'] or h['customer']
            address=ls[0].get('address') or h['address']
            taxid=ls[0].get('taxid') or h['taxid']
            method=ls[0].get('method') or h['method']
            assert method,(root,ls[0],h)
            src_lines=detail[target]
            action=a['action']; reverse=action=='Đảo điều chỉnh cũ'
            if reverse and a['case'] not in ('HH01','HH04','DC02','DC03','DC04','DC05','DC06'):
                tr={rate(x['rate']) for x in xml[target]['tax_breakdown'] if x['rate'] not in (None,'')}
                if not tr:tr={rate(x['rate']) for x in src_lines if x['rate'] not in (None,'')}
                assert len(tr)==1,(target,tr)
                lines=[{'name':f'Đảo khoản điều chỉnh HĐ {target[0]}/{target[1]} ngày {day_text(headers[target]["date"])}','unit':None,'qty':None,'price':None,'goods':a['money'][0],'rate':next(iter(tr)),'source':a['source_xml']}]
            elif reverse:
                assert len(src_lines)==1
                l=dict(src_lines[0]);xl=xml[target]['lines'][0]
                l.update(qty=neg(Decimal(str(xl['qty']))),price=Decimal(str(xl['price'])),goods=neg(Decimal(str(xl['goods']))))
                if a['case'].startswith('DC'):assert l['qty']==-1 and l['price']<0 and l['goods']>0
                lines=[l]
            else:
                lines=[]
                for i,l in enumerate(ls,1):
                    if (a['case']=='SL02' and i==45) or (a['case'] in ('SL07','SL08') and i==2):
                        assert all(l[c] in (None,0) for c in ('qty','price','goods')),(a['case'],i,l)
                        continue
                    l=dict(l)
                    l['qty']=neg(l['qty']);l['goods']=neg(l['goods'])
                    if action=='Giảm số lượng gốc':
                        assert l['price'] is None and l['goods'] in (None,0)
                        l['goods']=None
                    lines.append(l)
                assert lines
                if a['case']=='SL01':assert lines[5]['qty'] is None and lines[5]['price'] is None and lines[5]['goods']==-46296
            trates={rate(l['rate']) for l in lines if l.get('rate') not in (None,'')}
            assert len(trates)==1,(root,trates)
            reason=(f'Đảo điều chỉnh {target[0]}/{target[1]} ngày {day_text(headers[target]["date"])}; tham chiếu gốc {root[0]}/{root[1]} ngày {day_text(h["date"])}' if reverse else f'Điều chỉnh giảm '+('số lượng hàng hóa' if action=='Giảm số lượng gốc' else 'hàng bán')+f' HĐ {root[0]}/{root[1]} ngày {day_text(h["date"])} theo phương án kế toán đã chốt')
            start=last+len(metadata['new_rows'])+1
            for i,l in enumerate(lines):
                first=i==0
                row=[seq,'2026-10-12',buyer,address,taxid,buyer,ls[0].get('email'),method,None,*root,day_text(h['date']),None,reason,l['name'],l.get('unit'),l['qty'],l['price'],l['goods'],a['money'][0] if first else None,next(iter(trates)) if first else None,a['money'][1] if first else None,a['money'][2] if first else None,1 if l['goods']==0 else None]
                metadata['new_rows'].append(row)
            metadata['new_actions'].append({**a,'seq':seq,'start':start,'end':start+len(lines)-1,'sources':[l.get('source') for l in lines]})
        metadata['invoices']=seq;metadata['rows']=len(old)+len(metadata['new_rows'])
        metadata['roots']=len(existing)+len({tuple(a['root']) for a in metadata['new_actions']})
        for ident,decision in decisions.items():
            root=tuple(cases[ident]['root'])
            if root in main_rows:
                metadata['updates'].append({'case':ident,'root':root,**decision,'main_row':main_rows[root],'actions':[a for a in metadata['new_actions'] if tuple(a['root'])==root]})
        metadata['pending_rows']=[{'row':rn,'root':list(v[:2]),'case':next(ident for ident,c in cases.items() if tuple(c['root'])==tuple(v[:2]))} for rn,v in enumerate(cache['Cho KT xac nhan'].iter_rows(min_row=3,values_only=True),3) if v[0]] if kind=='adjusted' else []
        metadata['check11_rows']=[{'row':rn,'root':list(v[:2]),'case':next(ident for ident,c in cases.items() if tuple(c['root'])==tuple(v[:2]))} for rn,v in enumerate(cache['42 - Can check 11'].iter_rows(min_row=4,values_only=True),4) if v[0]] if kind=='adjusted' else []
    else:
        metadata['exception_rows']=[];metadata['source_rows']=[]
        for hold in extra_holds:
            root=tuple(hold['root']);h=headers[root]
            invs=[root]+[tuple(k) for k in hold['invoices'] if tuple(k)!=root]
            assert len(invs)==2
            v=[h['buyer'] or h['customer'],h['taxid']]
            for k in invs:
                hh=headers[k];v.extend([*k,hh['date'],*hh['money']])
                metadata['source_rows'].append([*root,*k,hh['date'],hh['buyer'] or hh['customer'],hh['taxid'],'Hóa đơn gốc' if k==root else 'Điều chỉnh; giữ nguyên cả chuỗi theo xác nhận',hh['status'],*hh['money'],hh['source']])
            v.extend([None]*6)
            cumulative=[sum(headers[k]['money'][i] for k in invs) if all(headers[k]['money'][i] is not None for k in invs) else None for i in range(3)]
            v.extend(cumulative+['Giữ nguyên toàn bộ chuỗi theo xác nhận điểm 6; không lập điều chỉnh mới.'])
            assert len(v)==24
            metadata['exception_rows'].append(v)
    w.close();cache.close();files.append(metadata)

assert sum(len(f['new_actions']) for f in files)==48
assert sum(len(f['new_rows']) for f in files)>48
assert [len(f['new_actions']) for f in files[:2]]==[6,42]
result={'date':'2026-10-12','files':files,'decisions':decisions,'cases':cases,'held_chains':plan['held_chains'],'protected':list(protected),'plan_sha256':sha(H/'confirmed-plan-data.json'),'source_hashes':audit['source_hashes'],'headers':{'/'.join(k):v for k,v in headers.items()},'details':{'/'.join(k):v for k,v in detail.items()}}
write(HERE/'edit-plan.json',result)
print(json.dumps({'files':[{'kind':f['kind'],'new_actions':len(f['new_actions']),'new_rows':len(f['new_rows']),'invoices':f.get('invoices'),'rows':f.get('rows'),'roots':f.get('roots')} for f in files],'new_exception_chains':len(extra_holds)},ensure_ascii=False))
