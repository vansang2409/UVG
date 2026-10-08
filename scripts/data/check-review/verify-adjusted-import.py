"""Xác minh sheet import nhóm đã điều chỉnh và chỉ lưu khi bảo toàn các sheet cũ."""
import sys,json,hashlib,shutil
from pathlib import Path
from datetime import datetime,date
from decimal import Decimal
from collections import defaultdict
import openpyxl
sys.stdout.reconfigure(encoding='utf8');r=Path(r'D:\UVG');h=r/'scripts/data/check-review';d=json.loads((h/'adjusted-import-data.json').read_text(encoding='utf8'));m=json.loads((h/'adjusted-import-meta.json').read_text(encoding='utf8'))
a=openpyxl.load_workbook(h/'before-adjusted-import.xlsx');b=openpyxl.load_workbook(h/'with-adjusted-import.xlsx');bd=openpyxl.load_workbook(h/'with-adjusted-import.xlsx',data_only=True,read_only=True)
assert b.sheetnames==a.sheetnames+['Import dieu chinh','Cho KT xac nhan']
for name in a.sheetnames:
 s,t=a[name],b[name];assert s.freeze_panes==t.freeze_panes and list(s.merged_cells.ranges)==list(t.merged_cells.ranges)
 for k in s.tables:assert s.tables[k].ref==t.tables[k].ref
 for rr,ss in zip(s.iter_rows(),t.iter_rows(max_row=s.max_row,max_col=s.max_column)):
  for c,e in zip(rr,ss):assert c.value==e.value and c.number_format==e.number_format,(name,c.coordinate,'old changed')
check=lambda v:Decimal(str(v)) if v is not None else None
s=b['Import dieu chinh'];assert s.max_row==40786 and s.max_column==24
byroot={tuple(x['key']):x for x in d['ready']};byseq={};cumulative=defaultdict(lambda:[Decimal(0)]*3);seqrows=defaultdict(list);promo=0;blankbuyers=0;errors=[]
for rr,e,info in zip(s.iter_rows(min_row=10,max_row=40786,max_col=24,values_only=True),m['rows'],m['meta']):
 v=list(rr);e[1]=datetime(2026,10,8);assert v==e,(info['row'],'metadata mismatch')
 key=tuple(info['root']);seq=v[0];assert tuple(v[9:11])==key and key in byroot
 assert v[2]==v[5] and v[1]==datetime(2026,10,8)
 assert len(str(v[10]))==8 and isinstance(v[10],str)
 if seq in byseq:assert byseq[seq]==key
 byseq[seq]=key;seqrows[seq].append(v)
 if info['first']:
  amounts=[check(v[i]) for i in [19,21,22]];source=d['headers']['/'.join(info['source_key'])]['money'];assert amounts==[-check(x) for x in source]
  assert amounts[0]+amounts[1]==amounts[2]
  for j,amt in enumerate(amounts):cumulative[key][j]+=amt
 else:assert all(v[i] is None for i in [19,20,21,22])
 if info['kind']=='Đảo khoản điều chỉnh cũ':assert v[16] is None and v[17] is None and v[23] is None and 'đảo' in v[13]
 else:assert v[16]!=0 and v[16] is not None and v[23]==(1 if v[18]==0 else None)
 promo+=v[23]==1;blankbuyers+=v[5] is None
assert len(byseq)==1600 and set(byseq)==set(range(1,1601)) and set(cumulative)==set(byroot)
for seq,rs in seqrows.items():assert sum(check(v[18]) for v in rs)==check(rs[0][19])
for key,x in byroot.items():
 assert all(check(x['prior'][j])+cumulative[key][j]==0 for j in range(3)),(key,'not zero')
 assert len({seq for seq,k in byseq.items() if k==key})==len(x['old'])
 source=d['headers']['/'.join(key)]['money'];slots=[]
 for op in x['old'][1:]+x['old'][:1]:slots.extend([-v for v in d['headers']['/'.join(op['key'])]['money']])
 # F4–F7 bảng chính vẫn khớp những khoản sẽ chuẩn bị.
 for j,op in enumerate(x['old'][1:]+x['old'][:1]):assert x['plan'][j*6+3:j*6+6]==[-v for v in d['headers']['/'.join(op['key'])]['money']]
expected_pending={tuple(x['key']) for x in d['pending']};actual_pending={tuple(rr[:2]) for rr in b['Cho KT xac nhan'].iter_rows(min_row=3,max_row=21,max_col=12,values_only=True)}
assert actual_pending==expected_pending and not actual_pending&set(byroot)
assert not {tuple(x['key']) for x in d['chains'] if x['state']=='zero'}&set(byroot)
for ss in bd:
 for rr in ss.iter_rows():
  for cell in rr:
   if cell.data_type=='e':errors.append((ss.title,cell.coordinate,cell.value))
assert not errors,errors[:5]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();assert all(sha(Path(p))==v for p,v in d['hashes'].items())
p=r/'outputs/HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx';assert sha(p)==sha(h/'before-adjusted-import.xlsx'),'Report changed concurrently'
shutil.copy2(h/'with-adjusted-import.xlsx',p);assert sha(p)==sha(h/'with-adjusted-import.xlsx')
print(json.dumps({'PASS':True,'chains_to_zero':len(byroot),'invoices':len(byseq),'lines':40777,'promo_lines':promo,'blank_buyer_lines':blankbuyers,'pending':len(actual_pending),'already_zero_omitted':d['zero'],'source_and_old_sheets_preserved':True,'saved':True}))
