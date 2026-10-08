from pathlib import Path
import openpyxl,hashlib,shutil
h=Path(r'D:\UVG\scripts\data\check-review'); p=Path(r'D:\UVG\outputs\HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx')
a=openpyxl.load_workbook(h/'before-rename.xlsx',data_only=False); b=openpyxl.load_workbook(h/'renamed.xlsx',data_only=False)
assert a.sheetnames==b.sheetnames
changes=[]
for s in a:
 t=b[s.title]
 assert (s.max_row,s.max_column)==(t.max_row,t.max_column)
 for row in s:
  for v in row:
   u=t.cell(v.row,v.column)
   expected=v.value.replace('sửa ','') if s.title=='Dieu chinh' and v.row==2 and isinstance(v.value,str) else v.value
   assert u.value==expected,(s.title,v.coordinate,v.value,u.value)
   if v.value!=u.value: changes.append(v.coordinate)
   assert v.number_format==u.number_format,(s.title,v.coordinate,'number format')
 assert s.freeze_panes==t.freeze_panes
assert len(changes)==18,changes
s=b['Dieu chinh']; assert not any('sửa' in str(c.value) for c in s[2])
assert [x.name for x in s.tables['DieuChinhChuoi'].tableColumns]==[c.value for c in s[2]]
assert hashlib.sha256(p.read_bytes()).digest()==hashlib.sha256((h/'before-rename.xlsx').read_bytes()).digest()
shutil.copyfile(h/'renamed.xlsx',p)
assert hashlib.sha256(p.read_bytes()).digest()==hashlib.sha256((h/'renamed.xlsx').read_bytes()).digest()
print('PASS: 18 header cells renamed; other values/formulas/number formats/freeze panes unchanged; table headers consistent; saved original file')
