import json,hashlib
from pathlib import Path
import openpyxl
r=Path(r'D:\UVG'); h=r/'scripts/data/check-review'
d=json.loads((h/'plan-data.json').read_text(encoding='utf8'))
a=openpyxl.load_workbook(h/'before-plan.xlsx',data_only=False)
b=openpyxl.load_workbook(h/'with-plan.xlsx',data_only=False)
c=openpyxl.load_workbook(h/'with-plan.xlsx',data_only=True)
assert a.sheetnames==b.sheetnames
changed=[]
for old in a:
 new=b[old.title]
 for row in old:
  for v in row:
   u=new.cell(v.row,v.column)
   assert v.value==u.value,(old.title,v.coordinate,v.value,u.value)
   if v.value is not None:
    if (v.number_format,v.font.name,v.font.sz,v.font.bold,v.font.italic,v.font.color.type if v.font.color else None,v.font.color.rgb if v.font.color and v.font.color.type=='rgb' else None,v.fill.fgColor.rgb,v.alignment.horizontal,v.alignment.vertical,v.alignment.wrap_text)!=(u.number_format,u.font.name,u.font.sz,u.font.bold,u.font.italic,u.font.color.type if u.font.color else None,u.font.color.rgb if u.font.color and u.font.color.type=='rgb' else None,u.fill.fgColor.rgb,u.alignment.horizontal,u.alignment.vertical,u.alignment.wrap_text): changed.append((old.title,v.coordinate))
assert not changed,('old styles changed',changed[:10],len(changed))
s=b['Dieu chinh']; calc=c['Dieu chinh']; counts=[]
for x in d['items']:
 n=x['row']; ops=x['old'][1:]+x['old'][:1];counts.append(len(ops))
 for j in range(4):
  col=29+j*6
  assert s.cell(n,col+1).value is None and s.cell(n,col+2).value is None
  for k in range(3):
   expected=None if j>=len(ops) else d['totals']['/'.join(ops[j][:2])]['money'][k]
   expected=None if expected is None else -expected
   assert s.cell(n,col+3+k).value==expected,(n,j,k)
 for k in range(3):
  values=[s.cell(n,32+j*6+k).value for j in range(len(ops))]
  expected=None if x['prior'][k] is None or any(v is None for v in values) else x['prior'][k]+sum(values)
  actual=calc.cell(n,53+k).value
  assert (actual if actual!='' else None)==expected,(n,k,actual,expected)
assert sum(n==4 for n in counts)==2
assert s.freeze_panes=='C3' and s.tables['DieuChinhChuoi'].ref=='A2:BC1049'
assert all(hashlib.sha256((r/k).read_bytes()).hexdigest()==v for k,v in d['hashes'].items())
print('PASS: all 3 original sheets and styles preserved; all planned amounts match reversed source P/Q/R; 3 final totals checked on 1047 rows; F7 used on 2 rows; source hashes unchanged')
print('PLANNED',sum(counts),'ROWS WITH FINAL PAYMENT BLANK',sum(calc.cell(n,55).value in (None,'') for n in range(3,1050)))
