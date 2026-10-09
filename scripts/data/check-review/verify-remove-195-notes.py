from pathlib import Path
import openpyxl,json,hashlib,shutil,copy
h=Path(r'D:\UVG\scripts\data\check-review');p=Path(r'D:\UVG\outputs\HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx')
d=json.loads((h/'remove-195-notes-data.json').read_text(encoding='utf-8'));sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
assert sha(p)==d['sha256']
a=openpyxl.load_workbook(h/'before-remove-195-notes.xlsx');b=openpyxl.load_workbook(h/'without-195-notes.xlsx')
assert a.sheetnames==b.sheetnames
allowed={f'AB{x["row"]}':x for x in d['changes']};changed=[];count=0;ac={};bc={}
def style(c,cache):
 if c.style_id not in cache:cache[c.style_id]=tuple(copy.copy(getattr(c,k)) for k in ['font','fill','border','alignment','protection'])
 return cache[c.style_id]
for name in a.sheetnames:
 s,t=a[name],b[name]
 assert s.max_row==t.max_row and s.max_column==t.max_column
 assert s.freeze_panes==t.freeze_panes and set(map(str,s.merged_cells.ranges))==set(map(str,t.merged_cells.ranges))
 assert {k:s.tables[k].ref for k in s.tables}=={k:t.tables[k].ref for k in t.tables}
 assert str(s.data_validations)==str(t.data_validations)
 for rr in s.iter_rows():
  for c in rr:
   e=t[c.coordinate];count+=1
   if c.value!=e.value:
    assert name=='Dieu chinh' and c.coordinate in allowed,(name,c.coordinate)
    assert e.value==allowed[c.coordinate]['note'];changed.append(c.coordinate)
   assert c.number_format==e.number_format and style(c,ac)==style(e,bc),(name,c.coordinate,'format')
 for n,dim in s.row_dimensions.items():
  target=next((x for x in d['changes'] if x['row']==n),None) if name=='Dieu chinh' else None
  assert t.row_dimensions[n].height==(target['height'] if target else dim.height)
assert set(changed)==set(allowed)
shutil.copy2(h/'without-195-notes.xlsx',p)
result={'passed':True,'notes_removed':195,'old_notes_retained':True,'plan_and_import_unchanged':True,'cells_checked':count,'sha256':sha(p)}
(h/'remove-195-notes-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result))
