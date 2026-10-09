from pathlib import Path
import openpyxl,json,datetime,itertools,shutil,hashlib
ROOT=Path(__file__).resolve().parents[3]
original=ROOT/'outputs/HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx'
staged=ROOT/'scripts/data/check-review/date-12885-updated.xlsx'
a=openpyxl.load_workbook(original,read_only=True,data_only=False);b=openpyxl.load_workbook(staged,read_only=True,data_only=False)
assert a.sheetnames==b.sheetnames
allowed={('Dieu chinh','I903'),('Dieu chinh','AB903'),('Chi tiet nguon','M1822')}
differences=[];checked=0
for sn in a.sheetnames:
 for oldrow,newrow in itertools.zip_longest(a[sn].iter_rows(),b[sn].iter_rows(),fillvalue=[]):
  for old,new in itertools.zip_longest(oldrow,newrow,fillvalue=None):
   ov=old.value if old else None;nv=new.value if new else None
   if ov is None and nv is None:continue
   checked+=1;coord=old.coordinate if old else new.coordinate
   if ov!=nv:differences.append({'sheet':sn,'cell':coord,'old':str(ov),'new':str(nv)})
assert all((x['sheet'],x['cell']) in allowed for x in differences),str(differences[:15])
assert b['Dieu chinh']['I903'].value==datetime.datetime(2025,8,18)
assert '18/08/2025' in b['Dieu chinh']['AB903'].value
assert 'người dùng chốt ngày đúng 18/08/2025' in b['Chi tiet nguon']['M1822'].value
assert b['Chi tiet nguon']['E1822'].value==datetime.datetime(2025,8,18)
a.close();b.close()
result={'date':'2026-10-09','checked_nonempty_cells':checked,'differences':differences,'only_authorized_cells_changed':True,'source_before_sha256':hashlib.sha256(original.read_bytes()).hexdigest(),'staged_sha256':hashlib.sha256(staged.read_bytes()).hexdigest()}
(ROOT/'scripts/data/check-review/date-12885-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(staged,original)
print(json.dumps({'checked_cells':checked,'changed_cells':[(x['sheet'],x['cell']) for x in differences],'replaced_original_after_verification':True},ensure_ascii=False))