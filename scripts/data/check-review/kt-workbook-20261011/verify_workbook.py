"""Independently read the exported accountant workbook and check source preservation."""
from collections import Counter
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import json

import openpyxl
from openpyxl.utils.cell import column_index_from_string

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
manifest = json.loads((HERE / 'review-manifest.json').read_text(encoding='utf-8'))
data = json.loads((HERE / 'review-data.json').read_text(encoding='utf-8'))
path = Path(manifest['output'])
formulas = openpyxl.load_workbook(path, data_only=False)
cached = openpyxl.load_workbook(path, data_only=True)
expected_sheets = ['Tong hop', *manifest['controls'], 'Da co can cu', 'Hoa don lien quan']
assert formulas.sheetnames == expected_sheets
assert len(manifest['cases']) == 29 and len(manifest['resolved']) == 7
assert len(manifest['related']) == 67
assert len({tuple(r['key']) for r in manifest['related']}) == 67
counts = Counter(c['group'] for c in manifest['cases'])
assert list(counts.values()) == [8, 8, 2, 5, 6]

def iso(value):
    assert isinstance(value, datetime), (value, type(value))
    return value.date().isoformat()

def key_text(key):
    return '/'.join(key) if key else None

for name, control in manifest['controls'].items():
    ws = formulas[name]
    case_rows = [c for c in manifest['cases'] if c['group'] == name]
    assert ws.freeze_panes == 'C8'
    assert len(ws.tables) == 1
    table = next(iter(ws.tables.values()))
    assert table.ref == f"A7:{control['status']}{control['end']}"
    assert table.autoFilter is not None
    assert table.tableStyleInfo.name == 'TableStyleLight1'
    for row, case in enumerate(case_rows, 8):
        assert ws.cell(row, 1).value == case['id']
        assert ws.cell(row, 2).value == key_text(case['root'])
        assert iso(ws.cell(row, 3).value) == case['date'][:10]
        assert ws.cell(row, 4).value == key_text(case['focus'])
        if case['focusDate']:
            assert iso(ws.cell(row, 5).value) == case['focusDate'][:10]
        else:
            assert ws.cell(row, 5).value is None
        start = column_index_from_string(control['first'])
        for col in range(start, start + 6):
            assert ws.cell(row, col).value is None
        status = ws[f"{control['status']}{row}"]
        assert status.data_type == 'f'
        assert cached[name][status.coordinate].value == 'Chưa phản hồi'
        assert ws.cell(row, start - 2).value  # Observed issue.
        assert ws.cell(row, start - 1).value  # Specific accountant question.

ws = formulas['Da co can cu']
for row, case in enumerate(manifest['resolved'], 8):
    assert [ws.cell(row, c).value for c in (1, 2, 4)] == [case['id'], key_text(case['root']), key_text(case['focus'])]
    assert isinstance(ws.cell(row, 3).value, datetime)
    assert isinstance(ws.cell(row, 5).value, datetime)

ws = formulas['Hoa don lien quan']
active_ids = {c['id'] for c in manifest['cases']}
active_evidence = 0
zero_blank = []
for row, record in enumerate(manifest['related'], 8):
    assert ws.cell(row, 1).value == record['id']
    assert ws.cell(row, 4).value == record['key'][0]
    invoice = ws.cell(row, 5)
    assert invoice.value == record['key'][1]
    assert invoice.data_type == 's' and invoice.number_format == '@'
    assert len(invoice.value) == 8 and invoice.value.isdigit()
    assert iso(ws.cell(row, 6).value) == record['date'][:10]
    assert [ws.cell(row, c).value for c in range(7, 10)] == record['money']
    assert [ws.cell(row, c).value for c in range(10, 13)] == record['xmlMoney']
    assert ws.cell(row, 14).value  # Traceable source filename/row.
    if record['money'][2] is None and record['xmlMoney'][2] == 0:
        zero_blank.append(record['key'])
    active_evidence += record['id'] in active_ids
assert active_evidence == 53
assert {r[1] for r in zero_blank} == {'00004203', '00004204', '00006108', '00009186'}

xml_map = {tuple(r['key']): r for r in data['xmls']}
for row, case in enumerate([c for c in manifest['cases'] if c['group'] == 'Lech tien'], 8):
    amounts = xml_map[tuple(case['focus'])]['money']
    assert [cached['Lech tien'].cell(row, c).value for c in range(6, 9)] == amounts
    assert cached['Lech tien'].cell(row, 9).value == amounts[0] + amounts[1] - amounts[2]
assert [cached['Tong hop'].cell(14, c).value for c in range(3, 7)] == [29, 29, 0, 0]
errors = [(ws.title, cell.coordinate, cell.value) for ws in cached for row in ws for cell in row if cell.data_type == 'e']
assert not errors, errors
for name, expected_hash in manifest['workbookHashes'].items():
    assert sha256((ROOT / 'outputs' / name).read_bytes()).hexdigest() == expected_hash
assert sum(len(ws.tables) for ws in formulas) == 8

result = {'passed': True, 'output': str(path), 'sheets': len(formulas.sheetnames),
          'active_cases': 29, 'reference_cases': 7, 'distinct_related_invoices': 67,
          'active_related_invoices': active_evidence, 'leading_zero_invoice_numbers_preserved': 67,
          'all_dates_checked': True, 'blank_zero_differences_preserved': zero_blank,
          'response_fields_blank': True, 'cached_formula_errors': 0,
          'summary_cached_values': [29, 29, 0, 0], 'source_workbooks_unchanged': True,
          'output_sha256': sha256(path.read_bytes()).hexdigest(), 'output_bytes': path.stat().st_size}
(HERE / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
formulas.close()
cached.close()
