"""Read source evidence for the accountant workbook; never write source Excel."""
from pathlib import Path
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from zipfile import ZipFile
import hashlib
import io
import json
import sys
import xml.etree.ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / 'scripts'))
from prepare_invoice_reports import rows, amount, invoice_day, invoice_key
from audit_pending_review import first, value

DEST = Path(__file__).resolve().parent
audit = json.loads((ROOT / 'scripts/data/check-review/pending-current-audit-20261011.json').read_text(encoding='utf-8'))
assert audit['passed'] and audit['summary']['pending_adjustments_with_xml'] == 20
hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT / 'outputs').glob('*.xlsx')}
path = ROOT / 'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx'
w = openpyxl.load_workbook(path, read_only=True, data_only=True)
extra = {}
for rn, v in enumerate(w['Kiem tra SL 0'].iter_rows(min_row=3, values_only=True), 3):
    if not v[0]:
        continue
    key = tuple(v[:2])
    rec = extra.setdefault(key, {'key': key, 'date': invoice_day(v[2]), 'lines': [], 'zero_lines': []})
    line = {'name': v[7], 'unit': v[8], 'qty': v[9], 'price': v[10], 'goods': v[11],
            'saved_sheet_row': rn, 'source': v[17]}
    rec['lines'].append(line)
    if v[9] in (None, 0):
        rec['zero_lines'].append(line)
for rn, v in enumerate(w['Chua dieu chinh'].iter_rows(min_row=3, values_only=True), 3):
    if tuple(v[:2]) in extra:
        extra[tuple(v[:2])]['saved_header'] = {'date': invoice_day(v[2]), 'money': list(v[3:6]),
                                             'customer': v[6], 'row': rn}
w.close()
assert len(extra) == 6 and sum(len(r['lines']) for r in extra.values()) == 27

current_detail = defaultdict(list)
for p in sorted((ROOT / 'data/Du lieu Meinvoice/Xuat_lai_2026-10-09').iterdir()):
    if p.suffix not in ('.xlsx', '.xls', '.zip') or '1791532943641' in p.name:
        continue
    if 'dieu_chinh_xuat' in p.name or 'thay_the_xuat' in p.name:
        continue
    if p.suffix == '.zip':
        with ZipFile(p) as z:
            books = [(n, z.read(n)) for n in z.namelist() if n.endswith(('.xlsx', '.xls'))]
    else:
        books = [(p.name, p.read_bytes())]
    for name, data in books:
        kind = 'detail' if 'chi_tiet' in name else 'total'
        for rn, v in rows(io.BytesIO(data)):
            if amount(v.get('A')) is None:
                continue
            key = invoice_key(v.get('B'), v.get('C'))
            if key not in extra:
                continue
            if kind == 'total':
                h = {'date': invoice_day(v.get('D')), 'money': [amount(v.get(c)) for c in ('P', 'Q', 'R')],
                     'evidence': {'file': name, 'row': rn}}
                if 'header' in extra[key]:
                    assert extra[key]['header']['date'] == h['date'] and extra[key]['header']['money'] == h['money']
                extra[key]['header'] = h
            else:
                current_detail[key].append({'name': v.get('L'), 'unit': v.get('M'), 'qty': amount(v.get('N')),
                    'price': amount(v.get('O')), 'goods': amount(v.get('P')), 'tax': amount(v.get('R')),
                    'payment': amount(v.get('S')), 'evidence': {'file': name, 'row': rn}})
for key, r in extra.items():
    assert r['header']['date'] == r['saved_header']['date']
    assert r['header']['money'] == [amount(v) for v in r['saved_header']['money']]
    r['current_details'] = current_detail[key]
    assert len(r['current_details']) == len(r['lines'])

targets = set(extra)
for c in audit['pending']:
    targets.update(tuple(i['key']) for i in c['invoices'])
for c in audit['review_11']:
    targets.update((tuple(c['root']), tuple(c['adjustment'])))
assert len(targets) == 67
xmls = {}
for p in sorted((ROOT / 'data/Bo sung/meInvoice').rglob('*.xml')):
    doc = ET.fromstring(p.read_bytes())
    hd = first(doc, 'DLHDon'); tt = first(hd, 'TTChung')
    key = invoice_key((value(tt, 'KHMSHDon') or '') + (value(tt, 'KHHDon') or ''), value(tt, 'SHDon'))
    if key not in targets:
        continue
    pay, related = first(hd, 'TToan'), first(tt, 'TTHDLQuan')
    lines = [{'name': value(l, 'THHDVu'), 'qty': amount(value(l, 'SLuong')),
              'price': amount(value(l, 'DGia')), 'goods': amount(value(l, 'ThTien')),
              'rate': value(l, 'TSuat'), 'nature': value(l, 'TChat')}
             for l in hd.iter() if l.tag.rsplit('}', 1)[-1] == 'HHDVu']
    tax = [{'rate': value(l, 'TSuat'), 'goods': amount(value(l, 'ThTien')), 'tax': amount(value(l, 'TThue'))}
           for l in pay.iter() if l.tag.rsplit('}', 1)[-1] == 'LTSuat']
    rec = {'key': key, 'date': invoice_day(value(tt, 'NLap')),
           'money': [amount(value(pay, f)) for f in ('TgTCThue', 'TgTThue', 'TgTTTBSo')],
           'related_note': value(related, 'GChu'), 'lines': lines, 'tax_breakdown': tax, 'file': p.name}
    if key in xmls:
        assert xmls[key]['date'] == rec['date'] and xmls[key]['money'] == rec['money'] and xmls[key]['lines'] == rec['lines']
    xmls[key] = rec

def encode(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    raise TypeError(type(v))

result = {'audit': audit, 'extra': list(extra.values()), 'xmls': list(xmls.values()), 'workbook_hashes': hashes}
out = DEST / 'review-data.json'
out.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=encode), encoding='utf-8')
assert hashes == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT / 'outputs').glob('*.xlsx')}
print(json.dumps({'extra_invoices': len(extra), 'related_invoices': len(targets), 'xml_available': len(xmls),
                  'source_workbooks_unchanged': True, 'output': str(out)}, ensure_ascii=False))
