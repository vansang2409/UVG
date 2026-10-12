"""Read and preserve the returned accountant responses, without applying them."""
from collections import Counter
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import json
import shutil

import openpyxl
from openpyxl.utils.cell import column_index_from_string

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
SOURCE = Path(r'C:\Users\thuong ngo\OneDrive\Desktop\JOB UVG\JOB HOA DON\UVG_CAC_VAN_DE_CAN_KE_TOAN_XAC_NHAN_20261011.xlsx')
manifest = json.loads((HERE / 'review-manifest.json').read_text(encoding='utf-8'))
digest = sha256(SOURCE.read_bytes()).hexdigest()
wb = openpyxl.load_workbook(SOURCE, data_only=False)
values = openpyxl.load_workbook(SOURCE, data_only=True)
base = openpyxl.load_workbook(manifest['output'], data_only=False)
assert wb.sheetnames == base.sheetnames

def encode(value):
    if isinstance(value, datetime):
        return value.isoformat()
    raise TypeError(type(value))

responses = []
allowed_cells = set()
for name, control in manifest['controls'].items():
    ws = values[name]
    first = column_index_from_string(control['first'])
    expected = [c for c in manifest['cases'] if c['group'] == name]
    for row, case in enumerate(expected, 8):
        assert ws.cell(row, 1).value == case['id']
        assert ws.cell(row, 2).value == '/'.join(case['root'])
        assert ws.cell(row, 4).value == ('/'.join(case['focus']) if case['focus'] else None)
        response = {ws.cell(7, col).value: ws.cell(row, col).value for col in range(first, first + 6)}
        assert all(ws.cell(row, col).value is not None for col in range(first, first + 3))
        responses.append({'id': case['id'], 'sheet': name, 'row': row,
                          'root': case['root'], 'root_date': ws.cell(row, 3).value,
                          'adjustment': case['focus'], 'adjustment_date': ws.cell(row, 5).value,
                          'response': response, 'cached_status': ws[control['status'] + str(row)].value})
        allowed_cells.update((name, ws.cell(row, col).coordinate) for col in range(first, first + 6))

known = []
for row, case in enumerate(manifest['resolved'], 8):
    ws = values['Da co can cu']
    assert ws.cell(row, 1).value == case['id']
    assert ws.cell(row, 8).value
    known.append({'id': case['id'], 'root': case['root'], 'adjustment': case['focus'],
                  'response': ws.cell(row, 8).value})
    allowed_cells.add(('Da co can cu', f'H{row}'))
allowed_cells.add(('Da co can cu', 'H7'))

changes = []
other_changes = []
for ws in wb:
    for row in ws:
        for cell in row:
            old = base[ws.title][cell.coordinate].value
            if old == cell.value:
                continue
            item = {'sheet': ws.title, 'cell': cell.coordinate, 'old': old, 'new': cell.value}
            changes.append(item)
            if (ws.title, cell.coordinate) not in allowed_cells:
                other_changes.append(item)
assert len(other_changes) == 8
assert all(isinstance(item['old'], str) and isinstance(item['new'], str)
           and item['old'].strip() == item['new'].strip() for item in other_changes)

money = []
for row in range(8, 16):
    ws = values['Lech tien']
    original = [ws.cell(row, col).value for col in (6, 7, 8)]
    corrected = [int(line.split(':', 1)[1].strip().replace(',', '')) for line in ws.cell(row, 14).value.splitlines()]
    assert corrected[0] + corrected[1] == corrected[2]
    money.append({'id': ws.cell(row, 1).value, 'adjustment': ws.cell(row, 4).value,
                  'source_money': original, 'accountant_corrected_money': corrected,
                  'reverse_source': [-n for n in original],
                  'residual_if_reverse_corrected': [a - b for a, b in zip(original, corrected)]})

issues = [
    {'id': 'R01', 'cases': [r['id'] for r in money],
     'finding': 'Eight corrected amounts balance, but reversing corrected values leaves a residual against the issued source. Clarify whether to correct first or reverse issued values, and the exact target invoice.'},
    {'id': 'R02', 'cases': ['DC03', 'DC04', 'DC05', 'DC06'],
     'finding': 'Repeated invoice 13038 in I10:I13 is outside the row pairs. I11:I13 also repeats root 3113. Expected pairs are 3113/13032, 4081/13033, 5555/13034, 5918/13035.'},
    {'id': 'R03', 'cases': ['DC02', 'DC03', 'DC04', 'DC05', 'DC06'],
     'finding': 'The old adjustment has positive quantity and negative unit price. A new -1 quantity at a positive price reduces money again. Specify the positive money reversal for the old adjustment separately from the negative money adjustment for the root.'},
    {'id': 'R04', 'cases': ['SL01'],
     'finding': 'Root 1C24TUV/555 dated 2024-10-31: response retains an explanatory line and calls for negative quantity, but supplies only -46296 goods amount. Quantity and treatment of explanation nature remain unspecified.'},
    {'id': 'R05', 'cases': ['SL08'],
     'finding': 'Root 1C25TUV/2797 dated 2025-06-01: source line is KM003, but response names KM002. Confirm that row 2 of invoice 2797, KM003, is intended.'},
    {'id': 'R06', 'cases': ['TT01', 'TT02'],
     'finding': 'Roots 6/16 dated 2025-01-07/2025-01-14: response supplies negative quantities, but does not settle blank root payment or specify the treatment of positive money adjustments 655/657.'},
    {'id': 'R07', 'cases': ['CC01', 'CC02', 'CC03', 'CC04', 'CC05', 'CC06'],
     'finding': 'Keep unchanged is stated at case level. Before preparing import, distinguish excluding both invoices of the pair from retaining the old adjustment while still adjusting the root. CC01 also differs from the earlier preparation decision for 3/508.'}
]

DEST = ROOT / 'data/Bo sung/Ke toan/2026-10-11' / digest[:12]
DEST.mkdir(parents=True, exist_ok=True)
saved = DEST / SOURCE.name
if saved.exists():
    assert sha256(saved.read_bytes()).hexdigest() == digest
else:
    shutil.copy2(SOURCE, saved)
assert sha256(saved.read_bytes()).hexdigest() == digest
provenance = {'source': str(SOURCE), 'preserved_file': saved.name, 'sha256': digest,
              'bytes': SOURCE.stat().st_size, 'received_date': '2026-10-11',
              'purpose': 'Accountant reply workbook. Responses are evidence; no import or source edits applied.'}
(DEST / 'PROVENANCE.json').write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding='utf-8')
result = {'source': str(SOURCE), 'source_sha256': digest, 'preserved_path': str(saved),
          'active_responses': len(responses), 'reference_responses': len(known),
          'cached_status_counts': dict(Counter(r['cached_status'] for r in responses)),
          'missing_basis_count': sum(r['response']['Căn cứ/chứng từ (KT điền)'] is None for r in responses),
          'all_case_identifiers_unchanged': True, 'original_evidence_unchanged_except_whitespace': True,
          'responses': responses, 'reference_responses_detail': known,
          'money_checks': money, 'issues': issues, 'changes': changes, 'applied_to_import': False}
(HERE / 'accountant-reply-review.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=encode), encoding='utf-8')
for name, expected_hash in manifest['workbookHashes'].items():
    assert sha256((ROOT / 'outputs' / name).read_bytes()).hexdigest() == expected_hash
assert sha256(SOURCE.read_bytes()).hexdigest() == digest
print(json.dumps({'source_sha256': digest, 'active_responses': len(responses), 'reference_responses': len(known),
                  'cached_status_counts': result['cached_status_counts'], 'money_arithmetic_passed': 8,
                  'issue_groups': len(issues), 'saved_source': str(saved), 'source_and_import_workbooks_unchanged': True}, ensure_ascii=False))
wb.close()
values.close()
base.close()
