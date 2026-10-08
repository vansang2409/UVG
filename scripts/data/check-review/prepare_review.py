import sys,json,re,hashlib
from pathlib import Path
from collections import Counter
from decimal import Decimal
import openpyxl
ROOT=Path(r'D:\UVG')
sys.path.insert(0,str(ROOT/'scripts'))
from prepare_invoice_reports import rows
items=[]; sources=[]
for name,group,offset in [('HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','Đã điều chỉnh ngoài ngoại lệ',0),('NGOAI_LE_F0_2024_T6_2025.xlsx','Ngoại lệ giữ nguyên',2)]:
 p=ROOT/'outputs'/name; sources.append(p); w=openpyxl.load_workbook(p,data_only=True); s=w.worksheets[0]; detail=list(w['Chi tiet nguon'].iter_rows(min_row=2,values_only=True))
 for row in s.iter_rows(min_row=3):
  v=[c.value for c in row]; flag=v[-1]
  if not flag: continue
  root=(v[offset],v[offset+1]); rel=[r for r in detail if (r[0],r[1])==root]
  cat='Lệch số học' if 'lệch thanh toán' in flag else 'Bảng tổng trống, chi tiết có 0' if 'bảng tổng trống' in flag else 'Quy ước riêng hóa đơn 574' if '574' in flag else 'Ngày giữa nguồn khác nhau' if 'ngày giữa nguồn' in flag else 'Chuỗi thay thế trong ngoại lệ'
  m=re.search(r'(1C\d\d\w+)/([0-9]{8})',flag); target=(m[1],m[2]) if m else root; inv=next(r for r in rel if (r[2],r[3])==target)
  items.append(dict(root='/'.join(root),group=group,target='/'.join(target),category=cat,day=inv[4].isoformat(),goods=inv[9],tax=inv[10],payment=inv[11],cumulative=v[-2],original_note=flag,chain='; '.join('/'.join((r[2],r[3])) for r in rel),provenance=f'{name} | {s.title} | dòng {row[0].row}; '+inv[12]))
items.append(dict(root='1C24TUV/00000003',group='Đã điều chỉnh ngoài ngoại lệ',target='1C25TUV/00000508',category='Bảng tổng khác bảng chi tiết',day='2025-04-26T00:00:00',goods=-356085,tax=-28487,payment=-384572,cumulative=11856288,original_note='Bất nhất bảng tổng và bảng chi tiết vừa phát hiện.',chain='1C24TUV/00000003; 1C25TUV/00000508',provenance='HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx | Dieu chinh | dòng 3; bảng kê tổng phần 1 | dòng 1299; bảng kê chi tiết 2025 | dòng 16434'))
zero_keys={tuple(x['target'].split('/')) for x in items if x['category']=='Bảng tổng trống, chi tiết có 0'}
checks={k:[] for k in zero_keys|{('1C25TUV','00000508')}}
for name in ['Bang_ke_chi_tiet_HD_da_su_dung_2024.xls','Bang_ke_chi_tiet_HD_da_su_dung_2025.xls']:
 p=ROOT/'data'/'Du lieu Meinvoice'/name; sources.append(p)
 for n,d in rows(p):
  key=(d.get('D'),d.get('B'))
  if key in checks: checks[key].append((n,d))
for x in items:
 key=tuple(x['target'].split('/')); cat=x['category']; x['difference']=None; x['status']='Chờ kiểm tra'
 if cat=='Lệch số học':
  x['difference']=x['payment']-(x['goods']+x['tax']); x['issue']=f"Thanh toán − (tiền hàng + thuế) = {x['difference']:+,} đồng."
  x['action']='Đối chiếu bản hóa đơn/XML để chốt tiền hàng, thuế và thanh toán đúng. Không tự coi lệch 1 đồng là do làm tròn.'
 elif cat=='Bảng tổng trống, chi tiết có 0':
  e=checks[key]; assert e and all(Decimal(d['U'])==0 for n,d in e)
  x['issue']=f'Bảng tổng trống thanh toán; đã kiểm lại {len(e)} dòng chi tiết, cột U đều ghi số 0. Cộng dồn thanh toán trong report vẫn trống.'
  x['action']='Đối chiếu bản hóa đơn/XML; xác nhận dùng thanh toán 0 từ chi tiết cho hóa đơn này, rồi mới cập nhật cộng dồn. Số 0 này không chứng minh cả chuỗi bằng 0.'
  x['status']='Có căn cứ 0 ở chi tiết, chờ chốt'; x['provenance']+='; bảng kê chi tiết '+('2024' if key[0]=='1C24TUV' else '2025')+' | dòng '+', '.join(str(n) for n,d in e)+' | cột U'
 elif cat=='Quy ước riêng hóa đơn 574':
  x['issue']='Thanh toán gốc trống; tạm tính 0 theo xác nhận trước. Cộng dồn chuỗi 1.839.000 đồng.'
  x['action']='Đối chiếu bản hóa đơn/XML để xác nhận thanh toán gốc. Quy ước riêng đã có, không áp dụng sang hóa đơn khác.'; x['status']='Đã chốt tạm, chờ căn cứ gốc'
 elif cat=='Ngày giữa nguồn khác nhau':
  x['issue']='12885: chi tiết và ảnh danh sách ghi 05/07/2025; bảng tổng/liên kết ghi 18/08/2025. Report dùng 05/07/2025 theo xác nhận. Chuỗi cộng dồn 0.'
  x['action']='Lấy XML/bản hóa đơn 12885 để đối chiếu ngày độc lập; giữ cả hai ngày nguồn và xác nhận đã có.'; x['status']='Đã xác nhận ngày, chờ XML'
 elif cat=='Chuỗi thay thế trong ngoại lệ':
  x['issue']='Chuỗi 60 → 61 → 1C25TUV/00002605; không cộng các bản bị thay thế. Cộng dồn thanh toán 2.739.000 đồng.'
  x['action']='Kiểm tra tham chiếu và bản hiện hành bằng XML. Thuộc ngoại lệ giữ nguyên, không đưa vào lô điều chỉnh ngoài ngoại lệ.'
 else:
  e=checks[key]; assert len(e)==1; n,d=e[0]; assert Decimal(d['P'])==0 and Decimal(d['T'])==-28487 and Decimal(d['U'])==-28487
  x['issue']='508: bảng tổng tiền hàng -356.085, thuế -28.487, thanh toán -384.572; chi tiết tiền hàng 0, thuế -28.487, thanh toán -28.487.'
  x['action']='Lấy XML/bản hóa đơn 508 để chốt số đúng trước khi đảo dấu. Chênh thanh toán giữa hai bảng 356.085 đồng; chưa kết luận lỗi nằm ở hóa đơn hay bảng kê.'
 x['result']=''; x['evidence']=''
order={'Bảng tổng khác bảng chi tiết':0,'Lệch số học':1,'Bảng tổng trống, chi tiết có 0':2,'Ngày giữa nguồn khác nhau':3,'Quy ước riêng hóa đơn 574':4,'Chuỗi thay thế trong ngoại lệ':5}
items.sort(key=lambda x:(order[x['category']],-abs(x['difference'] or 0),x['root']))
assert len(items)==19 and len({x['root'] for x in items})==19
assert sum(abs(x['difference']) for x in items if x['category']=='Lệch số học')==963901
payload={'items':items,'counts':dict(Counter(x['category'] for x in items)),'hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}}
(ROOT/'scripts/data/check-review/review.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'cases':len(items),'counts':payload['counts']},ensure_ascii=False))
