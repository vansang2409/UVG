from pathlib import Path
from datetime import datetime
import openpyxl,json,hashlib,shutil,copy
h=Path(r'D:\UVG\scripts\data\check-review');p=Path(r'D:\UVG\outputs\HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx')
d=json.loads((h/'split-42-data.json').read_text(encoding='utf-8'));sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
assert sha(p)==d['sha256']
a=openpyxl.load_workbook(h/'before-split-42.xlsx');b=openpyxl.load_workbook(h/'with-split-42.xlsx');cache=openpyxl.load_workbook(h/'with-split-42.xlsx',read_only=True,data_only=True)
assert b.sheetnames==a.sheetnames+list(d['groups'])
ac={};bc={};count=0
def style(c,cache):
 if c.style_id not in cache:cache[c.style_id]=tuple(copy.copy(getattr(c,k)) for k in ['font','fill','border','alignment','protection'])
 return cache[c.style_id]
for name in a.sheetnames:
 s,t=a[name],b[name]
 assert s.max_row==t.max_row and s.max_column==t.max_column
 assert s.freeze_panes==t.freeze_panes and set(map(str,s.merged_cells.ranges))==set(map(str,t.merged_cells.ranges))
 assert {k:s.tables[k].ref for k in s.tables}=={k:t.tables[k].ref for k in t.tables}
 assert str(s.data_validations)==str(t.data_validations)
 for rr in s.iter_rows():
  for c in rr:
   e=t[c.coordinate];count+=1
   assert c.value==e.value and c.number_format==e.number_format and style(c,ac)==style(e,bc),(name,c.coordinate)
 for n,dim in s.row_dimensions.items():assert dim.height==t.row_dimensions[n].height
ids=set();newcounts={}
for name,rows in d['groups'].items():
 s=b[name];last=len(rows)+3
 assert s.max_row==last and s.max_column==27
 assert [c.value for c in s[3]]==d['headers']
 assert len(s.tables)==1 and s.freeze_panes=='C4'
 for n,raw in enumerate(rows,4):
  expected=list(raw);expected[2]=datetime.fromisoformat(raw[2]);expected[5]=datetime.fromisoformat(raw[5])
  actual=[c.value for c in s[n]]
  assert actual[:22]==expected[:22] and actual[25:]==expected[25:]
  assert actual[22:25]==[f'=O{n}+P{n}',f'=Q{n}+R{n}',f'=S{n}+T{n}']
  key=tuple(actual[:2]);assert key not in ids;ids.add(key)
  assert isinstance(actual[1],str) and len(actual[1])==8 and isinstance(actual[4],str) and len(actual[4])==8
 for v in cache[name].iter_rows(min_row=4,max_row=last,min_col=23,max_col=25,values_only=True):assert list(v)==[0,0,0]
 newcounts[name]=len(rows)
assert len(ids)==42
cache.close();shutil.copy2(h/'with-split-42.xlsx',p)
result={'passed':True,'new_sheets':newcounts,'old_cells_preserved':count,'import_and_AB_unchanged':True,'all_42_cumulative_zero':True,'sha256':sha(p)}
(h/'split-42-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result))
