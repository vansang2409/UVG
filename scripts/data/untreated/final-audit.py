"""Rà soát workbook đang dùng với nguồn hiện tại; chỉ đọc, không ghi Excel."""
import sys,json,hashlib,re
from pathlib import Path
from collections import defaultdict,Counter
from datetime import datetime,date
from decimal import Decimal
import openpyxl
sys.stdout.reconfigure(encoding='utf8');r=Path(r'D:\UVG');sys.path.insert(0,str(r/'scripts'))
from prepare_invoice_reports import rows,invoice_key,invoice_day,amount,normal
p=r/'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx';src=r/'data/Du lieu Meinvoice';h=r/'scripts/data/untreated'
files=[p,*sorted((src/'Bang_ke_hoa_don_da_su_dung_2024_2025').glob('*.xlsx')),*[src/n for n in ['Bang_ke_chi_tiet_HD_da_su_dung_2024.xls','Bang_ke_chi_tiet_HD_da_su_dung_2025.xls','Bang_ke_chi_tiet_HD_da_su_dung_2025_MTT.xls']]]
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();hashes={str(f):sha(f) for f in files}
w=openpyxl.load_workbook(p,read_only=False,data_only=True);main={};issues=[]
def check(ok,msg):
 if not ok:issues.append(msg)
def norm(v):return None if v in (None,'') else str(v).strip()
def dec(v):return None if v in (None,'') else Decimal(str(v))
def rate(v):
 if v in (None,''):return None
 s=str(v);return Decimal(s.rstrip('%')) if re.fullmatch(r'\d+(\.\d+)?%?',s) else s
for n,rr in enumerate(w['Chua dieu chinh'].iter_rows(min_row=3,max_row=8412,max_col=16,values_only=True),3):
 k=invoice_key(*rr[:2]);check(k not in main,f'duplicate main {k}');main[k]=(n,list(rr))
check(len(main)==8410,'main count');check(w.sheetnames==['Chua dieu chinh','Huong dan','Import dieu chinh','Kiem tra SL 0'],'sheet names')
confirmed={invoice_key(*x['key'].split('/')):x for x in json.loads((h/'confirmed-payments.json').read_text(encoding='utf8'))}
excluded={invoice_key(*k.split('/')) for k in json.loads((h/'zero-exclusion-meta.json').read_text(encoding='utf8'))['keys']}
headers={};details=defaultdict(list)
for f in files[1:11]:
 for n,v in rows(f):
  if n<6 or amount(v.get('A')) is None:continue
  k=invoice_key(v.get('B'),v.get('C'))
  if k in main:
   check(k not in headers,f'duplicate total {k}');headers[k]={'money':[amount(v.get(c)) for c in ['P','Q','R']],'date':invoice_day(v.get('D')),'status':v.get('U'),'address':v.get('I'),'source':f'{f.name} | dòng {n}'}
for f in files[11:]:
 mt='_MTT' in f.name
 for n,v in rows(f):
  if n<7 or amount(v.get('A')) is None:continue
  k=invoice_key(v.get('F' if mt else 'D'),v.get('B'))
  if k not in main:continue
  cols=['S','T','U','V','W','Z','AA','AB','Y','H','I','J','K','N','D'] if mt else ['L','M','N','O','P','S','T','U','R','E','G','H',None,'I','C']
  a=[v.get(c) if c else None for c in cols]
  details[k].append({'name':a[0],'unit':a[1],'qty':amount(a[2]),'price':amount(a[3]),'goods':amount(a[4]),'rate':rate(a[5]),'tax':amount(a[6]),'pay':amount(a[7]),'discount':amount(a[8]),'customer':a[9],'taxid':a[10],'buyer':a[11],'email':a[12],'method':a[13],'date':invoice_day(a[14]),'address':headers[k]['address'] if mt else v.get('F'),'source':f'{f.name} | dòng {n}'})
check(set(headers)==set(main),'header coverage');check(set(details)==set(main),'detail coverage')
for k,(n,v) in main.items():
 hd=headers[k];ls=details[k];money=hd['money'][:]
 if k in confirmed:
  check(money==[Decimal(0),Decimal(0),None],f'confirmation source {k}');money[2]=Decimal(0);check(v[13]==confirmed[k]['note'],f'confirmation note {k}')
 check([dec(x) for x in v[3:6]]==money,f'main totals {k}')
 check(invoice_day(v[2])==hd['date'] and date(2024,1,1)<=hd['date']<=date(2025,6,30),f'main date {k}')
 check('huy' not in normal(hd['status']),f'cancelled main {k}')
 sums=[sum((l['goods'] or 0)-(l['discount'] or 0) for l in ls),sum(l['tax'] or 0 for l in ls),sum(l['pay'] or 0 for l in ls)]
 check(sums==money,f'live detail/header difference {k}');check(money[0]+money[1]==money[2],f'header arithmetic {k}')
 check(v[9:12]==[len(ls),sum(l['goods']!=0 for l in ls),sum(l['goods']==0 for l in ls)],f'main detail counts {k}')
for name,indices in [('HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx',(0,1)),('NGOAI_LE_F0_2024_T6_2025.xlsx',(2,3))]:
 other=openpyxl.load_workbook(r/'outputs'/name,read_only=True,data_only=True);keys={invoice_key(v[indices[0]],v[indices[1]]) for v in other.worksheets[0].iter_rows(min_row=3,values_only=True)};check(not(set(main)&keys),f'overlap {name}')
s=w['Import dieu chinh'];groups=defaultdict(list);seq_keys={};active=[]
for n,v in enumerate(s.iter_rows(min_row=10,max_col=24,values_only=True),10):
 if all(x is None for x in v):continue
 check(v[0] is not None,f'orphan import row {n}');k=invoice_key(v[9],v[10]);groups[k].append((n,list(v)));active.append((n,v));check(k not in excluded,f'excluded leaked {k}')
 if v[0] in seq_keys:check(seq_keys[v[0]]==k,f'seq reused {n}')
 seq_keys[v[0]]=k
check(set(groups)==set(main)-excluded,'import key set');check(len(groups)==8404 and len(active)==13543,'import counts');check(set(seq_keys)==set(range(1,8405)),'seq not continuous')
c=Counter();qty_product_differences=[]
for k,rs in groups.items():
 ls=details[k];check(len(rs)==len(ls),f'import line count {k}');check(len({v[0] for _,v in rs})==1,f'split sequence {k}');money=[dec(x) for x in main[k][1][3:6]]
 for j,((n,v),l) in enumerate(zip(rs,ls)):
  expected=[-l['qty'] if l['qty'] is not None else None,l['price'],-l['goods'] if l['goods'] is not None else None]
  check([dec(x) for x in v[16:19]]==expected,f'line qty/price/goods {k} row {n}')
  check(norm(v[14])==norm(l['name']) and norm(v[15])==norm(l['unit']),f'line name/unit {k} row {n}')
  check(norm(v[2])==norm(v[5])==norm(l['buyer']),f'buyer names {k} row {n}')
  check(norm(v[3])==norm(l['address']),f'address {k} row {n}')
  check(norm(v[4])==norm(l['taxid']) and norm(v[6])==norm(l['email']) and norm(v[7])==norm(l['method']),f'identity/method {k} row {n}')
  check(invoice_day(v[1])==date(2026,10,8) and invoice_day(v[11])==headers[k]['date']==l['date'],f'import dates {k} row {n}')
  check(v[16] not in (None,0),f'zero/blank quantity {k} row {n}');check(v[7] and v[14] and v[9] and v[10],f'required blanks {k} row {n}')
  check(v[23]==(1 if l['goods']==0 else None),f'promo {k} row {n}');c['promo_lines']+=v[23]==1;c['blank_buyer_lines']+=v[5] is None;c['positive_adjustment_qty']+=v[16]>0
  if l['qty'] is not None and l['price'] is not None and l['goods'] is not None and l['qty']*l['price']!=l['goods']:qty_product_differences.append({'key':list(k),'row':n,'source_qty':str(l['qty']),'source_price':str(l['price']),'source_goods':str(l['goods']),'difference':str(l['qty']*l['price']-l['goods'])})
  if j==0:check([dec(v[i]) for i in [19,21,22]]==[-x for x in money],f'header reversal {k}');check(rate(v[20])==l['rate'],f'VAT rate {k}')
  else:check(all(v[i] is None for i in [19,20,21,22]),f'repeated totals {k} row {n}')
 check(sum(dec(v[18]) for _,v in rs)==-money[0],f'import goods sum {k}')
rv=defaultdict(list)
for n,v in enumerate(w['Kiem tra SL 0'].iter_rows(min_row=3,max_row=29,max_col=19,values_only=True),3):rv[invoice_key(v[0],v[1])].append((n,list(v)))
check(set(rv)==excluded and sum(map(len,rv.values()))==27,'review completeness')
for k,rs in rv.items():
 ls=details[k];check(len(rs)==len(ls),f'review lines {k}')
 for j,((n,v),l) in enumerate(zip(rs,ls)):
  check([dec(x) for x in v[9:12]]==[l['qty'],l['price'],l['goods']],f'review quantities {k} row {n}')
  check(norm(v[7])==norm(l['name']) and norm(v[8])==norm(l['unit']),f'review names {k} row {n}')
  check(v[17]==l['source'] and v[16] and v[18],f'review evidence {k} row {n}')
  check(invoice_day(v[2])==l['date'],f'review dates {k}')
  check([dec(v[i]) for i in [12,14,15]]==([dec(x) for x in main[k][1][3:6]] if j==0 else [None]*3),f'review totals {k} row {n}')
errors=[]
for ss in w:
 for rr in ss.iter_rows():
  for cell in rr:
   if cell.data_type=='e':errors.append([ss.title,cell.coordinate,cell.value])
check(not errors,'Excel error cells');check(all(sha(Path(f))==v for f,v in hashes.items()),'files changed during audit')
result={'passed':not issues,'main_invoices':len(main),'import_invoices':len(groups),'import_lines':len(active),'excluded_invoices':len(rv),'review_lines':sum(map(len,rv.values())),'confirmed_zero_payments':len(confirmed),'counts':dict(c),'source_product_differences':len(qty_product_differences),'source_product_invoices':len({tuple(x['key']) for x in qty_product_differences}),'source_product_zero_goods':sum(Decimal(x['source_goods'])==0 for x in qty_product_differences),'source_product_max_abs':str(max((abs(Decimal(x['difference'])) for x in qty_product_differences),default=0)),'source_product_examples':sorted(qty_product_differences,key=lambda x:abs(Decimal(x['difference'])),reverse=True)[:8],'issues':issues,'product_difference_details':qty_product_differences,'read_only':True,'source_files_checked':len(files)-1}
(h/'final-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({k:v for k,v in result.items() if k!='product_difference_details'},ensure_ascii=False));assert not issues
