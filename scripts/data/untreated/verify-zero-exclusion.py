import openpyxl,json,hashlib,shutil
from pathlib import Path
from collections import OrderedDict
r=Path(r'D:\UVG');h=r/'scripts/data/untreated';a=openpyxl.load_workbook(h/'before-zero-exclusion.xlsx');b=openpyxl.load_workbook(h/'with-zero-exclusion.xlsx');m=json.loads((h/'zero-exclusion-meta.json').read_text(encoding='utf8'));keys=set(m['keys'])
assert b.sheetnames==a.sheetnames+['Kiem tra SL 0']
for name in ['Chua dieu chinh']:
 for row1,row2 in zip(a[name].iter_rows(),b[name].iter_rows()):
  for c,e in zip(row1,row2):assert c.value==e.value and c.number_format==e.number_format,(name,c.coordinate)
for n,row in enumerate(a['Huong dan'].iter_rows(),1):
 if n==13:continue
 for c in row:assert c.value==b['Huong dan'].cell(c.row,c.column).value
s=a['Import dieu chinh'];t=b['Import dieu chinh'];expected=[];origremoved=[];seq=OrderedDict()
for row in s.iter_rows(min_row=10,max_row=13579,max_col=24,values_only=True):
 v=list(row);k=str(v[9])+'/'+str(v[10])
 if k in keys:origremoved.append(v);continue
 if k not in seq:seq[k]=len(seq)+1
 v[0]=seq[k];expected.append(v)
assert len(expected)==13543 and len(seq)==8404 and len(origremoved)==27
actual=[list(v) for v in t.iter_rows(min_row=10,max_row=13552,max_col=24,values_only=True)]
assert actual==expected,'Unexpected import data changes'
assert not any(c.value is not None for rr in t.iter_rows(min_row=13553,max_row=13579,max_col=24) for c in rr)
assert all(v[16]!=0 and str(v[9])+'/'+str(v[10]) not in keys for v in actual)
for n,row in enumerate(s.iter_rows(min_row=1,max_row=9,max_col=24),1):
 for c in row:
  if c.coordinate!='A7':assert c.value==t.cell(c.row,c.column).value
review=b['Kiem tra SL 0'];rv=[list(v) for v in review.iter_rows(min_row=3,max_row=29,max_col=19,values_only=True)]
assert len(rv)==27 and {str(v[0])+'/'+str(v[1]) for v in rv}==keys
first=set();zero=0
for v,o in zip(rv,origremoved):
 k=str(o[9])+'/'+str(o[10]);isfirst=k not in first;first.add(k)
 assert v[:12]==[o[9],o[10],o[11],o[2],o[3],o[5],o[7],o[14],o[15],-o[16],o[17],-o[18]]
 assert v[12:16]==([-o[19],o[20],-o[21],-o[22]] if isfirst else [None]*4)
 assert v[16] and v[17] and v[18]
 zero+=v[9]==0
assert zero==6 and review.freeze_panes=='C3' and review.tables['KiemTraSoLuong0'].ref=='A2:S29'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();p=r/'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx'
assert sha(p)==sha(h/'before-zero-exclusion.xlsx'),'Report changed concurrently'
shutil.copy2(h/'with-zero-exclusion.xlsx',p);assert sha(p)==sha(h/'with-zero-exclusion.xlsx')
print(json.dumps({'PASS':True,'excluded_invoices':6,'review_lines':27,'import_invoices':8404,'import_lines':13543,'saved':True}))
