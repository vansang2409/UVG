import json,hashlib
from pathlib import Path
from datetime import datetime,date
import openpyxl
r=Path(r'D:\UVG');d=json.loads((r/'scripts/data/untreated/untreated.json').read_text(encoding='utf8'));p=r/'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx';w=openpyxl.load_workbook(p,data_only=True);s=w['Chua dieu chinh']
assert w.sheetnames==['Chua dieu chinh','Huong dan']
assert s.max_row==8412 and s.max_column==16
actual=set();missing=0;zero=0;year={}
for n,(g,row) in enumerate(zip(d['groups'],s.iter_rows(min_row=3,max_row=8412,max_col=16)),3):
 v=[c.value for c in row];inv=g['invoices'][0];key=tuple(v[:2]);assert key==tuple(g['root']);assert key not in actual;actual.add(key)
 assert v[2].date()==datetime.strptime(inv[2],'%d/%m/%Y').date()
 assert date(2024,1,1)<=v[2].date()<=date(2025,6,30)
 assert v[3:6]==inv[4:7],(key,'source money')
 assert v[9]==g['line_counts']['lines'] and v[10]==g['line_counts'].get('nonzero_goods',0) and v[11]==g['line_counts'].get('zero_goods',0)
 assert v[13]==(g['notes'] or None) and v[15]==g['detail_source']
 assert isinstance(v[1],str) and len(v[1])==8
 if v[5] is None:missing+=1
 if v[11]>0:zero+=1
 year[str(v[2].year)]=year.get(str(v[2].year),0)+1
assert missing==6 and zero==952 and year=={'2024':62,'2025':8348}
assert s.freeze_panes=='C3' and s.tables['ChuaDieuChinhChuoi'].ref=='A2:P8412'
assert sum(s.cell(n,10).value for n in range(3,8413))==13570
assert all(hashlib.sha256((r/k).read_bytes()).hexdigest()==v for k,v in d['hashes'].items())
assert all(hashlib.sha256((r/'outputs'/k).read_bytes()).hexdigest()==v for k,v in d['existing_report_hashes'].items())
print('PASS: 8410 unique F0, correct dates/money/source notes, 6 missing payments preserved, 952 roots with zero lines, 13570 detail rows, filter/freeze intact; source and 2 existing reports unchanged')
print('SIZE',p.stat().st_size)

