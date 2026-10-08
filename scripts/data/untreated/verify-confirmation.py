import json,hashlib,shutil
from pathlib import Path
import openpyxl
r=Path(r'D:\UVG'); h=r/'scripts/data/untreated'; dst=r/'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx'; before=h/'before-payment-confirmation.xlsx'; stage=h/'confirmed-payment.xlsx'
a=openpyxl.load_workbook(before); b=openpyxl.load_workbook(stage)
changes=json.loads((h/'confirmed-payments.json').read_text(encoding='utf8'))
assert len(changes)==6
allowed={('Chua dieu chinh',f'{col}{x["row"]}') for x in changes for col in ['F','M','N']}|{('Huong dan','A7'),('Huong dan','A9')}
assert a.sheetnames==b.sheetnames
actual=set()
for name in a.sheetnames:
 s,t=a[name],b[name]
 assert (s.max_row,s.max_column)==(t.max_row,t.max_column)
 assert s.freeze_panes==t.freeze_panes
 assert set(s.tables)==set(t.tables)
 for k in s.tables: assert s.tables[k].ref==t.tables[k].ref
 for row1,row2 in zip(s.iter_rows(),t.iter_rows()):
  for c,e in zip(row1,row2):
   if c.value!=e.value: actual.add((name,c.coordinate))
   assert c.number_format==e.number_format,(name,c.coordinate,'format')
assert actual==allowed,(actual-allowed,allowed-actual)
for x in changes:
 s=b['Chua dieu chinh'];n=x['row']
 assert '/'.join(str(s.cell(n,c).value) for c in [1,2])==x['key']
 assert s.cell(n,6).value==0 and s.cell(n,14).value==x['note']
d=json.loads((h/'untreated.json').read_text(encoding='utf8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
changed_sources=[k for k,v in d['hashes'].items() if sha(r/k)!=v]
assert changed_sources==['data/Du lieu Meinvoice/Bang_ke_chi_tiet_HD_da_su_dung_2024.xls']
from io import BytesIO
wsrc=openpyxl.load_workbook(BytesIO((r/changed_sources[0]).read_bytes()),read_only=True,data_only=True)
for rownum,num in [(1733,'00000744'),(6229,'00000671')]:
 vals=next(wsrc.worksheets[0].iter_rows(min_row=rownum,max_row=rownum,values_only=True))
 assert vals[1]==num and vals[3]=='1C24TUV' and vals[15]==0 and vals[19]==0 and vals[20]==0
source_snapshot={k:sha(r/k) for k in d['hashes']}
assert all(sha(r/'outputs'/k)==v for k,v in d['existing_report_hashes'].items())
assert sha(dst)==sha(before),'Report changed concurrently'
assert all(sha(r/k)==v for k,v in source_snapshot.items())
shutil.copy2(stage,dst)
assert sha(dst)==sha(stage)
print('PASS: exactly 6 payments, 6 classes, 6 notes and 2 guide cells changed; all other values, number formats, tables and freeze preserved. Current 2024 source rows rechecked as zero; no sources modified by this operation. Prior reports unchanged. Saved active report.')
