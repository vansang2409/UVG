import json,hashlib,shutil,sys
from pathlib import Path
from datetime import datetime
import openpyxl
sys.stdout.reconfigure(encoding='utf8');r=Path(r'D:\UVG');h=r/'scripts/data/untreated';p=r/'outputs/HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx'
a=openpyxl.load_workbook(h/'before-import.xlsx');b=openpyxl.load_workbook(h/'with-import.xlsx');data=json.loads((h/'import-data.json').read_text(encoding='utf8'));m=json.loads((h/'import-meta.json').read_text(encoding='utf8'))
assert b.sheetnames==a.sheetnames+['Import dieu chinh']
for name in a.sheetnames:
 s,t=a[name],b[name]
 assert s.freeze_panes==t.freeze_panes
 for k in s.tables:assert s.tables[k].ref==t.tables[k].ref
 for row1,row2 in zip(s.iter_rows(),t.iter_rows(max_row=s.max_row,max_col=s.max_column)):
  for c,e in zip(row1,row2):assert c.value==e.value and c.number_format==e.number_format,(name,c.coordinate)
s=b['Import dieu chinh'];assert s.max_row==13579 and s.max_column==24
by={i+1:x for i,x in enumerate(data['roots'])};groups={};promo=0
for row,expected,meta in zip(s.iter_rows(min_row=10,max_col=24,values_only=True),m['rows'],m['meta']):
 v=list(row);assert v[1]==datetime(2026,10,8)
 expected[1]=datetime(2026,10,8);assert v==expected,(meta['row'],'import mismatch')
 x=by[v[0]];assert v[9:11]==x['key']
 if meta['first']:assert v[19]==-x['values'][3] and v[21]==-x['values'][4] and v[22]==-x['values'][5]
 else:assert all(v[i] is None for i in [19,20,21,22])
 assert v[23]==(1 if v[18]==0 else None)
 groups.setdefault(v[0],[]).append(v);promo+=v[23]==1
assert len(groups)==8410 and len(m['rows'])==13570
for seq,vs in groups.items():assert abs(sum(v[18] for v in vs)-vs[0][19])<0.01
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(p)==sha(h/'before-import.xlsx'),'Report changed concurrently'
shutil.copy2(h/'with-import.xlsx',p);assert sha(p)==sha(h/'with-import.xlsx')
print(json.dumps({'PASS':True,'invoices':len(groups),'lines':len(m['rows']),'promo_lines':promo,'date':'08/10/2026','old_sheets_preserved':True,'saved':True}))
