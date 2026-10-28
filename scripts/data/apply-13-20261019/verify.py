"""Independently verify saved XLSX, unchanged scope, and eight chains returning to zero."""
from pathlib import Path
from collections import Counter,defaultdict
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from zipfile import ZipFile
from lxml import etree as E
from openpyxl.utils.cell import range_boundaries,get_column_letter
import json,sys
import openpyxl

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[2]
p=json.loads((HERE/'plan.json').read_text(encoding='utf8'))
src=json.loads((HERE/'sources.json').read_text(encoding='utf8'))
final='--final' in sys.argv
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
tag=lambda s:'{'+NS+'}'+s
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
def parts(z):
    wb=E.fromstring(z.read('xl/workbook.xml'));rs=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    targets={x.get('Id'):x.get('Target') for x in rs}
    return {s.get('name'):t.lstrip('/') if (t:=targets[s.get(RID)]).startswith('/') else 'xl/'+t for s in wb.find(tag('sheets'))}
def cellsig(c):
    return (tuple(sorted(c.attrib.items())),tuple((x.tag,tuple(sorted(x.attrib.items())),x.text,E.tostring(x)) for x in c))
def allowed(sp):
    out={c['cell'] for c in sp['changes']+sp['formulas']}
    for b in sp['blocks']:
        lo,r1,hi,r2=range_boundaries(b['range'])
        out.update(f'{get_column_letter(c)}{r}' for r in range(r1,r2+1) for c in range(lo,hi+1))
    return out
reports=[];groups_by_seq={};exception_roots=set();exception_keys=set();source_by_root=defaultdict(dict)
main_cache={};new_paths={}
for f in p['files']:
    path=ROOT/f['destination'] if final else HERE/'staged'/f['name']
    new_paths[f['kind']]=path
    before=HERE/'before'/f['name']
    assert sha256(before.read_bytes()).hexdigest()==f['before_sha256']
    preserved=0
    with ZipFile(before) as a,ZipFile(path) as b:
        assert b.testzip() is None
        assert set(a.namelist())==set(b.namelist())
        ap,bp=parts(a),parts(b);assert ap==bp
        for name in a.namelist():
            if name not in {ap[s] for s in f['sheets']} and not name.startswith('xl/tables/') and name!='xl/calcChain.xml':
                assert a.read(name)==b.read(name),('changed native resource',f['kind'],name)
        for sn,sp in f['sheets'].items():
            old=E.fromstring(a.read(ap[sn]));new=E.fromstring(b.read(bp[sn]));ok=allowed(sp)
            for child in old:
                if child.tag not in (tag('dimension'),tag('sheetData')):
                    assert E.tostring(child)==E.tostring(new.find(child.tag)),('sheet metadata',f['kind'],sn,child.tag)
            oldcells={c.get('r'):c for c in old.iter(tag('c'))};newcells={c.get('r'):c for c in new.iter(tag('c'))}
            replacement=sp.get('replace_data')
            for coord,c in oldcells.items():
                row=int(''.join(x for x in coord if x.isdigit()))
                if coord in ok or replacement and row>=replacement['start']:continue
                assert coord in newcells and cellsig(c)==cellsig(newcells[coord]),('unauthorized cell change',f['kind'],sn,coord)
                preserved+=1
            for coord,c in newcells.items():
                if coord not in oldcells:assert coord in ok,('unexpected new cell',f['kind'],sn,coord)
            if replacement:
                oldrows={int(x.get('r')):x for x in old.find(tag('sheetData'))}
                newrows={int(x.get('r')):x for x in new.find(tag('sheetData'))}
                assert max(newrows)==replacement['end']
                for target,donor in replacement['row_map'].items():
                    oldstyles={''.join(c0 for c0 in c.get('r') if c0.isalpha()):c.get('s') for c in oldrows[donor]}
                    for c in newrows[int(target)]:assert c.get('s')==oldstyles.get(''.join(x for x in c.get('r') if x.isalpha()))
            else:
                for coord,c in oldcells.items():
                    if coord in ok:assert c.get('s')==newcells[coord].get('s'),('changed cell style',f['kind'],sn,coord)
        # Table changes may resize their range and filter only.
        expected_tables={n:r for sp in f['sheets'].values() for n,r in sp['tables'].items()}
        for name in a.namelist():
            if not name.startswith('xl/tables/') or not name.endswith('.xml'):continue
            old=E.fromstring(a.read(name));new=E.fromstring(b.read(name));tn=old.get('name')
            if tn in expected_tables:
                assert new.get('ref')==expected_tables[tn]
                af=new.find(tag('autoFilter'));assert af is not None and af.get('ref')==expected_tables[tn]
                old.set('ref',new.get('ref'));old.find(tag('autoFilter')).set('ref',new.get('ref'))
            assert E.tostring(old)==E.tostring(new),('table metadata',f['kind'],tn)
    fw=openpyxl.load_workbook(path,read_only=True,data_only=False)
    cw=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if f['kind']=='compact':assert fw.sheetnames==['Import dieu chinh','F0, F1, F2, ...']
    # All authored scalar, block, and formula results must match the saved file.
    for sn,sp in f['sheets'].items():
        expected={c['cell']:c['value'] for c in sp['changes']}
        for block in sp['blocks']:
            lo,r1,hi,r2=range_boundaries(block['range'])
            rows=list(fw[sn].iter_rows(min_row=r1,max_row=r2,min_col=lo,max_col=hi,values_only=True))
            for rr,(actual,values) in enumerate(zip(rows,block['values']),r1):
                for i,(av,ev) in enumerate(zip(actual,values)):
                    coord=f'{get_column_letter(lo+i)}{rr}'
                    if any(c['cell']==coord for c in sp['formulas']):continue
                    if i in block['date_columns'] and ev:ev=datetime.fromisoformat(str(ev)[:10])
                    ev=None if ev=='' else ev
                    assert av==ev,('saved authored value',f['kind'],sn,coord,av,ev)
        for coord,ev in expected.items():assert fw[sn][coord].value==ev
        for c0 in sp['formulas']:
            assert fw[sn][c0['cell']].value==c0['formula']
            assert cw[sn][c0['cell']].value==c0['cache'],('saved cached formula',f['kind'],sn,c0)
    errors=[]
    for s in cw:
        for row in s:
            for cell in row:
                if cell.data_type=='e':errors.append((s.title,cell.coordinate,cell.value))
    assert not errors,errors[:5]
    if f['kind'] in ('adjusted','compact'):
        groups=defaultdict(list)
        for rn,r in enumerate(cw['Import dieu chinh'].iter_rows(min_row=10,max_col=24,values_only=True),10):
            if r[0] is not None:groups[r[0]].append((rn,r))
        assert len(groups)==2045 and sum(map(len,groups.values()))==41996
        assert set(groups)==set(range(1,2046))
        assert len({tuple(v[1][9:11]) for g in groups.values() for v in g})==1015
        for action in p['actions']:
            rows=groups[action['seq']]
            assert rows[0][0]==action['start'] and rows[-1][0]==action['end']
            assert [rows[0][1][i] for i in (19,21,22)]==action['money']
            assert sum(Decimal(str(v[18])) for _,v in rows if v[18] is not None)==Decimal(str(action['money'][0]))
            assert all(v[1]==datetime(2026,10,19) and v[6] is None and v[7]=='TM/CK' and v[2]==v[5] for _,v in rows)
            assert all(v[17] is None or v[17]>0 for _,v in rows)
            if action['target']==action['root']:
                original=src['details'][action['target']]
                assert len(original)==len(rows)
                for l,(_,v) in zip(original,rows):
                    assert Decimal(str(l['qty']))+Decimal(str(v[16]))==0
                    assert v[17]==l['price'] and v[18]==(-l['goods'] if l['goods'] is not None else None)
            else:assert len(rows)==1 and rows[0][1][16:18]==(None,None)
            assert [x+y for x,y in zip(action['source_money'],action['money'])]==[0,0,0]
        if f['kind']=='adjusted':
            groups_by_seq=groups
            for r in cw['Chi tiet nguon'].iter_rows(min_row=2,values_only=True):
                if r[0]:source_by_root['/'.join(r[:2])]['/'.join(r[2:4])]=list(r)
            assert len({k for rows in source_by_root.values() for k in rows})==2117
            for check in p['root_checks']:
                rn=check['row']
                main_cache[check['root']]=[cw['Dieu chinh'].cell(rn,c).value for c in (53,54,55)]
                assert main_cache[check['root']]==[0,0,0]
            states=Counter(cw['Cho KT xac nhan'].cell(r,7).value for r in range(3,22))
            assert states=={'Đã bổ sung import':19},states
    elif f['kind']=='exceptions':
        exception_roots={'/'.join(r[2:4]) for r in cw['Ngoai le'].iter_rows(min_row=3,values_only=True) if r[2]}
        exception_keys={'/'.join(r[2:4]) for r in cw['Chi tiet nguon'].iter_rows(min_row=2,values_only=True) if r[2]}
        assert len(exception_roots)==236 and len(exception_keys)==322
    elif f['kind']=='list72':
        assert cw['Tong hop']['B7'].value==10455 and cw['Tong hop']['B9'].value==13
        assert cw['Tong hop']['A9'].value=='Đã bổ sung vào bộ import'
        assert len([r for r in cw['Khong can dieu chinh'].iter_rows(min_row=7,values_only=True) if r[0]])==72
    fw.close();cw.close()
    reports.append({'name':f['name'],'destination':f['destination'],'sha256':sha256(path.read_bytes()).hexdigest(),'unchanged_cells_checked':preserved})

assert len(groups_by_seq)==2045
for check in p['root_checks']:
    original=[]
    for key in check['source_invoices']:
        money=source_by_root[check['root']][key][9:12]
        if key in p['retained_information_invoices']:
            assert src['xmls'][key]['money']==[0,0,0]
            money=[0 if v is None else v for v in money]
        if key=='1C24TUV/00000574':money[2]=0
        assert all(v is not None for v in money)
        original.append(money)
    actual_adjustments=[[groups_by_seq[seq][0][1][i] for i in (19,21,22)] for seq in check['sequences']]
    total=[sum(Decimal(str(m[i])) for m in original+actual_adjustments) for i in range(3)]
    assert total==[0,0,0],('chain not zero',check['root'],total)
    assert check['root'] not in exception_roots
assert not {a['target'] for a in p['actions']}&set(p['retained_information_invoices'])
assert not {a['target'] for a in p['actions']}&exception_keys
for name,h in p['source_hashes'].items():assert sha256((ROOT/name).read_bytes()).hexdigest()==h,('source changed',name)
for name,h in p['untreated_hashes'].items():assert sha256((ROOT/'outputs'/name).read_bytes()).hexdigest()==h,('untreated changed',name)
with ZipFile(new_paths['adjusted']) as a,ZipFile(new_paths['compact']) as b:
    ap,bp=parts(a),parts(b)
    for sn,cn in [('Import dieu chinh','Import dieu chinh'),('Dieu chinh','F0, F1, F2, ...')]:
        assert a.read(ap[sn])==b.read(bp[cn]),('two-sheet export differs',sn)
    assert a.read('xl/styles.xml')==b.read('xl/styles.xml')
result={'passed':True,'mode':'final' if final else 'staged','files':reports,'counts':p['counts'],
        'zero_goods_tax_payment_chains':main_cache,'new_quantities_reversed':True,
        'new_nonblank_unit_prices_positive':True,'amount_only_reversals':5,
        'source_blanks_preserved':True,'retained_information_invoices':p['retained_information_invoices'],
        'formula_errors':0,'main_and_two_sheet_views_identical':True,'sources_and_untreated_workbooks_unchanged':True,
        'software_imported':False,'signed':False,'issued':False}
(HERE/('final-verification.json' if final else 'staged-verification.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
