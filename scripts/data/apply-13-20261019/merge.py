"""Merge artifact-authored cells into original XLSX packages, preserving native parts."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from openpyxl.utils.cell import range_boundaries,get_column_letter,column_index_from_string
import hashlib,json,posixpath,re

HERE=Path(__file__).resolve().parent
p=json.loads((HERE/'plan.json').read_text(encoding='utf8'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
tag=lambda s:'{'+NS+'}'+s
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
sha=lambda path:hashlib.sha256(path.read_bytes()).hexdigest()
def resolve(z):
    w=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    targets={r.get('Id'):r.get('Target') for r in rels}
    return {s.get('name'):t.lstrip('/') if (t:=targets[s.get(RID)]).startswith('/') else 'xl/'+t for s in w.find(tag('sheets'))}
def serialise(tree):return E.tostring(tree,xml_declaration=True,encoding='utf-8',standalone=True)

(HERE/'staged').mkdir(exist_ok=True)
for f in p['files']:
    baseline=HERE/'before'/f['name']
    assert sha(baseline)==f['before_sha256']
    authored=HERE/'authored'/('HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx' if f['kind']=='compact' else f['name'])
    with ZipFile(baseline) as original,ZipFile(authored) as author:
        old_parts=resolve(original);new_parts=resolve(author)
        strings=list(E.fromstring(author.read('xl/sharedStrings.xml'))) if 'xl/sharedStrings.xml' in author.namelist() else []
        modified={}
        for sn,sp in f['sheets'].items():
            an=f.get('author_sheet_map',{}).get(sn,sn)
            old=E.fromstring(original.read(old_parts[sn]));new=E.fromstring(author.read(new_parts[an]))
            generated={c.get('r'):c for c in new.iter(tag('c'))}
            data=old.find(tag('sheetData'))
            old_rows={int(r.get('r')):deepcopy(r) for r in data}
            rows={int(r.get('r')):r for r in data}
            if 'replace_data' in sp:
                rd=sp['replace_data']
                for rn,row in list(rows.items()):
                    if rn>=rd['start']:data.remove(row);del rows[rn]
                for target,donor in rd['row_map'].items():
                    rn=int(target);row=deepcopy(old_rows[donor]);row.set('r',target)
                    for c in row:
                        if c.tag==tag('c'):
                            assert c.find(tag('f')) is None
                            c.set('r',re.sub(r'\d+$',target,c.get('r')))
                    data.append(row);rows[rn]=row
            expected={}
            for b in sp['blocks']:
                lo,start,hi,end=range_boundaries(b['range'])
                assert len(b['values'])==end-start+1
                for rn,values in enumerate(b['values'],start):
                    assert len(values)==hi-lo+1
                    for ci,v in enumerate(values,lo):expected[f'{get_column_letter(ci)}{rn}']={'value':v}
            for c in sp['changes']:expected[c['cell']]=c
            for c in sp['formulas']:expected[c['cell']]=c
            for coord,spec in expected.items():
                rn=int(re.search(r'\d+$',coord)[0]);letters=re.match(r'[A-Z]+',coord)[0]
                row=rows.get(rn)
                if row is None:
                    donor=old_rows.get(sp.get('append_template'))
                    row=E.Element(tag('row'),dict(donor.attrib) if donor is not None else {})
                    row.set('r',str(rn));data.append(row);rows[rn]=row
                oldc=next((c for c in row if c.get('r')==coord),None)
                style=oldc.get('s') if oldc is not None else None
                if style is None and 'append_template' in sp:
                    donor=old_rows.get(sp['append_template'])
                    dc=next((c for c in donor if re.match(r'[A-Z]+',c.get('r'))[0]==letters),None) if donor is not None else None
                    if dc is not None:style=dc.get('s')
                gen=generated.get(coord)
                blank='formula' not in spec and spec.get('value') is None
                if blank:
                    nc=E.Element(tag('c'),r=coord)
                else:
                    assert gen is not None,('missing authored cell',sn,coord)
                    nc=deepcopy(gen)
                    if nc.get('t')=='s':
                        si=strings[int(nc.findtext(tag('v')))];nc.set('t','inlineStr')
                        for c in list(nc):nc.remove(c)
                        inline=E.SubElement(nc,tag('is'))
                        for child in si:inline.append(deepcopy(child))
                    if 'formula' in spec:
                        assert nc.findtext(tag('f'))==spec['formula'].lstrip('='),(sn,coord)
                        assert float(nc.findtext(tag('v')))==spec['cache'],('authored cache',sn,coord,nc.findtext(tag('v')))
                if style is not None:nc.set('s',style)
                else:nc.attrib.pop('s',None)
                if oldc is not None:row.replace(oldc,nc)
                else:row.append(nc)
                row[:]=sorted(row,key=lambda c:column_index_from_string(re.match(r'[A-Z]+',c.get('r'))[0]))
            for rn,height in sp.get('row_heights',{}).items():
                if int(rn) in rows:rows[int(rn)].set('ht',str(height));rows[int(rn)].set('customHeight','1')
            data[:]=sorted(data,key=lambda r:int(r.get('r')))
            dim=old.find(tag('dimension'))
            maxrow=max(rows);maxcol=max((column_index_from_string(re.match(r'[A-Z]+',c.get('r'))[0]) for row in rows.values() for c in row),default=1)
            if dim is not None:dim.set('ref',f'A1:{get_column_letter(maxcol)}{maxrow}')
            modified[old_parts[sn]]=serialise(old)
            if sp.get('tables'):
                relpath=posixpath.dirname(old_parts[sn])+'/_rels/'+posixpath.basename(old_parts[sn])+'.rels'
                for rel in E.fromstring(original.read(relpath)):
                    if not rel.get('Type').endswith('/table'):continue
                    t=rel.get('Target');tp=t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(old_parts[sn]),t))
                    table=E.fromstring(original.read(tp));name=table.get('name')
                    if name not in sp['tables']:continue
                    table.set('ref',sp['tables'][name]);af=table.find(tag('autoFilter'))
                    if af is not None:af.set('ref',sp['tables'][name])
                    modified[tp]=serialise(table)
        if 'xl/calcChain.xml' in original.namelist():
            chain=E.fromstring(original.read('xl/calcChain.xml'))
            wb=E.fromstring(original.read('xl/workbook.xml'))
            sheet_ids={s.get('name'):s.get('sheetId') for s in wb.find(tag('sheets'))}
            known=set();sid=None
            for c in chain:
                sid=c.get('i',sid);known.add((sid,c.get('r')))
            added=False
            for sn,sp in f['sheets'].items():
                for c in sp['formulas']:
                    key=(sheet_ids[sn],c['cell'])
                    if key not in known:
                        E.SubElement(chain,tag('c'),r=c['cell'],i=sheet_ids[sn]);known.add(key);added=True
            if added:modified['xl/calcChain.xml']=serialise(chain)
        out=HERE/'staged'/f['name']
        with ZipFile(out,'w') as dest:
            for item in original.infolist():dest.writestr(deepcopy(item),modified.get(item.filename,original.read(item.filename)))
        with ZipFile(out) as check:
            assert check.testzip() is None
            assert set(check.namelist())==set(original.namelist())
            assert all(check.read(n)==original.read(n) for n in original.namelist() if n not in modified)
        print(f['kind'],'merged cells; preserved',len(original.namelist())-len(modified),'unchanged package entries')
