"""Chuẩn bị import trả lại theo mẫu; chờ người dùng map hóa đơn gốc."""
from pathlib import Path
from collections import defaultdict
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import hashlib,json,os,subprocess
from lxml import etree as E
from openpyxl import load_workbook
from match_import_2024_template import Q,xml,package,cell,cell_value,NS,REL,PKG,CT

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'scripts/data/amis-2024-linked-20261026'
OUT=Path(os.environ.get('UVG_OUTPUT_DIR',ROOT/'outputs/amis-2024-theo-btl-bhh-20261026'))
PENDING=set()

def prepare():
    p=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
    orders=json.loads((WORK/'return-order-map.json').read_text(encoding='utf8'))
    template=ROOT/'data/Du lieu Meinvoice/Mau HD Dieu chinh.xls'
    t=load_workbook(BytesIO(template.read_bytes()),read_only=True,data_only=True).active
    header=list(next(t.iter_rows(min_row=9,max_row=9,max_col=24,values_only=True)))+['Mã đơn hàng','Số chứng từ']
    instructions=[list(r)[0] for r in t.iter_rows(min_row=1,max_row=7,values_only=True)]
    instructions[6]='- Chờ map ký hiệu/số/ngày HĐ gốc và ngày điều chỉnh trước khi nhập. BTL00023, BTL00024 dùng đúng Giá ghi trả AMIS theo xác nhận 28/10/2026.'
    docs={d['doc']:d for d in p['docs']};groups=defaultdict(list)
    for r in p['links']:groups[r[0]].append(r)
    rows=[];audit=[]
    for btl,links in groups.items():
        if btl in PENDING:continue
        n=len({a['btl'] for a in audit})+1
        assert len({r[2] for r in links})==1
        sale=docs[links[0][2]]
        for j,r in enumerate(links):
            first=j==0
            reason=f'Trả lại hàng theo phiếu {btl}; chứng từ bán gốc {r[2]}'
            vals=[n,None,sale['buyer'],sale['address'] or None,sale['tax'] or None,sale['buyer'],None,'TM/CK',None,None,None,None,None,reason,r[5],r[6],-r[7],r[8],-r[9],-sum(x[9] for x in links) if first else None,r[10] if first else None,-sum(x[11] for x in links) if first else None,-sum(x[12] for x in links) if first else None,None,orders.get(r[2]),r[2]]
            assert len(vals)==26 and vals[16]<0 and vals[17]>=0 and vals[18]<0
            rows.append(vals);audit.append({'row':len(rows)+9,'btl':btl,'bhh':r[2],'sku':r[4],'source_rows':r[16],'source':r})
    pending=[r for r in p['return_drafts'] if r[0] in PENDING]
    assert len(rows)==219 and len({a['btl'] for a in audit})==163 and len(pending)==0
    assert sum(r[18] for r in rows)==-225865926 and sum(r[21] or 0 for r in rows)==-18069275 and sum(r[22] or 0 for r in rows)==-243935201
    result={'headers':header,'instructions':instructions,'rows':rows,'audit':audit,'pending_headers':['BTL','BHH bán gốc','Ngày chứng từ BTL','Mã hàng','Tên hàng','ĐVT','SL trả nguồn','Đơn giá trả nguồn','Tiền hàng trả nguồn','Thuế suất nguồn','VAT trả nguồn','Thanh toán trả nguồn','Ký hiệu HĐ gốc','Số HĐ gốc','Ngày HĐ gốc','Ngày HĐ điều chỉnh','Trạng thái','Dòng AMIS','Ghi chú chênh lệch'],'pending':pending,'template_sha256':hashlib.sha256(template.read_bytes()).hexdigest()}
    (WORK/'return-adjustment-import.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf8')
    return result

def merge_styles(dst,src):
    """Gộp style của sheet nhỏ; giữ nguyên mọi style workbook hiện tại."""
    maps={};num={}
    for tag in ['numFmts','fonts','fills','borders','cellStyleXfs','cellXfs']:
        a=dst.find(Q(tag));b=src.find(Q(tag))
        if b is None:continue
        if a is None:a=E.Element(Q(tag));dst.insert(0,a)
        if tag=='numFmts':
            nxt=max([163]+[int(x.get('numFmtId')) for x in a])+1
            for v in b:
                c=deepcopy(v);num[int(c.get('numFmtId'))]=nxt;c.set('numFmtId',str(nxt));nxt+=1;a.append(c)
        else:
            maps[tag]={}
            for i,v in enumerate(b):
                c=deepcopy(v)
                if tag in ['cellXfs','cellStyleXfs']:
                    for key,typ in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
                        if c.get(key) is not None:c.set(key,str(maps[typ][int(c.get(key))]))
                    if c.get('numFmtId') is not None:c.set('numFmtId',str(num.get(int(c.get('numFmtId')),int(c.get('numFmtId')))))
                    if tag=='cellXfs' and c.get('xfId') is not None:c.set('xfId',str(maps['cellStyleXfs'][int(c.get('xfId'))]))
                maps[tag][i]=len(a);a.append(c)
        a.set('count',str(len(a)))
    return maps['cellXfs']

def refresh_approval_notes(raw):
    """Đồng bộ ghi chú đã chốt; giữ nguyên toàn bộ số liệu và các sheet khác."""
    p=json.loads((WORK/'prepared.json').read_text(encoding='utf8'))
    with ZipFile(BytesIO(raw)) as z:
        wb=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        strings=[''.join(x.itertext()) for x in E.fromstring(z.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in z.namelist() else []
        changes={}
        for sh in wb.find(Q('sheets')):
            name=sh.get('name')
            if name not in ['Tong hop','Lien ket BTL-BHH']:continue
            path=next(r.get('Target') for r in rels if r.get('Id')==sh.get('{'+REL+'}id'));part=path.lstrip('/') if path.startswith('/') else 'xl/'+path
            s=E.fromstring(z.read(part));changed=False
            for row in s.find(Q('sheetData')):
                n=row.get('r');cells={c.get('r'):c for c in row};updates={}
                if name=='Tong hop' and cell_value(cells.get('A'+n),strings)=='Hai đơn lệch tiền trả':
                    updates['B'+n]='BTL00023 và BTL00024 đã được người dùng xác nhận ngày 28/10/2026: đưa vào import điều chỉnh theo Giá ghi trả AMIS; không tự tạo giảm giá bổ sung.'
                if name=='Lien ket BTL-BHH':
                    code=cell_value(cells.get('C'+n),strings)
                    if code in ['BHHS2024000629','BHHS2024005761']:
                        updates['R'+n]='Giá ghi trả đã xác nhận 28/10/2026';updates['S'+n]=p['link_issues'][code]
                for ref,value in updates.items():
                    old=cells.get(ref);style=old.get('s','0') if old is not None else '0'
                    if old is not None:row.remove(old)
                    row.append(cell(ref,value,style,True));changed=True
                if updates:
                    ordered=sorted(list(row),key=lambda c:(len(''.join(filter(str.isalpha,c.get('r')))),''.join(filter(str.isalpha,c.get('r')))))
                    for c in list(row):row.remove(c)
                    for c in ordered:row.append(c)
            if changed:changes[part]=xml(s)
    return package(raw,changes)

def remove_empty_pending_sheet(raw):
    """Bỏ tab chờ khi đã hết dữ liệu, không thay các sheet còn lại."""
    w=load_workbook(BytesIO(raw),read_only=True,data_only=True)
    if 'Cho KT tra lai' not in w.sheetnames:return raw
    assert not any(any(v is not None for v in row) for row in w['Cho KT tra lai'].iter_rows(min_row=2,values_only=True))
    w.close()
    with ZipFile(BytesIO(raw)) as z:
        wb=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'));ct=E.fromstring(z.read('[Content_Types].xml'))
        sh=next(s for s in wb.find(Q('sheets')) if s.get('name')=='Cho KT tra lai')
        rel=next(r for r in rels if r.get('Id')==sh.get('{'+REL+'}id'));path=rel.get('Target');part=path.lstrip('/') if path.startswith('/') else 'xl/'+path
        wb.find(Q('sheets')).remove(sh);rels.remove(rel)
        for item in list(ct):
            if item.get('PartName')=='/'+part:ct.remove(item)
        removed=[part];sr=str(Path(part).parent/'_rels'/(Path(part).name+'.rels')).replace('\\','/')
        if sr in z.namelist():removed.append(sr)
    return package(raw,{'xl/workbook.xml':xml(wb),'xl/_rels/workbook.xml.rels':xml(rels),'[Content_Types].xml':xml(ct)},removed)

def main():
    data=prepare();review=OUT/'KIEM_TRA_VA_LOAI_TRU_AMIS_2024.xlsx';sales=OUT/'IMPORT_HOA_DON_MOI_AMIS_2024.xlsx'
    sales_hash=hashlib.sha256(sales.read_bytes()).hexdigest()
    node=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    subprocess.run([str(node),str(ROOT/'scripts/build_return_adjustment_import.mjs')],check=True)
    raw=remove_empty_pending_sheet(review.read_bytes());small=(WORK/'_return-import.tmp.xlsx').read_bytes()
    with ZipFile(BytesIO(raw)) as z,ZipFile(BytesIO(small)) as src:
        wb=E.fromstring(z.read('xl/workbook.xml'));rels=E.fromstring(z.read('xl/_rels/workbook.xml.rels'));ct=E.fromstring(z.read('[Content_Types].xml'));styles=E.fromstring(z.read('xl/styles.xml'))
        sf=merge_styles(styles,E.fromstring(src.read('xl/styles.xml')))
        strings=[''.join(x.itertext()) for x in E.fromstring(src.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in src.namelist() else []
        changes={};remove=[];sheets=wb.find(Q('sheets'))
        for idx,name in enumerate(['Import dieu chinh']+(['Cho KT tra lai'] if data['pending'] else []),1):
            target=next((s for s in sheets if s.get('name') in ([name,'Tra lai cho dieu chinh'] if idx==1 else [name])),None)
            if target is None:
                sid=max(int(s.get('sheetId')) for s in sheets)+1;rid=f'UVG_Return_{sid}';part=f'xl/worksheets/sheet{sid}.xml'
                target=E.SubElement(sheets,Q('sheet'),name=name,sheetId=str(sid),attrib={'{'+REL+'}id':rid})
                E.SubElement(rels,'{'+PKG+'}Relationship',Id=rid,Type=REL+'/worksheet',Target='/'+part)
                E.SubElement(ct,'{'+CT+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
            else:
                target.set('name',name);rid=target.get('{'+REL+'}id');path=next(r.get('Target') for r in rels if r.get('Id')==rid);part=path.lstrip('/') if path.startswith('/') else 'xl/'+path
            sheet=E.fromstring(src.read(f'xl/worksheets/sheet{idx}.xml'))
            for c in sheet.findall('.//'+Q('c')):
                if c.get('s') is not None:c.set('s',str(sf[int(c.get('s'))]))
                if c.get('t')=='s':
                    v=c.find(Q('v'));text=strings[int(v.text)];c.remove(v);c.set('t','inlineStr');E.SubElement(E.SubElement(c,Q('is')),Q('t')).text=text
            # Không giữ bảng/quan hệ cũ của sheet nguồn vừa thay.
            oldrel=str(Path(part).parent/'_rels'/(Path(part).name+'.rels')).replace('\\','/')
            if oldrel in z.namelist():remove.append(oldrel)
            changes[part]=xml(sheet)
        changes.update({'xl/workbook.xml':xml(wb),'xl/_rels/workbook.xml.rels':xml(rels),'xl/styles.xml':xml(styles),'[Content_Types].xml':xml(ct)})
    result=refresh_approval_notes(package(raw,changes,remove))
    # Chỉ hai sheet ghi chú được phép đổi thêm ngoài các sheet import.
    with ZipFile(BytesIO(raw)) as old:
        bw=E.fromstring(old.read('xl/workbook.xml'));br=E.fromstring(old.read('xl/_rels/workbook.xml.rels'))
        note_parts=set()
        for sh in bw.find(Q('sheets')):
            if sh.get('name') in ['Tong hop','Lien ket BTL-BHH']:
                path=next(r.get('Target') for r in br if r.get('Id')==sh.get('{'+REL+'}id'));note_parts.add(path.lstrip('/') if path.startswith('/') else 'xl/'+path)
    with ZipFile(BytesIO(raw)) as a,ZipFile(BytesIO(result)) as b:
        for name in a.namelist():
            if name not in changes and name not in remove and name not in note_parts:assert a.read(name)==b.read(name),name
    w=load_workbook(BytesIO(result),read_only=True,data_only=True)
    assert 'Tra lai cho dieu chinh' not in w.sheetnames
    s=w['Import dieu chinh'];assert tuple(next(s.iter_rows(min_row=9,max_row=9,values_only=True)))==tuple(data['headers'])
    for got,expected in zip(s.iter_rows(min_row=10,values_only=True),data['rows']):assert list(got)==expected
    assert len(list(s.iter_rows(min_row=10,values_only=True)))==219
    assert 'Cho KT tra lai' not in w.sheetnames
    tmp=WORK/'_review-return.tmp.xlsx';tmp.write_bytes(result);os.replace(tmp,review)
    assert hashlib.sha256(sales.read_bytes()).hexdigest()==sales_hash
    report={'passed':True,'invoices':163,'lines':219,'goods':-225865926,'vat':-18069275,'payment':-243935201,'pending_btl':sorted(PENDING),'pending_lines':0,'confirmed_recorded_return_prices':['BTL00023','BTL00024'],'refs_blank':True,'adjustment_date_blank':True,'columns':'A:X đúng mẫu; Y Mã đơn hàng; Z Số chứng từ bán gốc','template_sha256':data['template_sha256'],'sales_unchanged_sha256':sales_hash,'review_sha256':hashlib.sha256(result).hexdigest(),'unrelated_parts_byte_identical':True}
    (WORK/'return-adjustment-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    format_check=WORK/'template-format-verification.json'
    if format_check.exists():
        v=json.loads(format_check.read_text(encoding='utf8'))
        for item in v.get('files',[]):
            file=Path(item['path'])
            if file==review:item['sha256']=hashlib.sha256(result).hexdigest()
        format_check.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
