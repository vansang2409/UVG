"""Audit supplemental detail for the 57 historical fallback roots; never edit workbooks."""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
from zipfile import ZipFile
import hashlib
import json
import xml.etree.ElementTree as ET
import openpyxl
from prepare_invoice_reports import rows, invoice_key, invoice_day, amount

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data/Bo sung/meInvoice/2026-10-10-57-F0'
EXPORT = SOURCE / 'Bang_ke_chi_tiet_hoa_don_da_su_dung_1791595024735.xls'
OUT = ROOT / 'scripts/data/check-review/supplement-57-20261010.json'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def norm(v):
    return ' '.join(str(v or '').split())


def first(e, name):
    return next((x for x in e.iter() if x.tag.rsplit('}', 1)[-1] == name), None)


def value(e, name):
    x = next((x for x in e if x.tag.rsplit('}', 1)[-1] == name), None)
    return x.text if x is not None else None


def line_counter(lines, zero_diagnostic=False):
    result = Counter()
    for name, unit, qty, price, goods in lines:
        numeric = (qty, price, goods)
        if zero_diagnostic:
            numeric = tuple(Decimal(0) if x is None else x for x in numeric)
        result[(norm(name), norm(unit), *numeric)] += 1
    return result


def main():
    prior = json.loads((ROOT / 'scripts/data/check-review/outputs-vs-new-data-20261009.json').read_text(encoding='utf-8'))
    exported = defaultdict(list)
    for rn, r in rows(EXPORT):
        if amount(r.get('A')) is None:
            continue
        k = invoice_key(r.get('B'), r.get('C'))
        if k:
            exported[k].append(r)
    xml = {}
    for p in sorted(SOURCE.glob('*.xml')):
        root = ET.fromstring(p.read_bytes())
        hd = first(root, 'DLHDon')
        tt, pay = first(hd, 'TTChung'), first(hd, 'TToan')
        k = invoice_key((value(tt, 'KHMSHDon') or '') + (value(tt, 'KHHDon') or ''), value(tt, 'SHDon'))
        assert k not in xml, k
        lines = []
        for l in hd.iter():
            if l.tag.rsplit('}', 1)[-1] != 'HHDVu':
                continue
            extra = {value(x, 'TTruong'): value(x, 'DLieu') for x in l.iter() if x.tag.rsplit('}', 1)[-1] == 'TTin'}
            lines.append({'name': value(l, 'THHDVu'), 'unit': value(l, 'DVTinh'), 'qty': amount(value(l, 'SLuong')), 'price': amount(value(l, 'DGia')), 'goods': amount(value(l, 'ThTien')), 'nature': value(l, 'TChat'), 'rate': value(l, 'TSuat'), 'tax': amount(extra.get('VATAmount')), 'discount': amount(value(l, 'STCKhau'))})
        xml[k] = {'path': str(p.relative_to(ROOT)), 'sha256': sha(p), 'date': invoice_day(value(tt, 'NLap')), 'money': [amount(value(pay, c)) for c in ['TgTCThue', 'TgTThue', 'TgTTTBSo']], 'lines': lines}
    records, workbooks = [], {}
    for name, previous in prior['outputs'].items():
        targets = {tuple(x['key']) for x in previous.get('import', {}).get('old_detail_fallback', [])}
        if not targets:
            continue
        p = ROOT / 'outputs' / name
        before = sha(p)
        w = openpyxl.load_workbook(p, read_only=True, data_only=True)
        headers = {}
        if 'Chi tiet nguon' in w.sheetnames:
            for v in w['Chi tiet nguon'].iter_rows(min_row=2, max_col=12, values_only=True):
                if v[2]:
                    headers[invoice_key(v[2], v[3])] = [amount(x) for x in v[9:12]]
        else:
            for v in w['Chua dieu chinh'].iter_rows(min_row=3, max_col=6, values_only=True):
                if v[0]:
                    headers[invoice_key(v[0], v[1])] = [amount(x) for x in v[3:6]]
        groups = defaultdict(list)
        for v in w['Import dieu chinh'].iter_rows(min_row=10, max_col=24, values_only=True):
            if v[0] is not None:
                groups[v[0]].append(v)
        for k in sorted(targets):
            candidates = [ls for ls in groups.values() if invoice_key(ls[0][9], ls[0][10]) == k and not (len(ls) == 1 and ls[0][16] is None)]
            assert len(candidates) == 1, (k, len(candidates))
            ls = candidates[0]
            actual = [(z[14], z[15], amount(z[16]), amount(z[17]), amount(z[18])) for z in ls]
            r = {'workbook': name, 'key': k, 'seq': ls[0][0], 'import_lines': len(ls)}
            if k in xml:
                s = xml[k]
                expected = [(l['name'], l['unit'], -l['qty'] if l['qty'] is not None else None, l['price'], -l['goods'] if l['goods'] is not None else None) for l in s['lines']]
                r.update(source_kind='XML', source=s['path'], source_sha256=s['sha256'], source_lines=len(expected), date_match=s['date'] == invoice_day(ls[0][11]), header_match=s['money'] == headers[k], line_goods_match_header=sum((l['goods'] for l in s['lines'] if l['goods'] is not None), Decimal(0)) == s['money'][0], header_arithmetic=s['money'][0] + s['money'][1] == s['money'][2], natures=dict(Counter(l['nature'] for l in s['lines'])), rates=dict(Counter(l['rate'] for l in s['lines'])), nonzero_discount=any(l['discount'] not in [None, 0] for l in s['lines']))
                r['line_tax_available'] = all(l['tax'] is not None for l in s['lines'])
                r['line_tax_match_header'] = sum((l['tax'] for l in s['lines']), Decimal(0)) == s['money'][1] if r['line_tax_available'] else None
            elif k in exported:
                src = exported[k]
                expected = [(z.get('L'), z.get('M'), -amount(z.get('N')) if amount(z.get('N')) is not None else None, amount(z.get('O')), -amount(z.get('P')) if amount(z.get('P')) is not None else None) for z in src]
                r.update(source_kind='Excel', source=str(EXPORT.relative_to(ROOT)), source_lines=len(src), date_match=all(invoice_day(z.get('D')) == invoice_day(ls[0][11]) for z in src), source_money=[sum((amount(z.get(c)) or Decimal(0) for z in src), Decimal(0)) for c in ['P', 'R', 'S']], header_match=[sum((amount(z.get(c)) or Decimal(0) for z in src), Decimal(0)) for c in ['P', 'R', 'S']] == headers[k])
            else:
                r['status'] = 'missing'
                records.append(r)
                continue
            a, b = line_counter(actual), line_counter(expected)
            r['exact_detail_match'] = a == b
            r['status'] = 'exact' if a == b else ('blank_vs_zero_only' if line_counter(actual, True) == line_counter(expected, True) else 'different')
            r['import_only'] = list((a - b).items())
            r['source_only'] = list((b - a).items())
            records.append(r)
        w.close()
        assert before == sha(p), name
        workbooks[name] = {'sha256': before, 'unchanged': True, 'targets': len(targets)}
    result = {'date': '2026-10-10', 'source_hashes': {str(p.relative_to(ROOT)): sha(p) for p in SOURCE.iterdir() if p.is_file()}, 'workbooks': workbooks, 'summary': dict(Counter(r['status'] for r in records)), 'source_kinds': dict(Counter(r.get('source_kind') for r in records)), 'records': records, 'limits': ['Blank and zero remain distinct; diagnostic comparison does not authorize replacement.', 'No cryptographic signature or live tax authority verification.', 'This audit does not modify workbooks or the 18-file base manifest.']}
    result['coverage_complete'] = len(records) == 57 and all(r['status'] != 'missing' for r in records)
    result['content_checks_passed'] = result['coverage_complete'] and all(r['status'] in ['exact', 'blank_vs_zero_only'] and r['date_match'] and r['header_match'] and r.get('line_goods_match_header', True) and r.get('line_tax_match_header') is not False and r.get('header_arithmetic', True) and not r.get('nonzero_discount', False) for r in records)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    print(json.dumps({k: v for k, v in result.items() if k not in ['source_hashes', 'records']}, ensure_ascii=False, indent=2))
    print(json.dumps([r for r in records if r['status'] == 'different' or not r.get('header_match', False) or not r.get('line_goods_match_header', True)], ensure_ascii=False, default=str))


if __name__ == '__main__':
    main()
