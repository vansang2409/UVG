"""Đọc lại file chuỗi đã xuất; so từng ô với nguồn và công thức với số dư."""
from pathlib import Path
from datetime import datetime,date
from collections import Counter
import json,sys,hashlib
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
ROOT=Path(__file__).resolve().parents[1]
DIR=ROOT/'outputs/chuoi-sau-xu-ly-20261030'
sys.stdout.reconfigure(encoding='utf-8')
p=json.loads((DIR/'prepared.json').read_text(encoding='utf-8'))
file=DIR/'CHUOI_F0_F1_SAU_XU_LY_20261029.xlsx'
wc=load_workbook(file,read_only=False,data_only=True)
wf=load_workbook(file,read_only=False,data_only=False)
assert wc.sheetnames==['Điều chỉnh','Thay thế']
counts=[];fields=0
for sp in p['sheets']:
 s=wc[sp['name']];f=wf[sp['name']];n=9+6*sp['stages'];last=9+len(sp['rows'])
 assert s.max_row==last and s.max_column==n
 assert s.freeze_panes=='E10'
 assert len(f.tables)==1
 assert list(f.tables.values())[0].ref==f'A9:{s.cell(9,n).column_letter}{last}'
 expected_formula_count=4*len(sp['rows'])+2
 assert sum(c.data_type=='f' for row in f for c in row)==expected_formula_count
 statuses=Counter();seen=set()
 for rn,record in enumerate(sp['rows'],10):
  expected=record['values'][:-1]
  for ci,v in enumerate(expected,1):
   if ci in [5,6,7,8]:continue
   got=s.cell(rn,ci).value
   if isinstance(got,(date,datetime)):got=to_excel(got)
   assert got==v,(sp['name'],rn,ci,got,v)
   assert f.cell(rn,ci).data_type!='e'
   fields+=1
  balance=[s.cell(rn,c).value for c in [5,6,7]]
  assert balance==record['expected'],(sp['name'],rn,balance)
  expectedstatus='KT loại trừ' if record['excluded'] else 'Đã về 0'
  assert s.cell(rn,8).value==expectedstatus
  statuses[expectedstatus]+=1
  root=tuple(record['keys'][0]);assert root not in seen;seen.add(root)
  for i,key in enumerate(record['keys']):
   assert isinstance(s.cell(rn,10+6*i).value,str) and len(s.cell(rn,10+6*i).value)==8
  if record['replacement']:
   assert len(record['keys'])==4
   assert record['keys'][0][0]=='1C24TUV' and record['keys'][-1][0]=='1C26MUV'
  assert all(s.cell(rn,c).data_type!='e' for c in range(1,n+1))
 assert s['E3'].value==sp['processed'] and s['H3'].value==sp['excluded']
 assert s.cell(9,n).value=='Ghi chú'
 assert not any(c.value=='Nguồn đối chiếu' for row in s for c in row)
 counts.append({'sheet':sp['name'],'rows':len(sp['rows']),'statuses':dict(statuses),'stages':sp['stages']})
for path,h in p['source_hashes'].items():
 assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==h,('Nguồn bị đổi',path)
result={'passed':True,'sheets':counts,'data_fields_compared':fields,'all_balances_and_formula_caches_correct':True,'formula_errors':0,'all_sources_unchanged':True,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()}
(DIR/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
wc.close();wf.close()
