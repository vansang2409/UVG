"""Đối chiếu số lượng theo ảnh kế toán, chỉ đọc nguồn và các workbook."""
from pathlib import Path
from collections import Counter
from datetime import date, datetime
import json, sys
from openpyxl import load_workbook
from prepare_invoice_reports import rows, invoice_key, invoice_day, amount
from filter_invoice_exceptions import is_company, tax_shaped, normal

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
headers = {}
for p in sorted((ROOT/'data/Du lieu Meinvoice/Bang_ke_hoa_don_da_su_dung_2024_2025').glob('*.xlsx')):
    for rn, r in rows(p):
        if rn < 6 or amount(r.get('A')) is None: continue
        k = invoice_key(r.get('B'), r.get('C'))
        if k:
            assert k not in headers
            headers[k] = {'key': list(k), 'date': str(invoice_day(r.get('D'))), 'status': r.get('U'), 'name': r.get('H') or r.get('L'), 'tax': r.get('J'), 'source': str(p.relative_to(ROOT)), 'row': rn}
period = {k for k,v in headers.items() if '2024-01-01'<=v['date']<='2025-06-30'}
scope_sets = {}; root_sets={}; source_rows={}; imports={}; source_roots={}
for label,fn,main in [('untreated','HOA_DON_CHUA_DIEU_CHINH_F0_2024_T6_2025.xlsx','Chua dieu chinh'),('adjusted','HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx','Dieu chinh'),('exceptions','NGOAI_LE_F0_2024_T6_2025.xlsx','Ngoai le')]:
    w=load_workbook(ROOT/'outputs'/fn,read_only=True,data_only=True)
    offset=2 if label=='exceptions' else 0
    roots={tuple(r[offset:offset+2]) for r in w[main].iter_rows(min_row=3,values_only=True) if r[offset] and r[offset+1]}
    root_sets[label]=roots
    if 'Chi tiet nguon' in w.sheetnames:
        sr=[r for r in w['Chi tiet nguon'].iter_rows(min_row=2,values_only=True) if r[0] and r[2]]
        keys={tuple(r[2:4]) for r in sr}
        for r in sr:source_roots.setdefault(tuple(r[2:4]),set()).add(tuple(r[:2]))
    else:
        sr=[];keys=roots
    source_rows[label]=sr;scope_sets[label]=keys
    if 'Import dieu chinh' in w.sheetnames:
        heads={}
        for rn,r in enumerate(w['Import dieu chinh'].values,1):
            if rn>=10 and isinstance(r[0],(int,float)):heads.setdefault(r[0],r)
        imports[label]=heads
all_source=set.union(*scope_sets.values())
all_roots=set.union(*root_sets.values())
late={k for k in all_source if k in headers and headers[k]['date']>'2025-06-30'}
not_covered=period-all_source
companies={k for k in period|late if k in headers and (is_company(headers[k]['name']) or tax_shaped(headers[k]['tax']))}
print('period',len(period),'status',dict(Counter(headers[k]['status'] for k in period)))
print('roots',len(all_roots),'source',len(all_source),'late',len(late),'groups', {k:len(v) for k,v in scope_sets.items()})
print('not_covered_period',len(not_covered),'status',dict(Counter(headers[k]['status'] for k in not_covered)))
print('direct_company_header',len(companies),'company_period',len(companies&period),'company_late',len(companies&late))
print('NOT_COVERED_EXAMPLES', [headers[k] for k in sorted(not_covered)[:3]])
print('IMPORTS', {k:len(v) for k,v in imports.items()})
for label, rs in source_rows.items():
    if not rs:continue
    print('SOURCE_TYPES',label,dict(Counter(str(r[7]).split(' cho ')[0] for r in rs)))
result={'period':len(period),'period_status':dict(Counter(headers[k]['status'] for k in period)),'roots':len(all_roots),'source_invoices':len(all_source),'late':len(late),'uncovered_period':[headers[k] for k in sorted(not_covered)],'source_group_counts':{k:len(v) for k,v in scope_sets.items()},'import_counts':{k:len(v) for k,v in imports.items()},'direct_company_count':len(companies)}
outside_edges = {}
for p in sorted((ROOT/'data/Du lieu Meinvoice').glob('*dieu_chinh_xuat*.xls')):
    for rn,r in rows(p):
        if amount(r.get('A')) is None: continue
        child=invoice_key(r.get('M'),r.get('O'))
        if child in not_covered:
            outside_edges.setdefault(child,{'invoice':list(child),'parent':list(invoice_key(r.get('B'),r.get('D'))),'parent_date':str(invoice_day(r.get('C'))),'source':str(p.relative_to(ROOT)),'row':rn})
assert len(outside_edges)==287
assert all(x['parent_date']<'2024-01-01' for x in outside_edges.values())
print('OUTSIDE_ROOT_YEARS',dict(Counter(x['parent_date'][:4] for x in outside_edges.values())))
print('COMPANY_LATE',[headers[k] for k in sorted(companies&late)])
print('EXCEPTION_MINUS_DIRECT',sorted(scope_sets['exceptions']-companies))
print('DIRECT_MINUS_EXCEPTION',sorted(companies-scope_sets['exceptions']))
replacements={k for k in period if 'thay the' in normal(headers[k]['status'])}
print('REPLACEMENT_OVERLAP_EXCEPTION',sorted(replacements&scope_sets['exceptions']))
wa=load_workbook(ROOT/'outputs/HOA_DON_DA_DIEU_CHINH_F0_2024_T6_2025.xlsx',read_only=True,data_only=True)
matched_roots={tuple(r[:2]) for r in wa['42 - Da khop 31'].iter_rows(min_row=4,values_only=True) if r[0]}
kept_roots={tuple(r[:2]) for r in wa['42 - Can check 11'].iter_rows(min_row=4,values_only=True) if r[0] and 'giu nguyen' in normal(r[25])}
matched_keys={tuple(r[2:4]) for r in source_rows['adjusted'] if tuple(r[:2]) in matched_roots}
kept_keys={tuple(r[2:4]) for r in source_rows['adjusted'] if tuple(r[:2]) in kept_roots}
print('MATCHED',len(matched_keys),'SAME_PRODUCT_KEPT',len(kept_keys))
cancelled={k for k in period if 'huy' in normal(headers[k]['status'])}
groups={'cancelled':cancelled,'replacement_outside_exception':replacements-scope_sets['exceptions'],'root_before_2024':set(outside_edges),'exceptions':scope_sets['exceptions'],'matched_31':matched_keys,'same_product_3':kept_keys,'name_only_9186':{('1C25TUV','00009186')}}
assert all(not a&b for i,a in enumerate(groups.values()) for b in list(groups.values())[i+1:])
universe=period|late
old_to_handle=universe-set.union(*groups.values())
assert len(old_to_handle)==sum(map(len,imports.values()))==10442
result.update(universe_count=len(universe),deduction_counts={k:len(v) for k,v in groups.items()},outside_scope_adjustments=list(outside_edges.values()),source_to_handle_count=len(old_to_handle),company_late=[headers[k] for k in sorted(companies&late)],replacement_exception_overlap=[headers[k] for k in sorted(replacements&scope_sets['exceptions'])],note='Ảnh KT có khoản trừ2 hóa đơn hủy nội bộ, nhưng chưa có danh sách định danh để đối chiếu khoản này; không tự trừ thêm vào nguồn. Phân loại theo các tập khóa không giao nhau, không sửa import.')
result['deduction_invoice_lists']={n:[headers.get(k,{'key':list(k)}) for k in sorted(v)] for n,v in groups.items()}
print('BRIDGE',len(universe),result['deduction_counts'],'=',len(old_to_handle))
out=ROOT/'scripts/data/count-difference-20261019';out.mkdir(exist_ok=True)
(out/'initial-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
