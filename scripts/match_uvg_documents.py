"""Đối chiếu ứng viên; không tự xác nhận liên kết hoặc lập hóa đơn."""
from pathlib import Path
import os
from collections import defaultdict, Counter
from datetime import datetime
from decimal import Decimal
from openpyxl import load_workbook
import re, csv

BASE=Path(os.environ.get('UVG_DATA_DIR', str(Path(__file__).resolve().parents[1] / 'data')))
OUT=Path(os.environ.get('UVG_OUTPUT_DIR', str(Path(__file__).resolve().parents[1] / 'outputs')))/'uvg-doi-chieu'
OUT.mkdir(parents=True,exist_ok=True)
def number(v):
    return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
def model(text):
    # Chỉ chuẩn hóa họ model và loại máy/lõi; không suy ra biến thể Pro/Light.
    t=str(text or '').upper()
    models=sorted(set(re.findall(r'\b(?:KA|KH|KM|KC)\s*[-_]?\s*\d{2,4}',t)))
    models=[re.sub(r'[\s_-]','',m) for m in models]
    if not models:
        m=re.search(r'(KA|KH|KM|KC)\d{2,4}',t)
        if m: models=[m.group(0)]
    machine_title=bool(re.match(r'^\s*(MÁY|ĐÈN)\b',t))
    isfilter=not machine_title and bool(re.search(r'FILTER|MÀNG LỌC|LÕI LỌC|BỘ LỌC|FT\s*[-:]|FT-',t))
    return ('FILTER:' if isfilter else 'MACHINE:')+','.join(models) if models else None
def signature(lines, strict=False):
    parts=[]
    for l in lines:
        if l['model'] is None: return None
        if strict and not l['complete']: return None
        parts.append((l['model'],l['qty'],l['sales'],l['vat']) if strict else (l['model'],l['qty']))
    return tuple(sorted(parts))
def csvwrite(name,headers,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f); writer.writerow(headers); writer.writerows(rows)

docs={}; mapping=Counter()
for p in sorted((BASE/'Du lieu Amis').glob('*.xlsx')):
    s=load_workbook(p,read_only=True,data_only=True).worksheets[0]
    for i,r in enumerate(s.values,1):
        if i<5 or not isinstance(r[1],datetime) or not r[2]: continue
        key=(r[1].year,str(r[2]))
        d=docs.setdefault(key,{'date':r[1],'file':p.name,'rows':[],'lines':[],'total':Decimal(0),'vat':Decimal(0),'sales':Decimal(0),'customer':str(r[9] or ''),'dates':set()})
        d['dates'].add(r[1]); d['rows'].append(i)
        m=model(str(r[12])+' '+str(r[13])); mapping[(str(r[12]),str(r[13]),m)]+=1
        d['lines'].append({'model':m,'qty':number(r[17]),'sales':number(r[20])-number(r[25]),'vat':number(r[32]),'complete':isinstance(r[20],(int,float)) and isinstance(r[32],(int,float))})
        d['total']+=number(r[34]); d['vat']+=number(r[32]); d['sales']+=number(r[20])-number(r[25])
print('AMIS loaded',len(docs),flush=True)
invoices={}; altered=set(); links=[]
for p in sorted((BASE/'Du lieu Meinvoice').glob('*.xls')):
    if not p.name.startswith(('Bang_ke','DS')): continue
    s=load_workbook(p.open('rb'),read_only=True,data_only=True).worksheets[0]
    for i,r in enumerate(s.values,1):
        if not isinstance(r[0],(int,float)): continue
        if p.name.startswith('Bang_ke') and i>=6:
            key=(str(r[1]),str(r[2]))
            invoices[key]={'date':datetime.strptime(r[3],'%d/%m/%Y'),'total':number(r[17]),'vat':number(r[16]),'sales':number(r[15]),'status':r[20],'file':p.name,'row':i,'complete':isinstance(r[17],(int,float)),'lines':[]}
        elif p.name.startswith('DS') and i>=7:
            old=(str(r[1]),str(r[3])); new=(str(r[12]),str(r[14])); altered.add(old)
            links.append((old,new,'adjust' if 'điều chỉnh' in p.name else 'replace'))
for p in sorted((BASE/'Du lieu Meinvoice').glob('Bảng kê chi tiết*.xls')):
    s=load_workbook(p.open('rb'),read_only=True,data_only=True).worksheets[0]
    for i,r in enumerate(s.values,1):
        if i<6 or not isinstance(r[0],(int,float)): continue
        key=(str(r[1]),str(r[2]))
        if key in invoices:
            invoices[key]['lines'].append({'model':model(r[11]),'qty':number(r[13]),'sales':number(r[15]),'vat':number(r[17]),'complete':isinstance(r[15],(int,float)) and isinstance(r[17],(int,float)),'file':p.name,'row':i})
print('MEINVOICE loaded',len(invoices),flush=True)
index=defaultdict(list); excluded=Counter(); eligible={}; bad_detail=[]
for key,h in invoices.items():
    if h['status'] not in ('Hoá đơn mới','Hoá đơn thay thế') or key in altered:
        excluded['Có lịch sử sửa hoặc không phải hóa đơn mới/thay thế hiện hành']+=1; continue
    if not h['complete'] or not h['lines']:
        excluded['Thiếu tổng tiền hoặc chi tiết']+=1; continue
    if sum((l['sales'] for l in h['lines']),Decimal(0))!=h['sales'] or sum((l['vat'] for l in h['lines']),Decimal(0))!=h['vat']:
        bad_detail.append(key); excluded['Tổng chi tiết không khớp bảng tổng']+=1; continue
    h['signature']=signature(h['lines'])
    h['strict_signature']=signature(h['lines'],True)
    eligible[key]=h
    index[(h['sales'],h['vat'],h['total'])].append(key)

group_index=defaultdict(list); strict_index=defaultdict(list); group_ids={}
for k,h in eligible.items():
    mk=(h['sales'],h['vat'],h['total'])
    gk=(mk,h['signature'])
    group_index[gk].append(k)
    strict_index[(mk,h['strict_signature'])].append(k)
for idx,gk in enumerate(group_index,1): group_ids[gk]=f'G{idx:05d}'
results={}; reverse=Counter(); pairs=[]
for key,d in docs.items():
    if not key[1].startswith(('BHH','BHDV')) or d['total']<=0 or len(d['dates'])!=1:
        results[key]={'status':'Trả lại/giảm giá hoặc chứng từ cần kiểm tra riêng','candidates':[],'strict':[]}; continue
    mk=(d['sales'],d['vat'],d['total']); raw=index.get(mk,[])
    sig=signature(d['lines']); strictsig=signature(d['lines'],True)
    candidates=group_index.get((mk,sig),[]) if sig is not None else []
    strict=strict_index.get((mk,strictsig),[]) if strictsig is not None else []
    for k in candidates: reverse[k]+=1
    results[key]={'candidates':candidates,'strict':strict,'raw_count':len(raw),'group':group_ids.get((mk,sig),'')}
for key,d in docs.items():
    z=results[key]
    if 'status' not in z:
        c=z['candidates']; strict=z['strict']
        if len(c)==1 and len(strict)==1 and reverse[c[0]]==1:
            gap=abs((d['date']-eligible[c[0]]['date']).days)
            z['status']='Ứng viên duy nhất: khớp từng dòng và cùng ngày' if gap==0 else 'Ứng viên duy nhất: khớp từng dòng, khác ngày'
        elif c: z['status']='Có ứng viên nhưng trùng lặp hoặc chỉ khớp mức tổng/model'
        elif z['raw_count']: z['status']='Chỉ khớp tiền: chưa khớp nhóm model và số lượng'
        else: z['status']='Chưa tìm thấy ứng viên trong tập hóa đơn đủ dữ liệu'
    for k in (z['candidates'] if z['status'].startswith('Ứng viên duy nhất') else []):
        h=eligible[k]
        pairs.append([key[0],key[1],d['date'].strftime('%d/%m/%Y'),k[0],k[1],h['date'].strftime('%d/%m/%Y'),(h['date']-d['date']).days,k in z['strict'],reverse[k],d['sales'],d['vat'],d['total'],d['file'],','.join(map(str,d['rows'])),h['file'],h['row'],z['status']])
summary=Counter(z['status'] for z in results.values())
csvwrite('AMIS_PHAN_NHOM_DOI_CHIEU.csv',['Năm','Số chứng từ','Ngày chứng từ','Tiền hàng sau chiết khấu','Thuế GTGT','Tổng thanh toán','Nhóm đối chiếu','Số ứng viên khớp tổng/model/số lượng','Số ứng viên khớp từng dòng','Số HĐ khớp bộ tiền trước kiểm tra model','File AMIS','Dòng AMIS','Mã nhóm ứng viên','Ghi chú'],[
    [k[0],k[1],d['date'].strftime('%d/%m/%Y'),d['sales'],d['vat'],d['total'],results[k]['status'],len(results[k]['candidates']),len(results[k]['strict']),results[k].get('raw_count',''),d['file'],','.join(map(str,d['rows'])),results[k].get('group',''),'Chưa xác nhận mã đơn và biến thể sản phẩm; chưa phải chỉ định xuất hóa đơn'] for k,d in docs.items()])
csvwrite('HOA_DON_THEO_NHOM_UNG_VIEN.csv',['Mã nhóm ứng viên','Ký hiệu HĐ','Số HĐ','Ngày HĐ','Tiền hàng','Thuế','Tổng tiền','File meInvoice','Dòng'],[[group_ids[gk],k[0],k[1],eligible[k]['date'].strftime('%d/%m/%Y'),eligible[k]['sales'],eligible[k]['vat'],eligible[k]['total'],eligible[k]['file'],eligible[k]['row']] for gk,keys in group_index.items() for k in keys])
csvwrite('UNG_VIEN_AMIS_MEINVOICE.csv',['Năm AMIS','Số chứng từ AMIS','Ngày AMIS','Ký hiệu HĐ','Số HĐ','Ngày HĐ','Lệch ngày HĐ-AMIS','Khớp từng dòng','Số chứng từ AMIS cùng ứng viên HĐ','Tiền hàng','Thuế','Tổng tiền','File AMIS','Dòng AMIS','File bảng tổng meInvoice','Dòng meInvoice','Nhóm đối chiếu'],pairs)
csvwrite('ANH_XA_MODEL_AMIS.csv',['Mã hàng AMIS','Tên hàng AMIS','Họ model suy ra','Số dòng','Giới hạn'],[[sku,name,m,count,'Chỉ phân biệt họ model và máy/lõi; chưa xác nhận biến thể Pro/Light hoặc SKU tương đương'] for (sku,name,m),count in mapping.items()])
report=['# UVG – đối chiếu chứng từ AMIS và hóa đơn meInvoice','',
'## Kết quả','',f'- Đã đọc {len(docs):,} chứng từ AMIS và {len(invoices):,} hóa đơn meInvoice trong bảng tổng.',
f'- {len(eligible):,} hóa đơn đủ điều kiện tìm ứng viên trực tiếp; hóa đơn đã bị sửa, thiếu dữ liệu hoặc lệch tổng chi tiết được giữ riêng.',
'- Chưa có liên kết nào được xác nhận bằng mã đơn hàng. Không có chứng từ nào được tự kết luận là cần xuất mới.','',
'| Nhóm | Số chứng từ AMIS |','|---|---:|']
for label,count in summary.items(): report.append(f'| {label} | {count:,} |')
report+=['','## Cách đối chiếu','',
'1. Gộp dòng AMIS theo năm và số chứng từ. Loại dòng Tổng cộng khỏi dữ liệu. Tính tiền hàng sau chiết khấu, thuế và thanh toán.',
'2. Với meInvoice, chỉ tìm ứng viên trực tiếp trong hóa đơn mới/thay thế chưa thấy liên kết sửa tiếp, có bảng tổng và chi tiết khớp. Hóa đơn gộp bị điều chỉnh cần phân tích chuỗi riêng, không bị coi là không tồn tại.',
'3. Yêu cầu khớp chính xác ba số: tiền hàng, thuế, tổng thanh toán. Sau đó so cùng số dòng, họ model máy/lõi và số lượng từng dòng. Nhóm khớp từng dòng còn yêu cầu tiền hàng và thuế từng dòng bằng nhau.',
'4. Kiểm tra số ứng viên theo cả hai chiều. Một hóa đơn có thể cùng khớp nhiều chứng từ; không tự chọn hóa đơn gần ngày nhất.',
'5. Cùng ngày hoặc ứng viên duy nhất chỉ làm tăng mức ưu tiên kiểm tra, không chứng minh cùng đơn hàng. Ngày AMIS là ngày chứng từ, không được tự coi là ngày giao hàng.',
'','## Giới hạn quan trọng','',
'- AMIS do khách hàng chốt là số liệu mục tiêu theo xác nhận của người dùng. Phép ghép chưa chứng minh tính pháp lý của phương án sửa.',
'- Chuẩn hóa sản phẩm mới tới họ model và loại máy/lõi. Các biến thể Pro/Light, hàng tặng hoặc cách tách giá máy/lõi có thể khác nhau; không tự suy ra tương đương SKU.',
'- Không ghép nhiều đơn thành hóa đơn gộp bằng việc chọn tổ hợp số tiền. Khi thiếu bảng kê đơn gốc, cách đó có nhiều đáp án và không tạo bằng chứng chắc chắn.',
'- Thiếu mã đơn sàn, CQT và lịch sử phát hành/sửa sau tháng 6/2025. Chưa tìm thấy ứng viên không đồng nghĩa chưa xuất hóa đơn.',
'- CSV là dữ liệu phân tích để kiểm tra, không phải file import meInvoice.','',
'## Hóa đơn giữ riêng khỏi bước ghép trực tiếp','']
for label,count in excluded.items(): report.append(f'- {label}: {count:,}.')
report+=['','## Ví dụ ứng viên duy nhất khớp từng dòng','', '| Chứng từ AMIS | Ngày AMIS | Hóa đơn | Ngày HĐ | Tổng thanh toán |','|---|---|---|---|---:|']
examples=0
for k,z in results.items():
    if z['status'].startswith('Ứng viên duy nhất') and examples<10:
        d=docs[k]; hk=z['candidates'][0]; h=eligible[hk]
        report.append(f'| {k[1]} | {d["date"]:%d/%m/%Y} | {hk[0]} / {hk[1]} | {h["date"]:%d/%m/%Y} | {d["total"]:,.0f} |'); examples+=1
report+=['','## Việc cần làm tiếp','',
'- Kiểm tra từng ứng viên duy nhất bằng mã đơn sàn hoặc chứng từ giao dịch gốc và bảng ánh xạ SKU được khách hàng xác nhận.',
'- Với hóa đơn gộp/đã sửa: tính số dư theo toàn bộ chuỗi, đối chiếu với nhóm giao dịch thực tế; chưa lập file giảm hoặc xuất lại khi chưa xác định được nhóm giao dịch.',
'- Bổ sung CQT và các lần sửa sau tháng 6/2025 trước khi chốt danh sách phát hành.','',
'## File kết quả','', '- AMIS_PHAN_NHOM_DOI_CHIEU.csv: một dòng cho mỗi chứng từ AMIS, có nguồn và dòng.',
'- UNG_VIEN_AMIS_MEINVOICE.csv: các cặp ứng viên duy nhất khớp từng dòng, có nguồn, dòng và độ lệch ngày.',
'- HOA_DON_THEO_NHOM_UNG_VIEN.csv: danh mục hóa đơn theo mã nhóm ứng viên. Nối với cột Mã nhóm ứng viên của bảng AMIS để xem đầy đủ các hóa đơn có thể khớp; không chọn một cặp khi nhóm có nhiều đáp án.',
'- ANH_XA_MODEL_AMIS.csv: họ model suy ra từ mã và tên hàng để rà soát.']
(OUT/'UVG_DOI_CHIEU_CHUNG_TU.md').write_text('\n'.join(report),encoding='utf-8')
assert sum(summary.values())==len(docs)
assert len(eligible)+sum(excluded.values())==len(invoices)
print('SUMMARY',dict(summary),flush=True)
print('ELIGIBLE',len(eligible),'EXCLUDED',dict(excluded),'PAIRS',len(pairs),'OUT',OUT,flush=True)
